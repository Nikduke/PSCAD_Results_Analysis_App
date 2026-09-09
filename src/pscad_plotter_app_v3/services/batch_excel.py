from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pscad_plotter_app_v3.models import (
    DEFAULT_TOV_WINDOW_COUNT,
    DEFAULT_TOV_WINDOW_S,
    MMElementRecord,
    PlotMode,
    PlotRequest,
    ProjectCatalog,
)
@dataclass(slots=True)
class BatchExcelError:
    row_number: int
    message: str
    sheet_name: str = ""


@dataclass(slots=True)
class BatchExcelImportResult:
    requests: list[PlotRequest]
    errors: list[BatchExcelError]


class BatchExcelService:
    """Import the MM waveform batches supported by this application."""

    MM_SHEET = "MM"
    TIME_RANGE_HEADERS = ["time_start_s", "time_end_s"]
    MM_HEADERS = [
        "case",
        "run",
        "element",
        "trace",
        "overview",
        "tov_windows",
        "tov_window_s",
        "tov_window_count",
        "limits",
        "legends_left",
        "excel_export",
        "annotate_max",
        "annotate_min",
        *TIME_RANGE_HEADERS,
    ]
    SHEET_DEFINITIONS = {MM_SHEET: (PlotMode.MM, MM_HEADERS)}

    def load_requests(self, path: str | Path, catalog: ProjectCatalog) -> BatchExcelImportResult:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if self.MM_SHEET not in workbook.sheetnames:
                return BatchExcelImportResult([], [BatchExcelError(1, "Missing required sheet 'MM'.", self.MM_SHEET)])
            rows = list(workbook[self.MM_SHEET].iter_rows(values_only=True))
        finally:
            workbook.close()

        if not rows:
            return BatchExcelImportResult([], [BatchExcelError(1, "Missing header row.", self.MM_SHEET)])

        header_map = self._header_map(rows[0])
        missing = [name for name in ("case", "run", "element") if name not in header_map]
        if missing:
            message = f"Missing required column(s): {', '.join(missing)}."
            return BatchExcelImportResult([], [BatchExcelError(1, message, self.MM_SHEET)])

        requests: list[PlotRequest] = []
        errors: list[BatchExcelError] = []
        for row_number, values in enumerate(rows[1:], start=2):
            if self._row_is_empty(values):
                continue
            row = {
                name: self._cell_text(values[position] if position < len(values) else None)
                for name, position in header_map.items()
            }
            try:
                requests.append(self._request_from_row(row, catalog))
            except ValueError as exc:
                errors.append(BatchExcelError(row_number, str(exc), self.MM_SHEET))
        return BatchExcelImportResult(requests, errors)

    def _request_from_row(self, row: dict[str, str], catalog: ProjectCatalog) -> PlotRequest:
        case_name = row.get("case", "").strip()
        run_number = self._parse_run(row.get("run", ""))
        if not case_name:
            raise ValueError("Missing case.")
        if case_name not in catalog.runs_by_case:
            raise ValueError(f"Unknown case '{case_name}'.")
        if run_number not in set(catalog.runs_by_case.get(case_name, [])):
            raise ValueError(f"Run {run_number} is not available for case '{case_name}'.")

        element = row.get("element", "").strip()
        if not element:
            raise ValueError("Missing MM element.")
        record = self._find_mm_record(catalog.mm_elements, case_name, run_number, element)
        if record is None:
            raise ValueError(f"MM element '{element}' is not available for {case_name} Run {run_number}.")
        return PlotRequest(
            mode=PlotMode.MM,
            case_name=case_name,
            run_numbers=[run_number],
            elements=[element],
            voltage_kv=record.voltage_kv,
            trace_type=self._parse_trace(row.get("trace", "")),
            show_three_phase_overview=self._parse_bool(row.get("overview", ""), default=True),
            show_tov_windows=self._parse_bool(row.get("tov_windows", ""), default=True),
            tov_window_s=self._parse_float(row.get("tov_window_s", ""), default=DEFAULT_TOV_WINDOW_S),
            tov_window_count=self._parse_int(row.get("tov_window_count", ""), default=DEFAULT_TOV_WINDOW_COUNT),
            show_limits=self._parse_bool(row.get("limits", ""), default=True),
            legends_left=self._parse_bool(row.get("legends_left", ""), default=False),
            excel_export=self._parse_bool(row.get("excel_export", ""), default=False),
            annotate_max=self._parse_bool(row.get("annotate_max", ""), default=False),
            annotate_min=self._parse_bool(row.get("annotate_min", ""), default=False),
            plot_variant=(row.get("plot_variant", "").strip() or None),
            **self._time_range_kwargs(row),
        )

    @staticmethod
    def _find_mm_record(
        records: list[MMElementRecord],
        case_name: str,
        run_number: int,
        element: str,
    ) -> MMElementRecord | None:
        return next(
            (
                record
                for record in records
                if record.case_name == case_name
                and record.element_name == element
                and run_number in record.available_runs
            ),
            None,
        )

    @staticmethod
    def _parse_run(value: str) -> int:
        try:
            run_number = int(float(value))
        except (TypeError, ValueError):
            raise ValueError(f"Invalid run '{value}'.")
        if run_number < 1:
            raise ValueError("Run must be 1 or greater.")
        return run_number

    @staticmethod
    def _parse_trace(value: str) -> str:
        token = value.strip()
        if not token:
            return "Both"
        normalized = token.lower().replace(" ", "")
        if normalized in {"both", "combined", "combinedlgp&llp", "combinedlgpllp"}:
            return "Both"
        if normalized == "lgp":
            return "LGp"
        if normalized == "llp":
            return "LLp"
        if normalized == "lgr":
            return "LGr"
        if normalized == "llr":
            return "LLr"
        raise ValueError(f"Invalid MM trace '{value}'. Use Both, LGp, LLp, LGr, or LLr.")

    @staticmethod
    def _parse_bool(value: str, *, default: bool) -> bool:
        token = value.strip().lower()
        if not token:
            return default
        if token in {"1", "true", "yes", "y", "on"}:
            return True
        if token in {"0", "false", "no", "n", "off"}:
            return False
        raise ValueError(f"Invalid boolean value '{value}'.")

    @staticmethod
    def _parse_float(value: str, *, default: float) -> float:
        token = value.strip()
        if not token:
            return default
        try:
            return float(token)
        except ValueError:
            raise ValueError(f"Invalid numeric value '{value}'.")

    @classmethod
    def _time_range_kwargs(cls, row: dict[str, str]) -> dict[str, float | None]:
        start_s = cls._parse_optional_float(row.get("time_start_s", ""), "time_start_s")
        end_s = cls._parse_optional_float(row.get("time_end_s", ""), "time_end_s")
        if start_s is not None and end_s is not None and end_s <= start_s:
            raise ValueError("time_end_s must be greater than time_start_s.")
        return {"time_start_s": start_s, "time_end_s": end_s}

    @staticmethod
    def _parse_optional_float(value: str, column_name: str) -> float | None:
        token = value.strip()
        if not token:
            return None
        try:
            return float(token)
        except ValueError:
            raise ValueError(f"Invalid numeric value in '{column_name}': '{value}'.")

    @staticmethod
    def _parse_int(value: str, *, default: int) -> int:
        token = value.strip()
        if not token:
            return default
        try:
            parsed = int(float(token))
        except ValueError:
            raise ValueError(f"Invalid integer value '{value}'.")
        if parsed < 1:
            raise ValueError("TOV window count must be 1 or greater.")
        return parsed

    @staticmethod
    def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
        aliases = {
            "case_name": "case",
            "run_number": "run",
            "show_three_phase_overview": "overview",
            "show_tov_windows": "tov_windows",
            "tov_windows_shown": "tov_window_count",
            "show_limits": "limits",
            "legend_left": "legends_left",
            "legend_side_left": "legends_left",
            "export_excel": "excel_export",
            "excel": "excel_export",
            "start_time_s": "time_start_s",
            "plot_start_s": "time_start_s",
            "start_s": "time_start_s",
            "end_time_s": "time_end_s",
            "plot_end_s": "time_end_s",
            "end_s": "time_end_s",
        }
        mapping: dict[str, int] = {}
        for index, value in enumerate(header_row):
            name = str(value or "").strip().lower().replace(" ", "_").replace("-", "_")
            if name:
                mapping[aliases.get(name, name)] = index
        return mapping

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "true" if value else "false"
        return str(value).strip()

    @staticmethod
    def _row_is_empty(values: tuple[Any, ...]) -> bool:
        return not any(str(value).strip() for value in values if value is not None)
