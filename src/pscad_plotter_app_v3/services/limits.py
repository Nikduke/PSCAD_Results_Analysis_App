from __future__ import annotations

from pathlib import Path
import math
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook

from pscad_plotter_app_v3.models import ProjectContext, VoltageLimitSet
from pscad_plotter_app_v3.services.json_store import read_json_file


class LimitService:
    """Load workbook MM limits and merge project-level overrides."""

    LIMITS_FILENAME = "limits.json"

    def load_effective_limits(self, context: ProjectContext) -> dict[str, VoltageLimitSet]:
        workbook_limits = self.load_workbook_limits(context.workbook_path)
        override_limits = self.load_override_limits(context.state_dir / self.LIMITS_FILENAME)
        merged = dict(workbook_limits)
        for key, value in override_limits.items():
            merged[key] = value
        return merged

    def load_workbook_limits(self, workbook_path: Path | None) -> dict[str, VoltageLimitSet]:
        workbook_data, _error = self._read_mm_blocks(workbook_path)
        if workbook_data is None:
            return {}
        headers, rows = workbook_data
        required = ["Un", "SDPF_LG", "SDPF_LL", "SIWL_LG", "SIWL_LL"]
        if any(column not in headers for column in required):
            return {}

        limits: dict[str, VoltageLimitSet] = {}
        for row in rows[1:]:
            if not row:
                continue
            try:
                voltage = float(row[headers["Un"]])
                sdpf_lg = float(row[headers["SDPF_LG"]])
                sdpf_ll = float(row[headers["SDPF_LL"]])
                siwl_lg = float(row[headers["SIWL_LG"]])
                siwl_ll = float(row[headers["SIWL_LL"]])
            except (TypeError, ValueError):
                continue
            key = f"{voltage:g}"
            limits.setdefault(
                key,
                VoltageLimitSet(
                    voltage_kv=voltage,
                    sdpf_lg=sdpf_lg,
                    sdpf_ll=sdpf_ll,
                    siwl_lg=siwl_lg,
                    siwl_ll=siwl_ll,
                    source="workbook",
                ),
            )
        return limits

    def load_workbook_limits_with_warnings(
        self,
        workbook_path: Path | None,
    ) -> tuple[dict[str, VoltageLimitSet], list[str]]:
        """Read SDPF rows with explicit validation for sustained-stress checks."""
        workbook_data, error = self._read_mm_blocks(workbook_path)
        if workbook_data is None:
            return {}, [error or "Could not read SDPF limits workbook."]
        headers, rows = workbook_data
        warnings: list[str] = []
        required = ["Un", "SDPF_LG", "SDPF_LL"]
        missing = [column for column in required if column not in headers]
        if missing:
            return {}, [f"Missing SDPF limit column(s): {', '.join(missing)}."]

        values_by_voltage: dict[str, list[tuple[float, float]]] = {}
        invalid_voltages: set[str] = set()
        for row_number, row in enumerate(rows[1:], start=2):
            if not row:
                continue
            try:
                voltage = float(row[headers["Un"]])
            except (TypeError, ValueError):
                warnings.append(f"Invalid voltage in MM_blocks row {row_number}; row skipped.")
                continue
            key = f"{voltage:g}"
            try:
                lg = float(row[headers["SDPF_LG"]])
                ll = float(row[headers["SDPF_LL"]])
            except (TypeError, ValueError):
                warnings.append(f"Invalid SDPF values in MM_blocks row {row_number}; row skipped.")
                invalid_voltages.add(key)
                continue
            if not all(math.isfinite(value) and value > 0 for value in (voltage, lg, ll)):
                warnings.append(f"Non-positive or non-finite SDPF values in MM_blocks row {row_number}; row skipped.")
                invalid_voltages.add(key)
                continue
            values_by_voltage.setdefault(key, []).append((lg, ll))

        limits: dict[str, VoltageLimitSet] = {}
        for key, values in values_by_voltage.items():
            if key in invalid_voltages:
                warnings.append(f"Invalid SDPF values for {key} kV; assessment skipped.")
                continue
            first = values[0]
            if any(abs(lg - first[0]) > 1e-9 or abs(ll - first[1]) > 1e-9 for lg, ll in values[1:]):
                warnings.append(f"Inconsistent SDPF values for {key} kV; assessment skipped.")
                continue
            voltage = float(key)
            limits[key] = VoltageLimitSet(
                voltage_kv=voltage,
                sdpf_lg=first[0],
                sdpf_ll=first[1],
                siwl_lg=math.nan,
                siwl_ll=math.nan,
                source="workbook",
            )
        if not limits and not warnings:
            warnings.append("No valid SDPF limits were found in MM_blocks.")
        return limits, warnings

    @staticmethod
    def _read_mm_blocks(
        workbook_path: Path | None,
    ) -> tuple[tuple[dict[str, int], list[tuple[Any, ...]]] | None, str | None]:
        if workbook_path is None or not workbook_path.exists():
            return None, "Input workbook with MM_blocks SDPF limits was not found."
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        except (OSError, BadZipFile, KeyError, ValueError):
            return None, f"Could not read SDPF limits workbook: {workbook_path.name}."
        try:
            if "MM_blocks" not in workbook.sheetnames:
                return None, "MM_blocks sheet is missing; SDPF assessment skipped."
            rows = list(workbook["MM_blocks"].iter_rows(values_only=True))
            if not rows:
                return None, "MM_blocks sheet is empty; SDPF assessment skipped."
            headers = {
                str(value).strip(): index
                for index, value in enumerate(rows[0])
                if value is not None and str(value).strip()
            }
            return (headers, rows), None
        finally:
            workbook.close()

    def load_override_limits(self, path: Path) -> dict[str, VoltageLimitSet]:
        if not path.exists():
            return {}
        payload = read_json_file(path)
        if not isinstance(payload, dict):
            return {}
        return {key: VoltageLimitSet.from_dict(value) for key, value in payload.items()}

    def get_limit_for_voltage(self, limits: dict[str, VoltageLimitSet], voltage_kv: float | None) -> VoltageLimitSet | None:
        if voltage_kv is None:
            return None
        return limits.get(f"{float(voltage_kv):g}")
