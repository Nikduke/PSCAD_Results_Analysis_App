from __future__ import annotations

from collections.abc import Callable, Iterable, Mapping
from contextlib import ExitStack, nullcontext
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import time
from typing import Any

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from results_analysis_app import resonance_checks, rms_analysis, storage, sustained_sdpf, sustained_sdpf_heatmap
from results_analysis_app.common import (
    CancelFn,
    LogFn,
    as_float,
    check_cancel as _cancel,
    log_message as _log,
)
from results_analysis_app.envelope_rows import nearest_rows, row_value
from results_analysis_app.excel import EXCEL_AUTOMATION_ERRORS, excel_app
from results_analysis_app.models import ScopeEntry
from results_analysis_app.project_config import DEFAULT_EVENT_TIMES
from pscad_plotter_app_v3.services.project_conventions import KNOWN_FAULT_LABELS


IMAGE_WIDTH_CM = 15.92
REPORT_HEADING_GREEN = RGBColor(112, 155, 50)
BUS_VOLTAGE_SLICER_SOURCE_TEXT = "Bus voltage [kV]"
BUS_VOLTAGE_SLICER_VALUE_OVERRIDES = {"22": "23"}
DASHBOARD_FILTER_SETTLE_S = 0.8
DASHBOARD_REPORT_TITLE_ORDER = {
    "initial voltages": 0,
    "initial reactive power": 1,
    "initial active power": 2,
}
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
JPEG_SIGNATURE = b"\xff\xd8"
WORD_2012_NAMESPACE = "http://schemas.microsoft.com/office/word/2012/wordml"
WORD_2016_CID_NAMESPACE = "http://schemas.microsoft.com/office/word/2016/wordml/cid"
REPORT_MANIFEST_VERSION = 2
SUSTAINED_REPORT_LAYOUT_VERSION = 3
RMS_REPORT_LAYOUT_VERSION = 2
LEGACY_REPORT_MANIFEST_FILENAME = ".report_manifest.json"

@dataclass(frozen=True)
class EnvelopeSummaryRow:
    event: str
    measurement: str
    time_s: float
    peak_kv: float
    rms_kv: float


@dataclass(frozen=True)
class FigureReference:
    bookmark: str
    fallback_label: str


@dataclass(frozen=True)
class TableReference:
    bookmark: str
    fallback_label: str


def _configure_report_styles(doc) -> None:
    """Apply the report typography without depending on a Word template."""
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri Light"
    normal.font.size = Pt(12)

    body = styles["Body Text"]
    body.font.name = "Calibri Light"
    body.font.size = Pt(12)
    body.paragraph_format.space_before = Pt(12)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.2
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.next_paragraph_style = body

    heading_1 = styles["Heading 1"]
    heading_1.font.name = "Calibri"
    heading_1.font.size = Pt(20)
    heading_1.font.bold = None
    heading_1.font.color.rgb = REPORT_HEADING_GREEN
    heading_1.paragraph_format.space_before = Pt(8)
    heading_1.paragraph_format.space_after = Pt(12)
    heading_1.paragraph_format.keep_with_next = True
    heading_1.paragraph_format.keep_together = True
    heading_1.next_paragraph_style = body

    heading_2 = styles["Heading 2"]
    heading_2.font.name = "Calibri Light"
    heading_2.font.size = Pt(16)
    heading_2.font.bold = None
    heading_2.font.color.rgb = REPORT_HEADING_GREEN
    heading_2.paragraph_format.space_before = Pt(18)
    heading_2.paragraph_format.space_after = Pt(6)
    heading_2.paragraph_format.keep_with_next = True
    heading_2.paragraph_format.keep_together = True
    heading_2.next_paragraph_style = body

    heading_3 = styles["Heading 3"]
    heading_3.font.name = "Calibri Light"
    heading_3.font.size = Pt(14)
    heading_3.font.bold = None
    heading_3.font.color.rgb = REPORT_HEADING_GREEN
    heading_3.paragraph_format.space_before = Pt(14)
    heading_3.paragraph_format.space_after = Pt(6)
    heading_3.paragraph_format.keep_with_next = True
    heading_3.paragraph_format.keep_together = True
    heading_3.next_paragraph_style = body

    caption = styles["Caption"]
    caption.font.name = "Calibri Light"
    caption.font.size = Pt(12)
    caption.font.bold = None
    caption.font.italic = True
    caption.font.color.rgb = REPORT_HEADING_GREEN
    caption.paragraph_format.space_before = Pt(12)
    caption.paragraph_format.space_after = Pt(14)
    caption.paragraph_format.line_spacing = 1.2
    caption.next_paragraph_style = styles["Normal"]

    bullet_style = styles["List Bullet"]
    bullet_style.base_style = normal
    bullet_style.next_paragraph_style = body
    bullet_style.font.name = "Calibri Light"
    bullet_style.font.size = Pt(12)
    bullet_style.paragraph_format.space_before = Pt(0)
    bullet_style.paragraph_format.space_after = Pt(0)
    bullet_style.paragraph_format.line_spacing = Pt(20)
    bullet_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    header = styles["Header"]
    header.font.name = "Calibri"
    header.font.size = Pt(10)
    header.font.bold = True
    header.font.color.rgb = REPORT_HEADING_GREEN

    footer = styles["Footer"]
    footer.font.name = "Arial"
    footer.font.size = Pt(10)

    try:
        plot_heading = styles["Plot Heading"]
    except KeyError:
        plot_heading = styles.add_style("Plot Heading", WD_STYLE_TYPE.PARAGRAPH)
    plot_heading.base_style = body
    plot_heading.next_paragraph_style = body
    plot_heading.font.name = "Calibri Light"
    plot_heading.font.size = Pt(14)
    plot_heading.font.color.rgb = REPORT_HEADING_GREEN
    plot_heading.paragraph_format.space_before = Pt(14)
    plot_heading.paragraph_format.space_after = Pt(6)
    plot_heading.paragraph_format.keep_with_next = True
    plot_heading.paragraph_format.keep_together = True
    outline_level = plot_heading._element.get_or_add_pPr().find(qn("w:outlineLvl"))
    if outline_level is None:
        outline_level = OxmlElement("w:outlineLvl")
        plot_heading._element.get_or_add_pPr().append(outline_level)
    outline_level.set(qn("w:val"), "2")


def _configure_report_header_footer(doc, header_text: str) -> None:
    """Add generic report header/footer content generated entirely in code."""
    section = doc.sections[0]
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.25)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.clear()
    header_paragraph.style = "Header"
    header_paragraph.add_run(header_text)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.style = "Footer"
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.add_run("Page ")
    _append_field(footer_paragraph, r"PAGE \* MERGEFORMAT", "1")

    paragraph_properties = footer_paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    top_border = OxmlElement("w:top")
    top_border.set(qn("w:val"), "single")
    top_border.set(qn("w:sz"), "4")
    top_border.set(qn("w:space"), "1")
    top_border.set(qn("w:color"), "000000")
    borders.append(top_border)
    paragraph_properties.append(borders)


def _add_numbered_caption(
    doc,
    reference: FigureReference | TableReference,
    count: int,
    label: str,
    bookmark_id_offset: int,
    text: str,
) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), str(count + bookmark_id_offset))
    bookmark_start.set(qn("w:name"), reference.bookmark)
    paragraph._p.append(bookmark_start)
    paragraph.add_run(f"{label} ").bold = True
    _append_field(paragraph, r"STYLEREF 1 \s", "1", bold=True)
    paragraph.add_run("-").bold = True
    _append_field(
        paragraph,
        rf"SEQ {label} \* ARABIC \s 1",
        str(count),
        bold=True,
    )
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), str(count + bookmark_id_offset))
    paragraph._p.append(bookmark_end)
    paragraph.add_run(f" – {text}")


@dataclass
class _FigureRegistry:
    count: int = 0

    def allocate(self) -> FigureReference:
        self.count += 1
        return FigureReference(
            bookmark=f"ReportFigure{self.count}",
            fallback_label=f"Figure 1-{self.count}",
        )

    def add_caption(self, doc, reference: FigureReference, text: str) -> None:
        _add_numbered_caption(doc, reference, self.count, "Figure", 1000, text)


@dataclass
class _TableRegistry:
    count: int = 0

    def allocate(self) -> TableReference:
        self.count += 1
        return TableReference(
            bookmark=f"ReportTable{self.count}",
            fallback_label=f"Table 1-{self.count}",
        )

    def add_caption(self, doc, reference: TableReference, text: str) -> None:
        _add_numbered_caption(doc, reference, self.count, "Table", 2000, text)


def _append_field(paragraph, instruction: str, fallback_text: str, bold: bool = False) -> None:
    begin_run = paragraph.add_run()
    begin_run.bold = bold
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    instruction_run.bold = bold
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    instruction_run._r.append(instruction_element)

    separator_run = paragraph.add_run()
    separator_run.bold = bold
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separator_run._r.append(separator)

    result_run = paragraph.add_run()
    result_run.bold = bold
    result = OxmlElement("w:t")
    result.text = fallback_text
    result_run._r.append(result)

    end_run = paragraph.add_run()
    end_run.bold = bold
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _create_heading_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    for number in list(numbering.findall(qn("w:num"))):
        numbering.remove(number)
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if element.get(qn("w:abstractNumId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    abstract.set(f"{{{WORD_2012_NAMESPACE}}}restartNumberingAfterBreak", "0")
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}")
    abstract.append(nsid)
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), "5902F2C2")
    abstract.append(template)
    for level in range(3):
        item = OxmlElement("w:lvl")
        item.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        item.append(start)
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "decimal")
        item.append(number_format)
        style_reference = OxmlElement("w:pStyle")
        style_reference.set(qn("w:val"), f"Heading{level + 1}")
        item.append(style_reference)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), ".".join(f"%{index}" for index in range(1, level + 2)))
        item.append(level_text)
        alignment = OxmlElement("w:lvlJc")
        alignment.set(qn("w:val"), "left")
        item.append(alignment)
        paragraph_properties = OxmlElement("w:pPr")
        indentation = OxmlElement("w:ind")
        if level == 0:
            indentation.set(qn("w:left"), "360")
            indentation.set(qn("w:hanging"), "360")
        elif level == 1:
            indentation.set(qn("w:left"), "0")
            indentation.set(qn("w:firstLine"), "0")
        else:
            indentation.set(qn("w:left"), "357")
            indentation.set(qn("w:hanging"), "357")
        paragraph_properties.append(indentation)
        item.append(paragraph_properties)
        abstract.append(item)
    for level in range(3, 9):
        item = OxmlElement("w:lvl")
        item.set(qn("w:ilvl"), str(level))
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "decimal")
        item.append(number_format)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "")
        item.append(level_text)
        alignment = OxmlElement("w:lvlJc")
        alignment.set(qn("w:val"), "left")
        item.append(alignment)
        abstract.append(item)
    numbering.insert(len(numbering.findall(qn("w:abstractNum"))), abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    number.set(f"{{{WORD_2016_CID_NAMESPACE}}}durableId", "1701663979")
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    for level in range(3):
        _set_numbering_properties(
            doc.styles[f"Heading {level + 1}"],
            None if level == 0 else level,
            num_id,
        )
    return num_id


def _create_bullet_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    existing = next(
        (
            number
            for number in numbering.findall(qn("w:num"))
            if number.get(qn("w:numId")) == "2"
        ),
        None,
    )
    if existing is not None:
        _set_numbering_properties(doc.styles["List Bullet"], None, 2)
        return 2

    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if element.get(qn("w:abstractNumId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    abstract.set(f"{{{WORD_2012_NAMESPACE}}}restartNumberingAfterBreak", "0")
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), "547D0181")
    abstract.append(nsid)
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), "1A2A351A")
    abstract.append(template)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "bullet")
    level.append(number_format)
    style_reference = OxmlElement("w:pStyle")
    style_reference.set(qn("w:val"), "ListBullet")
    level.append(style_reference)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "\u25e6")
    level.append(level_text)
    alignment = OxmlElement("w:lvlJc")
    alignment.set(qn("w:val"), "left")
    level.append(alignment)
    paragraph_properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:hint"), "default")
    run_properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(REPORT_HEADING_GREEN))
    run_properties.append(color)
    position = OxmlElement("w:position")
    position.set(qn("w:val"), "-4")
    run_properties.append(position)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "36")
    run_properties.append(size)
    level.append(run_properties)
    abstract.append(level)
    # Word repairs out-of-order numbering XML and can relink heading styles to this bullet list.
    numbering.insert(len(numbering.findall(qn("w:abstractNum"))), abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), "2")
    number.set(f"{{{WORD_2016_CID_NAMESPACE}}}durableId", "1656910545")
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    _set_numbering_properties(doc.styles["List Bullet"], None, 2)
    return 2


def _set_numbering_properties(target, level: int | None, num_id: int) -> None:
    element = target._element if hasattr(target, "_element") else target
    paragraph_properties = element.get_or_add_pPr()
    number_properties = paragraph_properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        insert_before = {
            qn("w:suppressLineNumbers"),
            qn("w:pBdr"),
            qn("w:shd"),
            qn("w:tabs"),
            qn("w:spacing"),
            qn("w:ind"),
            qn("w:contextualSpacing"),
            qn("w:jc"),
            qn("w:outlineLvl"),
            qn("w:rPr"),
        }
        insertion_index = next(
            (
                index
                for index, child in enumerate(paragraph_properties)
                if child.tag in insert_before
            ),
            len(paragraph_properties),
        )
        paragraph_properties.insert(insertion_index, number_properties)
    else:
        for child in list(number_properties):
            number_properties.remove(child)
    if level is not None:
        level_element = OxmlElement("w:ilvl")
        level_element.set(qn("w:val"), str(level))
        number_properties.append(level_element)
    number_element = OxmlElement("w:numId")
    number_element.set(qn("w:val"), str(num_id))
    number_properties.append(number_element)


def _apply_paragraph_numbering(paragraph, level: int, num_id: int) -> None:
    _set_numbering_properties(paragraph._p, level, num_id)


def _add_report_heading(doc, text: str, level: int):
    return doc.add_heading(text, level=level)


def _add_unumbered_plot_heading(doc, text: str) -> None:
    doc.add_paragraph(text, style="Plot Heading")


def _add_caption_reference_sentence(
    doc,
    before: str,
    reference: FigureReference | TableReference,
    after: str,
) -> None:
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.add_run(before)
    _append_field(paragraph, rf"REF {reference.bookmark} \h", reference.fallback_label, bold=True)
    paragraph.add_run(after)


def _add_figure_reference_sentence(doc, before: str, reference: FigureReference, after: str) -> None:
    _add_caption_reference_sentence(doc, before, reference, after)


def _add_table_reference_sentence(doc, before: str, reference: TableReference, after: str) -> None:
    _add_caption_reference_sentence(doc, before, reference, after)


def _add_centered_report_image(doc, image_path: Path) -> None:
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(IMAGE_WIDTH_CM))


ENVELOPE_SUMMARY_SHEETS = {"SFO": "LLp", "TOV": "LLp", "SA": "LGp"}
ENVELOPE_SUMMARY_RMS_EVENTS = {"TOV", "SA"}
SA_TOV_DURATION_MS = 300
PLOT_FAULT_LABEL_PATTERN = "|".join(
    re.escape(label)
    for label in sorted(
        (label for label in KNOWN_FAULT_LABELS if label.casefold() != "no fault"),
        key=len,
        reverse=True,
    )
)


def _safe_stem(text: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_.-]+", "_", text)
    value = re.sub(r"_+", "_", value)
    return value.strip("_.") or "figure"


def _report_json_value(value: Any) -> Any:
    if hasattr(value, "to_mapping") and callable(value.to_mapping):
        return _report_json_value(value.to_mapping())
    if isinstance(value, Path):
        return value.as_posix()
    if isinstance(value, Mapping):
        return {str(key): _report_json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_report_json_value(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _report_signature(value: Any) -> str:
    encoded = json.dumps(
        _report_json_value(value),
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
    ).encode("utf-8")
    return hashlib.blake2b(encoded, digest_size=16).hexdigest()


def _report_file_manifest_entry(project_root: Path, path: Path) -> dict[str, Any]:
    try:
        relative = path.resolve().relative_to(project_root.resolve()).as_posix()
    except (OSError, ValueError):
        relative = str(path.resolve())
    try:
        stat = path.stat()
    except OSError:
        return {"path": relative, "missing": True}
    return {
        "path": relative,
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


def _report_cache_key(project_root: Path, report_dir: Path) -> str:
    try:
        return report_dir.resolve().relative_to(project_root.resolve()).as_posix()
    except (OSError, ValueError):
        return report_dir.resolve().as_posix()


def _report_manifest_matches(
    project_root: Path,
    report_dir: Path,
    voltage: str,
    signature: str,
    output_path: Path,
) -> bool:
    cache = storage.load_project_analysis_cache(project_root)
    directories = cache.get("reports")
    directory = directories.get(_report_cache_key(project_root, report_dir)) if isinstance(directories, dict) else None
    record = directory.get(str(voltage)) if isinstance(directory, dict) else None
    if not isinstance(record, dict) or record.get("signature") != signature:
        return False
    output_record = record.get("output")
    return (
        isinstance(output_record, dict)
        and _report_file_manifest_entry(project_root, output_path) == output_record
    )


def _load_report_manifest_entries(
    project_root: Path,
    report_dir: Path,
) -> dict[str, dict[str, Any]]:
    cache = storage.load_project_analysis_cache(project_root)
    directories = cache.get("reports")
    reports = directories.get(_report_cache_key(project_root, report_dir)) if isinstance(directories, dict) else None
    if not isinstance(reports, dict):
        return {}
    return {
        str(key): value
        for key, value in reports.items()
        if isinstance(value, dict)
    }


def _write_report_manifest(
    project_root: Path,
    report_dir: Path,
    reports: dict[str, dict[str, Any]],
) -> None:
    cache = storage.load_project_analysis_cache(project_root)
    directories = cache.setdefault("reports", {})
    if not isinstance(directories, dict):
        directories = {}
        cache["reports"] = directories
    directories[_report_cache_key(project_root, report_dir)] = reports
    storage.save_project_analysis_cache(project_root, cache)
    try:
        (report_dir / LEGACY_REPORT_MANIFEST_FILENAME).unlink(missing_ok=True)
    except OSError:
        pass


def _report_manifest_payload(
    project_root: Path,
    scope: ScopeEntry,
    voltage: str,
    selected_events: list[str],
    dashboard_figure_ids: list[str],
    parsed_resonance: resonance_checks.ResonanceSettings,
    event_times: dict[str, float] | None,
    parsed_sustained: sustained_sdpf.SustainedSDPFSettings,
    parsed_sustained_ranking: sustained_sdpf.SustainedSDPFRankingSettings,
    sustained_heatmap_settings: Any,
    render_heatmaps: bool,
    image_cache: dict[Path, list[Path]],
    sustained_cache_valid: bool | None = None,
    rms_settings: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    source_paths: set[Path] = set()

    def add(path: Path) -> None:
        source_paths.add(path)

    envelope_dir = project_root / "Voltage_envelope" / scope.folder
    add(envelope_dir / f"MM_{voltage}.xlsx")
    add(envelope_dir / f"MM_{voltage}_with_combined_plot.xlsx")
    add(envelope_dir / resonance_checks.WORKBOOK_NAME)
    add(sustained_sdpf.result_path(project_root, scope.folder))
    add(sustained_sdpf.summary_path(project_root, scope.folder))
    parsed_rms = rms_analysis.normalize_rms_settings(rms_settings)
    if parsed_rms["enabled"]:
        add(project_root / "Results" / "MM results.csv")

    for event in selected_events:
        for image in _find_event_images(project_root, scope, event, voltage, image_cache):
            add(image)

    for check in parsed_resonance.effective_enabled_checks:
        for voltage_type in resonance_checks.VOLTAGE_TYPES:
            for image in _find_resonance_images(
                project_root,
                scope,
                check,
                voltage_type,
                voltage,
                image_cache,
            ):
                add(image)

    if parsed_rms["enabled"]:
        for quantity in parsed_rms["quantities"]:
            for image in _find_rms_images(
                project_root,
                scope,
                quantity,
                voltage,
                image_cache,
            ):
                add(image)

    heatmap_root = (
        project_root
        / "Plots"
        / "Generated"
        / scope.folder
        / sustained_sdpf_heatmap.HEATMAP_EVENT
    )
    if heatmap_root.is_dir():
        try:
            for path in heatmap_root.rglob("*"):
                if path.is_file():
                    add(path)
        except OSError:
            pass

    for figure_id in dashboard_figure_ids:
        try:
            workbook_name, _sheet_name, _chart_index, _title = _parse_dashboard_figure_id(figure_id)
        except ValueError:
            continue
        add(project_root / "Dashboards" / workbook_name)

    payload = {
        "version": REPORT_MANIFEST_VERSION,
        "scope": {
            "name": scope.name,
            "mode": scope.mode,
            "tokens": list(scope.tokens),
        },
        "voltage": str(voltage),
        "events": list(selected_events),
        "dashboard_figure_ids": list(dashboard_figure_ids),
        "resonance": parsed_resonance.to_mapping(),
        "event_times": dict(event_times or {}),
        "sustained": parsed_sustained.to_mapping(),
        "sustained_ranking": parsed_sustained_ranking.to_mapping(),
        "sustained_heatmaps": sustained_heatmap_settings,
        "sustained_cache_valid": sustained_cache_valid,
        "rms": parsed_rms,
        "render_heatmaps": bool(render_heatmaps),
        "sources": sorted(
            (_report_file_manifest_entry(project_root, path) for path in source_paths),
            key=lambda item: str(item["path"]).casefold(),
        ),
    }
    if parsed_rms["enabled"]:
        payload["rms_result_version"] = rms_analysis.RMS_RESULT_VERSION
        payload["rms_report_layout_version"] = RMS_REPORT_LAYOUT_VERSION
    if parsed_sustained.enabled:
        payload["sustained_sdpf_result_version"] = sustained_sdpf.RESULT_VERSION
        payload["sustained_sdpf_report_layout_version"] = SUSTAINED_REPORT_LAYOUT_VERSION
    return payload


def _is_report_image(path: Path) -> bool:
    try:
        if not path.is_file() or path.stat().st_size <= 0:
            return False
        with path.open("rb") as handle:
            header = handle.read(8)
    except OSError:
        return False
    suffix = path.suffix.casefold()
    if suffix == ".png":
        return header.startswith(PNG_SIGNATURE)
    if suffix in {".jpg", ".jpeg"}:
        return header.startswith(JPEG_SIGNATURE)
    return False


def _dashboard_slicer_value(voltage: str) -> str:
    value = str(voltage).strip()
    return BUS_VOLTAGE_SLICER_VALUE_OVERRIDES.get(value, value)


def _dashboard_figure_uses_saved_voltage_filter(title: str) -> bool:
    return " ".join(str(title).split()).casefold() in {
        "initial voltages",
        "initial active power",
        "initial reactive power",
    }


def _slicer_cache_touches_sheet(slicer_cache, sheet_name: str) -> bool:
    wanted = sheet_name.strip().casefold()
    try:
        count = slicer_cache.Slicers.Count
    except EXCEL_AUTOMATION_ERRORS:
        return False

    for index in range(1, count + 1):
        try:
            slicer = slicer_cache.Slicers(index)
            if slicer.Shape.TopLeftCell.Worksheet.Name.strip().casefold() == wanted:
                return True
        except EXCEL_AUTOMATION_ERRORS:
            continue
    return False


def _find_bus_voltage_slicer_cache(workbook, sheet_name: str):
    fallback = None
    try:
        count = workbook.SlicerCaches.Count
    except EXCEL_AUTOMATION_ERRORS:
        return None

    for index in range(1, count + 1):
        slicer_cache = workbook.SlicerCaches(index)
        try:
            source_name = str(slicer_cache.SourceName)
        except EXCEL_AUTOMATION_ERRORS:
            source_name = ""
        try:
            cache_name = str(slicer_cache.Name)
        except EXCEL_AUTOMATION_ERRORS:
            cache_name = ""

        combined = f"{source_name} {cache_name}".casefold()
        if (
            BUS_VOLTAGE_SLICER_SOURCE_TEXT.casefold() not in combined
            and "bus_voltage" not in combined
        ):
            continue

        if _slicer_cache_touches_sheet(slicer_cache, sheet_name):
            return slicer_cache
        if fallback is None:
            fallback = slicer_cache
    return fallback


def _set_dashboard_voltage_filter(
    excel,
    workbook,
    sheet_name: str,
    voltage: str,
    log: LogFn | None,
) -> bool:
    slicer_cache = _find_bus_voltage_slicer_cache(workbook, sheet_name)
    if slicer_cache is None:
        _log(log, f"Dashboard voltage slicer not found: {sheet_name}; exporting current view.")
        return False

    slicer_value = _dashboard_slicer_value(voltage)
    try:
        source_name = str(slicer_cache.SourceName)
        slicer_cache.VisibleSlicerItemsList = (f"{source_name}.&[{slicer_value}]",)
    except EXCEL_AUTOMATION_ERRORS as exc:
        _log(log, f"Dashboard voltage slicer update failed: {sheet_name} | {voltage} kV | {exc}")
        return False

    try:
        excel.CalculateUntilAsyncQueriesDone()
    except EXCEL_AUTOMATION_ERRORS as exc:
        _log(log, f"Dashboard query wait skipped: {sheet_name} | {voltage} kV | {exc}")
    time.sleep(DASHBOARD_FILTER_SETTLE_S)
    _log(log, f"Dashboard voltage slicer set: {sheet_name} | {voltage} kV")
    return True


def _excel_context(get_excel: Callable[[], object] | None):
    return nullcontext(get_excel()) if get_excel is not None else excel_app()


def _find_event_images(
    project_root: Path,
    scope: ScopeEntry,
    event: str,
    voltage: str,
    image_cache: dict[Path, list[Path]] | None = None,
) -> list[Path]:
    event_dir = project_root / "Plots" / "Generated" / scope.folder / event
    return _find_voltage_images(event_dir, voltage, image_cache)


def _find_resonance_images(
    project_root: Path,
    scope: ScopeEntry,
    check: str,
    voltage_type: str,
    voltage: str,
    image_cache: dict[Path, list[Path]] | None = None,
) -> list[Path]:
    folder, type_folder = resonance_checks.report_folder(check, voltage_type)
    event_dir = project_root / "Plots" / "Generated" / scope.folder / folder / type_folder
    return _find_voltage_images(event_dir, voltage, image_cache)


def _find_rms_images(
    project_root: Path,
    scope: ScopeEntry,
    quantity: str,
    voltage: str,
    image_cache: dict[Path, list[Path]] | None = None,
) -> list[Path]:
    event_dir = rms_analysis.rms_output_dir(project_root, scope.folder, quantity)
    return _find_voltage_images(event_dir, voltage, image_cache)


def _find_voltage_images(
    event_dir: Path,
    voltage: str,
    image_cache: dict[Path, list[Path]] | None = None,
) -> list[Path]:
    if not event_dir.is_dir():
        return []
    voltage_patterns = (
        re.compile(rf"MM_{re.escape(voltage)}(?:_|\.|\b)", re.IGNORECASE),
        re.compile(rf"(?:^|[_ -]){re.escape(voltage)}(?:[_ -]?kV|[_ .-])", re.IGNORECASE),
    )
    if image_cache is None:
        image_paths = [
            path
            for pattern in ("*.png", "*.jpg", "*.jpeg")
            for path in sorted(event_dir.rglob(pattern))
        ]
    else:
        image_paths = image_cache.get(event_dir)
        if image_paths is None:
            image_paths = [
                path
                for pattern in ("*.png", "*.jpg", "*.jpeg")
                for path in sorted(event_dir.rglob(pattern))
            ]
            image_cache[event_dir] = image_paths
    return [
        path
        for path in image_paths
        if any(regex.search(path.name) for regex in voltage_patterns)
    ]


def _envelope_summary_rows(
    workbook_path: Path,
    event_times: dict[str, float] | None = None,
    events: Iterable[str] | None = None,
    measurement_overrides: dict[str, str] | None = None,
) -> list[EnvelopeSummaryRow]:
    if not workbook_path.is_file():
        return []

    import openpyxl

    if events is None:
        selected_events = ["SFO", "TOV"]
    else:
        selected_events = []
        for event in events:
            event_name = str(event).strip().upper()
            if event_name in ENVELOPE_SUMMARY_SHEETS and event_name not in selected_events:
                selected_events.append(event_name)
    selected_times = [
        (event, float((event_times or {}).get(event, DEFAULT_EVENT_TIMES[event])))
        for event in selected_events
        if event in DEFAULT_EVENT_TIMES
    ]
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        selected_by_sheet: dict[str, dict[str, float]] = {}
        for event_name, event_time in selected_times:
            sheet_name = (measurement_overrides or {}).get(
                event_name,
                ENVELOPE_SUMMARY_SHEETS[event_name],
            )
            if sheet_name not in workbook.sheetnames:
                continue
            selected_by_sheet.setdefault(sheet_name, {})[event_name] = event_time

        matched: dict[str, tuple[str, dict[str, int], tuple[Any, ...]]] = {}
        for sheet_name, targets in selected_by_sheet.items():
            headers, rows = nearest_rows(workbook[sheet_name], targets)
            for event_name, row in rows.items():
                if row is not None:
                    matched[event_name] = (sheet_name, headers, row)

        output: list[EnvelopeSummaryRow] = []
        for event_name, _event_time in selected_times:
            match = matched.get(event_name)
            if match is None:
                continue
            sheet_name, headers, best_row = match

            peak = as_float(row_value(best_row, headers, "Max_all", "Max"))
            time_s = as_float(row_value(best_row, headers, "Time (s)"))
            if peak is None or time_s is None:
                continue
            output.append(
                EnvelopeSummaryRow(
                    event=event_name,
                    measurement=sheet_name,
                    time_s=time_s,
                    peak_kv=peak,
                    rms_kv=abs(peak) / math.sqrt(2.0),
                )
            )
        return output
    finally:
        workbook.close()


def _format_float(value: float, decimals: int = 3) -> str:
    text = f"{value:.{decimals}f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


def _add_subscripted_kv_unit(paragraph, subscript: str) -> None:
    paragraph.add_run(" kV")
    run = paragraph.add_run(subscript)
    run.font.subscript = True


def _add_envelope_summary_list(doc, rows: list[EnvelopeSummaryRow]) -> None:
    if not rows:
        return
    bullet_num_id = _create_bullet_numbering(doc)
    for row in rows:
        paragraph = doc.add_paragraph(style="List Bullet")
        _apply_paragraph_numbering(paragraph, level=0, num_id=bullet_num_id)
        paragraph.add_run(f"{row.event}: {_format_float(row.peak_kv, 1)}")
        _add_subscripted_kv_unit(paragraph, "peak")
        if row.event in ENVELOPE_SUMMARY_RMS_EVENTS:
            paragraph.add_run(f" ({_format_float(row.rms_kv, 1)}")
            _add_subscripted_kv_unit(paragraph, "RMS")
            paragraph.add_run(")")
    doc.add_paragraph(style="Body Text")


def _report_dashboard_title(title: str) -> str:
    text = " ".join(str(title).split())
    return re.sub(r"\bRMS overvoltages\b", "RMS voltage rise", text, flags=re.IGNORECASE)


def _dashboard_figure_caption(title: str, voltage: str) -> str:
    normalized = " ".join(str(title).split())
    key = normalized.casefold()
    fixed_captions = {
        "initial voltages": "Initial voltage across voltage levels in the OWF",
        "initial reactive power": "Initial reactive power at the POC",
        "initial active power": "Initial active power at the PoC",
    }
    if key in fixed_captions:
        return fixed_captions[key]
    if "instantaneous overvoltages" in key:
        return f"Overvoltage results at the {voltage} kV side"
    return f"{_report_dashboard_title(normalized)} at {voltage} kV"


def _add_dashboard_figure_sentence(doc, title: str, voltage: str, reference: FigureReference) -> None:
    normalized = " ".join(str(title).split())
    key = normalized.casefold()
    if key == "initial voltages":
        before = "The initial voltage conditions of all the cases are shown in the "
    elif key == "initial reactive power":
        before = "The initial reactive power conditions of all the cases are shown in the "
    elif key == "initial active power":
        before = "The initial active power conditions of all the cases are shown in the "
    elif "instantaneous overvoltages" in key:
        before = f"The {voltage} kV side overvoltage peaks obtained for the assessed switching instants are shown in the "
    elif "RMS voltage" in key:
        before = f"The {_report_dashboard_title(normalized)} results at the PoC are shown in the "
    else:
        before = f"The {_report_dashboard_title(normalized)} results at {voltage} kV are shown in the "
    _add_figure_reference_sentence(doc, before, reference, ".")


def _add_time_domain_figure_sentence(doc, event: str, voltage: str, reference: FigureReference) -> None:
    _add_figure_reference_sentence(
        doc,
        "",
        reference,
        f" shows the time-domain plot for the highest {_report_event_name(event)} candidate recorded at the {voltage} kV level.",
    )


def _report_event_name(event: str) -> str:
    return str(event).strip().upper()


def _time_domain_figure_caption(event: str, voltage: str) -> str:
    if _report_event_name(event) == "SA":
        return f"Time-domain plot of the highest TOV at {voltage} kV for surge arrester selection"
    return f"Time-domain plots of the highest {_report_event_name(event)} at {voltage} kV"


def _add_sa_figure_sentence(
    doc,
    voltage: str,
    reference: FigureReference,
    tov_row: EnvelopeSummaryRow | None,
) -> None:
    _add_figure_reference_sentence(
        doc,
        "",
        reference,
        f" shows the highest TOV line-ground value found at {voltage} kV, which will be used for surge arrester selection.",
    )
    if tov_row is None:
        return
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.add_run("This TOV value is ")
    paragraph.add_run(_format_float(tov_row.rms_kv, 1))
    _add_subscripted_kv_unit(paragraph, "RMS")
    paragraph.add_run(f" and was assumed to last for {SA_TOV_DURATION_MS} ms.")


def _add_analysis_figure_sentence(
    doc,
    label: str,
    voltage_type: str,
    voltage: str,
    reference: FigureReference,
) -> None:
    _add_figure_reference_sentence(
        doc,
        f"The {label} {voltage_type} results recorded at {voltage} kV are shown in the ",
        reference,
        ".",
    )


def _load_sustained_sdpf_report_results(
    project_root: Path,
    scope: ScopeEntry,
    voltage: str,
    settings: sustained_sdpf.SustainedSDPFSettings,
    payload: Mapping[str, Any] | None = None,
    shared_manifest_current: bool | None = None,
    cache_validation: sustained_sdpf.SustainedSDPFCacheValidation | None = None,
) -> tuple[
    dict[
        str,
        dict[str, dict[str, sustained_sdpf.SustainedSDPFResult]],
    ],
    sustained_sdpf.SustainedSDPFCacheValidation,
]:
    saved_payload = (
        payload
        if isinstance(payload, Mapping)
        else sustained_sdpf.load_results(project_root, scope.folder)
    )
    validation = cache_validation or sustained_sdpf.validate_result_cache(
        saved_payload,
        project_root,
        settings,
        shared_manifest_current=shared_manifest_current,
    )
    if not validation.valid:
        return {}, validation
    return sustained_sdpf.current_representatives_for_voltage(
        saved_payload,
        voltage,
        project_root,
        settings,
        shared_manifest_current=validation.shared_manifest_current,
        cache_validation=validation,
    )


def _sustained_image_for_result(
    images: Iterable[Path],
    result: sustained_sdpf.SustainedSDPFResult,
) -> Path | None:
    prefix = (
        f"Case: {result.case} | Run: {int(result.run)} | "
        f"Element: {result.mm_name}"
    )
    for image in images:
        if _plot_heading_from_image_path(image).startswith(prefix):
            return image
    return None


_SUSTAINED_SELECTION_LABELS = {
    sustained_sdpf.HIGHEST_VOLTAGE_SUSTAINED_SELECTION: "Highest sustained Vₜ",
    sustained_sdpf.CUMULATIVE_STRESS_SELECTION: "Cumulative stress",
    sustained_sdpf.CONTINUOUS_DURATION_SELECTION: "Longest duration",
}


def _sustained_population_title(population: str) -> str:
    if population == sustained_sdpf.ACTUAL_SDPF_POPULATION:
        return "Actual SDPF (threshold: SDPF)"
    return "Safety-margin-only (threshold: SDPF/1.15)"


def _sustained_selection_text(metrics: Iterable[str]) -> str:
    labels = [
        _SUSTAINED_SELECTION_LABELS[metric]
        for metric in metrics
        if metric in _SUSTAINED_SELECTION_LABELS
    ]
    return "; ".join(labels)


def _sustained_selection_groups(
    selections: Mapping[str, sustained_sdpf.SustainedSDPFResult],
    population: str,
) -> list[tuple[sustained_sdpf.SustainedSDPFResult, tuple[str, ...]]]:
    grouped: dict[tuple[Any, ...], list[str]] = {}
    for metric in sustained_sdpf.RANKING_SELECTIONS:
        result = selections.get(metric)
        if not isinstance(result, sustained_sdpf.SustainedSDPFResult):
            continue
        phase = result.governing
        event = sustained_sdpf.selected_event(phase, population, metric)
        grouped.setdefault(
            (
                *sustained_sdpf._result_identity(result),
                str(phase.measurement).casefold(),
                str(phase.phase).casefold(),
                round(float(event.start_s), 12) if event is not None else None,
                round(float(event.end_s), 12) if event is not None else None,
            ),
            [],
        ).append(metric)
    return [
        (selections[metrics[0]], tuple(metrics))
        for metrics in grouped.values()
    ]


def _report_metric_text(value: float | None, decimals: int = 2) -> str:
    if value is None:
        return "—"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    return _format_float(number, decimals) if math.isfinite(number) else "—"


def _sustained_table_rows(
    groups: Iterable[tuple[sustained_sdpf.SustainedSDPFResult, tuple[str, ...]]],
    population: str,
) -> list[list[str]]:
    rows: list[list[str]] = []
    for result, metrics in groups:
        phase = result.governing
        metric = metrics[0] if metrics else sustained_sdpf.CUMULATIVE_STRESS_SELECTION
        (
            _threshold,
            _threshold_peak,
            area,
            qualification_duration_ms,
            _duration_ms,
            sustained_t_peak,
            sustained_t_percent,
            _event_start,
            _event_end,
        ) = sustained_sdpf._population_metric_values(phase, population, metric)
        ratio_percent = (
            float(sustained_t_percent)
            if sustained_t_percent is not None
            else None
        )
        rows.append(
            [
                _sustained_selection_text(metrics),
                f"{result.case} / {int(result.run)}",
                str(result.fault_type or "—"),
                str(result.mm_name or "—"),
                f"{phase.measurement} {phase.phase}",
                f"{_report_metric_text(sustained_t_peak)} / {_report_metric_text(ratio_percent, 1)}%",
                _report_metric_text(area),
                _report_metric_text(qualification_duration_ms),
            ]
        )
    return rows


def _add_sustained_selection_table(
    doc,
    groups: Iterable[tuple[sustained_sdpf.SustainedSDPFResult, tuple[str, ...]]],
    population: str,
    voltage: str,
    table_registry: _TableRegistry,
) -> None:
    reference = table_registry.allocate()
    _add_table_reference_sentence(
        doc,
        "Selected cases are summarized in ",
        reference,
        ".",
    )
    table_registry.add_caption(
        doc,
        reference,
        f"Selected Sustained SDPF cases — {_sustained_population_title(population)} at {voltage} kV",
    )
    headers = (
        "Selection criteria",
        "Case / Run",
        "Fault",
        "MM element",
        "Path",
        "Vₜ (kVₚₑₐₖ / %SDPF)",
        "Area (pu·ms)",
        "Episode duration (ms)",
    )
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in _sustained_table_rows(groups, population):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True


def _add_sustained_selected_plot(
    doc,
    metrics: Iterable[str],
    population: str,
    voltage: str,
    image_path: Path | None,
    figure_registry: _FigureRegistry,
    rendered_images: set[Path],
    check_cancel: CancelFn | None = None,
) -> None:
    if image_path is None:
        return
    _cancel(check_cancel)
    if image_path in rendered_images:
        return
    rendered_images.add(image_path)
    criteria_text = _sustained_selection_text(metrics)
    reference = figure_registry.allocate()
    _add_unumbered_plot_heading(doc, _plot_heading_from_image_path(image_path))
    _add_figure_reference_sentence(
        doc,
        "",
        reference,
        f" shows the selected Sustained SDPF case ({criteria_text}).",
    )
    _add_centered_report_image(doc, image_path)
    figure_registry.add_caption(
        doc,
        reference,
        f"Selected Sustained SDPF case — {_sustained_population_title(population)}; "
        f"{criteria_text}; {voltage} kV",
    )


def _add_sustained_population_subsection(
    doc,
    population: str,
    selections: Mapping[str, sustained_sdpf.SustainedSDPFResult],
    voltage: str,
    image_paths: Mapping[str, Path | None],
    figure_registry: _FigureRegistry,
    check_cancel: Callable[[], None] | None = None,
    table_registry: _TableRegistry | None = None,
) -> None:
    groups = _sustained_selection_groups(selections, population)
    if not groups:
        return
    _add_report_heading(
        doc,
        f"{_sustained_population_title(population)} selected cases",
        level=3,
    )
    table_registry = table_registry or _TableRegistry()
    _add_sustained_selection_table(doc, groups, population, voltage, table_registry)

    plot_groups = [
        (
            metrics,
            next(
                (
                    image_paths.get(metric)
                    for metric in metrics
                    if image_paths.get(metric) is not None
                ),
                None,
            ),
        )
        for _result, metrics in groups
    ]
    if any(image_path is not None for _metrics, image_path in plot_groups):
        _add_report_heading(doc, "Selected time-domain plots", level=4)
    rendered_images: set[Path] = set()
    for metrics, image_path in plot_groups:
        _add_sustained_selected_plot(
            doc,
            metrics,
            population,
            voltage,
            image_path,
            figure_registry,
            rendered_images,
            check_cancel,
        )


def _add_sustained_sdpf_section(
    doc,
    result: sustained_sdpf.SustainedSDPFResult,
    voltage: str,
    image_path: Path | None,
    figure_registry: _FigureRegistry,
    check_cancel: CancelFn | None = None,
    heatmap_groups: Iterable[tuple[str, Iterable[Path]]] = (),
    heatmap_settings_by_name: Mapping[str, sustained_sdpf_heatmap.HeatmapSettings] | None = None,
    selection_results_by_population: Mapping[str, Any] | None = None,
    image_paths_by_population: Mapping[str, Any] | None = None,
    ranking_settings: sustained_sdpf.SustainedSDPFRankingSettings | None = None,
    table_registry: _TableRegistry | None = None,
) -> None:
    _add_report_heading(doc, "Sustained SDPF", level=2)
    parsed_ranking = ranking_settings or sustained_sdpf.SustainedSDPFRankingSettings()
    criteria_text = _sustained_selection_text(parsed_ranking.enabled_selections()) or "the enabled criteria"
    duration_ms = result.duration_s * 1000.0
    doc.add_paragraph(
        f"Candidates require peak-envelope persistence for at least {_format_float(duration_ms, 2)} ms. "
        f"Actual SDPF and SDPF/1.15 safety-margin populations are ranked separately by {criteria_text}.",
        style="Body Text",
    )
    table_registry = table_registry or _TableRegistry()
    raw_selections = selection_results_by_population or {}
    population_selections: dict[str, dict[str, sustained_sdpf.SustainedSDPFResult]] = {
        population: {}
        for population in sustained_sdpf.RANKING_POPULATIONS
    }
    if any(population in raw_selections for population in sustained_sdpf.RANKING_POPULATIONS):
        for population in sustained_sdpf.RANKING_POPULATIONS:
            raw_population = raw_selections.get(population, {})
            if isinstance(raw_population, Mapping):
                population_selections[population] = {
                    metric: selected
                    for metric, selected in raw_population.items()
                    if metric in sustained_sdpf.RANKING_SELECTIONS
                    and isinstance(selected, sustained_sdpf.SustainedSDPFResult)
                }
    else:
        # Keep direct callers that omit the detailed population structure
        # usable for a concise report section.
        fallback_population = sustained_sdpf.result_population(result)
        if fallback_population is not None:
            population_selections[fallback_population] = {
                metric: selected
                for metric, selected in raw_selections.items()
                if metric in {
                    sustained_sdpf.CUMULATIVE_STRESS_SELECTION,
                    sustained_sdpf.CONTINUOUS_DURATION_SELECTION,
                }
                and isinstance(selected, sustained_sdpf.SustainedSDPFResult)
            }
        if not any(population_selections.values()):
            fallback_population = fallback_population or sustained_sdpf.SAFETY_MARGIN_ONLY_POPULATION
            population_selections[fallback_population] = {
                sustained_sdpf.CUMULATIVE_STRESS_SELECTION: result,
                sustained_sdpf.CONTINUOUS_DURATION_SELECTION: result,
            }

    raw_image_paths = image_paths_by_population or {}
    population_images: dict[str, dict[str, Path | None]] = {
        population: {}
        for population in sustained_sdpf.RANKING_POPULATIONS
    }
    if any(population in raw_image_paths for population in sustained_sdpf.RANKING_POPULATIONS):
        for population in sustained_sdpf.RANKING_POPULATIONS:
            raw_population = raw_image_paths.get(population, {})
            if isinstance(raw_population, Mapping):
                population_images[population] = {
                    metric: path
                    for metric, path in raw_population.items()
                    if metric in sustained_sdpf.RANKING_SELECTIONS
                }
    else:
        fallback_population = sustained_sdpf.result_population(result) or sustained_sdpf.SAFETY_MARGIN_ONLY_POPULATION
        population_images[fallback_population] = {
            metric: raw_image_paths.get(metric)
            for metric in sustained_sdpf.RANKING_SELECTIONS
        }
        if image_path is not None:
            for metric in sustained_sdpf.RANKING_SELECTIONS:
                population_images[fallback_population].setdefault(metric, image_path)

    enabled = set(parsed_ranking.enabled_selections())
    for population in sustained_sdpf.RANKING_POPULATIONS:
        selected = {
            metric: selected_result
            for metric, selected_result in population_selections[population].items()
            if metric in enabled
        }
        if not selected:
            continue
        _add_sustained_population_subsection(
            doc,
            population,
            selected,
            voltage,
            population_images[population],
            figure_registry,
            check_cancel,
            table_registry,
        )
    groups = [
        (str(name), [path for path in paths if _is_report_image(path)])
        for name, paths in heatmap_groups
    ]
    groups = [(name, paths) for name, paths in groups if paths]
    if not groups:
        return
    _add_report_heading(doc, "Incidence heatmaps", level=3)
    doc.add_paragraph(
        "Top: actual SDPF incidence; bottom: inclusive SDPF/1.15 safety-margin incidence, "
        "counted per eligible Run × MM cell.",
        style="Body Text",
    )
    for set_name, paths in groups:
        heatmap_settings = (heatmap_settings_by_name or {}).get(set_name)
        report_set_name = (
            sustained_sdpf_heatmap.display_heatmap_set_name(set_name, heatmap_settings)
            or "Heatmap"
        )
        _add_report_heading(doc, report_set_name, level=4)
        for heatmap_path in paths:
            _cancel(check_cancel)
            reference = figure_registry.allocate()
            panel_heading = _heatmap_plot_heading_from_image_path(heatmap_path, heatmap_settings)
            if panel_heading:
                _add_unumbered_plot_heading(doc, panel_heading)
            _add_figure_reference_sentence(doc, "", reference, " shows Sustained SDPF incidence.")
            _add_centered_report_image(doc, heatmap_path)
            caption = f"Sustained SDPF incidence — {report_set_name}"
            if panel_heading:
                caption += f", {panel_heading}"
            figure_registry.add_caption(doc, reference, f"{caption} at {voltage} kV")


def _analysis_figure_caption(label: str, voltage_type: str, voltage: str) -> str:
    return f"{label} {voltage_type} results at {voltage} kV"


def _plot_heading_from_image_path(path: Path) -> str:
    stem = path.stem
    if stem.endswith("_3Ph") or stem.endswith("_1Ch"):
        stem = stem.rsplit("_", 1)[0]
    # Group labels can contain underscores that must remain part of Element.
    match = re.match(
        rf"(?P<case>.+?)_(?P<element>MM_\d+(?:\.\d+)?(?:_[^_]+)*?)(?:_(?P<fault>{PLOT_FAULT_LABEL_PATTERN}))?_(?P<run>\d{{3,}})_(?P<trace>LGp_LLp|LGr|LLr|LGp|LLp|[^_]+)(?:_|$)",
        stem,
        flags=re.IGNORECASE,
    )
    if match:
        parts = [
            f"Case: {match.group('case')}",
            f"Run: {int(match.group('run'))}",
            f"Element: {match.group('element')}",
        ]
        fault = match.group("fault")
        if fault:
            parts.append(f"Fault: {fault}")
        trace = match.group("trace")
        if trace:
            parts.append(f"Trace: {trace.replace('_', ' & ') if trace == 'LGp_LLp' else trace}")
        return " | ".join(parts)
    return path.stem


def _rms_variant_from_image_path(path: Path) -> str:
    match = re.search(r"_(max|min)(?:_3Ph)?$", path.stem, re.IGNORECASE)
    return match.group(1).lower() if match else ""


_RMS_REPORT_HEADINGS = {
    ("LG", "max"): "Line-to-Ground RMS Voltage Rise",
    ("LG", "min"): "Line-to-Ground RMS Voltage Dip",
    ("LL", "max"): "Line-to-Line RMS Voltage Rise",
    ("LL", "min"): "Line-to-Line RMS Voltage Dip",
}


def _rms_report_heading(quantity: str, variant: str) -> str:
    return _RMS_REPORT_HEADINGS.get((quantity, variant), "RMS Voltage")


def _rms_report_calculation(selection: rms_analysis.RMSSelection | None) -> str | None:
    if selection is None:
        return None
    reference = rms_analysis.rms_reference_kv(selection)
    change_percent = rms_analysis.rms_change_percent(selection)
    if reference is None or change_percent is None:
        return None
    if selection.variant == "min":
        value_word = "lowest"
        change_word = "voltage dip"
    elif selection.variant == "max":
        value_word = "highest"
        change_word = "voltage rise"
    else:
        return None
    return (
        f"The {value_word} {selection.quantity} RMS voltage is "
        f"{_format_float(selection.value_kv, 1)} kV, corresponding to a "
        f"{change_word} of {_format_float(change_percent, 1)}%."
    )


def _add_rms_report_content(
    doc,
    project_root: Path,
    scope: ScopeEntry,
    voltage: str,
    rms_settings: Mapping[str, Any],
    image_cache: dict[Path, list[Path]],
    figure_registry: _FigureRegistry,
    check_cancel: CancelFn | None = None,
    selections_by_key: Mapping[tuple[str, str, str], rms_analysis.RMSSelection] | None = None,
) -> int:
    grouped: dict[str, list[Path]] = {}
    for quantity in rms_analysis.RMS_QUANTITIES:
        if quantity not in rms_settings.get("quantities", []):
            continue
        images = [
            image
            for image in _find_rms_images(project_root, scope, quantity, voltage, image_cache)
            if _is_report_image(image)
        ]
        if images:
            grouped[quantity] = sorted(images, key=lambda path: path.name.casefold())
    if not grouped:
        return 0
    _add_report_heading(doc, "RMS", level=2)
    doc.add_paragraph(
        "Selected RMS voltage traces show the governing maximum and minimum values "
        "for the configured MM elements.",
        style="Body Text",
    )
    count = 0
    voltage_key = rms_analysis.normalize_voltage(voltage)
    for quantity in rms_analysis.RMS_QUANTITIES:
        images = grouped.get(quantity, [])
        if not images:
            continue
        image_order = {
            _rms_variant_from_image_path(image_path): image_path
            for image_path in images
        }
        ordered_variants = [
            variant for variant in rms_analysis.RMS_REPORT_VARIANT_ORDER
            if variant in image_order
        ]
        ordered_variants.extend(
            variant for variant in image_order
            if variant not in ordered_variants
        )
        for variant in ordered_variants:
            image_path = image_order[variant]
            _cancel(check_cancel)
            _add_report_heading(doc, _rms_report_heading(quantity, variant), level=3)
            reference = figure_registry.allocate()
            _add_unumbered_plot_heading(doc, _plot_heading_from_image_path(image_path))
            selection = (selections_by_key or {}).get((voltage_key, quantity, variant))
            calculation = _rms_report_calculation(selection)
            if calculation:
                doc.add_paragraph(calculation, style="Body Text")
            _add_figure_reference_sentence(
                doc,
                f"The selected {quantity} RMS voltage trace at {voltage} kV is shown in the ",
                reference,
                ".",
            )
            _add_centered_report_image(doc, image_path)
            figure_registry.add_caption(
                doc,
                reference,
                f"{_rms_report_heading(quantity, variant)} at {voltage} kV",
            )
            count += 1
    return count


def _heatmap_plot_heading_from_image_path(
    path: Path,
    settings: sustained_sdpf_heatmap.HeatmapSettings | None = None,
) -> str:
    """Return a semantic report heading without exposing the PNG filename."""
    stem = path.stem
    combined = re.fullmatch(r"MM_[^_]+_faceted_(?P<page>\d+)_heatmap", stem, flags=re.IGNORECASE)
    if combined:
        return "Combined panels"

    separate = re.fullmatch(r"MM_[^_]+_(?P<body>.+)_heatmap", stem, flags=re.IGNORECASE)
    if not separate:
        return "Heatmap"
    body = separate.group("body")
    part = None
    part_match = re.fullmatch(r"(?P<value>.+)_(?P<part>\d{2})", body)
    if part_match:
        body = part_match.group("value")
        part = int(part_match.group("part"))

    if settings is not None and settings.split_by and body.casefold() != "all":
        value = sustained_sdpf_heatmap.display_group_value(settings.split_by, body)
        heading = f"Split {settings.split_by}: {value}"
    else:
        heading = "" if body.casefold() == "all" else "Heatmap"
    if part is not None:
        heading = f"{heading} · Panel {part}" if heading else f"Panel {part}"
    return heading


def _export_envelope_chart(
    project_root: Path,
    scope: ScopeEntry,
    voltage: str,
    export_dir: Path,
    log: LogFn | None = None,
    get_excel: Callable[[], object] | None = None,
    check_cancel: CancelFn | None = None,
) -> Path | None:
    base_workbook_path = (
        project_root
        / "Voltage_envelope"
        / scope.folder
        / f"MM_{voltage}.xlsx"
    )
    workbook_path = (
        project_root
        / "Voltage_envelope"
        / scope.folder
        / f"MM_{voltage}_with_combined_plot.xlsx"
    )
    if not workbook_path.is_file():
        return None
    if (
        base_workbook_path.is_file()
        and workbook_path.stat().st_mtime_ns < base_workbook_path.stat().st_mtime_ns
    ):
        _log(
            log,
            f"Envelope chart is older than its data and was skipped: "
            f"{scope.folder} / {workbook_path.name}. Rebuild envelope charts.",
        )
        return None
    _cancel(check_cancel)
    export_dir.mkdir(parents=True, exist_ok=True)
    output_path = export_dir / f"Envelope_{scope.folder}_{voltage}_kV.png"
    try:
        output_path.unlink(missing_ok=True)
    except OSError:
        pass
    try:
        with _excel_context(get_excel) as excel:
            workbook = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=True)
            try:
                worksheet = workbook.Worksheets("LLp")
                if worksheet.ChartObjects().Count < 1:
                    return None
                chart_object = worksheet.ChartObjects(1)
                _log(log, f"Exporting envelope chart: {scope.folder} | {voltage} kV")
                for _attempt in range(2):
                    _cancel(check_cancel)
                    try:
                        chart_object.Activate()
                    except EXCEL_AUTOMATION_ERRORS:
                        pass
                    time.sleep(0.5)
                    ok = chart_object.Chart.Export(str(output_path), "PNG")
                    if ok and _is_report_image(output_path):
                        return output_path
                    try:
                        output_path.unlink(missing_ok=True)
                    except OSError:
                        pass
                _log(log, f"Excel did not export a valid envelope chart: {output_path.name}")
                return None
            finally:
                workbook.Close(SaveChanges=False)
    except EXCEL_AUTOMATION_ERRORS as exc:
        _log(log, f"Envelope chart export skipped: {exc}")
        return None


def _parse_dashboard_figure_id(figure_id: str) -> tuple[str, str, int, str]:
    parts = figure_id.split("|", 3)
    if len(parts) != 4:
        raise ValueError(f"Unsupported dashboard figure id: {figure_id}")
    workbook, sheet, chart_index, title = [part.strip() for part in parts]
    return workbook, sheet, int(chart_index), title


def _export_dashboard_figures(
    project_root: Path,
    figure_ids: Iterable[str],
    voltage: str,
    export_dir: Path,
    log: LogFn | None = None,
    get_excel: Callable[[], object] | None = None,
    get_dashboard_workbook: Callable[[Path, str], object] | None = None,
    check_cancel: CancelFn | None = None,
) -> list[tuple[str, Path]]:
    selected_ids = [figure_id for figure_id in figure_ids if figure_id]
    if not selected_ids:
        return []
    export_dir.mkdir(parents=True, exist_ok=True)
    exported: list[tuple[str, Path]] = []
    try:
        with _excel_context(get_excel) as excel:
            grouped: dict[str, list[tuple[str, int, str, str]]] = {}
            for figure_id in selected_ids:
                _cancel(check_cancel)
                try:
                    workbook_name, sheet_name, chart_index, title = _parse_dashboard_figure_id(figure_id)
                except ValueError as exc:
                    _log(log, str(exc))
                    continue
                grouped.setdefault(workbook_name, []).append((sheet_name, chart_index, title, figure_id))

            for workbook_name, figures in grouped.items():
                _cancel(check_cancel)
                workbook_path = project_root / "Dashboards" / workbook_name
                if not workbook_path.is_file():
                    _log(log, f"Dashboard workbook not found, skipping: {workbook_name}")
                    continue
                use_cached_workbook = get_dashboard_workbook is not None and not any(
                    _dashboard_figure_uses_saved_voltage_filter(figure[2])
                    for figure in figures
                )

                saved_filter_figures = [
                    figure
                    for figure in figures
                    if _dashboard_figure_uses_saved_voltage_filter(figure[2])
                ]
                voltage_filter_figures = [
                    figure
                    for figure in figures
                    if not _dashboard_figure_uses_saved_voltage_filter(figure[2])
                ]
                for figures_to_export, apply_voltage_filter in (
                    (saved_filter_figures, False),
                    (voltage_filter_figures, True),
                ):
                    if not figures_to_export:
                        continue
                    if apply_voltage_filter and use_cached_workbook:
                        _log(log, f"Opening dashboard for report figures: {workbook_name}")
                        workbook = get_dashboard_workbook(project_root, workbook_name)
                        close_workbook = False
                    else:
                        _log(log, f"Opening dashboard for saved figure state: {workbook_name}")
                        workbook = excel.Workbooks.Open(
                            str(workbook_path),
                            UpdateLinks=0,
                            ReadOnly=True,
                        )
                        close_workbook = True
                    try:
                        filtered_sheets: set[str] = set()
                        for sheet_name, chart_index, title, figure_id in figures_to_export:
                            _cancel(check_cancel)
                            try:
                                sheet_key = sheet_name.strip().casefold()
                                if apply_voltage_filter and sheet_key not in filtered_sheets:
                                    _set_dashboard_voltage_filter(excel, workbook, sheet_name, str(voltage), log)
                                    filtered_sheets.add(sheet_key)
                                worksheet = workbook.Worksheets(sheet_name)
                                chart_object = worksheet.ChartObjects(chart_index)
                                figure_title = title or f"Chart {chart_index}"
                                file_name = _safe_stem(
                                    f"{voltage}_kV_{workbook_name}_{sheet_name}_{chart_index}_{figure_title}.png"
                                )
                                output_path = export_dir / file_name
                                try:
                                    output_path.unlink(missing_ok=True)
                                except OSError:
                                    pass
                                _log(log, f"Exporting dashboard figure: {voltage} kV | {figure_title}")
                                ok = chart_object.Chart.Export(str(output_path), "PNG")
                                if not ok or not _is_report_image(output_path):
                                    _log(log, f"Excel did not export dashboard figure: {figure_id}")
                                    try:
                                        output_path.unlink(missing_ok=True)
                                    except OSError:
                                        pass
                                    continue
                                exported.append((figure_title, output_path))
                            except EXCEL_AUTOMATION_ERRORS as exc:
                                _log(log, f"Dashboard figure skipped: {figure_id} | {exc}")
                    finally:
                        if close_workbook:
                            workbook.Close(SaveChanges=False)
    except EXCEL_AUTOMATION_ERRORS as exc:
        _log(log, f"Dashboard figure export skipped: {exc}")

    return exported


def _add_sustained_report_content(
    doc,
    root: Path,
    scope: ScopeEntry,
    voltage: str,
    parsed_sustained: sustained_sdpf.SustainedSDPFSettings,
    sustained_payload: Mapping[str, Any],
    shared_manifest_current: bool | None,
    sustained_cache_validation: sustained_sdpf.SustainedSDPFCacheValidation | None,
    image_cache: dict[Path, list[Path]],
    figure_registry: _FigureRegistry,
    check_cancel: Callable[[], None] | None,
    log: LogFn | None,
    heatmap_settings_value: Any,
    render_heatmaps: bool,
    ranking_settings: sustained_sdpf.SustainedSDPFRankingSettings,
) -> int:
    """Add Sustained SDPF figures for one report voltage."""
    _cancel(check_cancel)
    sustained_results, result_validation = _load_sustained_sdpf_report_results(
        root,
        scope,
        voltage,
        parsed_sustained,
        sustained_payload,
        shared_manifest_current,
        sustained_cache_validation,
    )
    sustained_result = next(
        (
            sustained_results.get(population, {}).get(
                sustained_sdpf.CUMULATIVE_STRESS_SELECTION
            )
            for population in sustained_sdpf.RANKING_POPULATIONS
            if sustained_results.get(population, {}).get(
                sustained_sdpf.CUMULATIVE_STRESS_SELECTION
            ) is not None
        ),
        next(
            (
                selected
                for population in sustained_sdpf.RANKING_POPULATIONS
                for selected in sustained_results.get(population, {}).values()
                if selected is not None
            ),
            None,
        ),
    )
    if sustained_result is None:
        if result_validation is not None and not result_validation.valid:
            if sustained_cache_validation is None or sustained_cache_validation.valid:
                _log(
                    log,
                    f"Sustained SDPF report result skipped for {voltage} kV: "
                    f"{result_validation.reason}.",
                )
            sustained_sdpf_heatmap.clear_heatmaps(root, scope.folder)
        elif render_heatmaps:
            sustained_sdpf_heatmap.clear_heatmaps(root, scope.folder)
        return 0

    sustained_images = _find_event_images(
        root,
        scope,
        sustained_sdpf.SUSTAINED_SDPF,
        voltage,
        image_cache,
    )
    image_paths_by_population = {
        population: {
            selection: _sustained_image_for_result(sustained_images, selected_result)
            for selection, selected_result in selections.items()
        }
        for population, selections in sustained_results.items()
    }
    heatmap_groups: list[tuple[str, list[Path]]] = []
    heatmap_settings = sustained_sdpf_heatmap.heatmap_sets_from_mapping(
        heatmap_settings_value
    )
    heatmap_enabled = any(item.settings.enabled for item in heatmap_settings)
    if heatmap_enabled and render_heatmaps:
        heatmap_groups = sustained_sdpf_heatmap.generate_heatmap_sets(
            root,
            scope.folder,
            voltage,
            heatmap_settings,
            parsed_sustained,
            log=log,
            check_cancel=check_cancel,
            payload=sustained_payload,
            shared_manifest_current=shared_manifest_current,
        )
    elif heatmap_enabled:
        heatmap_groups = sustained_sdpf_heatmap.heatmap_paths_for_sets(
            root,
            scope.folder,
            voltage,
            heatmap_settings,
        )
    else:
        sustained_sdpf_heatmap.clear_heatmaps(root, scope.folder)

    cumulative_image = next(
        (
            image_paths_by_population.get(population, {}).get(
                sustained_sdpf.CUMULATIVE_STRESS_SELECTION
            )
            for population in sustained_sdpf.RANKING_POPULATIONS
            if image_paths_by_population.get(population, {}).get(
                sustained_sdpf.CUMULATIVE_STRESS_SELECTION
            ) is not None
        ),
        None,
    )
    _add_sustained_sdpf_section(
        doc,
        sustained_result,
        voltage,
        cumulative_image,
        figure_registry,
        check_cancel,
        heatmap_groups=heatmap_groups,
        heatmap_settings_by_name={
            item.name: item.settings for item in heatmap_settings
        },
        selection_results_by_population=sustained_results,
        image_paths_by_population=image_paths_by_population,
        ranking_settings=ranking_settings,
    )
    return len({
        path
        for paths in image_paths_by_population.values()
        for path in paths.values()
        if path is not None
    }) + sum(len(paths) for _name, paths in heatmap_groups)


def build_reports_from_existing_plots(
    project_roots: Iterable[str | Path],
    scopes: Iterable[ScopeEntry],
    voltages: Iterable[str],
    events: Iterable[str],
    dashboard_figure_ids_by_project: Mapping[str, Iterable[str]] | None = None,
    resonance_settings: dict[str, object] | None = None,
    event_times: dict[str, float] | None = None,
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    sustained_sdpf_settings: dict[str, object] | None = None,
    sustained_sdpf_heatmap_settings_by_project: dict[str, object] | None = None,
    sustained_sdpf_ranking_settings_by_project: dict[str, object] | None = None,
    render_heatmaps: bool = True,
    sustained_payloads_by_project_scope: Mapping[str, Mapping[str, Mapping[str, Any]]] | None = None,
    sustained_cache_validations_by_project_scope: Mapping[
        str, Mapping[str, sustained_sdpf.SustainedSDPFCacheValidation]
    ] | None = None,
    rms_settings_by_project: Mapping[str, Mapping[str, Any]] | None = None,
    resonance_settings_by_project: Mapping[str, Mapping[str, Any]] | None = None,
    sustained_sdpf_settings_by_project: Mapping[str, Mapping[str, Any]] | None = None,
) -> list[Path]:
    """Build draft Word reports from already generated plot image files."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to build Word reports.") from exc

    written: list[Path] = []
    selected_scopes = list(scopes)
    selected_events = [
        str(event)
        for event in events
        if str(event) not in rms_analysis.RMS_BATCH_EVENTS.values()
    ]
    selected_voltages = list(voltages)
    parsed_sustained_ranking = sustained_sdpf.SustainedSDPFRankingSettings()
    _cancel(check_cancel)

    with ExitStack() as stack:
        excel = None
        dashboard_workbooks: dict[tuple[Path, str], object] = {}

        def get_excel():
            nonlocal excel
            _cancel(check_cancel)
            if excel is None:
                excel = stack.enter_context(excel_app())
            return excel

        def get_dashboard_workbook(project_root: Path, workbook_name: str):
            _cancel(check_cancel)
            key = (project_root, workbook_name)
            workbook = dashboard_workbooks.get(key)
            if workbook is None:
                workbook_path = project_root / "Dashboards" / workbook_name
                workbook = get_excel().Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=True)
                dashboard_workbooks[key] = workbook
                stack.callback(lambda wb=workbook: wb.Close(SaveChanges=False))
            return workbook

        for project_root in project_roots:
            _cancel(check_cancel)
            root = Path(project_root).resolve()
            parsed_resonance = resonance_checks.ResonanceSettings.from_mapping(
                (resonance_settings_by_project or {}).get(str(root), resonance_settings)
            )
            parsed_sustained = sustained_sdpf.SustainedSDPFSettings.from_mapping(
                (sustained_sdpf_settings_by_project or {}).get(str(root), sustained_sdpf_settings)
            )
            selected_dashboard_figure_ids = list(
                (dashboard_figure_ids_by_project or {}).get(str(root), ())
            )
            parsed_sustained_ranking = sustained_sdpf.SustainedSDPFRankingSettings.from_mapping(
                (sustained_sdpf_ranking_settings_by_project or {}).get(str(root))
            )
            parsed_rms = rms_analysis.normalize_rms_settings(
                (rms_settings_by_project or {}).get(str(root))
            )
            rms_selections_by_key: dict[
                tuple[str, str, str], rms_analysis.RMSSelection
            ] = {}
            if parsed_rms["enabled"]:
                rms_rows = rms_analysis.load_mm_results(root)
                rms_selections_by_key = {
                    (
                        selection.voltage_key,
                        selection.quantity,
                        selection.variant,
                    ): selection
                    for selection in rms_analysis.select_rms_rows(
                        rms_rows,
                        selected_elements=parsed_rms["elements"],
                        selected_quantities=parsed_rms["quantities"],
                        selected_voltages=selected_voltages,
                    )
                }
            for scope in selected_scopes:
                _cancel(check_cancel)
                report_dir = root / "Reports" / scope.folder
                report_dir.mkdir(parents=True, exist_ok=True)
                manifest_entries = _load_report_manifest_entries(root, report_dir)
                if not parsed_sustained.enabled:
                    sustained_sdpf_heatmap.clear_heatmaps(root, scope.folder)
                project_payloads = (sustained_payloads_by_project_scope or {}).get(str(root), {})
                sustained_payload = (
                    project_payloads.get(scope.folder)
                    if isinstance(project_payloads, Mapping)
                    else None
                )
                if parsed_sustained.enabled and not isinstance(sustained_payload, Mapping):
                    sustained_payload = sustained_sdpf.load_results(root, scope.folder)
                if not parsed_sustained.enabled:
                    sustained_payload = {}
                project_cache_validations = (
                    (sustained_cache_validations_by_project_scope or {}).get(str(root), {})
                )
                sustained_cache_validation = (
                    project_cache_validations.get(scope.folder)
                    if parsed_sustained.enabled
                    and isinstance(project_cache_validations, Mapping)
                    else None
                )
                if parsed_sustained.enabled and sustained_cache_validation is None:
                    sustained_cache_validation = sustained_sdpf.validate_result_cache(
                        sustained_payload,
                        root,
                        parsed_sustained,
                    )
                shared_manifest_current = (
                    sustained_cache_validation.shared_manifest_current
                    if sustained_cache_validation is not None
                    else None
                )
                if (
                    sustained_cache_validation is not None
                    and not sustained_cache_validation.valid
                ):
                    _log(
                        log,
                        f"Sustained SDPF report data unavailable for {scope.folder}: "
                        f"{sustained_cache_validation.reason}. "
                        "Rebuild envelope data/checks before rebuilding reports.",
                    )
                for temporary_dir in (
                    report_dir / "_envelope_exports",
                    report_dir / "_dashboard_exports",
                ):
                    stack.callback(shutil.rmtree, temporary_dir, ignore_errors=True)
                image_cache: dict[Path, list[Path]] = {}
                for voltage in selected_voltages:
                    _cancel(check_cancel)
                    _log(log, f"Building report: {root.name} | {scope.folder} | {voltage} kV")
                    report_payload = _report_manifest_payload(
                        root,
                        scope,
                        str(voltage),
                        selected_events,
                        selected_dashboard_figure_ids,
                        parsed_resonance,
                        event_times,
                        parsed_sustained,
                        parsed_sustained_ranking,
                        (sustained_sdpf_heatmap_settings_by_project or {}).get(str(root)),
                        render_heatmaps,
                        image_cache,
                        sustained_cache_valid=(
                            sustained_cache_validation.valid
                            if sustained_cache_validation is not None
                            else None
                        ),
                        rms_settings=parsed_rms,
                    )
                    report_signature = _report_signature(report_payload)
                    output_path = report_dir / f"Voltage_{voltage}_kV.docx"
                    sustained_cache_invalid = (
                        parsed_sustained.enabled
                        and sustained_cache_validation is not None
                        and not sustained_cache_validation.valid
                    )
                    if not sustained_cache_invalid and _report_manifest_matches(
                        root,
                        report_dir,
                        str(voltage),
                        report_signature,
                        output_path,
                    ):
                        written.append(output_path)
                        try:
                            (report_dir / LEGACY_REPORT_MANIFEST_FILENAME).unlink(missing_ok=True)
                        except OSError:
                            pass
                        _log(log, f"Skipping unchanged report: {output_path.name}")
                        continue
                    doc = Document()
                    section = doc.sections[0]
                    section.page_width = Cm(21.0)
                    section.page_height = Cm(29.7)
                    section.left_margin = Cm(2.54)
                    section.right_margin = Cm(2.54)
                    section.top_margin = Cm(2.54)
                    section.bottom_margin = Cm(2.54)

                    _configure_report_styles(doc)
                    _configure_report_header_footer(
                        doc,
                        f"PSCAD Results Analysis - {root.name} - {scope.name}",
                    )

                    _create_heading_numbering(doc)
                    figure_registry = _FigureRegistry()
                    _add_report_heading(doc, f"{voltage} kV Voltage Assessment", level=1)
                    envelope_workbook = root / "Voltage_envelope" / scope.folder / f"MM_{voltage}.xlsx"
                    sa_tov_rows = []
                    if any(_report_event_name(event) == "SA" for event in selected_events):
                        sa_tov_rows = _envelope_summary_rows(
                            envelope_workbook,
                            event_times,
                            events=["TOV"],
                            measurement_overrides={"TOV": "LGp"},
                        )
                    sa_tov_row = sa_tov_rows[0] if sa_tov_rows else None

                    dashboard_figures = _export_dashboard_figures(
                        root,
                        selected_dashboard_figure_ids,
                        str(voltage),
                        report_dir / "_dashboard_exports" / f"Voltage_{voltage}_kV",
                        log,
                        get_excel,
                        get_dashboard_workbook,
                        check_cancel,
                    )
                    _cancel(check_cancel)
                    dashboard_figures.sort(
                        key=lambda figure: DASHBOARD_REPORT_TITLE_ORDER.get(
                            figure[0].strip().casefold(),
                            len(DASHBOARD_REPORT_TITLE_ORDER),
                        )
                    )

                    envelope_chart = _export_envelope_chart(
                        root,
                        scope,
                        str(voltage),
                        report_dir / "_envelope_exports",
                        log,
                        get_excel,
                        check_cancel,
                    )
                    _cancel(check_cancel)
                    if envelope_chart is not None:
                        _add_report_heading(
                            doc,
                            "Voltage Envelope",
                            level=2,
                        )
                        reference = figure_registry.allocate()
                        _add_figure_reference_sentence(
                            doc,
                            "",
                            reference,
                            " shows that the highest overvoltages are:",
                        )
                        _add_envelope_summary_list(
                            doc,
                            _envelope_summary_rows(envelope_workbook, event_times, selected_events),
                        )
                        _add_centered_report_image(doc, envelope_chart)
                        figure_registry.add_caption(
                            doc,
                            reference,
                            f"Representative overvoltages envelope at {voltage} kV",
                        )

                    if dashboard_figures:
                        _add_report_heading(
                            doc,
                            "Dashboard Figures",
                            level=2,
                        )
                        for figure_title, image_path in dashboard_figures:
                            _cancel(check_cancel)
                            reference = figure_registry.allocate()
                            _add_unumbered_plot_heading(doc, _report_dashboard_title(figure_title))
                            _add_dashboard_figure_sentence(doc, figure_title, str(voltage), reference)
                            _add_centered_report_image(doc, image_path)
                            figure_registry.add_caption(
                                doc,
                                reference,
                                _dashboard_figure_caption(figure_title, str(voltage)),
                            )
                    elif selected_dashboard_figure_ids:
                        doc.add_paragraph(
                            "Selected dashboard figures could not be exported.",
                            style="Body Text",
                        )

                    plot_count = 0

                    for event in selected_events:
                        _cancel(check_cancel)
                        images = _find_event_images(root, scope, event, voltage, image_cache)
                        if not images:
                            continue
                        _add_report_heading(
                            doc,
                            _report_event_name(event),
                            level=2,
                        )
                        for image_path in images:
                            _cancel(check_cancel)
                            if not _is_report_image(image_path):
                                _log(log, f"Skipping invalid plot image: {image_path}")
                                continue
                            reference = figure_registry.allocate()
                            _add_unumbered_plot_heading(doc, _plot_heading_from_image_path(image_path))
                            if _report_event_name(event) == "SA":
                                _add_sa_figure_sentence(doc, str(voltage), reference, sa_tov_row)
                            else:
                                _add_time_domain_figure_sentence(doc, event, str(voltage), reference)
                            _add_centered_report_image(doc, image_path)
                            figure_registry.add_caption(
                                doc,
                                reference,
                                _time_domain_figure_caption(event, str(voltage)),
                            )
                            plot_count += 1

                    if parsed_rms["enabled"]:
                        plot_count += _add_rms_report_content(
                            doc,
                            root,
                            scope,
                            str(voltage),
                            parsed_rms,
                            image_cache,
                            figure_registry,
                            check_cancel,
                            rms_selections_by_key,
                        )

                    if parsed_sustained.enabled:
                        plot_count += _add_sustained_report_content(
                            doc,
                            root,
                            scope,
                            str(voltage),
                            parsed_sustained,
                            sustained_payload,
                            shared_manifest_current,
                            sustained_cache_validation,
                            image_cache,
                            figure_registry,
                            check_cancel,
                            log,
                            (sustained_sdpf_heatmap_settings_by_project or {}).get(str(root)),
                            render_heatmaps,
                            parsed_sustained_ranking,
                        )

                    for check in parsed_resonance.effective_enabled_checks:
                        _cancel(check_cancel)
                        for voltage_type in resonance_checks.VOLTAGE_TYPES:
                            _cancel(check_cancel)
                            images = _find_resonance_images(root, scope, check, voltage_type, voltage, image_cache)
                            if not images:
                                continue
                            label = resonance_checks.CHECK_DEFINITIONS[check]["label"]
                            _add_report_heading(
                                doc,
                                f"{label} - {voltage_type}",
                                level=2,
                            )
                            for image_path in images:
                                _cancel(check_cancel)
                                if not _is_report_image(image_path):
                                    _log(log, f"Skipping invalid plot image: {image_path}")
                                    continue
                                reference = figure_registry.allocate()
                                _add_unumbered_plot_heading(doc, _plot_heading_from_image_path(image_path))
                                _add_analysis_figure_sentence(
                                    doc,
                                    label,
                                    voltage_type,
                                    str(voltage),
                                    reference,
                                )
                                _add_centered_report_image(doc, image_path)
                                figure_registry.add_caption(
                                    doc,
                                    reference,
                                    _analysis_figure_caption(label, voltage_type, str(voltage)),
                                )
                                plot_count += 1

                    if plot_count == 0:
                        doc.add_paragraph(
                            "No generated plot images were found for this report.",
                            style="Body Text",
                        )

                    _cancel(check_cancel)
                    doc.save(output_path)
                    _cancel(check_cancel)
                    written.append(output_path)
                    manifest_entries[str(voltage)] = {
                        "signature": report_signature,
                        "output": _report_file_manifest_entry(root, output_path),
                    }
                    _write_report_manifest(root, report_dir, manifest_entries)
                    _log(log, f"Wrote report: {output_path}")

                for temporary_dir in (
                    report_dir / "_envelope_exports",
                    report_dir / "_dashboard_exports",
                ):
                    shutil.rmtree(temporary_dir, ignore_errors=True)

    return written
