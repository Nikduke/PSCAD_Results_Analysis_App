from __future__ import annotations

from collections.abc import Iterable, Mapping
from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
from concurrent.futures.process import BrokenProcessPool
from dataclasses import dataclass, fields, is_dataclass, replace
from enum import Enum
import math
import multiprocessing
from pathlib import Path
import shutil
import time
from typing import Any
import uuid

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from results_analysis_app import resonance_checks, rms_analysis, storage, sustained_sdpf, sustained_sdpf_heatmap
from results_analysis_app.background import OperationCancelled
from results_analysis_app.common import (
    CancelFn,
    LogFn,
    as_float,
    check_cancel as _cancel,
    log_message as _log,
    save_workbook_atomic,
)
from results_analysis_app.envelope_rows import nearest_rows, row_value
from results_analysis_app.models import ScopeEntry
from pscad_plotter_app_v3.models import DEFAULT_TOV_WINDOW_S
from results_analysis_app.project_config import DEFAULT_EVENT_TIMES


LEGACY_PLOT_MANIFEST_FILENAME = ".plot_manifest.json"
# Kept only as the name of the legacy per-folder file that is removed when a
# folder is next rendered. New stage metadata lives in the project cache.
PLOT_MANIFEST_VERSION = 3
_DIRECTORY_REPLACE_ATTEMPTS = 3
_DIRECTORY_REPLACE_DELAY_S = 0.05


@dataclass(slots=True)
class _PlotBatchPlan:
    project_root: Path
    batch_file: Path
    event_name: str
    output_dir: Path
    stage_dir: Path
    jobs: list[Any]

EVENT_DEFINITIONS = {
    "TOV": {"source_sheet": "LLp", "trace": "Both"},
    "SFO": {"source_sheet": "LLp", "trace": "Both"},
    "SA": {"source_sheet": "LGp", "trace": "LGp"},
}


def _as_int_if_possible(value: Any) -> Any:
    number = as_float(value)
    if number is None:
        return value
    if number.is_integer():
        return int(number)
    return number


def _batch_points_from_sheet(
    workbook,
    sheet_name: str,
    event_times: dict[str, float],
) -> dict[str, dict[str, Any] | None]:
    output = {event_name: None for event_name in event_times}
    if sheet_name not in workbook.sheetnames:
        return output
    headers, matches = nearest_rows(workbook[sheet_name], event_times)
    for event_name, row in matches.items():
        if row is None:
            continue
        output[event_name] = {
            "case": row_value(row, headers, "Case_all", "Case"),
            "run": row_value(row, headers, "Run_all", "Run"),
            "element": row_value(row, headers, "MM_name_all", "MM_name"),
            "fault": row_value(row, headers, "Fault_type_all", "Fault_type"),
        }
    return output


def _read_batch_points(
    workbook_path: Path,
    selected_events: Iterable[str],
    event_times: dict[str, float],
) -> dict[str, dict[str, Any] | None]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        by_sheet: dict[str, dict[str, float]] = {}
        for event_name in selected_events:
            by_sheet.setdefault(EVENT_DEFINITIONS[event_name]["source_sheet"], {})[event_name] = event_times[event_name]
        output: dict[str, dict[str, Any] | None] = {event_name: None for event_name in selected_events}
        for sheet_name, sheet_event_times in by_sheet.items():
            output.update(_batch_points_from_sheet(wb, sheet_name, sheet_event_times))
        return output
    finally:
        wb.close()


def _create_batch_workbook(path: Path, mm_rows: list[dict[str, Any]]) -> None:
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    try:
        for index, (sheet_name, (_mode, headers)) in enumerate(BatchExcelService.SHEET_DEFINITIONS.items()):
            ws = wb.active if index == 0 else wb.create_sheet(sheet_name)
            ws.title = sheet_name
            ws.append(headers)
            ws.freeze_panes = "A2"
            if headers:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(1, col_idx)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(str(header)) + 2)
        header_to_col = {
            str(wb["MM"].cell(1, col).value): col
            for col in range(1, wb["MM"].max_column + 1)
        }
        for row_idx, row_data in enumerate(mm_rows, start=2):
            for key, value in row_data.items():
                col = header_to_col.get(key)
                if col is not None:
                    wb["MM"].cell(row_idx, col).value = value
        save_workbook_atomic(wb, path)
    finally:
        wb.close()


def _sustained_plot_row(
    result: sustained_sdpf.SustainedSDPFResult,
    *,
    excel_waveform_exports_enabled: bool = True,
) -> dict[str, Any]:
    plot_row: dict[str, Any] = {
        "case": result.case,
        "run": result.run,
        "element": result.mm_name,
        "trace": "Both",
        "overview": True,
        "tov_windows": True,
        "tov_window_count": 4,
        "limits": True,
        "excel_export": bool(excel_waveform_exports_enabled),
    }
    return plot_row


def _sustained_plot_rows(
    results: Iterable[tuple[str, sustained_sdpf.SustainedSDPFResult]],
    *,
    excel_waveform_exports_enabled: bool = True,
) -> list[dict[str, Any]]:
    """Build deterministic, duplicate-free MM rows from selected results."""
    ordered_results = sorted(
        results,
        key=lambda item: (
            sustained_sdpf.normalized_voltage_key(item[0]),
            str(item[1].case).casefold(),
            int(item[1].run),
            str(item[1].mm_name).casefold(),
        ),
    )
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str, int, str]] = set()
    for voltage_key, result in ordered_results:
        identity = (
            sustained_sdpf.normalized_voltage_key(voltage_key),
            str(result.case),
            int(result.run),
            str(result.mm_name),
        )
        if identity in seen:
            continue
        seen.add(identity)
        rows.append(
            _sustained_plot_row(
                result,
                excel_waveform_exports_enabled=excel_waveform_exports_enabled,
            )
        )
    return rows


def _current_sustained_plot_results(
    project_root: Path,
    scope_folder: str,
    settings: sustained_sdpf.SustainedSDPFSettings,
    ranking_settings: sustained_sdpf.SustainedSDPFRankingSettings | None = None,
    *,
    payload: Mapping[str, Any] | None = None,
    voltage_keys: Iterable[str] | None = None,
    log: LogFn | None = None,
    cache_validation: sustained_sdpf.SustainedSDPFCacheValidation | None = None,
) -> dict[str, dict[str, dict[str, sustained_sdpf.SustainedSDPFResult]]]:
    saved_payload = payload if isinstance(payload, Mapping) else sustained_sdpf.load_results(
        project_root,
        scope_folder,
    )
    validation = cache_validation or sustained_sdpf.validate_result_cache(
        saved_payload,
        project_root,
        settings,
    )
    if not validation.valid:
        _log(
            log,
            f"Sustained SDPF cache unavailable for {scope_folder}: {validation.reason}. "
            "Rebuild envelope data/checks before creating or rendering Sustained SDPF plots.",
        )
        return {}
    keys = (
        [str(value) for value in voltage_keys]
        if voltage_keys is not None
        else [str(value) for value in saved_payload.get("results", {})]
    )
    current: dict[
        str,
        dict[str, dict[str, sustained_sdpf.SustainedSDPFResult]],
    ] = {}
    parsed_ranking = ranking_settings or sustained_sdpf.SustainedSDPFRankingSettings()
    enabled = set(parsed_ranking.enabled_selections())
    for voltage in keys:
        selections, result_validation = sustained_sdpf.current_representatives_for_voltage(
            saved_payload,
            voltage,
            project_root,
            settings,
            shared_manifest_current=validation.shared_manifest_current,
            cache_validation=validation,
        )
        if not selections:
            if not result_validation.valid:
                _log(
                    log,
                    f"Sustained SDPF result skipped for {voltage} kV: "
                    f"{result_validation.reason}.",
                )
            continue
        filtered = {
            population: {
                metric: result
                for metric, result in metrics.items()
                if metric in enabled
            }
            for population, metrics in selections.items()
        }
        filtered = {population: metrics for population, metrics in filtered.items() if metrics}
        if filtered:
            current[sustained_sdpf.normalized_voltage_key(voltage)] = filtered
    return current


def _clear_sustained_outputs(project_root: Path, scope_folder: str) -> None:
    """Remove stale Sustained SDPF plot and heatmap outputs."""
    output_dir = _desired_output_dir(project_root, scope_folder, sustained_sdpf.SUSTAINED_SDPF)
    _clear_plot_cache_entry(project_root, output_dir)
    _remove_generated_directory(
        output_dir,
        project_root / "Plots" / "Generated" / scope_folder,
    )
    sustained_sdpf_heatmap.clear_heatmaps(project_root, scope_folder)


def _clear_rms_outputs(project_root: Path, scope_folder: str, quantity: str) -> None:
    output_dir = rms_analysis.rms_output_dir(project_root, scope_folder, quantity)
    _clear_plot_cache_entry(project_root, output_dir)
    _remove_generated_directory(
        output_dir,
        project_root / "Plots" / "Generated" / scope_folder,
    )
    rms_analysis.rms_batch_path(project_root, scope_folder, quantity).unlink(missing_ok=True)


def _voltage_key(value: Any) -> str:
    return sustained_sdpf.normalized_voltage_key(value)


def _sustained_request_matches(
    request: Any,
    result: sustained_sdpf.SustainedSDPFResult,
) -> bool:
    expected = _sustained_plot_row(result, excel_waveform_exports_enabled=request.excel_export)
    if request.case_name != expected["case"]:
        return False
    if request.run_numbers != [int(expected["run"])] or request.elements != [expected["element"]]:
        return False
    if request.trace_type != expected["trace"]:
        return False
    if not request.show_three_phase_overview or not request.show_tov_windows:
        return False
    if not math.isclose(
        float(request.tov_window_s),
        DEFAULT_TOV_WINDOW_S,
        rel_tol=0.0,
        abs_tol=1e-12,
    ):
        return False
    if request.tov_window_count != expected["tov_window_count"] or not request.show_limits:
        return False
    for key in ("time_start_s", "time_end_s"):
        actual = getattr(request, key)
        wanted = expected.get(key)
        if actual is None or wanted is None:
            if actual != wanted:
                return False
        elif not math.isclose(float(actual), float(wanted), rel_tol=0.0, abs_tol=1e-12):
            return False
    return _voltage_key(request.voltage_kv) == _voltage_key(result.voltage)


def _expected_sustained_plot_results(
    project_root: Path,
    scope_folder: str,
    settings: sustained_sdpf.SustainedSDPFSettings,
    ranking_settings: sustained_sdpf.SustainedSDPFRankingSettings,
    payload: Mapping[str, Any] | None,
    voltage_keys: Iterable[str] | None,
    log: LogFn | None,
    cache_validation: sustained_sdpf.SustainedSDPFCacheValidation | None = None,
) -> dict[tuple[str, str, int, str], sustained_sdpf.SustainedSDPFResult]:
    expected: dict[tuple[str, str, int, str], sustained_sdpf.SustainedSDPFResult] = {}
    for voltage_key, population_selections in _current_sustained_plot_results(
        project_root,
        scope_folder,
        settings,
        ranking_settings,
        payload=payload,
        voltage_keys=voltage_keys,
        log=log,
        cache_validation=cache_validation,
    ).items():
        for selections in population_selections.values():
            for result in selections.values():
                expected.setdefault(
                    (
                        voltage_key,
                        str(result.case),
                        int(result.run),
                        str(result.mm_name),
                    ),
                    result,
                )
    return expected


def _sustained_batch_is_current(
    requests: Iterable[Any],
    expected_results: Mapping[
        tuple[str, str, int, str], sustained_sdpf.SustainedSDPFResult
    ],
) -> bool:
    request_keys: set[tuple[str, str, int, str]] = set()
    for request in requests:
        request_key = (
            _voltage_key(request.voltage_kv),
            request.case_name,
            int(request.run_numbers[0]) if request.run_numbers else 0,
            request.elements[0] if request.elements else "",
        )
        result = expected_results.get(request_key)
        if request_key in request_keys or result is None or not _sustained_request_matches(
            request, result
        ):
            return False
        request_keys.add(request_key)
    return request_keys == set(expected_results)


def _build_mm_plot_jobs(
    requests: Iterable[Any],
    stage_dir: Path,
    limit_service: Any,
    event_limits: Mapping[str, Any],
    check_cancel: CancelFn | None,
    log: LogFn | None,
) -> list[Any]:
    from pscad_plotter_app_v3.models import PlotMode
    from pscad_plotter_app_v3.services.batching import build_mm_jobs

    jobs: list[Any] = []
    for request in requests:
        _cancel(check_cancel)
        request.output_dir = str(stage_dir)
        if request.mode is not PlotMode.MM:
            raise RuntimeError(
                f"Unsupported batch mode in analysis app: {request.mode.value}. "
                "Only MM voltage-envelope batches are supported."
            )
        limit = limit_service.get_limit_for_voltage(event_limits, request.voltage_kv)
        if request.mode is PlotMode.MM and request.show_limits and limit is None:
            request.show_limits = False
            _log(
                log,
                f"No MM limits found; plotting without limits: "
                f"{request.case_name} | {request.elements[0] if request.elements else ''}",
            )
        jobs.extend(build_mm_jobs(request, limit))
    return jobs


def create_plot_batches(
    project_root: Path,
    scopes: Iterable[ScopeEntry],
    voltages: Iterable[str],
    events: Iterable[str],
    event_times: dict[str, float] | None = None,
    resonance_settings: dict[str, Any] | None = None,
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    sustained_sdpf_settings: dict[str, Any] | None = None,
    sustained_sdpf_heatmap_settings_by_project: dict[str, Any] | None = None,
    sustained_sdpf_ranking_settings: dict[str, Any] | None = None,
    sustained_payloads_by_scope: Mapping[str, Mapping[str, Any]] | None = None,
    excel_waveform_exports_enabled: bool = True,
    sustained_cache_validations_by_scope: Mapping[
        str, sustained_sdpf.SustainedSDPFCacheValidation
    ] | None = None,
    rms_settings: Mapping[str, Any] | None = None,
) -> list[Path]:
    outputs: list[Path] = []
    selected_events = [
        str(event)
        for event in events
        if str(event) not in rms_analysis.RMS_BATCH_EVENTS.values()
    ]
    selected_voltages = list(voltages)
    selected_event_times = {
        event: float((event_times or {}).get(event, DEFAULT_EVENT_TIMES[event]))
        for event in selected_events
    }
    parsed_rms = rms_analysis.normalize_rms_settings(rms_settings) if rms_settings is not None else None
    rms_batch_rows_by_event: dict[str, list[dict[str, object]]] | None = None
    if parsed_rms is not None:
        if parsed_rms["enabled"]:
            catalog = _load_rms_catalog(project_root)
            selections = rms_analysis.select_rms_rows(
                catalog.mm_results,
                selected_elements=parsed_rms["elements"],
                selected_quantities=parsed_rms["quantities"],
                selected_voltages=selected_voltages,
            )
            rms_batch_rows_by_event = rms_analysis.rms_batch_rows(
                selections,
                excel_export=excel_waveform_exports_enabled,
            )
        else:
            rms_batch_rows_by_event = {
                event: [] for event in rms_analysis.RMS_BATCH_EVENTS.values()
            }
    for scope in scopes:
        scenario_dir = project_root / "Voltage_envelope" / scope.folder
        rows_by_event: dict[str, list[dict[str, Any]]] = {
            event_name: [] for event_name in selected_events
        }
        for voltage in selected_voltages:
            _cancel(check_cancel)
            workbook = scenario_dir / f"MM_{voltage}.xlsx"
            if not workbook.is_file():
                continue
            points = _read_batch_points(workbook, selected_events, selected_event_times)
            for event_name, point in points.items():
                if point is None or not point.get("case") or not point.get("element"):
                    _log(log, f"No plot point found: {scope.folder} | {event_name} | {voltage} kV")
                    continue
                event = EVENT_DEFINITIONS[event_name]
                rows_by_event[event_name].append(
                    {
                        "case": point["case"],
                        "run": _as_int_if_possible(point["run"]),
                        "element": point["element"],
                        "trace": event["trace"],
                        "overview": True,
                        "tov_windows": True,
                        "tov_window_count": 4,
                        "limits": True,
                        "excel_export": bool(excel_waveform_exports_enabled),
                    }
                )
        for event_name, mm_rows in rows_by_event.items():
            _cancel(check_cancel)
            output_path = project_root / "Plots" / "Plot_batch" / f"batch_paste_{scope.folder}_{event_name}.xlsx"
            _create_batch_workbook(output_path, mm_rows)
            outputs.append(output_path)
            _log(log, f"Wrote batch: {output_path.name} | MM rows={len(mm_rows)}")

        # RMS is explicitly project-specific.  ``None`` means this action did
        # not request RMS work (for example, event-only batch creation), while
        # a mapping, including a disabled mapping, reconciles stale RMS files.
        if parsed_rms is not None:
            for quantity in rms_analysis.RMS_QUANTITIES:
                event_name = rms_analysis.RMS_BATCH_EVENTS[quantity]
                batch_path = rms_analysis.rms_batch_path(project_root, scope.folder, quantity)
                quantity_enabled = quantity in parsed_rms.get("quantities", [])
                rows = (
                    rms_batch_rows_by_event.get(event_name, [])
                    if parsed_rms["enabled"] and quantity_enabled and rms_batch_rows_by_event is not None
                    else []
                )
                if rows:
                    _create_batch_workbook(batch_path, rows)
                    outputs.append(batch_path)
                    _log(log, f"Wrote RMS batch: {batch_path.name} | MM rows={len(rows)}")
                else:
                    _clear_rms_outputs(project_root, scope.folder, quantity)
                    # _clear_rms_outputs removes the batch itself; leave an
                    # empty workbook only when there are no selected elements
                    # if the user explicitly enabled RMS quantities.
                    if parsed_rms["enabled"] and quantity_enabled:
                        _create_batch_workbook(batch_path, [])
                        outputs.append(batch_path)
                        _log(log, f"Wrote empty RMS batch: {batch_path.name}")

        # ``None`` means this action did not request Sustained SDPF work.  In
        # particular, event-only batch creation must not delete a valid
        # Sustained batch or its heatmaps. An explicit disabled settings
        # mapping still means "turn Sustained SDPF off" and clears them.
        if sustained_sdpf_settings is not None:
            expected_settings = sustained_sdpf.SustainedSDPFSettings.from_mapping(
                sustained_sdpf_settings
            )
            sustained_enabled = expected_settings.enabled
            heatmap_sets = sustained_sdpf_heatmap.heatmap_sets_from_mapping(
                (sustained_sdpf_heatmap_settings_by_project or {}).get(str(project_root.resolve()))
            )
            heatmap_enabled = any(heatmap_set.settings.enabled for heatmap_set in heatmap_sets)
            sustained_event = sustained_sdpf.SUSTAINED_SDPF
            sustained_path = project_root / "Plots" / "Plot_batch" / f"batch_paste_{scope.folder}_{sustained_event}.xlsx"
            if not sustained_enabled:
                if sustained_path.is_file():
                    sustained_path.unlink()
                    _log(log, f"Removed obsolete Sustained SDPF batch: {sustained_path.name}")
                sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
            else:
                if not heatmap_enabled:
                    sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
                payload = (sustained_payloads_by_scope or {}).get(scope.folder)
                if not isinstance(payload, Mapping):
                    payload = sustained_sdpf.load_results(project_root, scope.folder)
                validation = (sustained_cache_validations_by_scope or {}).get(scope.folder)
                if validation is None:
                    validation = sustained_sdpf.validate_result_cache(
                        payload,
                        project_root,
                        expected_settings,
                    )
                if not validation.valid:
                    _log(
                        log,
                        f"Sustained SDPF batch not created for {scope.folder}: "
                        f"{validation.reason}. Rebuild envelope data/checks first.",
                    )
                    if sustained_path.is_file():
                        sustained_path.unlink(missing_ok=True)
                    _clear_sustained_outputs(project_root, scope.folder)
                else:
                    current_results = _current_sustained_plot_results(
                        project_root,
                        scope.folder,
                        expected_settings,
                        sustained_sdpf.SustainedSDPFRankingSettings.from_mapping(
                            sustained_sdpf_ranking_settings
                        ),
                        payload=payload,
                        voltage_keys=selected_voltages,
                        log=log,
                        cache_validation=validation,
                    )
                    sustained_rows = _sustained_plot_rows(
                        (
                            (voltage_key, result)
                            for voltage_key, population_selections in current_results.items()
                            for selections in population_selections.values()
                            for result in selections.values()
                        ),
                        excel_waveform_exports_enabled=excel_waveform_exports_enabled,
                    )
                    _create_batch_workbook(sustained_path, sustained_rows)
                    outputs.append(sustained_path)
                    if not sustained_rows:
                        _clear_sustained_outputs(project_root, scope.folder)
                        _log(
                            log,
                            f"No qualifying Sustained SDPF governing result for {scope.folder}; "
                            "the sustained plot batch is empty by design.",
                        )
                    _log(log, f"Wrote Sustained SDPF batch: {sustained_path.name} | MM rows={len(sustained_rows)}")

        if resonance_settings is None:
            continue
        parsed_resonance = resonance_checks.ResonanceSettings.from_mapping(resonance_settings)
        selected_resonance_events = resonance_checks.selected_plot_events(parsed_resonance)
        selected_resonance_event_set = set(selected_resonance_events)
        for check in resonance_checks.CHECK_DEFINITIONS:
            for voltage_type in resonance_checks.VOLTAGE_TYPES:
                stale_event = resonance_checks.plot_event_name(check, voltage_type)
                if stale_event in selected_resonance_event_set:
                    continue
                stale_path = (
                    project_root
                    / "Plots"
                    / "Plot_batch"
                    / f"batch_paste_{scope.folder}_{stale_event}.xlsx"
                )
                if stale_path.is_file():
                    stale_path.unlink()
                    _log(log, f"Removed obsolete resonance batch: {stale_path.name}")
        resonance_rows = resonance_checks.create_plot_batch_rows(
            project_root,
            scope.folder,
            parsed_resonance.effective_enabled_checks,
            excel_waveform_exports_enabled=excel_waveform_exports_enabled,
        )
        for event_name in selected_resonance_events:
            _cancel(check_cancel)
            mm_rows = resonance_rows.get(event_name, [])
            output_path = project_root / "Plots" / "Plot_batch" / f"batch_paste_{scope.folder}_{event_name}.xlsx"
            _create_batch_workbook(output_path, mm_rows)
            outputs.append(output_path)
            _log(log, f"Wrote resonance batch: {output_path.name} | MM rows={len(mm_rows)}")
    return outputs


def _load_embedded_plotter_session(
    project_root: Path,
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
):
    import pscad_plotter_app_v3
    from pscad_plotter_app_v3.services.exporter import ExcelExporter
    from pscad_plotter_app_v3.services.limits import LimitService
    from pscad_plotter_app_v3.services.project import (
        CatalogCache,
        ProjectDiscoveryService,
        ResultsCatalogService,
        RunAvailabilityService,
    )
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    version = getattr(pscad_plotter_app_v3, "__version__", "embedded")
    _log(log, f"Using embedded plotter engine: pscad_plotter_app_v3 {version}")

    discovery_service = ProjectDiscoveryService()
    run_service = RunAvailabilityService()
    results_service = ResultsCatalogService()
    limit_service = LimitService()

    context = discovery_service.discover(project_root)
    context.state_dir.mkdir(parents=True, exist_ok=True)
    cache = CatalogCache(context.state_dir)
    try:
        _log(log, "plotter: Indexing runs...")
        run_index = run_service.build_index(context)
        _log(log, "plotter: Loading MM results...")
        catalog = results_service.build_base_catalog(context, run_index, cache)
        limits = limit_service.load_effective_limits(context)
        renderer = MatplotlibRenderer(run_index, check_cancel)
        exporter = ExcelExporter(renderer)
        return catalog, limits, renderer, exporter
    finally:
        cache.close()


def _load_rms_catalog(project_root: Path):
    """Load the shared MM catalog without constructing a renderer."""
    from pscad_plotter_app_v3.services.project import (
        CatalogCache,
        ProjectDiscoveryService,
        ResultsCatalogService,
        RunAvailabilityService,
    )

    context = ProjectDiscoveryService().discover(project_root)
    context.state_dir.mkdir(parents=True, exist_ok=True)
    cache = CatalogCache(context.state_dir)
    try:
        run_index = RunAvailabilityService().build_index(context)
        return ResultsCatalogService().build_base_catalog(context, run_index, cache)
    finally:
        cache.close()


def _apply_sustained_sdpf_plot_limit_overrides(
    limits: dict[str, Any],
    overrides: dict[str, dict[str, float]] | None,
) -> dict[str, Any]:
    """Apply Sustained SDPF RMS overrides while preserving the SIWL limits."""
    if not overrides:
        return limits

    adjusted = dict(limits)
    for raw_voltage, values in overrides.items():
        if not isinstance(values, dict):
            continue
        try:
            voltage_key = f"{float(raw_voltage):g}"
        except (TypeError, ValueError):
            continue
        current = limits.get(voltage_key)
        if current is None:
            continue
        updates: dict[str, float] = {}
        for setting_key, limit_field in (("LGp", "sdpf_lg"), ("LLp", "sdpf_ll")):
            try:
                value = float(values[setting_key])
            except (KeyError, TypeError, ValueError):
                continue
            if math.isfinite(value) and value > 0:
                updates[limit_field] = value
        if not updates:
            continue
        adjusted[voltage_key] = replace(current, **updates, source="manual")
    return adjusted


def _plot_job_sort_key(job: Any) -> tuple[str, int, str, str, str, str, str, str]:
    return (
        str(getattr(job, "case_name", "")).casefold(),
        int(getattr(job, "run_number", 0)),
        str(getattr(job, "group_label", "")).casefold(),
        str(getattr(job, "trace_type", "") or "").casefold(),
        str(getattr(job, "plot_variant", "") or "").casefold(),
        str(getattr(job, "time_start_s", "")),
        str(getattr(job, "time_end_s", "")),
        str(getattr(job, "output_dir", "")),
    )


def _plot_job_groups(
    planned_jobs: list[tuple[_PlotBatchPlan, int, Any]],
    group_size: int,
) -> list[tuple[tuple[_PlotBatchPlan, int, Any], ...]]:
    """Keep nearby jobs from one Case/Run in the same renderer process."""
    groups: list[tuple[tuple[_PlotBatchPlan, int, Any], ...]] = []
    current: list[tuple[_PlotBatchPlan, int, Any]] = []
    current_key: tuple[str, int] | None = None
    for item in planned_jobs:
        job = item[2]
        key = (
            str(getattr(job, "case_name", "")).casefold(),
            int(getattr(job, "run_number", 0)),
        )
        if current and (key != current_key or len(current) >= group_size):
            groups.append(tuple(current))
            current = []
        current_key = key
        current.append(item)
    if current:
        groups.append(tuple(current))
    return groups


def _cleanup_plot_stages(plans: Iterable[_PlotBatchPlan]) -> None:
    for plan in plans:
        shutil.rmtree(plan.stage_dir, ignore_errors=True)


def _plot_job_context(plan: _PlotBatchPlan, job: Any | None) -> str:
    if job is None:
        return plan.event_name
    return f"{job.case_name} | {job.group_label} | run {job.run_number}"


def _plot_failure(plan: _PlotBatchPlan, job: Any | None, error: Exception) -> RuntimeError:
    return RuntimeError(
        f"Plot rendering failed for {plan.batch_file.name}: "
        f"{_plot_job_context(plan, job)} | {error}"
    )


_PLOT_JOB_MANIFEST_FIELDS = (
    "mode",
    "case_name",
    "run_number",
    "group_label",
    "trace_type",
    "show_three_phase_overview",
    "show_limits",
    "show_tov_windows",
    "tov_window_s",
    "tov_window_count",
    "legends_left",
    "excel_export",
    "time_start_s",
    "time_end_s",
    "voltage_kv",
    "limits",
    "annotate_max",
    "annotate_min",
    "plot_variant",
)


def _manifest_value(value: Any) -> Any:
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, Path):
        return value.as_posix()
    if is_dataclass(value) and not isinstance(value, type):
        return {
            field.name: _manifest_value(getattr(value, field.name))
            for field in fields(value)
        }
    if isinstance(value, dict):
        return {str(key): _manifest_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_manifest_value(item) for item in value]
    if value is None or isinstance(value, (bool, float, int, str)):
        return value
    return str(value)


def _plot_job_manifest_payload(job: Any) -> dict[str, Any]:
    if is_dataclass(job) and not isinstance(job, type):
        payload = {
            field.name: getattr(job, field.name)
            for field in fields(job)
            if field.name != "output_dir"
        }
    else:
        payload = {
            name: getattr(job, name, None)
            for name in _PLOT_JOB_MANIFEST_FIELDS
        }
    return _manifest_value(payload)


def _plot_manifest_path(project_root: Path, path: Path) -> str:
    resolved_root = project_root.resolve()
    resolved_path = path.resolve()
    try:
        return resolved_path.relative_to(resolved_root).as_posix()
    except ValueError:
        return resolved_path.as_posix()


def _plot_source_manifest(
    project_root: Path,
    renderer: Any,
    jobs: Iterable[Any],
) -> list[dict[str, Any]]:
    run_index = getattr(renderer, "run_index", {})
    source_paths: dict[str, Path] = {}
    special_entries: dict[str, dict[str, Any]] = {}
    scanned_inf_paths: set[str] = set()
    for job in jobs:
        case_name = str(getattr(job, "case_name", ""))
        run_number = int(getattr(job, "run_number", 0))
        try:
            inf_path = run_index.get(case_name, {}).get(run_number)
        except AttributeError:
            inf_path = None
        if inf_path is None:
            marker = f"missing:{case_name}:{run_number}"
            special_entries[marker] = {"path": marker, "missing": True}
            continue

        inf_path = Path(inf_path)
        if not inf_path.is_file():
            relative = _plot_manifest_path(project_root, inf_path)
            special_entries[f"missing:{relative}"] = {"path": relative, "missing": True}
            continue

        resolved_inf = str(inf_path.resolve())
        if resolved_inf in scanned_inf_paths:
            continue
        scanned_inf_paths.add(resolved_inf)
        source_paths[resolved_inf] = inf_path
        try:
            candidates = list(inf_path.parent.iterdir())
        except OSError:
            relative = _plot_manifest_path(project_root, inf_path.parent)
            special_entries[f"directory:{relative}"] = {
                "path": relative,
                "directory_error": True,
            }
            continue
        stem = inf_path.stem.casefold()
        for candidate in candidates:
            candidate_name = candidate.name.casefold()
            if (
                candidate.is_file()
                and candidate.suffix.casefold() == ".out"
                and (
                    candidate_name.startswith(f"{stem}_")
                    or candidate_name.startswith("statistic")
                )
            ):
                source_paths[str(candidate.resolve())] = candidate

    entries: list[dict[str, Any]] = []
    for path in sorted(source_paths.values(), key=lambda item: str(item).casefold()):
        relative = _plot_manifest_path(project_root, path)
        try:
            stat = path.stat()
        except OSError:
            entries.append({"path": relative, "missing": True})
        else:
            entries.append(
                {
                    "path": relative,
                    "size": int(stat.st_size),
                    "mtime_ns": int(stat.st_mtime_ns),
                }
            )
    entries.extend(special_entries.values())
    return sorted(entries, key=lambda item: str(item["path"]).casefold())


def _plot_signature(plan: _PlotBatchPlan, renderer: Any) -> str:
    payload = {
        "version": PLOT_MANIFEST_VERSION,
        "sustained_sdpf_result_version": (
            sustained_sdpf.RESULT_VERSION
            if plan.event_name == sustained_sdpf.SUSTAINED_SDPF
            else None
        ),
        "sources": _plot_source_manifest(plan.project_root, renderer, plan.jobs),
        "jobs": [
            _plot_job_manifest_payload(job)
            for job in sorted(plan.jobs, key=_plot_job_sort_key)
        ],
    }
    return sustained_sdpf.make_signature(payload)


def _plot_final_output_path(plan: _PlotBatchPlan, path: Path) -> Path | None:
    try:
        relative = path.resolve().relative_to(plan.stage_dir.resolve())
    except (OSError, ValueError):
        return None
    return plan.output_dir / relative


def _plot_output_records(plan: _PlotBatchPlan, outputs: Iterable[Any]) -> list[dict[str, Any]]:
    output_list = list(outputs)
    if len(output_list) != len(plan.jobs) or any(output is None for output in output_list):
        raise RuntimeError(f"Plot cache output count does not match {plan.batch_file.name}.")
    records: list[dict[str, Any]] = []
    for output in output_list:
        for raw_path in filter(None, (output.png_path, output.excel_path)):
            final_path = _plot_final_output_path(plan, Path(raw_path))
            if final_path is None or not final_path.is_file():
                raise RuntimeError(f"Plot output is missing after commit: {raw_path}")
            try:
                relative = final_path.resolve().relative_to(plan.project_root.resolve()).as_posix()
                stat = final_path.stat()
            except (OSError, ValueError) as exc:
                raise RuntimeError(f"Could not record plot output: {final_path}") from exc
            records.append(
                {
                    "path": relative,
                    "size": int(stat.st_size),
                    "mtime_ns": int(stat.st_mtime_ns),
                }
            )
    return sorted(records, key=lambda item: str(item["path"]).casefold())


def _safe_project_cache_path(project_root: Path, value: Any) -> Path | None:
    if not isinstance(value, str) or not value:
        return None
    relative = Path(value)
    if relative.is_absolute() or ".." in relative.parts:
        return None
    candidate = project_root / relative
    try:
        candidate.resolve().relative_to(project_root.resolve())
    except (OSError, ValueError):
        return None
    return candidate


def _plot_cache_key(project_root: Path, output_dir: Path) -> str:
    return _plot_manifest_path(project_root, output_dir)


def _clear_plot_cache_entry(project_root: Path, output_dir: Path) -> None:
    cache = storage.load_project_analysis_cache(project_root)
    entries = cache.get("plots")
    if not isinstance(entries, dict) or _plot_cache_key(project_root, output_dir) not in entries:
        return
    entries.pop(_plot_cache_key(project_root, output_dir), None)
    if not entries:
        cache.pop("plots", None)
    storage.save_project_analysis_cache(project_root, cache)


def _write_plot_manifest(
    plan: _PlotBatchPlan,
    renderer: Any,
    outputs: Iterable[Any],
) -> None:
    cache = storage.load_project_analysis_cache(plan.project_root)
    entries = cache.setdefault("plots", {})
    if not isinstance(entries, dict):
        entries = {}
        cache["plots"] = entries
    entries[_plot_cache_key(plan.project_root, plan.output_dir)] = {
        "version": PLOT_MANIFEST_VERSION,
        "signature": _plot_signature(plan, renderer),
        "outputs": _plot_output_records(plan, outputs),
    }
    storage.save_project_analysis_cache(plan.project_root, cache)
    try:
        (plan.output_dir / LEGACY_PLOT_MANIFEST_FILENAME).unlink(missing_ok=True)
    except OSError:
        pass


def _plot_manifest_matches(plan: _PlotBatchPlan, renderer: Any) -> bool:
    output_dir = plan.output_dir
    if not output_dir.is_dir():
        return False
    cache = storage.load_project_analysis_cache(plan.project_root)
    entries = cache.get("plots")
    entry = entries.get(_plot_cache_key(plan.project_root, output_dir)) if isinstance(entries, dict) else None
    if not isinstance(entry, dict):
        return False
    if entry.get("version") != PLOT_MANIFEST_VERSION or entry.get("signature") != _plot_signature(plan, renderer):
        return False
    records = entry.get("outputs")
    if not isinstance(records, list):
        return False
    expected_paths: set[str] = set()
    for record in records:
        if not isinstance(record, dict):
            return False
        path = _safe_project_cache_path(plan.project_root, record.get("path"))
        if path is None or not path.is_file():
            return False
        try:
            relative = path.resolve().relative_to(plan.project_root.resolve()).as_posix()
        except (OSError, ValueError):
            return False
        if relative != record.get("path"):
            return False
        try:
            stat = path.stat()
        except OSError:
            return False
        if stat.st_size != record.get("size") or stat.st_mtime_ns != record.get("mtime_ns"):
            return False
        expected_paths.add(relative)

    actual_paths = {
        _plot_manifest_path(plan.project_root, path)
        for path in output_dir.glob("*.png")
        if path.is_file()
    }
    excel_dir = output_dir / "Excel"
    if excel_dir.is_dir():
        actual_paths.update(
            _plot_manifest_path(plan.project_root, path)
            for path in excel_dir.glob("*.xlsx")
            if path.is_file()
        )
    return actual_paths == expected_paths


def _log_plot_output(log: LogFn | None, output: Any) -> None:
    _log(log, f"plotter: Saved {output.png_path}")
    if output.excel_path is not None:
        _log(log, f"plotter: Saved {output.excel_path}")


def _render_plot_plans_sequential(
    plans: list[_PlotBatchPlan],
    renderer: Any,
    exporter: Any,
    log: LogFn | None,
    check_cancel: CancelFn | None,
) -> None:
    from pscad_plotter_app_v3.services.plot_execution import execute_plot_job

    current_plan: _PlotBatchPlan | None = None
    current_job = None
    try:
        for plan in plans:
            current_plan = plan
            current_job = None
            outputs: list[Any] = []
            for job in plan.jobs:
                current_job = job
                _cancel(check_cancel)
                _log(log, f"plotter: Rendering {job.case_name} | {job.group_label} | run {job.run_number}")
                output = execute_plot_job(renderer, exporter, job)
                outputs.append(output)
                _log_plot_output(log, output)
            _cancel(check_cancel)
            _replace_generated_directory(plan.stage_dir, plan.output_dir)
            _write_plot_manifest(plan, renderer, outputs)
    except OperationCancelled:
        _cleanup_plot_stages(plans)
        raise
    except Exception as exc:
        _cleanup_plot_stages(plans)
        if current_plan is not None:
            raise _plot_failure(current_plan, current_job, exc) from exc
        raise


def _render_plot_plans_parallel(
    plans: list[_PlotBatchPlan],
    renderer: Any,
    exporter: Any,
    log: LogFn | None,
    check_cancel: CancelFn | None,
) -> None:
    from pscad_plotter_app_v3.services.plot_execution import (
        PLOT_JOB_GROUP_SIZE,
        PLOT_PROCESS_QUEUE_MULTIPLIER,
        automatic_plot_worker_count,
        execute_plot_job_group_in_process,
        initialize_plot_process,
    )

    planned_jobs = sorted(
        ((plan, index, job) for plan in plans for index, job in enumerate(plan.jobs)),
        key=lambda item: _plot_job_sort_key(item[2]),
    )
    job_groups = _plot_job_groups(planned_jobs, PLOT_JOB_GROUP_SIZE)
    worker_count = min(
        automatic_plot_worker_count(len(planned_jobs)),
        len(job_groups),
    )
    if worker_count <= 1:
        _render_plot_plans_sequential(plans, renderer, exporter, log, check_cancel)
        return

    _log(
        log,
        f"Parallel plot rendering: {len(planned_jobs)} job(s) in {len(job_groups)} "
        f"{worker_count} process(es); PNG and requested Excel exports run together",
    )
    executor = None
    futures: dict[Any, tuple[tuple[_PlotBatchPlan, int, Any], ...]] = {}
    pending_groups = iter(job_groups)
    pending_exhausted = False
    completed_by_plan = {id(plan): 0 for plan in plans}
    outputs_by_plan: dict[int, list[Any | None]] = {
        id(plan): [None] * len(plan.jobs)
        for plan in plans
    }
    current_plan: _PlotBatchPlan | None = None
    current_job = None
    aborted = False
    pool_failure: BrokenProcessPool | None = None
    try:
        try:
            run_index = getattr(renderer, "run_index")
            executor = ProcessPoolExecutor(
                max_workers=worker_count,
                mp_context=multiprocessing.get_context("spawn"),
                initializer=initialize_plot_process,
                initargs=(run_index,),
            )
        except Exception as exc:
            _log(log, f"Parallel plotting unavailable; continuing sequentially: {exc}")
            _render_plot_plans_sequential(plans, renderer, exporter, log, check_cancel)
            return

        while futures or not pending_exhausted:
            while not pending_exhausted and len(futures) < worker_count * PLOT_PROCESS_QUEUE_MULTIPLIER:
                try:
                    group = next(pending_groups)
                except StopIteration:
                    pending_exhausted = True
                    break
                current_plan = group[0][0]
                current_job = group[0][2]
                _cancel(check_cancel)
                for _plan, _job_index, job in group:
                    _log(log, f"plotter: Rendering {job.case_name} | {job.group_label} | run {job.run_number}")
                future = executor.submit(
                    execute_plot_job_group_in_process,
                    tuple(item[2] for item in group),
                )
                futures[future] = group

            if not futures:
                break
            _cancel(check_cancel)
            completed, _pending = wait(futures, return_when=FIRST_COMPLETED)
            for future in completed:
                group = futures.pop(future)
                current_plan = group[0][0]
                current_job = group[0][2]
                try:
                    outputs = tuple(future.result())
                except Exception:
                    raise
                if len(outputs) != len(group):
                    raise RuntimeError("Plot worker returned an unexpected output count.")
                for (plan, job_index, _job), output in zip(group, outputs):
                    outputs_by_plan[id(plan)][job_index] = output
                    _log_plot_output(log, output)
                    completed_by_plan[id(plan)] += 1
                    if completed_by_plan[id(plan)] == len(plan.jobs):
                        _replace_generated_directory(plan.stage_dir, plan.output_dir)
                        _write_plot_manifest(plan, renderer, outputs_by_plan[id(plan)])

    except OperationCancelled:
        aborted = True
        raise
    except BrokenProcessPool as exc:
        aborted = True
        pool_failure = exc
    except Exception as exc:
        aborted = True
        if current_plan is not None:
            raise _plot_failure(current_plan, current_job, exc) from exc
        raise
    finally:
        if executor is not None:
            if aborted:
                _terminate_plot_executor(executor)
            else:
                executor.shutdown(wait=True)
        if aborted:
            _cleanup_plot_stages(plans)
    if pool_failure is not None:
        _log(log, f"Parallel plot worker failed; retrying sequentially: {pool_failure}")
        _render_plot_plans_sequential(plans, renderer, exporter, log, check_cancel)


def _terminate_plot_executor(executor: ProcessPoolExecutor) -> None:
    """Stop active plotting workers promptly after cancellation or failure."""
    from pscad_plotter_app_v3.services.plot_execution import terminate_process_executor

    terminate_process_executor(executor)


def _render_plot_plans(
    plans: list[_PlotBatchPlan],
    renderer: Any,
    exporter: Any,
    log: LogFn | None,
    check_cancel: CancelFn | None,
) -> None:
    from pscad_plotter_app_v3.services.plot_execution import MIN_PARALLEL_PLOT_JOBS

    total_jobs = sum(len(plan.jobs) for plan in plans)
    if total_jobs < MIN_PARALLEL_PLOT_JOBS:
        _render_plot_plans_sequential(plans, renderer, exporter, log, check_cancel)
        return
    _render_plot_plans_parallel(plans, renderer, exporter, log, check_cancel)


def _log_batch_import_errors(
    batch_file: Path,
    import_result: Any,
    log: LogFn | None,
) -> None:
    for error in import_result.errors:
        _log(
            log,
            f"Batch import error: {batch_file.name} | "
            f"{error.sheet_name} row {error.row_number}: {error.message}",
        )


def _render_plot_batches_direct(
    project_root: Path,
    scopes: Iterable[ScopeEntry],
    events: Iterable[str],
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    sustained_sdpf_limit_overrides: dict[str, dict[str, float]] | None = None,
    sustained_sdpf_settings: dict[str, Any] | None = None,
    sustained_sdpf_ranking_settings: dict[str, Any] | None = None,
    sustained_payloads_by_scope: Mapping[str, Mapping[str, Any]] | None = None,
    sustained_voltage_keys: Iterable[str] | None = None,
    excel_waveform_exports_enabled: bool = True,
    sustained_cache_validations_by_scope: Mapping[
        str, sustained_sdpf.SustainedSDPFCacheValidation
    ] | None = None,
    rms_settings: Mapping[str, Any] | None = None,
) -> None:
    _cancel(check_cancel)
    selected_scopes = list(scopes)
    selected_events = list(events)
    parsed_rms = rms_analysis.normalize_rms_settings(rms_settings) if rms_settings is not None else None
    if parsed_rms is not None and parsed_rms["enabled"]:
        for event_name in rms_analysis.RMS_BATCH_EVENTS.values():
            if event_name not in selected_events:
                selected_events.append(event_name)
    elif parsed_rms is not None:
        for scope in selected_scopes:
            for quantity in rms_analysis.RMS_QUANTITIES:
                _clear_rms_outputs(project_root, scope.folder, quantity)
    catalog, limits, renderer, exporter = _load_embedded_plotter_session(
        project_root,
        log,
        check_cancel,
    )

    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services.limits import LimitService

    batch_service = BatchExcelService()
    limit_service = LimitService()
    batch_dir = project_root / "Plots" / "Plot_batch"
    plans: list[_PlotBatchPlan] = []
    selected_sustained_voltage_keys = (
        tuple(sustained_voltage_keys)
        if sustained_voltage_keys is not None
        else None
    )
    try:
        for scope in selected_scopes:
            for event_name in selected_events:
                _cancel(check_cancel)
                batch_file = batch_dir / f"batch_paste_{scope.folder}_{event_name}.xlsx"
                if not batch_file.is_file():
                    continue

                output_dir = _desired_output_dir(project_root, scope.folder, event_name)
                _log(log, f"Preparing plots with embedded plotter: {batch_file.name}")
                import_result = batch_service.load_requests(batch_file, catalog)
                _log_batch_import_errors(batch_file, import_result, log)
                expected_results: dict[
                    tuple[str, str, int, str], sustained_sdpf.SustainedSDPFResult
                ] = {}
                if event_name == sustained_sdpf.SUSTAINED_SDPF and sustained_sdpf_settings is not None:
                    request_voltage_keys = tuple(
                        request.voltage_kv for request in import_result.requests
                    )
                    voltage_keys = (
                        selected_sustained_voltage_keys
                        if selected_sustained_voltage_keys is not None
                        else request_voltage_keys or None
                    )
                    expected_results = _expected_sustained_plot_results(
                        project_root,
                        scope.folder,
                        sustained_sdpf.SustainedSDPFSettings.from_mapping(
                            sustained_sdpf_settings
                        ),
                        sustained_sdpf.SustainedSDPFRankingSettings.from_mapping(
                            sustained_sdpf_ranking_settings
                        ),
                        (sustained_payloads_by_scope or {}).get(scope.folder),
                        voltage_keys,
                        log,
                        cache_validation=(sustained_cache_validations_by_scope or {}).get(
                            scope.folder
                        ),
                    )
                    batch_is_current = not import_result.errors and _sustained_batch_is_current(
                        import_result.requests,
                        expected_results,
                    )
                    if not batch_is_current:
                        if expected_results:
                            _log(
                                log,
                                f"Rebuilding stale or incomplete Sustained SDPF batch: {batch_file.name}.",
                            )
                            _clear_sustained_outputs(project_root, scope.folder)
                            _create_batch_workbook(
                                batch_file,
                                _sustained_plot_rows(
                                    (
                                        (key[0], result)
                                        for key, result in expected_results.items()
                                    ),
                                    excel_waveform_exports_enabled=excel_waveform_exports_enabled,
                                ),
                            )
                            import_result = batch_service.load_requests(batch_file, catalog)
                            if import_result.errors:
                                _log_batch_import_errors(batch_file, import_result, log)
                                raise RuntimeError(f"Batch import failed for {batch_file.name}.")
                            if not _sustained_batch_is_current(
                                import_result.requests,
                                expected_results,
                            ):
                                raise RuntimeError(
                                    f"Rebuilt Sustained SDPF batch is incomplete: {batch_file.name}."
                                )
                        else:
                            _log(
                                log,
                                f"No current Sustained SDPF selections for {scope.folder}; "
                                f"removing stale batch {batch_file.name}.",
                            )
                            batch_file.unlink(missing_ok=True)
                            _clear_sustained_outputs(project_root, scope.folder)
                            continue
                elif import_result.errors:
                    raise RuntimeError(f"Batch import failed for {batch_file.name}.")
                if not import_result.requests:
                    _log(log, f"No plot rows to render: {batch_file.name}")
                    _clear_plot_cache_entry(project_root, output_dir)
                    _remove_generated_directory(
                        output_dir,
                        project_root / "Plots" / "Generated" / scope.folder,
                    )
                    continue

                stage_dir = output_dir.parent / f".{output_dir.name}.render-{uuid.uuid4().hex}"
                jobs = []
                event_limits = (
                    _apply_sustained_sdpf_plot_limit_overrides(
                        limits,
                        sustained_sdpf_limit_overrides,
                    )
                    if event_name == sustained_sdpf.SUSTAINED_SDPF
                    else limits
                )
                jobs = _build_mm_plot_jobs(
                    import_result.requests,
                    stage_dir,
                    limit_service,
                    event_limits,
                    check_cancel,
                    log,
                )

                if not jobs:
                    _log(log, f"No plot jobs to render: {batch_file.name}")
                    shutil.rmtree(stage_dir, ignore_errors=True)
                    _clear_plot_cache_entry(project_root, output_dir)
                    _remove_generated_directory(
                        output_dir,
                        project_root / "Plots" / "Generated" / scope.folder,
                    )
                    continue

                _log(log, f"Queued {len(jobs)} plot(s): {scope.folder} | {event_name}")
                plan = _PlotBatchPlan(
                    project_root=project_root,
                    batch_file=batch_file,
                    event_name=event_name,
                    output_dir=output_dir,
                    stage_dir=stage_dir,
                    jobs=jobs,
                )
                if _plot_manifest_matches(plan, renderer):
                    _log(log, f"Skipping unchanged plots: {batch_file.name}")
                    try:
                        (output_dir / LEGACY_PLOT_MANIFEST_FILENAME).unlink(missing_ok=True)
                    except OSError:
                        pass
                    continue
                plans.append(plan)

        if not plans:
            return
        total_jobs = sum(len(plan.jobs) for plan in plans)
        _log(log, f"Rendering {total_jobs} plot(s) across {len(plans)} batch(es)")
        _render_plot_plans(plans, renderer, exporter, log, check_cancel)
    except OperationCancelled:
        _cleanup_plot_stages(plans)
        raise
    except Exception:
        _cleanup_plot_stages(plans)
        raise


def _replace_generated_directory(stage_dir: Path, output_dir: Path) -> None:
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    backup_dir = output_dir.parent / f".{output_dir.name}.previous-{uuid.uuid4().hex}"
    had_previous = output_dir.exists()
    if had_previous:
        _replace_path_with_retry(output_dir, backup_dir)
    try:
        _replace_path_with_retry(stage_dir, output_dir)
    except Exception:
        if had_previous and backup_dir.exists():
            _replace_path_with_retry(backup_dir, output_dir)
        raise
    shutil.rmtree(backup_dir, ignore_errors=True)


def _replace_path_with_retry(source: Path, target: Path) -> None:
    """Retry transient Windows file-lock failures during directory commits."""
    for attempt in range(_DIRECTORY_REPLACE_ATTEMPTS):
        try:
            source.replace(target)
            return
        except PermissionError:
            if attempt + 1 == _DIRECTORY_REPLACE_ATTEMPTS:
                raise
            time.sleep(_DIRECTORY_REPLACE_DELAY_S * (attempt + 1))


def _remove_generated_directory(output_dir: Path, scope_root: Path) -> None:
    if output_dir.exists():
        shutil.rmtree(output_dir)
    parent = output_dir.parent
    while parent != scope_root.parent and parent != scope_root:
        try:
            parent.rmdir()
        except OSError:
            break
        parent = parent.parent

def _desired_output_dir(project_root: Path, scope_folder: str, event_name: str) -> Path:
    if event_name == sustained_sdpf.SUSTAINED_SDPF:
        return project_root / "Plots" / "Generated" / scope_folder / sustained_sdpf.SUSTAINED_SDPF
    for quantity, rms_event in rms_analysis.RMS_BATCH_EVENTS.items():
        if event_name == rms_event:
            return rms_analysis.rms_output_dir(project_root, scope_folder, quantity)
    resonance_dir = resonance_checks.output_dir_for_event(project_root, scope_folder, event_name)
    if resonance_dir is not None:
        return resonance_dir
    return project_root / "Plots" / "Generated" / scope_folder / event_name


def render_plot_batches(
    project_root: Path,
    scopes: Iterable[ScopeEntry],
    events: Iterable[str],
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    sustained_sdpf_settings: dict[str, Any] | None = None,
    sustained_sdpf_heatmap_settings: Any | None = None,
    event_times: dict[str, float] | None = None,
    sustained_sdpf_limit_overrides: dict[str, dict[str, float]] | None = None,
    sustained_sdpf_ranking_settings: dict[str, Any] | None = None,
    sustained_payloads_by_scope: Mapping[str, Mapping[str, Any]] | None = None,
    sustained_voltage_keys: Iterable[str] | None = None,
    excel_waveform_exports_enabled: bool = True,
    sustained_cache_validations_by_scope: Mapping[
        str, sustained_sdpf.SustainedSDPFCacheValidation
    ] | None = None,
    rms_settings: Mapping[str, Any] | None = None,
) -> None:
    selected_scopes = list(scopes)
    selected_events = list(events)
    if rms_settings is not None and rms_analysis.normalize_rms_settings(rms_settings)["enabled"]:
        selected_events.extend(
            event
            for event in rms_analysis.RMS_BATCH_EVENTS.values()
            if event not in selected_events
        )
    _render_plot_batches_direct(
        project_root,
        selected_scopes,
        selected_events,
        log,
        check_cancel,
        sustained_sdpf_limit_overrides=sustained_sdpf_limit_overrides,
        sustained_sdpf_settings=sustained_sdpf_settings,
        sustained_sdpf_ranking_settings=sustained_sdpf_ranking_settings,
        sustained_payloads_by_scope=sustained_payloads_by_scope,
        sustained_voltage_keys=sustained_voltage_keys,
        excel_waveform_exports_enabled=excel_waveform_exports_enabled,
        sustained_cache_validations_by_scope=sustained_cache_validations_by_scope,
        rms_settings=rms_settings,
    )
    if sustained_sdpf_settings is not None or sustained_sdpf.SUSTAINED_SDPF in selected_events:
        render_sustained_heatmaps(
            project_root,
            selected_scopes,
            sustained_sdpf_settings=sustained_sdpf_settings,
            sustained_sdpf_heatmap_settings=sustained_sdpf_heatmap_settings,
            event_times=event_times,
            log=log,
            check_cancel=check_cancel,
            sustained_payloads_by_scope=sustained_payloads_by_scope,
            sustained_cache_validations_by_scope=sustained_cache_validations_by_scope,
        )


def render_sustained_heatmaps(
    project_root: Path,
    scopes: Iterable[ScopeEntry],
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    sustained_sdpf_settings: dict[str, Any] | None = None,
    sustained_sdpf_heatmap_settings: Any | None = None,
    event_times: dict[str, float] | None = None,
    sustained_payloads_by_scope: Mapping[str, Mapping[str, Any]] | None = None,
    sustained_cache_validations_by_scope: Mapping[
        str, sustained_sdpf.SustainedSDPFCacheValidation
    ] | None = None,
) -> None:
    """Regenerate Sustained SDPF heatmaps from saved analysis results."""
    selected_scopes = list(scopes)
    sustained_settings = sustained_sdpf.SustainedSDPFSettings.from_mapping(sustained_sdpf_settings)
    if not sustained_settings.enabled:
        for scope in selected_scopes:
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
        return
    heatmap_sets = sustained_sdpf_heatmap.heatmap_sets_from_mapping(sustained_sdpf_heatmap_settings)
    if not any(heatmap_set.settings.enabled for heatmap_set in heatmap_sets):
        for scope in selected_scopes:
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
        return
    for scope in selected_scopes:
        payload = (sustained_payloads_by_scope or {}).get(scope.folder)
        if not isinstance(payload, Mapping):
            payload = sustained_sdpf.load_results(project_root, scope.folder)
        if not payload:
            result_file = sustained_sdpf.result_path(project_root, scope.folder)
            if result_file.is_file():
                _log(
                    log,
                    f"Sustained SDPF cache is missing or obsolete for {scope.folder}; "
                    "run the envelope analysis before rebuilding heatmaps.",
                )
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
            continue
        observations = payload.get("observations", {})
        voltages = [str(key) for key in observations] if isinstance(observations, dict) else []
        if not voltages:
            result_mapping = payload.get("results", {})
            voltages = [str(key) for key in result_mapping] if isinstance(result_mapping, dict) else []
        if not voltages:
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
            continue
        sustained_sdpf_heatmap.clear_obsolete_heatmap_voltages(
            project_root,
            scope.folder,
            voltages,
        )
        validation = (sustained_cache_validations_by_scope or {}).get(scope.folder)
        if validation is None:
            validation = sustained_sdpf.validate_result_cache(
                payload,
                project_root,
                sustained_settings,
            )
        if not validation.valid:
            _log(
                log,
                f"Sustained SDPF heatmaps unavailable for {scope.folder}: "
                f"{validation.reason}. Rebuild envelope data/checks first.",
            )
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)
            continue
        shared_manifest_current = validation.shared_manifest_current
        for voltage in voltages:
            _cancel(check_cancel)
            sustained_sdpf_heatmap.generate_heatmap_sets(
                project_root,
                scope.folder,
                voltage,
                heatmap_sets,
                sustained_settings,
                log=log,
                check_cancel=check_cancel,
                payload=payload,
                shared_manifest_current=shared_manifest_current,
            )
