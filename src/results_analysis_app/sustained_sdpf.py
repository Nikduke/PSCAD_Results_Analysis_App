from __future__ import annotations

import hashlib
import json
import math
import re
from collections.abc import Iterable, Mapping
from dataclasses import asdict, dataclass, replace
from itertools import pairwise
from pathlib import Path
from typing import Any

import numpy as np

from results_analysis_app.common import as_bool, save_workbook_atomic
from results_analysis_app.project_config import normalize_voltage

SUSTAINED_SDPF = "Sustained_SDPF"
RESULT_FILENAME = "Sustained_SDpf.json"
SUMMARY_FILENAME = "Sustained_SDpf_summary.xlsx"
# Version 19 persists compact event descriptors for metric-specific report
# provenance.  The project screening rule remains the absolute peak per
# complete cycle; polarity-specific peaks are used by the full-wave ranking
# envelope.
# Older result caches must be rebuilt because the persisted result schema
# deliberately no longer carries flat aliases or embedded duplicate results.
RESULT_VERSION = 19
# The workbook format has its own version because summary-only changes must
# invalidate the envelope-stage output without changing the engineering-result
# schema/version.
SUMMARY_WORKBOOK_VERSION = 4
SOURCE_MANIFEST_VERSION = 1
DEFAULT_SUSTAINED_DURATION_MS = 30.0
# CIGRE/TB 913 expresses the safety factor as 1.15.  The usable threshold is
# therefore the SDPF limit divided by 1.15; this is not the same calculation as
# subtracting 15 percent (0.85 × limit).
SDPF_SAFETY_FACTOR = 1.15
NUMERIC_TOLERANCE = 1e-12
CUMULATIVE_STRESS_SELECTION = "cumulative_stress"
CONTINUOUS_DURATION_SELECTION = "continuous_duration"
HIGHEST_VOLTAGE_SUSTAINED_SELECTION = "highest_voltage_sustained"
ACTUAL_SDPF_POPULATION = "actual_sdpf"
SAFETY_MARGIN_ONLY_POPULATION = "safety_margin_only"
RANKING_SELECTIONS = (
    HIGHEST_VOLTAGE_SUSTAINED_SELECTION,
    CUMULATIVE_STRESS_SELECTION,
    CONTINUOUS_DURATION_SELECTION,
)
RANKING_POPULATIONS = (
    ACTUAL_SDPF_POPULATION,
    SAFETY_MARGIN_ONLY_POPULATION,
)


@dataclass(frozen=True, slots=True)
class SustainedSDPFSettings:
    enabled: bool = False
    duration_ms: float = DEFAULT_SUSTAINED_DURATION_MS

    @classmethod
    def from_mapping(cls, data: Mapping[str, Any] | None) -> SustainedSDPFSettings:
        data = data if isinstance(data, Mapping) else {}
        return cls(
            enabled=as_bool(data.get("enabled", data.get("sustained_sdpf_enabled", False))),
            duration_ms=normalize_duration_ms(
                data.get(
                    "duration_ms",
                    data.get("minimum_sustained_duration_ms", DEFAULT_SUSTAINED_DURATION_MS),
                )
            ),
        )

    @classmethod
    def from_session(cls, session: Any) -> SustainedSDPFSettings:
        return cls.from_mapping(
            {
                "enabled": getattr(session, "sustained_sdpf_enabled", False),
                "duration_ms": getattr(
                    session,
                    "sustained_sdpf_duration_ms",
                    DEFAULT_SUSTAINED_DURATION_MS,
                ),
            }
        )

    def effective_duration(self, frequency_hz: float | None = None) -> float:
        """Return the configured physical persistence duration in seconds."""
        del frequency_hz
        return self.duration_ms / 1000.0

    def required_cycles(self, frequency_hz: float) -> int:
        """Return the minimum complete-cycle coverage implied by the duration."""
        try:
            frequency = float(frequency_hz)
        except (TypeError, ValueError):
            return 0
        if not math.isfinite(frequency) or frequency <= 0:
            return 0
        return required_cycles_for_duration(self.effective_duration(), frequency)

    def to_mapping(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True, slots=True)
class SustainedSDPFRankingSettings:
    """Project-level controls for representative output only.

    These settings deliberately live outside ``SustainedSDPFSettings`` so
    changing them cannot invalidate the engineering result cache.
    """

    highest_voltage_sustained: bool = True
    cumulative_stress: bool = True
    continuous_duration: bool = True

    @classmethod
    def from_mapping(cls, data: Mapping[str, Any] | None) -> "SustainedSDPFRankingSettings":
        data = data if isinstance(data, Mapping) else {}
        return cls(
            highest_voltage_sustained=as_bool(
                data.get(
                    HIGHEST_VOLTAGE_SUSTAINED_SELECTION,
                    data.get("highest_voltage_sustained_for_t", True),
                ),
                True,
            ),
            cumulative_stress=as_bool(
                data.get(CUMULATIVE_STRESS_SELECTION, data.get("worst_cumulative_stress", True)),
                True,
            ),
            continuous_duration=as_bool(
                data.get(
                    CONTINUOUS_DURATION_SELECTION,
                    data.get("longest_continuous_duration", True),
                ),
                True,
            ),
        )

    def to_mapping(self) -> dict[str, bool]:
        return {
            HIGHEST_VOLTAGE_SUSTAINED_SELECTION: bool(self.highest_voltage_sustained),
            CUMULATIVE_STRESS_SELECTION: bool(self.cumulative_stress),
            CONTINUOUS_DURATION_SELECTION: bool(self.continuous_duration),
        }

    def enabled_selections(self) -> tuple[str, ...]:
        enabled = {
            HIGHEST_VOLTAGE_SUSTAINED_SELECTION: self.highest_voltage_sustained,
            CUMULATIVE_STRESS_SELECTION: self.cumulative_stress,
            CONTINUOUS_DURATION_SELECTION: self.continuous_duration,
        }
        return tuple(selection for selection in RANKING_SELECTIONS if enabled[selection])


def normalize_duration_ms(value: Any, default: float = DEFAULT_SUSTAINED_DURATION_MS) -> float:
    """Return a positive finite Sustained SDPF duration in milliseconds."""
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return default
    if not math.isfinite(numeric) or numeric <= 0:
        return default
    return numeric


def required_cycles_for_duration(duration_s: float, frequency_hz: float) -> int:
    """Return the complete-cycle coverage required by a physical duration."""
    try:
        duration = float(duration_s)
        frequency = float(frequency_hz)
    except (TypeError, ValueError):
        return 0
    if not math.isfinite(duration) or duration <= 0:
        return 0
    if not math.isfinite(frequency) or frequency <= 0:
        return 0
    return max(1, int(math.ceil(duration * frequency - NUMERIC_TOLERANCE)))


def normalize_required_cycles(value: Any, default: int = 0) -> int:
    """Normalize a derived complete-cycle coverage count."""
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return default
    if not math.isfinite(numeric) or numeric < 0 or not numeric.is_integer():
        return default
    return int(numeric)


def _finite_float(data: Mapping[str, Any], key: str, default: float = 0.0) -> float:
    """Read one persisted numeric value and reject non-finite cache data."""
    try:
        value = float(data.get(key, default))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid Sustained SDPF value: {key}") from exc
    if not math.isfinite(value):
        raise ValueError(f"Invalid Sustained SDPF value: {key}")
    return value


def _finite_int(data: Mapping[str, Any], key: str, default: int = 0) -> int:
    value = _finite_float(data, key, default)
    if not value.is_integer():
        raise ValueError(f"Invalid Sustained SDPF value: {key}")
    return int(value)


@dataclass(frozen=True, slots=True)
class SDPFVoltageLimits:
    voltage_kv: float
    sdpf_lg_rms: float
    sdpf_ll_rms: float
    source: str = "workbook"

    def rms(self, measurement: str) -> float:
        return self.sdpf_lg_rms if measurement == "LGp" else self.sdpf_ll_rms

    def peak(self, measurement: str) -> float:
        return self.rms(measurement) * math.sqrt(2.0)

    def margin_rms(self, measurement: str) -> float:
        return self.rms(measurement) / SDPF_SAFETY_FACTOR

    def margin_peak(self, measurement: str) -> float:
        return self.peak(measurement) / SDPF_SAFETY_FACTOR

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def resolve_project_limits(
    project_root: str | Path,
    *,
    workbook_path: str | Path | None = None,
    overrides_path: str | Path | None = None,
    manual_overrides: Mapping[str, Mapping[str, float]] | None = None,
) -> tuple[dict[str, SDPFVoltageLimits], list[str]]:
    """Reuse the embedded plotter's effective MM limit source."""
    from pscad_plotter_app_v3.services.limits import LimitService
    from pscad_plotter_app_v3.services.project import ProjectDiscoveryService

    if workbook_path is None or overrides_path is None:
        try:
            context = ProjectDiscoveryService().discover(project_root)
        except (OSError, ValueError):
            return {}, ["Could not inspect the project while resolving SDPF limits."]
        workbook_path = workbook_path or context.workbook_path
        overrides_path = overrides_path or context.state_dir / LimitService.LIMITS_FILENAME
    workbook_path = Path(workbook_path) if workbook_path is not None else None
    overrides_path = Path(overrides_path) if overrides_path is not None else None
    service = LimitService()
    workbook_limits, warnings = service.load_workbook_limits_with_warnings(workbook_path)
    try:
        overrides = service.load_override_limits(overrides_path)
    except (OSError, TypeError, ValueError, KeyError):
        warnings.append("Invalid embedded plotter SDPF overrides; workbook limits used.")
        overrides = {}
    effective = dict(workbook_limits)
    effective.update(overrides)
    resolved: dict[str, SDPFVoltageLimits] = {}
    for key, limit in effective.items():
        try:
            voltage = float(limit.voltage_kv)
            lg = float(limit.sdpf_lg)
            ll = float(limit.sdpf_ll)
        except (AttributeError, TypeError, ValueError):
            warnings.append(f"Invalid SDPF limits for {key} kV; assessment skipped.")
            continue
        if not all(math.isfinite(value) and value > 0 for value in (voltage, lg, ll)):
            warnings.append(f"Invalid SDPF limits for {key} kV; assessment skipped.")
            continue
        resolved[f"{voltage:g}"] = SDPFVoltageLimits(voltage, lg, ll, str(getattr(limit, "source", "workbook")))
    return apply_limit_overrides(resolved, manual_overrides), warnings


def apply_limit_overrides(
    limits: Mapping[str, SDPFVoltageLimits],
    manual_overrides: Mapping[str, Mapping[str, float]] | None = None,
) -> dict[str, SDPFVoltageLimits]:
    """Apply session overrides to already-loaded workbook limits."""
    resolved = dict(limits)
    for raw_voltage, raw_values in (manual_overrides or {}).items():
        if not isinstance(raw_values, Mapping):
            continue
        try:
            voltage_key = f"{float(raw_voltage):g}"
        except (TypeError, ValueError):
            continue
        base = resolved.get(voltage_key)
        if base is None:
            continue
        values = {
            "LGp": base.sdpf_lg_rms,
            "LLp": base.sdpf_ll_rms,
        }
        for measurement, default_value in values.items():
            try:
                value = float(raw_values.get(measurement, default_value))
            except (TypeError, ValueError):
                continue
            if math.isfinite(value) and value > 0:
                values[measurement] = value
        resolved[voltage_key] = SDPFVoltageLimits(
            base.voltage_kv,
            values["LGp"],
            values["LLp"],
            "manual",
        )
    return resolved


@dataclass(frozen=True, slots=True)
class PhaseStressResult:
    measurement: str
    phase: str
    sdpf_rms_kv: float
    sdpf_peak_kv: float
    margin_rms_kv: float
    margin_peak_kv: float
    # Actual peak-envelope durations of the longest qualifying event at each
    # threshold.  These are not nominal cycle-count durations.
    longest_margin_s: float
    longest_sdpf_s: float
    # Peak and RMS measured in the most severe qualifying peak-envelope event.
    # RMS is retained as a supporting diagnostic; peak persistence is the
    # sustained-TOV registration criterion.
    sustained_peak_kv: float
    sustained_rms_kv: float
    sustained_ratio: float
    classification: str
    # ``margin_exceeded`` and ``sdpf_exceeded`` are sustained peak flags: the
    # peak envelope persists above the relevant limit for the configured
    # physical duration and every relevant complete cycle contains a peak at
    # or above that limit. The RMS flags are supporting diagnostics and do not
    # qualify a sustained TOV on their own. ``peak_kv`` is a raw diagnostic max.
    peak_kv: float = 0.0
    margin_exceeded: bool = False
    sdpf_exceeded: bool = False
    margin_rms_exceeded: bool = False
    sdpf_rms_exceeded: bool = False
    # Full-wave exposure-severity areas for the largest qualifying event at
    # each threshold. Absolute values are kV*ms; normalized values are pu*ms.
    margin_excess_area_kv_ms: float = 0.0
    margin_excess_area_norm_ms: float = 0.0
    sdpf_excess_area_kv_ms: float = 0.0
    sdpf_excess_area_norm_ms: float = 0.0
    # Full-wave envelope duration used only for the independent duration
    # selection. The existing ``longest_*_s`` fields remain the qualification
    # event durations and are intentionally not redefined here.
    margin_longest_continuous_s: float = 0.0
    sdpf_longest_continuous_s: float = 0.0
    margin_event_start_s: float = 0.0
    margin_event_end_s: float = 0.0
    sdpf_event_start_s: float = 0.0
    sdpf_event_end_s: float = 0.0
    # Highest level sustained by the same qualification peak envelope for the
    # configured physical duration. Values are retained independently for the
    # margin and actual SDPF thresholds; ratios always use the actual SDPF
    # peak limit.
    margin_sustained_t_peak_kv: float = 0.0
    margin_sustained_t_rms_kv: float = 0.0
    margin_sustained_t_ratio: float = 0.0
    sdpf_sustained_t_peak_kv: float = 0.0
    sdpf_sustained_t_rms_kv: float = 0.0
    sdpf_sustained_t_ratio: float = 0.0
    # Compact event descriptors keep report duration, area, and V_T tied to
    # the same physical run when different events win different metrics.
    margin_events: tuple[_CycleRun, ...] = ()
    sdpf_events: tuple[_CycleRun, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> PhaseStressResult:
        if not isinstance(data, Mapping):
            raise TypeError("Sustained SDPF phase data must be a mapping")
        measurement = str(data.get("measurement", "")).strip()
        phase = str(data.get("phase", "")).strip()
        if not measurement or not phase:
            raise ValueError("Sustained SDPF phase identity is missing")
        return cls(
            measurement=measurement,
            phase=phase,
            sdpf_rms_kv=_finite_float(data, "sdpf_rms_kv"),
            sdpf_peak_kv=_finite_float(data, "sdpf_peak_kv"),
            margin_rms_kv=_finite_float(data, "margin_rms_kv"),
            margin_peak_kv=_finite_float(data, "margin_peak_kv"),
            longest_margin_s=_finite_float(data, "longest_margin_s"),
            longest_sdpf_s=_finite_float(data, "longest_sdpf_s"),
            sustained_peak_kv=_finite_float(data, "sustained_peak_kv"),
            sustained_rms_kv=_finite_float(data, "sustained_rms_kv"),
            sustained_ratio=_finite_float(data, "sustained_ratio"),
            classification=str(data.get("classification", "")),
            peak_kv=_finite_float(data, "peak_kv"),
            margin_exceeded=as_bool(data.get("margin_exceeded"), False),
            sdpf_exceeded=as_bool(data.get("sdpf_exceeded"), False),
            margin_rms_exceeded=as_bool(data.get("margin_rms_exceeded"), False),
            sdpf_rms_exceeded=as_bool(data.get("sdpf_rms_exceeded"), False),
            margin_excess_area_kv_ms=_finite_float(data, "margin_excess_area_kv_ms"),
            margin_excess_area_norm_ms=_finite_float(data, "margin_excess_area_norm_ms"),
            sdpf_excess_area_kv_ms=_finite_float(data, "sdpf_excess_area_kv_ms"),
            sdpf_excess_area_norm_ms=_finite_float(data, "sdpf_excess_area_norm_ms"),
            margin_longest_continuous_s=_finite_float(
                data,
                "margin_longest_continuous_s",
            ),
            sdpf_longest_continuous_s=_finite_float(
                data,
                "sdpf_longest_continuous_s",
            ),
            margin_event_start_s=_finite_float(data, "margin_event_start_s"),
            margin_event_end_s=_finite_float(data, "margin_event_end_s"),
            sdpf_event_start_s=_finite_float(data, "sdpf_event_start_s"),
            sdpf_event_end_s=_finite_float(data, "sdpf_event_end_s"),
            margin_sustained_t_peak_kv=_finite_float(
                data,
                "margin_sustained_t_peak_kv",
            ),
            margin_sustained_t_rms_kv=_finite_float(
                data,
                "margin_sustained_t_rms_kv",
            ),
            margin_sustained_t_ratio=_finite_float(
                data,
                "margin_sustained_t_ratio",
            ),
            sdpf_sustained_t_peak_kv=_finite_float(
                data,
                "sdpf_sustained_t_peak_kv",
            ),
            sdpf_sustained_t_rms_kv=_finite_float(
                data,
                "sdpf_sustained_t_rms_kv",
            ),
            sdpf_sustained_t_ratio=_finite_float(
                data,
                "sdpf_sustained_t_ratio",
            ),
            margin_events=_cycle_runs_from_dict(data, "margin_events"),
            sdpf_events=_cycle_runs_from_dict(data, "sdpf_events"),
        )


@dataclass(frozen=True, slots=True)
class SustainedSDPFResult:
    scope_folder: str
    voltage: str
    case: str
    run: int
    mm_name: str
    fault_type: str
    duration_s: float
    governing: PhaseStressResult
    phases: tuple[PhaseStressResult, ...]
    signature: str = ""
    # Derived complete-cycle coverage for the physical duration. This is a
    # result diagnostic, not the user-configured persistence setting.
    cycle_coverage: int = 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "scope_folder": self.scope_folder,
            "voltage": self.voltage,
            "case": self.case,
            "run": self.run,
            "mm_name": self.mm_name,
            "fault_type": self.fault_type,
            "duration_s": self.duration_s,
            "phases": [phase.to_dict() for phase in self.phases],
            "cycle_coverage": self.cycle_coverage,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> SustainedSDPFResult:
        if not isinstance(data, Mapping):
            raise TypeError("Sustained SDPF result data must be a mapping")
        raw_phases = data.get("phases", [])
        if not isinstance(raw_phases, (list, tuple)):
            raise ValueError("Invalid Sustained SDPF phase list")
        if any(not isinstance(item, Mapping) for item in raw_phases):
            raise ValueError("Invalid Sustained SDPF phase entry")
        phases = tuple(PhaseStressResult.from_dict(item) for item in raw_phases)
        duration_s = _finite_float(data, "duration_s")
        if phases:
            # Do not trust a persisted governing copy when the full fixed-path
            # population is available. Recompute it so a stale or hand-edited
            # governing row cannot change ranking or the selected plot.
            governing = _select_governing_phase(phases)
        else:
            raw_governing = data.get("governing")
            if not isinstance(raw_governing, Mapping):
                raise ValueError("Sustained SDPF result has no governing phase")
            governing = PhaseStressResult.from_dict(raw_governing)
        return cls(
            scope_folder=str(data.get("scope_folder", "")),
            voltage=str(data.get("voltage", "")),
            case=str(data.get("case", "")),
            run=_finite_int(data, "run"),
            mm_name=str(data.get("mm_name", "")),
            fault_type=str(data.get("fault_type", "")),
            duration_s=duration_s,
            governing=governing,
            phases=phases,
            signature=str(data.get("signature", "")),
            cycle_coverage=normalize_required_cycles(data.get("cycle_coverage", 0)),
        )


def _finite_segments(time: np.ndarray, values: np.ndarray) -> list[tuple[np.ndarray, np.ndarray]]:
    """Return sorted finite waveform segments, preserving gaps as boundaries."""
    time = np.asarray(time, dtype=float)
    values = np.asarray(values, dtype=float)
    if time.shape != values.shape:
        raise ValueError("Time and amplitude arrays must have the same shape.")
    finite_time = np.isfinite(time)
    if finite_time.sum() < 2:
        return []
    time = time[finite_time]
    values = values[finite_time]
    order = np.argsort(time, kind="stable")
    time = time[order]
    values = values[order]
    finite_values = np.isfinite(values)
    segments: list[tuple[np.ndarray, np.ndarray]] = []
    boundaries = np.flatnonzero(np.diff(np.r_[False, finite_values, False]))
    for segment_start, segment_end in boundaries.reshape(-1, 2):
        if segment_end - segment_start >= 2:
            segments.append((time[segment_start:segment_end], values[segment_start:segment_end]))
    return segments


@dataclass(frozen=True, slots=True)
class _CycleTrack:
    starts: np.ndarray
    times: tuple[np.ndarray, ...]
    values: tuple[np.ndarray, ...]


def _full_cycle_tracks(
    time: np.ndarray,
    values: np.ndarray,
    frequency_hz: float,
) -> list[_CycleTrack]:
    """Extract complete cycles on one deterministic global time grid.

    Cycle numbering is anchored once to the first finite sample and is not
    restarted after a NaN or sampling gap. Partial and poorly covered cycles
    are discarded, while signed samples are retained for RMS diagnostics.
    """
    if not math.isfinite(frequency_hz) or frequency_hz <= 0:
        return []
    period = 1.0 / float(frequency_hz)
    if period <= 0:
        return []
    segments = _finite_segments(time, values)
    if not segments:
        return []
    reference_time = float(segments[0][0][0])
    tracks: list[_CycleTrack] = []
    for segment_time, segment_values in segments:
        if len(segment_time) < 2:
            continue
        differences = np.diff(segment_time)
        positive_differences = differences[differences > NUMERIC_TOLERANCE]
        median_dt = (
            float(np.median(positive_differences))
            if len(positive_differences)
            else period
        )
        gap_limit = max(3.0 * median_dt, NUMERIC_TOLERANCE * 100.0)
        gap_breaks = np.flatnonzero(differences > gap_limit) + 1
        boundaries = np.r_[0, gap_breaks, len(segment_time)]
        for start, end in pairwise(boundaries):
            segment_time_part = segment_time[start:end]
            segment_values_part = segment_values[start:end]
            if len(segment_time_part) < 2:
                continue
            cycle_index = np.floor(
                (segment_time_part - reference_time) / period + NUMERIC_TOLERANCE
            ).astype(np.int64)
            unique_cycles = np.unique(cycle_index)
            valid_cycles: list[int] = []
            valid_times: list[np.ndarray] = []
            valid_values: list[np.ndarray] = []
            boundary_tolerance = max(1.5 * median_dt, NUMERIC_TOLERANCE * 100.0)
            for cycle in unique_cycles:
                cycle_start_index = int(np.searchsorted(cycle_index, cycle, side="left"))
                cycle_end_index = int(np.searchsorted(cycle_index, cycle, side="right"))
                times_for_cycle = segment_time_part[cycle_start_index:cycle_end_index]
                values_for_cycle = segment_values_part[cycle_start_index:cycle_end_index]
                if len(values_for_cycle) < 2:
                    continue
                cycle_start = reference_time + int(cycle) * period
                cycle_end = cycle_start + period
                if times_for_cycle[0] > cycle_start + boundary_tolerance:
                    continue
                if times_for_cycle[-1] < cycle_end - boundary_tolerance:
                    continue
                local_differences = np.diff(times_for_cycle)
                if len(local_differences) and np.max(local_differences) > gap_limit:
                    continue
                valid_cycles.append(int(cycle))
                valid_times.append(times_for_cycle)
                valid_values.append(values_for_cycle)
            if not valid_values:
                continue
            # A missing/invalid cycle is a data boundary. Never bridge it in
            # a persistence run, even when the raw time gap is small.
            block_start = 0
            for block_end in range(1, len(valid_cycles) + 1):
                if (
                    block_end < len(valid_cycles)
                    and valid_cycles[block_end] == valid_cycles[block_end - 1] + 1
                ):
                    continue
                block_cycles = valid_cycles[block_start:block_end]
                tracks.append(
                    _CycleTrack(
                        starts=reference_time + np.asarray(block_cycles, dtype=float) * period,
                        times=tuple(valid_times[block_start:block_end]),
                        values=tuple(valid_values[block_start:block_end]),
                    )
                )
                block_start = block_end
    return tracks


@dataclass(frozen=True, slots=True)
class _CycleRun:
    start_s: float
    end_s: float
    max_peak_kv: float
    max_rms_kv: float
    excess_area_kv_s: float = 0.0
    excess_area_norm_s: float = 0.0
    longest_continuous_s: float = 0.0
    sustained_t_peak_kv: float = 0.0

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> _CycleRun:
        if not isinstance(data, Mapping):
            raise TypeError("Sustained SDPF event data must be a mapping")
        return cls(
            start_s=_finite_float(data, "start_s"),
            end_s=_finite_float(data, "end_s"),
            max_peak_kv=_finite_float(data, "max_peak_kv"),
            max_rms_kv=_finite_float(data, "max_rms_kv"),
            excess_area_kv_s=_finite_float(data, "excess_area_kv_s"),
            excess_area_norm_s=_finite_float(data, "excess_area_norm_s"),
            longest_continuous_s=_finite_float(data, "longest_continuous_s"),
            sustained_t_peak_kv=_finite_float(data, "sustained_t_peak_kv"),
        )

    @property
    def duration_s(self) -> float:
        return max(0.0, self.end_s - self.start_s)


def _cycle_runs_from_dict(
    data: Mapping[str, Any],
    key: str,
) -> tuple[_CycleRun, ...]:
    raw_events = data.get(key, ())
    if raw_events in (None, ()):
        return ()
    if not isinstance(raw_events, (list, tuple)):
        raise ValueError(f"Invalid Sustained SDPF event list: {key}")
    if any(not isinstance(item, Mapping) for item in raw_events):
        raise ValueError(f"Invalid Sustained SDPF event entry: {key}")
    return tuple(_CycleRun.from_dict(item) for item in raw_events)


@dataclass(frozen=True, slots=True)
class _CycleMetrics:
    starts: np.ndarray
    peaks: np.ndarray
    peak_times: np.ndarray
    positive_peaks: np.ndarray
    positive_peak_times: np.ndarray
    negative_peaks: np.ndarray
    negative_peak_times: np.ndarray
    rms: np.ndarray


def _cycle_metrics(
    cycle_tracks: Iterable[_CycleTrack],
) -> list[_CycleMetrics]:
    """Calculate peak and RMS once for each complete-cycle track."""
    metrics: list[_CycleMetrics] = []
    for track in cycle_tracks:
        cycle_starts = track.starts
        cycle_values = track.values
        if len(cycle_values) == 0:
            continue
        peaks_list: list[float] = []
        peak_times_list: list[float] = []
        positive_peaks_list: list[float] = []
        positive_peak_times_list: list[float] = []
        negative_peaks_list: list[float] = []
        negative_peak_times_list: list[float] = []
        sums: list[float] = []
        counts: list[int] = []
        for times_for_cycle, values_for_cycle in zip(track.times, cycle_values):
            absolute_values = np.abs(values_for_cycle)
            peak_index = int(np.argmax(absolute_values))
            peaks_list.append(float(absolute_values[peak_index]))
            peak_times_list.append(float(times_for_cycle[peak_index]))
            positive_index = int(np.argmax(values_for_cycle))
            positive_value = float(values_for_cycle[positive_index])
            if positive_value > NUMERIC_TOLERANCE:
                positive_peaks_list.append(positive_value)
                positive_peak_times_list.append(float(times_for_cycle[positive_index]))
            else:
                positive_peaks_list.append(float("nan"))
                positive_peak_times_list.append(float("nan"))
            negative_index = int(np.argmin(values_for_cycle))
            negative_value = float(values_for_cycle[negative_index])
            if negative_value < -NUMERIC_TOLERANCE:
                negative_peaks_list.append(-negative_value)
                negative_peak_times_list.append(float(times_for_cycle[negative_index]))
            else:
                negative_peaks_list.append(float("nan"))
                negative_peak_times_list.append(float("nan"))
            sums.append(float(np.sum(values_for_cycle ** 2)))
            counts.append(len(values_for_cycle))
        peaks = np.asarray(peaks_list)
        peak_times = np.asarray(peak_times_list)
        sums_array = np.asarray(sums)
        counts_array = np.asarray(counts, dtype=float)
        rms = np.sqrt(
            np.divide(
                sums_array,
                counts_array,
                out=np.zeros_like(sums_array),
                where=counts_array > 0,
            )
        )
        metrics.append(
            _CycleMetrics(
                starts=cycle_starts,
                peaks=peaks,
                peak_times=peak_times,
                positive_peaks=np.asarray(positive_peaks_list),
                positive_peak_times=np.asarray(positive_peak_times_list),
                negative_peaks=np.asarray(negative_peaks_list),
                negative_peak_times=np.asarray(negative_peak_times_list),
                rms=rms,
            )
        )
    return metrics


def _shared_cycle_metrics(
    times: tuple[np.ndarray, ...],
    values: tuple[np.ndarray, ...],
    frequency_hz: float,
    *,
    include_tracks: bool = False,
) -> tuple[list[_CycleMetrics], ...] | tuple[tuple[list[_CycleMetrics], ...], tuple[list[_CycleTrack], ...]] | None:
    """Reuse one cycle index for finite phases on one common time grid.

    PSCAD bus channels normally share one strictly increasing time vector.
    In that safe fast path the cycle boundaries are identical for every
    phase, so indexing the waveform once per phase is unnecessary.  Any
    non-finite, unsorted, duplicated, or mismatched input falls back to the
    general per-phase extractor so missing data cannot be hidden by sharing.
    """
    if not times or len(times) != len(values):
        return None
    base_time = np.asarray(times[0], dtype=float)
    if (
        base_time.ndim != 1
        or len(base_time) < 2
        or not np.all(np.isfinite(base_time))
        or not np.all(np.diff(base_time) > NUMERIC_TOLERANCE)
    ):
        return None
    waveforms: list[np.ndarray] = []
    for raw_time, raw_values in zip(times, values):
        phase_time = np.asarray(raw_time, dtype=float)
        phase_values = np.asarray(raw_values, dtype=float)
        if (
            phase_time.shape != base_time.shape
            or phase_values.shape != base_time.shape
            or not np.array_equal(phase_time, base_time)
            or not np.all(np.isfinite(phase_values))
        ):
            return None
        waveforms.append(phase_values)
    tracks = _full_cycle_tracks(base_time, waveforms[0], frequency_hz)
    if not tracks:
        return tuple([] for _ in waveforms)

    indexed_tracks: list[tuple[_CycleTrack, tuple[np.ndarray, ...]]] = []
    for track in tracks:
        indices: list[np.ndarray] = []
        for cycle_time in track.times:
            cycle_indices = np.searchsorted(base_time, cycle_time, side="left")
            if not np.array_equal(base_time[cycle_indices], cycle_time):
                return None
            indices.append(cycle_indices)
        indexed_tracks.append((track, tuple(indices)))

    metrics_by_phase: list[list[_CycleMetrics]] = []
    tracks_by_phase: list[list[_CycleTrack]] = []
    for waveform in waveforms:
        phase_tracks = [
            _CycleTrack(
                starts=track.starts,
                times=track.times,
                values=tuple(waveform[cycle_indices] for cycle_indices in indices),
            )
            for track, indices in indexed_tracks
        ]
        tracks_by_phase.append(phase_tracks)
        metrics_by_phase.append(_cycle_metrics(phase_tracks))
    if include_tracks:
        return tuple(metrics_by_phase), tuple(tracks_by_phase)
    return tuple(metrics_by_phase)


def _peak_qualifying_mask(
    metrics: _CycleMetrics,
    threshold: float,
) -> np.ndarray:
    """Mark complete cycles whose absolute peak reaches ``threshold``."""
    return metrics.peaks + NUMERIC_TOLERANCE >= threshold


def _threshold_crossing_time(
    time_a: float,
    value_a: float,
    time_b: float,
    value_b: float,
    threshold: float,
) -> float:
    """Linearly interpolate a peak-envelope threshold crossing."""
    if time_b <= time_a:
        return float(time_a)
    denominator = value_b - value_a
    if abs(denominator) <= NUMERIC_TOLERANCE:
        return float(time_a)
    fraction = (threshold - value_a) / denominator
    return float(time_a + min(1.0, max(0.0, fraction)) * (time_b - time_a))


def _peak_envelope_points(
    metrics: _CycleMetrics,
    start: int,
    end: int,
    threshold: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Return the clipped linear peak-envelope points for one run."""
    peak_times = metrics.peak_times
    peaks = metrics.peaks
    times: list[float] = []
    values: list[float] = []
    if start > 0 and peaks[start - 1] + NUMERIC_TOLERANCE < threshold:
        times.append(
            _threshold_crossing_time(
                float(peak_times[start - 1]),
                float(peaks[start - 1]),
                float(peak_times[start]),
                float(peaks[start]),
                threshold,
            )
        )
        values.append(float(threshold))
    times.extend(float(value) for value in peak_times[start:end])
    values.extend(float(value) for value in peaks[start:end])
    if end < len(peaks) and peaks[end] + NUMERIC_TOLERANCE < threshold:
        times.append(
            _threshold_crossing_time(
                float(peak_times[end - 1]),
                float(peaks[end - 1]),
                float(peak_times[end]),
                float(peaks[end]),
                threshold,
            )
        )
        values.append(float(threshold))
    return np.asarray(times, dtype=float), np.asarray(values, dtype=float)


def _threshold_segment_interval(
    time_a: float,
    value_a: float,
    time_b: float,
    value_b: float,
    threshold: float,
) -> tuple[float, float] | None:
    """Return the interval of one linear segment at/above a threshold."""
    if time_b <= time_a:
        return None
    tolerance = NUMERIC_TOLERANCE
    above_a = value_a + tolerance >= threshold
    above_b = value_b + tolerance >= threshold
    if not above_a and not above_b:
        return None
    if above_a and above_b:
        return float(time_a), float(time_b)
    crossing = _threshold_crossing_time(
        time_a,
        value_a,
        time_b,
        value_b,
        threshold,
    )
    if above_a:
        return float(time_a), max(float(time_a), crossing)
    return min(crossing, float(time_b)), float(time_b)


def _longest_continuous_above(
    times: np.ndarray,
    values: np.ndarray,
    threshold: float,
) -> float:
    """Return the longest continuous above-threshold interval of a linear envelope."""
    if (
        len(times) < 2
        or len(times) != len(values)
        or not math.isfinite(threshold)
    ):
        return 0.0
    longest = 0.0
    current_start: float | None = None
    current_end: float | None = None
    for time_a, value_a, time_b, value_b in zip(
        times[:-1],
        values[:-1],
        times[1:],
        values[1:],
    ):
        interval = _threshold_segment_interval(
            float(time_a),
            float(value_a),
            float(time_b),
            float(value_b),
            threshold,
        )
        if interval is None:
            if current_start is not None and current_end is not None:
                longest = max(longest, current_end - current_start)
            current_start = None
            current_end = None
            continue
        interval_start, interval_end = interval
        if interval_end - interval_start <= NUMERIC_TOLERANCE:
            continue
        if current_end is not None and interval_start <= current_end + NUMERIC_TOLERANCE:
            current_end = max(current_end, interval_end)
        else:
            if current_start is not None and current_end is not None:
                longest = max(longest, current_end - current_start)
            current_start = interval_start
            current_end = interval_end
    if current_start is not None and current_end is not None:
        longest = max(longest, current_end - current_start)
    return max(0.0, longest)


def _maximum_sustained_envelope_level(
    times: np.ndarray,
    values: np.ndarray,
    duration_s: float,
) -> float:
    """Return the highest piecewise-linear envelope level held for ``duration_s``."""
    if (
        len(times) < 2
        or len(times) != len(values)
        or not math.isfinite(duration_s)
        or duration_s <= 0
    ):
        return 0.0
    finite = np.isfinite(times) & np.isfinite(values)
    times = np.asarray(times[finite], dtype=float)
    values = np.asarray(values[finite], dtype=float)
    if len(times) < 2:
        return 0.0
    span = float(times[-1] - times[0])
    if span + NUMERIC_TOLERANCE < duration_s:
        return 0.0
    low = float(np.min(values))
    high = float(np.max(values))
    if low < 0.0:
        low = 0.0
    if _longest_continuous_above(times, values, high) + NUMERIC_TOLERANCE >= duration_s:
        return high
    for _ in range(48):
        midpoint = 0.5 * (low + high)
        if _longest_continuous_above(times, values, midpoint) + NUMERIC_TOLERANCE >= duration_s:
            low = midpoint
        else:
            high = midpoint
    return max(0.0, low)


def _full_wave_envelope_points(
    metrics: _CycleMetrics,
    start: int,
    end: int,
) -> tuple[np.ndarray, np.ndarray]:
    """Return chronological positive/negative peak-magnitude points.

    The absolute cycle peak remains the qualification metric.  These two
    polarity-independent points per cycle are used only for the ranking
    envelope, so a lower opposite-polarity peak cannot artificially fill the
    area or continuous-duration measure.
    """
    times: list[float] = []
    values: list[float] = []
    for index in range(start, end):
        positive_time = float(metrics.positive_peak_times[index])
        positive_peak = float(metrics.positive_peaks[index])
        if math.isfinite(positive_time) and math.isfinite(positive_peak):
            times.append(positive_time)
            values.append(positive_peak)
        negative_time = float(metrics.negative_peak_times[index])
        negative_peak = float(metrics.negative_peaks[index])
        if math.isfinite(negative_time) and math.isfinite(negative_peak):
            times.append(negative_time)
            values.append(negative_peak)
    if not times:
        return np.empty(0, dtype=float), np.empty(0, dtype=float)
    order = np.argsort(np.asarray(times), kind="stable")
    return np.asarray(times, dtype=float)[order], np.asarray(values, dtype=float)[order]


def _full_wave_segment_metrics(
    time_a: float,
    value_a: float,
    time_b: float,
    value_b: float,
    threshold: float,
) -> tuple[float, float, float, float] | None:
    """Return area and above-limit interval for one linear envelope segment."""
    if time_b <= time_a:
        return None
    excess_a = value_a - threshold
    excess_b = value_b - threshold
    duration = time_b - time_a
    if excess_a <= 0.0 and excess_b <= 0.0:
        return 0.0, 0.0, 0.0, 0.0
    if excess_a >= 0.0 and excess_b >= 0.0:
        return (
            0.5 * (excess_a + excess_b) * duration,
            time_a,
            time_b,
            duration,
        )
    crossing = _threshold_crossing_time(
        time_a,
        value_a,
        time_b,
        value_b,
        threshold,
    )
    if excess_a > 0.0:
        interval = max(0.0, crossing - time_a)
        return 0.5 * excess_a * interval, time_a, crossing, interval
    interval = max(0.0, time_b - crossing)
    return 0.5 * excess_b * interval, crossing, time_b, interval


def _envelope_metrics(
    times: np.ndarray,
    values: np.ndarray,
    threshold: float,
) -> tuple[float, float, float]:
    """Return area, normalized area, and longest continuous duration."""
    if threshold <= 0 or not math.isfinite(threshold):
        return 0.0, 0.0, 0.0
    finite = np.isfinite(times) & np.isfinite(values)
    times = np.asarray(times[finite], dtype=float)
    values = np.asarray(values[finite], dtype=float)
    if len(times) < 2:
        return 0.0, 0.0, 0.0
    area_kv_s = 0.0
    longest_continuous_s = 0.0
    current_start: float | None = None
    current_end: float | None = None
    for time_a, value_a, time_b, value_b in zip(
        times[:-1],
        values[:-1],
        times[1:],
        values[1:],
    ):
        segment = _full_wave_segment_metrics(
            float(time_a),
            float(value_a),
            float(time_b),
            float(value_b),
            threshold,
        )
        if segment is None:
            continue
        segment_area, segment_start, segment_end, segment_duration = segment
        area_kv_s += segment_area
        if segment_duration <= NUMERIC_TOLERANCE:
            continue
        if (
            current_end is not None
            and segment_start <= current_end + NUMERIC_TOLERANCE
        ):
            current_end = max(current_end, segment_end)
        else:
            if current_start is not None and current_end is not None:
                longest_continuous_s = max(
                    longest_continuous_s,
                    current_end - current_start,
                )
            current_start = segment_start
            current_end = segment_end
    if current_start is not None and current_end is not None:
        longest_continuous_s = max(
            longest_continuous_s,
            current_end - current_start,
        )
    return area_kv_s, area_kv_s / threshold, longest_continuous_s


def _full_wave_envelope_metrics(
    metrics: _CycleMetrics,
    start: int,
    end: int,
    threshold: float,
) -> tuple[float, float, float]:
    """Return full-wave area, normalized area, and longest continuous duration."""
    times, values = _full_wave_envelope_points(metrics, start, end)
    return _envelope_metrics(times, values, threshold)


def _threshold_crossing_bounds(
    times: np.ndarray,
    values: np.ndarray,
    threshold: float,
) -> tuple[float, float] | None:
    """Return first/last absolute-waveform threshold crossings in one run."""
    finite = np.isfinite(times) & np.isfinite(values)
    times = np.asarray(times[finite], dtype=float)
    values = np.abs(np.asarray(values[finite], dtype=float))
    if len(times) == 0 or len(times) != len(values):
        return None
    order = np.argsort(times, kind="stable")
    times = times[order]
    values = values[order]
    qualifying = values + NUMERIC_TOLERANCE >= threshold
    indices = np.flatnonzero(qualifying)
    if len(indices) == 0:
        return None
    first = int(indices[0])
    last = int(indices[-1])
    if first == 0:
        start = float(times[0])
    else:
        start = _threshold_crossing_time(
            float(times[first - 1]),
            float(values[first - 1]),
            float(times[first]),
            float(values[first]),
            threshold,
        )
    if last == len(times) - 1:
        end = float(times[-1])
    else:
        end = _threshold_crossing_time(
            float(times[last]),
            float(values[last]),
            float(times[last + 1]),
            float(values[last + 1]),
            threshold,
        )
    return start, max(start, end)


def _run_envelope_points(
    metrics: _CycleMetrics,
    start: int,
    end: int,
    threshold: float,
    start_s: float,
    end_s: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Build the absolute qualification envelope with exact boundaries.

    The qualification and ``V_T`` calculations use the maximum absolute peak
    per complete cycle.  The separate full-wave envelope remains responsible
    for area and continuous-duration ranking, where a lower opposite-polarity
    peak creates a real below-limit gap.
    """
    peak_times, peaks = _peak_envelope_points(metrics, start, end, threshold)
    times = np.r_[float(start_s), peak_times, float(end_s)]
    values = np.r_[float(threshold), peaks, float(threshold)]
    finite = np.isfinite(times) & np.isfinite(values)
    times = np.asarray(times[finite], dtype=float)
    values = np.asarray(values[finite], dtype=float)
    if len(times) == 0:
        return times, values
    order = np.argsort(times, kind="stable")
    times = times[order]
    values = values[order]
    unique_times: list[float] = []
    unique_values: list[float] = []
    for time, value in zip(times, values):
        if unique_times and time <= unique_times[-1] + NUMERIC_TOLERANCE:
            unique_values[-1] = max(unique_values[-1], float(value))
        else:
            unique_times.append(float(time))
            unique_values.append(float(value))
    return np.asarray(unique_times, dtype=float), np.asarray(unique_values, dtype=float)


def _threshold_cycle_runs(
    cycle_metrics: Iterable[_CycleMetrics],
    minimum_duration_s: float,
    frequency_hz: float,
    threshold: float,
    cycle_tracks: Iterable[_CycleTrack] | None = None,
) -> list[_CycleRun]:
    """Return qualifying absolute-peak runs at one threshold.

    A complete cycle qualifies when its maximum absolute peak reaches the
    threshold.  Raw cycle tracks, when available, refine the physical event
    start/end to the first and last waveform crossings; polarity-specific peak
    points are used only for full-wave area and continuous-duration ranking.
    """
    if (
        not math.isfinite(minimum_duration_s)
        or minimum_duration_s <= 0
        or not math.isfinite(frequency_hz)
        or frequency_hz <= 0
        or not math.isfinite(threshold)
        or threshold <= 0
    ):
        return []
    period = 1.0 / float(frequency_hz)
    required_cycles = required_cycles_for_duration(minimum_duration_s, frequency_hz)
    if required_cycles <= 0:
        return []
    metric_tracks = tuple(cycle_tracks) if cycle_tracks is not None else None
    runs: list[_CycleRun] = []
    for track_index, metrics in enumerate(cycle_metrics):
        cycle_starts = metrics.starts
        peaks = metrics.peaks
        qualifying = _peak_qualifying_mask(metrics, threshold)
        run_start: int | None = None
        track = (
            metric_tracks[track_index]
            if metric_tracks is not None and track_index < len(metric_tracks)
            else None
        )

        def finish(run_end: int) -> None:
            nonlocal run_start
            if run_start is None:
                return
            run_length = run_end - run_start
            if run_length < required_cycles:
                run_start = None
                return
            if track is not None and run_end <= len(track.times):
                first_cycle_times = track.times[run_start]
                last_cycle_times = track.times[run_end - 1]
                first_bounds = _threshold_crossing_bounds(
                    first_cycle_times,
                    track.values[run_start],
                    threshold,
                )
                last_bounds = _threshold_crossing_bounds(
                    last_cycle_times,
                    track.values[run_end - 1],
                    threshold,
                )
                if first_bounds is None or last_bounds is None:
                    run_start = None
                    return
                start_s = float(first_bounds[0])
                end_s = float(last_bounds[1])
                envelope_times, envelope_values = _run_envelope_points(
                    metrics,
                    run_start,
                    run_end,
                    threshold,
                    start_s,
                    end_s,
                )
            else:
                envelope_times, envelope_values = _peak_envelope_points(
                    metrics,
                    run_start,
                    run_end,
                    threshold,
                )
                if len(envelope_times) == 0:
                    run_start = None
                    return
                start_s = float(envelope_times[0])
                end_s = max(start_s, float(envelope_times[-1]))
            if len(envelope_times) == 0:
                run_start = None
                return
            if end_s - start_s + NUMERIC_TOLERANCE >= minimum_duration_s:
                (
                    excess_area_kv_s,
                    excess_area_norm_s,
                    longest_continuous_s,
                ) = _full_wave_envelope_metrics(
                    metrics,
                    run_start,
                    run_end,
                    threshold,
                )
                runs.append(
                    _CycleRun(
                        start_s=start_s,
                        end_s=end_s,
                        max_peak_kv=float(np.max(peaks[run_start:run_end])),
                        max_rms_kv=float(np.max(metrics.rms[run_start:run_end])),
                        excess_area_kv_s=excess_area_kv_s,
                        excess_area_norm_s=excess_area_norm_s,
                        longest_continuous_s=longest_continuous_s,
                        sustained_t_peak_kv=_maximum_sustained_envelope_level(
                            envelope_times,
                            envelope_values,
                            minimum_duration_s,
                        ),
                    )
                )
            run_start = None

        for index, is_qualifying in enumerate(qualifying):
            if index > 0 and cycle_starts[index] - cycle_starts[index - 1] > period * (1.0 + 1e-9):
                finish(index)
            if is_qualifying:
                if run_start is None:
                    run_start = index
            else:
                finish(index)
        finish(len(qualifying))
    return runs


def _has_rms_cycle_run(
    cycle_metrics: Iterable[_CycleMetrics],
    minimum_duration_s: float,
    frequency_hz: float,
    threshold: float,
) -> bool:
    """Return whether one threshold has a qualifying cycle run.

    This is intentionally a boolean-only path for RMS diagnostics. RMS flags
    are useful for explaining a result, but they do not need the peak-envelope
    timestamps, area integration, or ``_CycleRun`` objects used by ranking.
    """
    if (
        not math.isfinite(minimum_duration_s)
        or minimum_duration_s <= 0
        or not math.isfinite(frequency_hz)
        or frequency_hz <= 0
        or not math.isfinite(threshold)
        or threshold <= 0
    ):
        return False
    period = 1.0 / float(frequency_hz)
    required_cycles = required_cycles_for_duration(minimum_duration_s, frequency_hz)
    if required_cycles <= 0:
        return False
    for metrics in cycle_metrics:
        cycle_starts = metrics.starts
        qualifying = metrics.rms + NUMERIC_TOLERANCE >= threshold
        run_start: int | None = None

        def qualifies(run_end: int, start: int | None) -> bool:
            if start is None:
                return False
            run_length = run_end - start
            return (
                run_length >= required_cycles
                and run_length * period + NUMERIC_TOLERANCE >= minimum_duration_s
            )

        for index, is_qualifying in enumerate(qualifying):
            if (
                index > 0
                and cycle_starts[index] - cycle_starts[index - 1]
                > period * (1.0 + 1e-9)
            ):
                if qualifies(index, run_start):
                    return True
                run_start = None
            if is_qualifying:
                if run_start is None:
                    run_start = index
            elif qualifies(index, run_start):
                return True
            elif not is_qualifying:
                run_start = None
        if qualifies(len(qualifying), run_start):
            return True
    return False


def _most_severe_run(runs: Iterable[_CycleRun]) -> _CycleRun | None:
    materialized = list(runs)
    if not materialized:
        return None
    return max(
        materialized,
        key=lambda run: (
            run.excess_area_norm_s,
            -run.start_s,
        ),
    )


def _highest_sustained_t_run(runs: Iterable[_CycleRun]) -> _CycleRun | None:
    materialized = list(runs)
    if not materialized:
        return None
    return max(
        materialized,
        key=lambda run: (
            run.sustained_t_peak_kv,
            -run.start_s,
        ),
    )


def _phase_result_from_metrics(
    cycle_metrics: Iterable[_CycleMetrics],
    amplitude_peak: np.ndarray,
    measurement: str,
    phase: str,
    sdpf_rms: float,
    minimum_duration_s: float,
    frequency_hz: float = 60.0,
    cycle_tracks: Iterable[_CycleTrack] | None = None,
) -> PhaseStressResult:
    cycle_metrics = tuple(cycle_metrics)
    sdpf_peak = float(sdpf_rms) * math.sqrt(2.0)
    margin_rms = float(sdpf_rms) / SDPF_SAFETY_FACTOR
    margin_peak = sdpf_peak / SDPF_SAFETY_FACTOR
    effective_frequency = float(frequency_hz) if math.isfinite(frequency_hz) and frequency_hz > 0 else 0.0
    try:
        minimum_duration = float(minimum_duration_s)
    except (TypeError, ValueError):
        minimum_duration = 0.0
    if not math.isfinite(minimum_duration):
        minimum_duration = 0.0
    finite_peaks = np.asarray(amplitude_peak, dtype=float)
    finite_peaks = np.abs(finite_peaks[np.isfinite(finite_peaks)])
    peak_kv = float(np.max(finite_peaks)) if len(finite_peaks) else 0.0
    margin_peak_runs = _threshold_cycle_runs(
        cycle_metrics,
        minimum_duration,
        effective_frequency,
        margin_peak,
        cycle_tracks,
    )
    sdpf_peak_runs = _threshold_cycle_runs(
        cycle_metrics,
        minimum_duration,
        effective_frequency,
        sdpf_peak,
        cycle_tracks,
    )
    # Keep RMS persistence as a separate diagnostic.  It is intentionally not
    # allowed to register a sustained TOV: the project rule is peak-based.
    margin_rms_exceeded = _has_rms_cycle_run(
        cycle_metrics, minimum_duration, effective_frequency, margin_rms
    )
    sdpf_rms_exceeded = _has_rms_cycle_run(
        cycle_metrics, minimum_duration, effective_frequency, sdpf_rms
    )
    margin_exceeded = bool(margin_peak_runs)
    sdpf_exceeded = bool(sdpf_peak_runs)
    margin_duration = max((run.duration_s for run in margin_peak_runs), default=0.0)
    sdpf_duration = max((run.duration_s for run in sdpf_peak_runs), default=0.0)
    margin_longest_continuous = max(
        (run.longest_continuous_s for run in margin_peak_runs),
        default=0.0,
    )
    sdpf_longest_continuous = max(
        (run.longest_continuous_s for run in sdpf_peak_runs),
        default=0.0,
    )
    margin_run = _most_severe_run(margin_peak_runs)
    sdpf_run = _most_severe_run(sdpf_peak_runs)
    margin_t_run = _highest_sustained_t_run(margin_peak_runs)
    sdpf_t_run = _highest_sustained_t_run(sdpf_peak_runs)
    governing_run = sdpf_run or margin_run
    sustained_peak = governing_run.max_peak_kv if governing_run is not None else 0.0
    sustained_rms = governing_run.max_rms_kv if governing_run is not None else 0.0
    ratio = sustained_peak / sdpf_peak if sdpf_peak > 0 else 0.0
    duration_label = f"at least {minimum_duration * 1000.0:g} ms of peak-envelope persistence"
    if sdpf_exceeded:
        classification = f"SDPF exceeded for {duration_label} (peak)"
    elif margin_exceeded:
        classification = (
            f"SDPF safety margin exceeded; SDPF not exceeded for {duration_label} "
            "(peak)"
        )
    else:
        classification = f"No sustained SDPF safety-margin exceedance for {duration_label} (peak)"
    return PhaseStressResult(
        measurement=measurement,
        phase=phase,
        sdpf_rms_kv=float(sdpf_rms),
        sdpf_peak_kv=sdpf_peak,
        margin_rms_kv=margin_rms,
        margin_peak_kv=margin_peak,
        longest_margin_s=margin_duration,
        longest_sdpf_s=sdpf_duration,
        sustained_peak_kv=sustained_peak,
        sustained_rms_kv=sustained_rms,
        sustained_ratio=ratio,
        classification=classification,
        peak_kv=peak_kv,
        margin_exceeded=margin_exceeded,
        sdpf_exceeded=sdpf_exceeded,
        margin_rms_exceeded=margin_rms_exceeded,
        sdpf_rms_exceeded=sdpf_rms_exceeded,
        margin_excess_area_kv_ms=max(
            (run.excess_area_kv_s * 1000.0 for run in margin_peak_runs),
            default=0.0,
        ),
        margin_excess_area_norm_ms=max(
            (run.excess_area_norm_s * 1000.0 for run in margin_peak_runs),
            default=0.0,
        ),
        sdpf_excess_area_kv_ms=max(
            (run.excess_area_kv_s * 1000.0 for run in sdpf_peak_runs),
            default=0.0,
        ),
        sdpf_excess_area_norm_ms=max(
            (run.excess_area_norm_s * 1000.0 for run in sdpf_peak_runs),
            default=0.0,
        ),
        margin_longest_continuous_s=margin_longest_continuous,
        sdpf_longest_continuous_s=sdpf_longest_continuous,
        margin_event_start_s=margin_run.start_s if margin_run is not None else 0.0,
        margin_event_end_s=margin_run.end_s if margin_run is not None else 0.0,
        sdpf_event_start_s=sdpf_run.start_s if sdpf_run is not None else 0.0,
        sdpf_event_end_s=sdpf_run.end_s if sdpf_run is not None else 0.0,
        margin_sustained_t_peak_kv=(
            margin_t_run.sustained_t_peak_kv if margin_t_run is not None else 0.0
        ),
        margin_sustained_t_rms_kv=(
            margin_t_run.sustained_t_peak_kv / math.sqrt(2.0)
            if margin_t_run is not None
            else 0.0
        ),
        margin_sustained_t_ratio=(
            margin_t_run.sustained_t_peak_kv / sdpf_peak
            if margin_t_run is not None and sdpf_peak > 0
            else 0.0
        ),
        sdpf_sustained_t_peak_kv=(
            sdpf_t_run.sustained_t_peak_kv if sdpf_t_run is not None else 0.0
        ),
        sdpf_sustained_t_rms_kv=(
            sdpf_t_run.sustained_t_peak_kv / math.sqrt(2.0)
            if sdpf_t_run is not None
            else 0.0
        ),
        sdpf_sustained_t_ratio=(
            sdpf_t_run.sustained_t_peak_kv / sdpf_peak
            if sdpf_t_run is not None and sdpf_peak > 0
            else 0.0
        ),
        margin_events=tuple(margin_peak_runs),
        sdpf_events=tuple(sdpf_peak_runs),
    )


def analyze_phase_amplitude(
    time: np.ndarray,
    amplitude_peak: np.ndarray,
    measurement: str,
    phase: str,
    sdpf_rms: float,
    minimum_duration_s: float,
    frequency_hz: float = 60.0,
) -> PhaseStressResult:
    """Analyze one phase using its complete-cycle peak envelope."""
    effective_frequency = (
        float(frequency_hz)
        if math.isfinite(frequency_hz) and frequency_hz > 0
        else 0.0
    )
    cycle_tracks = _full_cycle_tracks(time, amplitude_peak, effective_frequency)
    cycle_metrics = _cycle_metrics(cycle_tracks)
    return _phase_result_from_metrics(
        cycle_metrics,
        amplitude_peak,
        measurement,
        phase,
        sdpf_rms,
        minimum_duration_s,
        effective_frequency,
        cycle_tracks,
    )


def analyze_waveform(
    time: np.ndarray,
    values: np.ndarray,
    measurement: str,
    phase: str,
    sdpf_rms: float,
    minimum_duration_s: float,
    frequency_hz: float,
) -> PhaseStressResult:
    return analyze_phase_amplitude(
        time,
        values,
        measurement,
        phase,
        sdpf_rms,
        minimum_duration_s,
        frequency_hz,
    )


def analyze_bus_phase_results(
    measurement: str,
    times: Iterable[np.ndarray],
    values: Iterable[np.ndarray],
    signals: Iterable[str],
    limits: SDPFVoltageLimits,
    minimum_duration_s: float,
    frequency_hz: float,
) -> list[PhaseStressResult]:
    """Analyze all phases, sharing cycle indexing when their grid is common."""
    time_values = tuple(times)
    waveform_values = tuple(values)
    phase_count = min(len(time_values), len(waveform_values))
    if phase_count == 0:
        return []
    labels = phase_labels(measurement, tuple(signals))
    sdpf_rms = limits.rms(measurement)
    effective_frequency = (
        float(frequency_hz)
        if math.isfinite(frequency_hz) and frequency_hz > 0
        else 0.0
    )
    shared_data = _shared_cycle_metrics(
        time_values[:phase_count],
        waveform_values[:phase_count],
        effective_frequency,
        include_tracks=True,
    )
    results: list[PhaseStressResult] = []
    for index in range(phase_count):
        phase = labels[index] if index < len(labels) else f"Phase {index + 1}"
        waveform = np.asarray(waveform_values[index], dtype=float)
        if shared_data is None:
            result = analyze_phase_amplitude(
                time_values[index],
                waveform,
                measurement,
                phase,
                sdpf_rms,
                minimum_duration_s,
                effective_frequency,
            )
        else:
            shared_metrics, shared_tracks = shared_data
            result = _phase_result_from_metrics(
                shared_metrics[index],
                waveform,
                measurement,
                phase,
                sdpf_rms,
                minimum_duration_s,
                effective_frequency,
                shared_tracks[index],
            )
        results.append(result)
    return results


def analyze_bus_waveforms(
    measurement: str,
    times: Iterable[np.ndarray],
    values: Iterable[np.ndarray],
    signals: Iterable[str],
    limits: SDPFVoltageLimits,
    case: str,
    run: int,
    mm_name: str,
    minimum_duration_s: float,
    frequency_hz: float,
) -> list[dict[str, Any]]:
    return [
        {
            "case": str(case),
            "run": int(run),
            "mm_name": str(mm_name),
            **result.to_dict(),
        }
        for result in analyze_bus_phase_results(
            measurement,
            times,
            values,
            signals,
            limits,
            minimum_duration_s,
            frequency_hz,
        )
    ]


def phase_labels(measurement: str, signals: Iterable[str]) -> list[str]:
    labels: list[str] = []
    fallback_lg = ("A-G", "B-G", "C-G")
    fallback_ll = ("A-B", "B-C", "C-A")
    is_lg = str(measurement).casefold() == "lgp"
    for index, signal in enumerate(signals):
        text = str(signal).upper()
        marker = re.search(r"(?:^|[-_])(LGp|LLp)(?:[-_]|$)", text, flags=re.IGNORECASE)
        suffix = (
            text[marker.end():]
            if marker is not None
            else re.split(r"[-_]", text)[-1]
        )
        tokens = re.findall(r"[ABC]{1,2}", suffix)
        if is_lg:
            phase = next((token for token in reversed(tokens) if len(token) == 1), None)
            labels.append(f"{phase or fallback_lg[min(index, 2)][0]}-G")
        else:
            pair = next(
                (
                    token
                    for token in reversed(tokens)
                    if token in {"AB", "BC", "CA"}
                ),
                None,
            )
            if pair is None and len(tokens) >= 2:
                candidate = "".join(tokens[-2:])
                pair = candidate if candidate in {"AB", "BC", "CA"} else None
            labels.append(
                f"{pair[0]}-{pair[1]}" if pair else fallback_ll[min(index, 2)]
            )
    return labels


def classify_rows(
    scope_folder: str,
    voltage: str,
    case: str,
    run: int,
    mm_name: str,
    rows: Iterable[PhaseStressResult],
    duration_s: float,
    fault_type: str = "",
    signature: str = "",
    cycle_coverage: int = 0,
) -> SustainedSDPFResult | None:
    phases = tuple(row for row in rows if isinstance(row, PhaseStressResult))
    if not phases:
        return None
    governing = _select_governing_phase(phases)
    return SustainedSDPFResult(
        scope_folder=scope_folder,
        voltage=str(voltage),
        case=str(case),
        run=int(run),
        mm_name=str(mm_name),
        fault_type=str(fault_type or ""),
        duration_s=float(duration_s),
        governing=governing,
        phases=tuple(sorted(phases, key=lambda row: (row.measurement, row.phase))),
        signature=signature,
        cycle_coverage=normalize_required_cycles(cycle_coverage),
    )


def result_population(result: SustainedSDPFResult) -> str | None:
    """Return the mutually exclusive representative population of a result."""
    phases = result.phases or (result.governing,)
    if any(phase.sdpf_exceeded for phase in phases):
        return ACTUAL_SDPF_POPULATION
    if any(phase.margin_exceeded for phase in phases):
        return SAFETY_MARGIN_ONLY_POPULATION
    return None


def _selection_key(
    item: tuple[SustainedSDPFResult, PhaseStressResult],
    metric: str,
    population: str,
) -> tuple[Any, ...]:
    result, phase = item
    if metric == HIGHEST_VOLTAGE_SUSTAINED_SELECTION:
        value = (
            phase.sdpf_sustained_t_ratio
            if population == ACTUAL_SDPF_POPULATION
            else phase.margin_sustained_t_ratio
        )
    elif metric == CUMULATIVE_STRESS_SELECTION:
        value = (
            phase.sdpf_excess_area_norm_ms
            if population == ACTUAL_SDPF_POPULATION
            else phase.margin_excess_area_norm_ms
        )
    else:
        value = (
            phase.sdpf_longest_continuous_s
            if population == ACTUAL_SDPF_POPULATION
            else phase.margin_longest_continuous_s
        )
    return (
        value,
        str(result.case).casefold(),
        int(result.run),
        str(result.mm_name).casefold(),
        str(phase.measurement).casefold(),
        str(phase.phase).casefold(),
    )


def _selected_candidate(
    candidates: list[tuple[SustainedSDPFResult, PhaseStressResult]],
    metric: str,
    population: str,
) -> SustainedSDPFResult | None:
    if not candidates:
        return None
    result, phase = max(
        candidates,
        key=lambda item: _selection_key(item, metric, population),
    )
    return result if phase == result.governing else replace(result, governing=phase)


def select_representatives(
    results: Iterable[SustainedSDPFResult],
) -> dict[str, dict[str, SustainedSDPFResult | None]]:
    """Select each ranking criterion independently within each population."""
    materialized = [
        result for result in results if isinstance(result, SustainedSDPFResult)
    ]
    candidates_by_population: dict[
        str,
        list[tuple[SustainedSDPFResult, PhaseStressResult]],
    ] = {population: [] for population in RANKING_POPULATIONS}
    for result in materialized:
        population = result_population(result)
        if population is None:
            continue
        phases = result.phases or (result.governing,)
        candidates_by_population[population].extend(
            (result, phase)
            for phase in phases
            if (
                phase.sdpf_exceeded
                if population == ACTUAL_SDPF_POPULATION
                else phase.margin_exceeded and not phase.sdpf_exceeded
            )
        )
    return {
        population: {
            metric: _selected_candidate(
                candidates_by_population[population],
                metric,
                population,
            )
            for metric in RANKING_SELECTIONS
        }
        for population in RANKING_POPULATIONS
    }


def select_governing(results: Iterable[SustainedSDPFResult]) -> SustainedSDPFResult | None:
    """Select the cumulative-stress governing result.

    Qualification is already complete before this function runs. Short
    events remain excluded from both sustained selections and stay with the
    existing SFO/AFO analysis.
    """
    materialized = list(results)
    representatives = select_representatives(materialized)
    return (
        representatives[ACTUAL_SDPF_POPULATION][CUMULATIVE_STRESS_SELECTION]
        or representatives[SAFETY_MARGIN_ONLY_POPULATION][CUMULATIVE_STRESS_SELECTION]
    )


def candidate_category(row: PhaseStressResult) -> int:
    """Return the sustained-TOV severity for one fixed phase/pair.

    ``2`` is an SDPF-limit candidate, ``1`` is a safety-margin-only candidate,
    and ``0`` is not a candidate. Qualification is decided by the explicit
    peak-envelope flags; duration and diagnostic ratios are not fallbacks.
    """
    if row.sdpf_exceeded:
        return 2
    if row.margin_exceeded:
        return 1
    return 0


def _select_governing_phase(
    phases: Iterable[PhaseStressResult],
) -> PhaseStressResult:
    """Choose the worst fixed phase, preferring sustained candidates."""
    phase_values = tuple(phases)
    if not phase_values:
        raise ValueError("At least one phase is required")
    qualified = tuple(
        row for row in phase_values if candidate_category(row) > 0
    )
    return max(
        qualified or phase_values,
        key=_phase_governing_key,
    )


def _candidate_ranking_values(
    row: PhaseStressResult,
) -> tuple[int, float, float]:
    """Return category, normalized area, and continuous duration."""
    category = candidate_category(row)
    if row.sdpf_exceeded:
        area = row.sdpf_excess_area_norm_ms
        event_duration = row.sdpf_longest_continuous_s
    elif row.margin_exceeded:
        area = row.margin_excess_area_norm_ms
        event_duration = row.margin_longest_continuous_s
    else:
        area = 0.0
        event_duration = 0.0
    return category, area, event_duration


def _phase_governing_key(row: PhaseStressResult) -> tuple[Any, ...]:
    """Sort fixed phases by qualification and normalized excess exposure."""
    category, area, _event_duration = _candidate_ranking_values(row)
    return (
        category,
        area,
        -len(row.phase),
        row.measurement,
        row.phase,
    )


def _result_rank_key(result: SustainedSDPFResult) -> tuple[Any, ...]:
    return (
        *_phase_governing_key(result.governing),
        str(result.case).casefold(),
        int(result.run),
        str(result.mm_name).casefold(),
    )


def rank_results(results: Iterable[SustainedSDPFResult]) -> list[SustainedSDPFResult]:
    """Return all fixed LG/LL result rows in the report's severity order."""
    return sorted(
        [result for result in results if isinstance(result, SustainedSDPFResult)],
        key=_result_rank_key,
        reverse=True,
    )


_POPULATION_LABELS = {
    ACTUAL_SDPF_POPULATION: "Actual SDPF",
    SAFETY_MARGIN_ONLY_POPULATION: "Safety-margin-only",
}

_SELECTION_LABELS = {
    HIGHEST_VOLTAGE_SUSTAINED_SELECTION: "Highest sustained T",
    CUMULATIVE_STRESS_SELECTION: "Cumulative stress",
    CONTINUOUS_DURATION_SELECTION: "Longest continuous duration",
}


def _population_label(population: str | None) -> str:
    return _POPULATION_LABELS.get(population, "No qualifying sustained TOV")


def _summary_limit(
    limits_by_voltage: Mapping[str, SDPFVoltageLimits] | None,
    voltage: str,
) -> SDPFVoltageLimits | None:
    limits = limits_by_voltage or {}
    direct = limits.get(str(voltage))
    if direct is not None:
        return direct
    try:
        return limits.get(f"{float(voltage):g}")
    except (TypeError, ValueError):
        return None


def _summary_voltage(value: str) -> float | str:
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)


def _summary_voltage_sort_key(value: str) -> tuple[int, float | str]:
    parsed = _summary_voltage(value)
    return (0, parsed) if isinstance(parsed, float) else (1, parsed)


def _qualifying_path_labels(
    result: SustainedSDPFResult,
) -> tuple[str, ...]:
    return tuple(
        f"{phase.measurement}:{phase.phase}"
        for phase in result.phases
        if candidate_category(phase) > 0
    )


def selected_event(
    phase: PhaseStressResult,
    population: str | None,
    metric: str = CUMULATIVE_STRESS_SELECTION,
) -> _CycleRun | None:
    """Return the event that supplies one population/selection metric."""
    if population == ACTUAL_SDPF_POPULATION:
        events = phase.sdpf_events
    elif population == SAFETY_MARGIN_ONLY_POPULATION:
        events = phase.margin_events
    else:
        return None
    if not events:
        return None
    if metric == HIGHEST_VOLTAGE_SUSTAINED_SELECTION:
        return max(events, key=lambda event: (event.sustained_t_peak_kv, -event.start_s))
    if metric == CONTINUOUS_DURATION_SELECTION:
        return max(events, key=lambda event: (event.longest_continuous_s, -event.start_s))
    return max(events, key=lambda event: (event.excess_area_norm_s, -event.start_s))


def _population_metric_values(
    phase: PhaseStressResult,
    population: str | None,
    metric: str = CUMULATIVE_STRESS_SELECTION,
) -> tuple[
    str,
    float | None,
    float | None,
    float | None,
    float | None,
    float | None,
    float | None,
    float | None,
    float | None,
]:
    """Return compact threshold metrics; the V_T/SDPF value is already a percent."""
    event = selected_event(phase, population, metric)
    if population == ACTUAL_SDPF_POPULATION:
        if event is not None:
            ratio = (
                event.sustained_t_peak_kv / phase.sdpf_peak_kv * 100.0
                if phase.sdpf_peak_kv > 0
                else 0.0
            )
            return (
                "SDPF",
                float(phase.sdpf_peak_kv),
                float(event.excess_area_norm_s * 1000.0),
                float(event.duration_s * 1000.0),
                float(event.longest_continuous_s * 1000.0),
                float(event.sustained_t_peak_kv),
                float(ratio),
                float(event.start_s),
                float(event.end_s),
            )
        return (
            "SDPF",
            float(phase.sdpf_peak_kv),
            float(phase.sdpf_excess_area_norm_ms),
            float(phase.longest_sdpf_s * 1000.0),
            float(phase.sdpf_longest_continuous_s * 1000.0),
            float(phase.sdpf_sustained_t_peak_kv),
            float(phase.sdpf_sustained_t_ratio * 100.0),
            float(phase.sdpf_event_start_s),
            float(phase.sdpf_event_end_s),
        )
    if population == SAFETY_MARGIN_ONLY_POPULATION:
        if event is not None:
            ratio = (
                event.sustained_t_peak_kv / phase.sdpf_peak_kv * 100.0
                if phase.sdpf_peak_kv > 0
                else 0.0
            )
            return (
                "SDPF/1.15",
                float(phase.margin_peak_kv),
                float(event.excess_area_norm_s * 1000.0),
                float(event.duration_s * 1000.0),
                float(event.longest_continuous_s * 1000.0),
                float(event.sustained_t_peak_kv),
                float(ratio),
                float(event.start_s),
                float(event.end_s),
            )
        return (
            "SDPF/1.15",
            float(phase.margin_peak_kv),
            float(phase.margin_excess_area_norm_ms),
            float(phase.longest_margin_s * 1000.0),
            float(phase.margin_longest_continuous_s * 1000.0),
            float(phase.margin_sustained_t_peak_kv),
            float(phase.margin_sustained_t_ratio * 100.0),
            float(phase.margin_event_start_s),
            float(phase.margin_event_end_s),
        )
    return ("", None, None, None, None, None, None, None, None)


def _summary_metric_values(
    result: SustainedSDPFResult,
    phase: PhaseStressResult,
    rank: int | None,
    limits_by_voltage: Mapping[str, SDPFVoltageLimits] | None,
) -> list[Any]:
    population = result_population(result)
    (
        threshold,
        threshold_peak,
        normalized_area,
        qualification_duration_ms,
        continuous_duration_ms,
        sustained_t_peak,
        sustained_t_ratio,
        event_start,
        event_end,
    ) = _population_metric_values(phase, population)
    limit = _summary_limit(limits_by_voltage, result.voltage)
    qualifying_paths = _qualifying_path_labels(result)
    return [
        rank if rank is not None else "",
        _summary_voltage(result.voltage),
        _population_label(population),
        str(result.case),
        int(result.run),
        str(result.mm_name),
        str(result.fault_type or ""),
        str(phase.measurement),
        str(phase.phase),
        threshold,
        threshold_peak,
        float(phase.sdpf_peak_kv),
        float(phase.margin_peak_kv),
        normalized_area,
        qualification_duration_ms,
        continuous_duration_ms,
        sustained_t_peak,
        sustained_t_ratio,
        float(phase.sustained_peak_kv) if population is not None else None,
        event_start,
        event_end,
        len(qualifying_paths),
        "; ".join(qualifying_paths),
        str(limit.source) if limit is not None else "",
    ]


QUALIFICATION_DURATION_HEADER = "Qualification duration (ms)"
FULL_WAVE_DURATION_HEADER = "Full-wave continuous duration (ranking, ms)"


SUMMARY_HEADERS = [
    "Population rank",
    "Voltage (kV)",
    "Population",
    "Case",
    "Run",
    "MM element",
    "Fault type",
    "Measurement",
    "Phase",
    "Threshold",
    "Threshold (kVpeak)",
    "SDPF limit (kVpeak)",
    "Safety margin (kVpeak)",
    "Normalized excess area (pu*ms)",
    QUALIFICATION_DURATION_HEADER,
    FULL_WAVE_DURATION_HEADER,
    "Highest sustained T, VT (kVpeak)",
    "VT / SDPF (%)",
    "Sustained peak (kVpeak)",
    "Event start (s)",
    "Event end (s)",
    "Qualifying path count",
    "Qualifying paths",
    "Limit source",
]

REPRESENTATIVE_HEADERS = [
    "Voltage (kV)",
    "Population",
    "Selection criterion",
    "Case",
    "Run",
    "MM element",
    "Measurement",
    "Phase",
    "Selected metric",
    "Metric value",
    "Limit source",
]


def _selection_metric_value(
    phase: PhaseStressResult,
    metric: str,
    population: str,
) -> tuple[str, float]:
    values = _population_metric_values(phase, population, metric)
    if metric == HIGHEST_VOLTAGE_SUSTAINED_SELECTION:
        return "VT / SDPF (%)", float(values[6])
    if metric == CUMULATIVE_STRESS_SELECTION:
        return "Normalized excess area (pu*ms)", float(values[2])
    return FULL_WAVE_DURATION_HEADER, float(values[4])


def _representative_rows(
    observations_by_voltage: Mapping[str, Iterable[SustainedSDPFResult]],
    limits_by_voltage: Mapping[str, SDPFVoltageLimits] | None,
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for raw_voltage in sorted(observations_by_voltage, key=_summary_voltage_sort_key):
        results = [
            result
            for result in observations_by_voltage[raw_voltage]
            if isinstance(result, SustainedSDPFResult)
        ]
        representatives = select_representatives(results)
        for population in RANKING_POPULATIONS:
            for metric in RANKING_SELECTIONS:
                result = representatives[population][metric]
                if result is None:
                    continue
                phase = result.governing
                metric_name, metric_value = _selection_metric_value(
                    phase,
                    metric,
                    population,
                )
                limit = _summary_limit(limits_by_voltage, result.voltage)
                rows.append(
                    [
                        _summary_voltage(result.voltage),
                        _population_label(population),
                        _SELECTION_LABELS[metric],
                        str(result.case),
                        int(result.run),
                        str(result.mm_name),
                        str(phase.measurement),
                        str(phase.phase),
                        metric_name,
                        metric_value,
                        str(limit.source) if limit is not None else "",
                    ]
                )
    return rows


def _summary_context(
    scope_folder: str,
    observations: Mapping[str, Iterable[SustainedSDPFResult]],
    settings: SustainedSDPFSettings | None,
    frequency_hz: float | None,
) -> list[tuple[str, Any]]:
    results = [
        result
        for voltage_results in observations.values()
        for result in voltage_results
        if isinstance(result, SustainedSDPFResult)
    ]
    duration_values = sorted({round(result.duration_s * 1000.0, 9) for result in results})
    if settings is not None:
        duration_value: Any = float(settings.duration_ms)
    elif len(duration_values) == 1:
        duration_value = duration_values[0]
    elif duration_values:
        duration_value = ", ".join(f"{value:g}" for value in duration_values)
    else:
        duration_value = ""
    try:
        frequency_value = float(frequency_hz) if frequency_hz is not None else None
    except (TypeError, ValueError):
        frequency_value = None
    if frequency_value is not None and (not math.isfinite(frequency_value) or frequency_value <= 0):
        frequency_value = None
    cycle_values = sorted(
        {
            int(result.cycle_coverage)
            for result in results
            if result.cycle_coverage > 0
        }
    )
    if not cycle_values and settings is not None and frequency_value is not None:
        required_cycles = settings.required_cycles(frequency_value)
        if required_cycles > 0:
            cycle_values = [required_cycles]
    cycles_value: Any = (
        cycle_values[0]
        if len(cycle_values) == 1
        else ", ".join(str(value) for value in cycle_values)
        if cycle_values
        else ""
    )
    return [
        ("Scope", str(scope_folder)),
        ("Result version", RESULT_VERSION),
        ("Minimum persistence (ms)", duration_value),
        ("Frequency (Hz)", frequency_value if frequency_value is not None else ""),
        ("Derived complete-cycle coverage", cycles_value),
    ]


def _summary_sheet(
    workbook,
    title: str,
    headers: list[str],
    rows: Iterable[list[Any]],
    table_name: str,
    context: Iterable[tuple[str, Any]] = (),
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    worksheet = workbook.create_sheet(title)
    materialized = list(rows)
    context_rows = list(context)
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="Arial", bold=True, size=14, color="1F4E78")
    for row_index, (label, value) in enumerate(context_rows, start=2):
        label_cell = worksheet.cell(row_index, 1, label)
        label_cell.font = Font(name="Arial", bold=True)
        value_cell = worksheet.cell(row_index, 2, value)
        value_cell.font = Font(name="Arial")
    header_row = len(context_rows) + 3
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row_index, row in enumerate(materialized, start=header_row + 1):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(row_index, column_index, value)
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(headers))
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column_index)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(
            12,
            min(34, len(str(header)) + 2),
        )
    worksheet.column_dimensions["A"].width = max(
        worksheet.column_dimensions["A"].width or 0,
        max((len(str(label)) for label, _value in context_rows), default=0) + 2,
    )
    worksheet.column_dimensions["B"].width = max(
        worksheet.column_dimensions["B"].width or 0,
        max((len(str(value)) for _label, value in context_rows), default=0) + 2,
    )
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial")
    if materialized:
        table = Table(
            displayName=table_name,
            ref=f"A{header_row}:{last_column}{header_row + len(materialized)}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        # A table owns its AutoFilter.  Setting a second worksheet-level
        # AutoFilter makes Excel repair/remove the table on open.
        worksheet.auto_filter.ref = f"A{header_row}:{last_column}{header_row}"


def write_summary_workbook(
    project_root: str | Path,
    scope_folder: str,
    observations_by_voltage: Mapping[str, Iterable[SustainedSDPFResult]],
    limits_by_voltage: Mapping[str, SDPFVoltageLimits] | None = None,
    *,
    settings: SustainedSDPFSettings | None = None,
    frequency_hz: float | None = None,
) -> Path:
    """Write a compact ranking workbook for manual review and selection."""
    from openpyxl import Workbook

    path = summary_path(project_root, scope_folder)
    path.parent.mkdir(parents=True, exist_ok=True)
    observations: dict[str, list[SustainedSDPFResult]] = {
        str(voltage): [
            result
            for result in results
            if isinstance(result, SustainedSDPFResult)
        ]
        for voltage, results in observations_by_voltage.items()
    }
    ranked_by_voltage = {
        voltage: rank_results(results)
        for voltage, results in observations.items()
    }
    ranked_rows: list[list[Any]] = []
    for voltage in sorted(ranked_by_voltage, key=_summary_voltage_sort_key):
        ranked = ranked_by_voltage[voltage]
        population_ranks = {population: 0 for population in RANKING_POPULATIONS}
        for result in ranked:
            population = result_population(result)
            rank: int | None = None
            if population is not None:
                population_ranks[population] += 1
                rank = population_ranks[population]
            ranked_rows.append(
                _summary_metric_values(result, result.governing, rank, limits_by_voltage)
            )
    representative_rows = _representative_rows(observations, limits_by_voltage)
    summary_context = _summary_context(scope_folder, observations, settings, frequency_hz)
    workbook = Workbook()
    try:
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        _summary_sheet(
            workbook,
            "Ranked cases",
            SUMMARY_HEADERS,
            ranked_rows,
            "RankedCases",
            summary_context,
        )
        _summary_sheet(
            workbook,
            "Representative selections",
            REPRESENTATIVE_HEADERS,
            representative_rows,
            "RepresentativeSelections",
        )
        save_workbook_atomic(workbook, path)
    finally:
        workbook.close()
    return path


def make_signature(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, sort_keys=True, default=str, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def result_path(project_root: str | Path, scope_folder: str) -> Path:
    return Path(project_root) / "Voltage_envelope" / scope_folder / RESULT_FILENAME


def summary_path(project_root: str | Path, scope_folder: str) -> Path:
    return Path(project_root) / "Voltage_envelope" / scope_folder / SUMMARY_FILENAME


def invalidate_results(project_root: str | Path) -> None:
    """Remove app-owned Sustained SDPF metadata after an exclusion/settings edit."""
    root = Path(project_root) / "Voltage_envelope"
    try:
        paths = root.glob(f"*/{RESULT_FILENAME}")
        for path in paths:
            path.unlink(missing_ok=True)
            path.with_name(SUMMARY_FILENAME).unlink(missing_ok=True)
    except OSError:
        return


def load_results(project_root: str | Path, scope_folder: str) -> dict[str, Any]:
    path = result_path(project_root, scope_folder)
    try:
        # Do not load an obsolete result cache. The persisted schema is tied
        # to the peak-envelope persistence rule and is regenerated when the
        # method version changes.
        with path.open("r", encoding="utf-8") as handle:
            header = handle.read(256)
        if not re.search(rf'"version"\s*:\s*{RESULT_VERSION}(?:\D|$)', header):
            return {}
        with path.open("r", encoding="utf-8") as handle:
            value = json.load(handle)
    except (OSError, ValueError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _compact_signature_inputs(
    signature_inputs: Mapping[str, Mapping[str, Any]] | None,
    project_root: Path,
) -> tuple[dict[str, dict[str, Any]], dict[str, Any] | None]:
    """Store source freshness as one fingerprint plus compact source roots.

    Source file size/mtime triples are needed while a result is being built,
    but retaining every triple in the result makes large projects needlessly
    large. The persisted result therefore keeps the fingerprint and the
    selected case directories needed to recompute it. Any full input manifest
    is compacted before it is written.
    """

    def compact_entry(entry: Any) -> list[Any] | None:
        if isinstance(entry, Mapping):
            raw_path = entry.get("path")
            if raw_path in (None, ""):
                return None
            path = str(raw_path)
            if bool(entry.get("missing", False)):
                return [path, None, None]
            try:
                return [path, int(entry["size"]), int(entry["mtime_ns"])]
            except (KeyError, TypeError, ValueError):
                return None
        if isinstance(entry, (list, tuple)) and len(entry) >= 3:
            path = str(entry[0])
            if not path:
                return None
            try:
                size = None if entry[1] is None else int(entry[1])
                mtime_ns = None if entry[2] is None else int(entry[2])
            except (TypeError, ValueError):
                return None
            return [path, size, mtime_ns]
        return None

    def project_path(raw_path: Any) -> Path:
        path = Path(str(raw_path))
        return path if path.is_absolute() else project_root / path

    def relative_root(raw_path: Any) -> str:
        return _source_relative_path(project_root, project_path(raw_path).parent)

    def normalize_entry(entry: list[Any]) -> list[Any]:
        entry[0] = _source_relative_path(project_root, project_path(entry[0]))
        return entry

    normalized: dict[str, dict[str, Any]] = {}
    compact_by_voltage: dict[str, list[list[Any]] | None] = {}
    roots_by_voltage: dict[str, list[str]] = {}
    static_by_voltage: dict[str, bool] = {}
    extra_files_by_voltage: dict[str, list[str]] = {}
    for raw_voltage, raw_inputs in (signature_inputs or {}).items():
        if not isinstance(raw_inputs, Mapping):
            continue
        inputs = dict(raw_inputs)
        files = inputs.get("files")
        raw_roots = inputs.pop("source_roots", None)
        if isinstance(files, list):
            compact_files = [
                normalize_entry(item)
                for item in (compact_entry(entry) for entry in files)
                if item is not None
            ]
            compact_files.sort(key=lambda item: str(item[0]).casefold())
            compact_by_voltage[str(raw_voltage)] = compact_files
            static_by_voltage[str(raw_voltage)] = any(
                str(item[0]).casefold() == "pscad_log.txt"
                or str(item[0]).casefold().startswith("input_data_pscad")
                or str(item[0]).casefold() == "plots/.plottool_v3/limits.json"
                for item in compact_files
            )
            extra_files_by_voltage[str(raw_voltage)] = sorted(
                {
                    str(item[0])
                    for item in compact_files
                    if not (
                        str(item[0]).casefold().endswith((".inf", ".out"))
                        or str(item[0]).casefold() == "pscad_log.txt"
                        or str(item[0]).casefold().startswith("input_data_pscad")
                        or str(item[0]).casefold() == "plots/.plottool_v3/limits.json"
                    )
                },
                key=str.casefold,
            )
            roots = (
                [str(item) for item in raw_roots if str(item)]
                if isinstance(raw_roots, list)
                else sorted({relative_root(item[0]) for item in compact_files})
            )
            roots_by_voltage[str(raw_voltage)] = sorted(dict.fromkeys(roots))
            inputs.pop("files", None)
        else:
            compact_by_voltage[str(raw_voltage)] = None
            extra_files_by_voltage[str(raw_voltage)] = []
        normalized[str(raw_voltage)] = inputs
    file_values = [value for value in compact_by_voltage.values() if value is not None]
    shared_files = file_values[0] if file_values and all(value == file_values[0] for value in file_values) else None
    root_values = [roots_by_voltage[key] for key, value in compact_by_voltage.items() if value is not None]
    shared_roots = root_values[0] if root_values and all(value == root_values[0] for value in root_values) else None
    shared_static = list(static_by_voltage.values())
    extra_values = [
        extra_files_by_voltage[key]
        for key, value in compact_by_voltage.items()
        if value is not None
    ]
    shared_extra_files = (
        extra_values[0]
        if extra_values and all(value == extra_values[0] for value in extra_values)
        else None
    )
    if (
        shared_files is not None
        and shared_roots is not None
        and shared_extra_files is not None
        and len(file_values) == len(compact_by_voltage)
        and shared_static
        and all(value == shared_static[0] for value in shared_static)
    ):
        normalized = {
            voltage: inputs
            for voltage, inputs in normalized.items()
        }
        manifest = {
            "version": SOURCE_MANIFEST_VERSION,
            "roots": shared_roots,
            "include_static": shared_static[0],
            "fingerprint": make_signature({"files": shared_files}),
        }
        if shared_extra_files:
            manifest["extra_files"] = shared_extra_files
        return normalized, manifest
    for voltage, inputs in normalized.items():
        files = compact_by_voltage.get(voltage)
        roots = roots_by_voltage.get(voltage)
        if files is not None and roots is not None:
            manifest = {
                "version": SOURCE_MANIFEST_VERSION,
                "roots": roots,
                "include_static": static_by_voltage.get(voltage, False),
                "fingerprint": make_signature({"files": files}),
            }
            extra_files = extra_files_by_voltage.get(voltage, [])
            if extra_files:
                manifest["extra_files"] = extra_files
            inputs["source_manifest"] = manifest
    return normalized, None


def _result_identity(result: SustainedSDPFResult) -> tuple[str, int, str]:
    return str(result.case), int(result.run), str(result.mm_name)


def _selection_descriptor(
    result: SustainedSDPFResult,
    population: str,
    metric: str,
    *,
    canonical: SustainedSDPFResult | None,
    extra_results: list[dict[str, Any]],
    extra_index_by_identity: dict[tuple[str, int, str], int],
) -> dict[str, Any]:
    """Return a compact reference to one selected result and fixed path."""
    identity = _result_identity(result)
    if canonical is not None and identity == _result_identity(canonical):
        source = "canonical"
        index = 0
    else:
        source = "representative"
        index = extra_index_by_identity.get(identity)
        if index is None:
            selected_phases = tuple(
                phase
                for phase in (result.phases or (result.governing,))
                if phase.measurement == result.governing.measurement
                and phase.phase == result.governing.phase
            )
            compact_result = replace(
                result,
                phases=selected_phases or (result.governing,),
            ).to_dict()
            index = len(extra_results)
            extra_index_by_identity[identity] = index
            extra_results.append(compact_result)
        else:
            existing = extra_results[index]
            existing_phases = existing.get("phases", [])
            if not isinstance(existing_phases, list):
                existing_phases = []
            path_key = (str(result.governing.measurement), str(result.governing.phase))
            if not any(
                isinstance(raw_phase, Mapping)
                and (
                    str(raw_phase.get("measurement", "")),
                    str(raw_phase.get("phase", "")),
                )
                == path_key
                for raw_phase in existing_phases
            ):
                existing_phases.append(result.governing.to_dict())
                existing["phases"] = existing_phases
    return {
        "source": source,
        "index": int(index),
        "population": population,
        "metric": metric,
        "case": str(result.case),
        "run": int(result.run),
        "mm_name": str(result.mm_name),
        "measurement": str(result.governing.measurement),
        "phase": str(result.governing.phase),
    }


def _build_representative_cache(
    canonical: SustainedSDPFResult | None,
    representatives: Mapping[str, Mapping[str, SustainedSDPFResult | None]],
) -> tuple[dict[str, dict[str, dict[str, Any] | None]], list[dict[str, Any]]]:
    extra_results: list[dict[str, Any]] = []
    extra_index_by_identity: dict[tuple[str, int, str], int] = {}
    selections: dict[str, dict[str, dict[str, Any] | None]] = {}
    for population in RANKING_POPULATIONS:
        population_selections: dict[str, dict[str, Any] | None] = {}
        for metric in RANKING_SELECTIONS:
            result = representatives.get(population, {}).get(metric)
            population_selections[metric] = (
                _selection_descriptor(
                    result,
                    population,
                    metric,
                    canonical=canonical,
                    extra_results=extra_results,
                    extra_index_by_identity=extra_index_by_identity,
                )
                if result is not None
                else None
            )
        selections[population] = population_selections
    return selections, extra_results


def save_results(
    project_root: str | Path,
    scope_folder: str,
    settings: SustainedSDPFSettings,
    results_by_voltage: dict[str, SustainedSDPFResult | None],
    warnings: list[str] | None = None,
    signature_inputs: dict[str, dict[str, Any]] | None = None,
    observations_by_voltage: dict[str, Iterable[SustainedSDPFResult]] | None = None,
) -> Path:
    path = result_path(project_root, scope_folder)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Keep one canonical result for the heatmap's existing payload contract.
    # Representative selections use compact references into a small pool of
    # selected fixed-path results, so changing ranking controls never rereads
    # waveform files or stores a second copy of every observation.
    from results_analysis_app import sustained_sdpf_heatmap

    materialized_observations = {
        str(voltage): list(results)
        for voltage, results in (observations_by_voltage or {}).items()
    }
    compact_by_voltage = {
        str(voltage): [
            observation.to_mapping()
            for observation in sustained_sdpf_heatmap.compact_observations(results)
        ]
        for voltage, results in materialized_observations.items()
    }
    selected_by_voltage: dict[str, SustainedSDPFResult | None] = {}
    representative_selections: dict[str, dict[str, dict[str, Any] | None]] = {}
    representative_results: dict[str, list[dict[str, Any]]] = {}
    for raw_voltage, raw_results in materialized_observations.items():
        representatives = select_representatives(raw_results)
        actual_area = representatives[ACTUAL_SDPF_POPULATION][CUMULATIVE_STRESS_SELECTION]
        margin_area = representatives[SAFETY_MARGIN_ONLY_POPULATION][CUMULATIVE_STRESS_SELECTION]
        area_result = results_by_voltage.get(raw_voltage) or actual_area or margin_area
        selected_by_voltage[raw_voltage] = area_result
        selections, extras = _build_representative_cache(area_result, representatives)
        representative_selections[raw_voltage] = selections
        representative_results[raw_voltage] = extras
    for raw_voltage, result in results_by_voltage.items():
        selected_by_voltage.setdefault(str(raw_voltage), result)
        voltage_key = str(raw_voltage)
        if voltage_key not in representative_selections:
            representatives = {
                population: {
                    metric: result
                    for metric in RANKING_SELECTIONS
                }
                if result is not None and result_population(result) == population
                else {metric: None for metric in RANKING_SELECTIONS}
                for population in RANKING_POPULATIONS
            }
            selections, extras = _build_representative_cache(result, representatives)
            representative_selections[voltage_key] = selections
            representative_results[voltage_key] = extras
    compact_signature_inputs, shared_manifest = _compact_signature_inputs(
        signature_inputs,
        Path(project_root),
    )
    persisted_selections = representative_selections
    payload = {
        "version": RESULT_VERSION,
        "scope_folder": scope_folder,
        "settings": settings.to_mapping(),
        "results": {
            str(voltage): result.to_dict() if result is not None else None
            for voltage, result in selected_by_voltage.items()
        },
        "representative_results": representative_results,
        "selections": persisted_selections,
        "observations": compact_by_voltage,
        "signature_inputs": compact_signature_inputs,
        "warnings": list(warnings or []),
    }
    if shared_manifest is not None:
        payload["source_manifest"] = shared_manifest
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8") as handle:
        # This cache is machine-generated and can contain thousands of source
        # files.  Compact JSON keeps it readable while avoiding a large amount
        # of whitespace on every rebuild.
        json.dump(payload, handle, separators=(",", ":"))
        handle.write("\n")
    temp.replace(path)
    return path


def _source_relative_path(root: Path, path: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except (OSError, ValueError):
        try:
            return path.resolve(strict=False).relative_to(root.resolve(strict=False)).as_posix()
        except (OSError, ValueError):
            return path.resolve(strict=False).as_posix()


def _compact_source_entries(
    project_root: Path,
    roots: Iterable[Any],
    include_static: bool = True,
    extra_files: Iterable[Any] = (),
) -> list[list[Any]]:
    """Recreate the small metadata input used by a compact source fingerprint."""
    candidates: list[Path] = []
    if include_static:
        candidates.extend(
            [
                *sorted(project_root.glob("Input_Data_PSCAD*.xlsx")),
                project_root / "PSCAD_log.txt",
                project_root / "Plots" / ".plottool_v3" / "limits.json",
            ]
        )
    for raw_root in roots:
        source_root = Path(str(raw_root))
        if not source_root.is_absolute():
            source_root = project_root / source_root
        try:
            candidates.extend(sorted(source_root.glob("*.inf")))
            candidates.extend(sorted(source_root.glob("*.out")))
        except OSError:
            continue
    for raw_path in extra_files:
        path = Path(str(raw_path))
        if not path.is_absolute():
            path = project_root / path
        candidates.append(path)

    entries: list[list[Any]] = []
    seen: set[str] = set()
    for path in candidates:
        relative = _source_relative_path(project_root, path)
        if relative in seen:
            continue
        seen.add(relative)
        try:
            stat = path.stat()
        except OSError:
            entries.append([relative, None, None])
        else:
            entries.append([relative, int(stat.st_size), int(stat.st_mtime_ns)])
    return sorted(entries, key=lambda item: str(item[0]).casefold())


def _compact_source_manifest_current(
    manifest: Mapping[str, Any],
    project_root: str | Path | None,
) -> bool:
    if manifest.get("version") != SOURCE_MANIFEST_VERSION:
        return False
    roots = manifest.get("roots")
    fingerprint = manifest.get("fingerprint")
    if not isinstance(roots, list) or not isinstance(fingerprint, str) or project_root is None:
        return False
    entries = _compact_source_entries(
        Path(project_root).resolve(),
        roots,
        include_static=bool(manifest.get("include_static", True)),
        extra_files=manifest.get("extra_files", ()),
    )
    return make_signature({"files": entries}) == fingerprint


def source_manifest_current(
    payload: Mapping[str, Any],
    project_root: str | Path | None = None,
) -> bool:
    """Validate one shared source manifest without rereading waveform data."""
    if "version" in payload and payload.get("version") != RESULT_VERSION:
        return False
    shared_manifest = payload.get("source_manifest")
    return (
        isinstance(shared_manifest, Mapping)
        and _compact_source_manifest_current(shared_manifest, project_root)
    )


def result_inputs_current(
    payload: Mapping[str, Any],
    voltage: str,
    project_root: str | Path | None = None,
    *,
    shared_manifest_current: bool | None = None,
) -> bool:
    """Check the recorded source-file manifest without rereading waveform data.

    A saved result normally contains one shared manifest for all voltage
    levels.  Callers processing several voltages can validate that manifest
    once and pass the result here; legacy payloads with voltage-specific file
    lists continue to validate those lists independently.
    """
    if "version" in payload and payload.get("version") != RESULT_VERSION:
        return False
    inputs = payload.get("signature_inputs", {}).get(str(voltage))
    if not isinstance(inputs, dict):
        # Results written by the first implementation have no manifest and are
        # intentionally treated as stale rather than trusted silently.
        return False
    if shared_manifest_current is not None:
        return bool(shared_manifest_current)
    source_manifest = inputs.get("source_manifest")
    if isinstance(source_manifest, Mapping):
        return _compact_source_manifest_current(source_manifest, project_root)
    return source_manifest_current(payload, project_root)


@dataclass(frozen=True, slots=True)
class SustainedSDPFCacheValidation:
    """Explain whether persisted Sustained SDPF data can be reused."""

    valid: bool
    reason: str = ""
    shared_manifest_current: bool | None = None


def validate_result_cache(
    payload: Mapping[str, Any] | None,
    project_root: str | Path,
    settings: SustainedSDPFSettings | None = None,
    *,
    shared_manifest_current: bool | None = None,
) -> SustainedSDPFCacheValidation:
    """Validate reusable result metadata without reading waveform data.

    A valid cache is not the same as a JSON file that can be parsed.  The
    persisted method version, requested settings, and source fingerprint must
    all agree before a later batch, heatmap, or report may consume it.
    """
    if not isinstance(payload, Mapping) or not payload:
        return SustainedSDPFCacheValidation(False, "result cache is missing or obsolete")

    version = payload.get("version")
    if version != RESULT_VERSION:
        return SustainedSDPFCacheValidation(
            False,
            f"result cache version {version!r} is unsupported; expected {RESULT_VERSION}",
        )

    if settings is not None and payload.get("settings") != settings.to_mapping():
        return SustainedSDPFCacheValidation(
            False,
            "saved analysis settings do not match the current Sustained SDPF settings",
        )

    if not isinstance(payload.get("results"), Mapping):
        return SustainedSDPFCacheValidation(False, "result cache has no valid results mapping")

    if shared_manifest_current is None:
        try:
            shared_manifest_current = source_manifest_current(payload, project_root)
        except (OSError, RuntimeError, TypeError, ValueError):
            shared_manifest_current = False
    if not shared_manifest_current:
        return SustainedSDPFCacheValidation(
            False,
            "recorded source files or their fingerprint are missing or stale",
            shared_manifest_current=False,
        )

    return SustainedSDPFCacheValidation(
        True,
        shared_manifest_current=True,
    )


def normalized_voltage_key(value: Any) -> str:
    """Return the canonical persisted key for a voltage-level value."""
    normalized = normalize_voltage(value)
    return normalized or str(value).strip()


def _persisted_result_for_descriptor(
    descriptor: Mapping[str, Any],
    canonical: SustainedSDPFResult | None,
    extra_results: list[SustainedSDPFResult],
    population: str,
    metric: str,
    voltage_key: str,
    scope_folder: str,
    expected_duration_s: float | None,
) -> SustainedSDPFResult | None:
    if (
        descriptor.get("population") != population
        or descriptor.get("metric") != metric
    ):
        return None
    source = str(descriptor.get("source", ""))
    if source == "canonical":
        result = canonical
    elif source == "representative":
        try:
            result = extra_results[int(descriptor.get("index", -1))]
        except (IndexError, TypeError, ValueError):
            return None
    else:
        return None
    if result is None:
        return None
    if normalized_voltage_key(result.voltage) != voltage_key:
        return None
    if scope_folder and str(result.scope_folder) != scope_folder:
        return None
    if expected_duration_s is not None and not math.isclose(
        float(result.duration_s), expected_duration_s, rel_tol=0.0, abs_tol=NUMERIC_TOLERANCE
    ):
        return None
    try:
        expected_identity = (
            str(descriptor.get("case", "")),
            int(float(descriptor.get("run", 0))),
            str(descriptor.get("mm_name", "")),
        )
    except (TypeError, ValueError):
        return None
    if _result_identity(result) != expected_identity:
        return None
    measurement = str(descriptor.get("measurement", "")).strip()
    phase_name = str(descriptor.get("phase", "")).strip()
    phases = result.phases or (result.governing,)
    phase = next(
        (
            item
            for item in phases
            if item.measurement == measurement and item.phase == phase_name
        ),
        None,
    )
    if phase is None:
        return None
    if population == ACTUAL_SDPF_POPULATION and not phase.sdpf_exceeded:
        return None
    if population == SAFETY_MARGIN_ONLY_POPULATION and (
        not phase.margin_exceeded or phase.sdpf_exceeded
    ):
        return None
    return replace(result, governing=phase)


def current_representatives_for_voltage(
    payload: Mapping[str, Any] | None,
    voltage: Any,
    project_root: str | Path,
    settings: SustainedSDPFSettings | None = None,
    *,
    shared_manifest_current: bool | None = None,
    cache_validation: SustainedSDPFCacheValidation | None = None,
) -> tuple[
    dict[str, dict[str, SustainedSDPFResult]],
    SustainedSDPFCacheValidation,
]:
    """Load all population-specific representative selections for one voltage."""
    validation = cache_validation or validate_result_cache(
        payload,
        project_root,
        settings,
        shared_manifest_current=shared_manifest_current,
    )
    if not validation.valid:
        return {}, validation

    voltage_key = normalized_voltage_key(voltage)
    if not result_inputs_current(
        payload,
        voltage_key,
        project_root,
        shared_manifest_current=validation.shared_manifest_current,
    ):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"recorded source signature for {voltage_key} kV is missing or stale",
            shared_manifest_current=validation.shared_manifest_current,
        )

    result_mapping = payload.get("results")
    raw_result = result_mapping.get(voltage_key) if isinstance(result_mapping, Mapping) else None
    if raw_result is not None and not isinstance(raw_result, Mapping):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted governing result for {voltage_key} kV is malformed",
            shared_manifest_current=validation.shared_manifest_current,
        )
    try:
        canonical = SustainedSDPFResult.from_dict(raw_result) if isinstance(raw_result, Mapping) else None
    except (TypeError, ValueError, KeyError):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted governing result for {voltage_key} kV is malformed",
            shared_manifest_current=validation.shared_manifest_current,
        )
    raw_extra = payload.get("representative_results", {})
    raw_extra_for_voltage = (
        raw_extra.get(voltage_key)
        if isinstance(raw_extra, Mapping)
        else None
    )
    if raw_extra_for_voltage is None:
        raw_extra_for_voltage = []
    if not isinstance(raw_extra_for_voltage, list):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted representative pool for {voltage_key} kV is malformed",
            shared_manifest_current=validation.shared_manifest_current,
        )
    extra_results: list[SustainedSDPFResult] = []
    try:
        if any(not isinstance(item, Mapping) for item in raw_extra_for_voltage):
            raise ValueError("Invalid representative pool entry")
        extra_results = [
            SustainedSDPFResult.from_dict(item)
            for item in raw_extra_for_voltage
        ]
    except (TypeError, ValueError, KeyError):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted representative pool for {voltage_key} kV is malformed",
            shared_manifest_current=validation.shared_manifest_current,
        )
    raw_voltage_selections = payload.get("selections", {})
    raw_selections = (
        raw_voltage_selections.get(voltage_key)
        if isinstance(raw_voltage_selections, Mapping)
        else None
    )
    if not isinstance(raw_selections, Mapping):
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted selections for {voltage_key} kV are missing",
            shared_manifest_current=validation.shared_manifest_current,
        )
    missing_populations = [
        population for population in RANKING_POPULATIONS if population not in raw_selections
    ]
    if missing_populations:
        return {}, SustainedSDPFCacheValidation(
            False,
            f"persisted selections for {voltage_key} kV are incomplete",
            shared_manifest_current=validation.shared_manifest_current,
        )
    expected_duration_s = settings.effective_duration() if settings is not None else None
    scope_folder = str(payload.get("scope_folder", ""))
    representatives: dict[str, dict[str, SustainedSDPFResult]] = {}
    for population in RANKING_POPULATIONS:
        raw_population = raw_selections.get(population)
        if not isinstance(raw_population, Mapping):
            return {}, SustainedSDPFCacheValidation(
                False,
                f"persisted {population} selections for {voltage_key} kV are malformed",
                shared_manifest_current=validation.shared_manifest_current,
            )
        missing_metrics = [metric for metric in RANKING_SELECTIONS if metric not in raw_population]
        if missing_metrics:
            return {}, SustainedSDPFCacheValidation(
                False,
                f"persisted {population} selections for {voltage_key} kV are incomplete",
                shared_manifest_current=validation.shared_manifest_current,
            )
        selected: dict[str, SustainedSDPFResult] = {}
        for metric in RANKING_SELECTIONS:
            raw_selection = raw_population.get(metric)
            if raw_selection is None:
                continue
            if not isinstance(raw_selection, Mapping):
                return {}, SustainedSDPFCacheValidation(
                    False,
                    f"persisted {population}/{metric} selection for {voltage_key} kV is malformed",
                    shared_manifest_current=validation.shared_manifest_current,
                )
            selected_result = _persisted_result_for_descriptor(
                raw_selection,
                canonical,
                extra_results,
                population,
                metric,
                voltage_key,
                scope_folder,
                expected_duration_s,
            )
            if selected_result is None:
                return {}, SustainedSDPFCacheValidation(
                    False,
                    f"persisted {population}/{metric} selection for {voltage_key} kV is malformed",
                    shared_manifest_current=validation.shared_manifest_current,
                )
            selected[metric] = selected_result
        if selected:
            representatives[population] = selected
    return representatives, validation
