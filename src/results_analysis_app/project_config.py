from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
import math
import re
from typing import Any
from zipfile import BadZipFile


DEFAULT_EVENT_TIMES = {"TOV": 0.03, "SFO": 0.004, "SA": 0.1}


@dataclass(frozen=True)
class VoltageConfig:
    voltage: str
    bus_prefix: str
    um: float


@dataclass(frozen=True)
class ProjectTiming:
    frequency: float | None = None
    final_duration: float | None = None


def normalize_voltage(value: Any) -> str:
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return str(value).strip()
    if not math.isfinite(number) or number <= 0:
        return ""
    return str(int(number)) if number.is_integer() else f"{number:g}"


def input_data_workbook(project_root: str | Path) -> Path | None:
    root = Path(project_root).resolve()
    matches = sorted(root.glob("Input_Data_PSCAD*.xlsx"))
    return matches[0] if matches else None


def _positive_number(value: Any) -> float | None:
    if isinstance(value, str):
        match = re.search(r"\d+(?:\.\d+)?", value)
        value = match.group(0) if match else value
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number > 0 else None


def load_project_timing(project_root: str | Path) -> ProjectTiming:
    workbook_path = input_data_workbook(project_root)
    if workbook_path is None:
        return ProjectTiming()

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except (OSError, BadZipFile, KeyError, ValueError):
        return ProjectTiming()

    try:
        if "Input_Data" not in workbook.sheetnames:
            return ProjectTiming()
        sheet = workbook["Input_Data"]
        values: dict[str, float] = {}
        for label, value in sheet.iter_rows(
            min_col=1,
            max_col=2,
            values_only=True,
        ):
            key = str(label or "").strip().casefold()
            if key not in {"frequency", "final duration"}:
                continue
            number = _positive_number(value)
            if number is not None:
                values[key] = number
        return ProjectTiming(
            frequency=values.get("frequency") or _positive_number(sheet["B16"].value),
            final_duration=values.get("final duration"),
        )
    finally:
        workbook.close()


def load_project_frequency(project_root: str | Path) -> float | None:
    return load_project_timing(project_root).frequency


def discover_voltage_prefixes(project_root: str | Path) -> dict[str, str]:
    root = Path(project_root).resolve()
    return discover_voltage_prefixes_from_files(sorted((root / "Case_folder").rglob("*.inf")))


def discover_voltage_prefixes_from_files(inf_paths: Iterable[str | Path]) -> dict[str, str]:
    prefixes: dict[str, str] = {}
    for raw_path in inf_paths:
        path = Path(raw_path)
        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
        except OSError:
            continue
        for match in re.finditer(r'Group="(MM_(\d+(?:\.\d+)?))[^"]*"', text, flags=re.IGNORECASE):
            voltage = normalize_voltage(match.group(2))
            if voltage:
                prefixes.setdefault(voltage, match.group(1))
    return prefixes


def load_voltage_configs(
    project_root: str | Path,
    um_overrides: dict[str, float] | None = None,
    discovered_prefixes: dict[str, str] | None = None,
) -> dict[str, VoltageConfig]:
    workbook_path = input_data_workbook(project_root)
    configs: dict[str, VoltageConfig] = {}
    if workbook_path is not None:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        except (OSError, BadZipFile, KeyError, ValueError):
            workbook = None

        if workbook is not None:
            try:
                if "MM_blocks" in workbook.sheetnames:
                    sheet = workbook["MM_blocks"]
                    header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
                    if header_row is not None:
                        headers = {
                            str(cell.value).strip().casefold(): index
                            for index, cell in enumerate(header_row, start=1)
                            if cell.value is not None
                        }
                        group_col = headers.get("group")
                        un_col = headers.get("un")
                        um_col = headers.get("um")
                        if group_col is not None and um_col is not None:
                            for row in sheet.iter_rows(min_row=2):
                                group = str(row[group_col - 1].value or "").strip()
                                match = re.match(r"^(MM_(\d+(?:\.\d+)?))", group, flags=re.IGNORECASE)
                                voltage = normalize_voltage(row[un_col - 1].value) if un_col is not None else ""
                                if not voltage and match:
                                    voltage = normalize_voltage(match.group(2))
                                if not voltage:
                                    continue
                                try:
                                    um = float(row[um_col - 1].value)
                                except (TypeError, ValueError):
                                    continue
                                if not math.isfinite(um) or um <= 0:
                                    continue
                                bus_prefix = match.group(1) if match else f"MM_{voltage}"
                                configs.setdefault(voltage, VoltageConfig(voltage, bus_prefix, um))
            finally:
                workbook.close()

    overrides = um_overrides or {}
    if overrides:
        prefixes = discovered_prefixes if discovered_prefixes is not None else discover_voltage_prefixes(project_root)
        override_configs: dict[str, VoltageConfig] = {}
        for raw_voltage, raw_um in overrides.items():
            voltage = normalize_voltage(raw_voltage)
            try:
                um = float(raw_um)
            except (TypeError, ValueError):
                continue
            if voltage and math.isfinite(um) and um > 0:
                bus_prefix = configs.get(voltage, VoltageConfig(voltage, prefixes.get(voltage, f"MM_{voltage}"), um)).bus_prefix
                override_configs[voltage] = VoltageConfig(voltage, bus_prefix, um)
        return {**configs, **override_configs}

    return configs


def available_voltage_keys(
    project_root: str | Path,
    inf_paths: Iterable[str | Path] | None = None,
    um_overrides: dict[str, float] | None = None,
) -> list[str]:
    root = Path(project_root).resolve()
    paths = sorted(Path(path) for path in inf_paths) if inf_paths is not None else sorted((root / "Case_folder").rglob("*.inf"))
    prefixes = discover_voltage_prefixes_from_files(paths)
    configs = load_voltage_configs(root, um_overrides=um_overrides, discovered_prefixes=prefixes)
    found = set(prefixes) | set(configs)
    return sorted(found, key=lambda value: float(value) if _is_number(value) else value)


def _is_number(value: str) -> bool:
    try:
        float(value)
    except ValueError:
        return False
    return True
