from __future__ import annotations

from collections.abc import Callable, Iterable
from concurrent.futures import Executor, ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from contextlib import ExitStack
from dataclasses import dataclass, field, fields, is_dataclass, replace
import hashlib
import json
import shutil
from pathlib import Path
from typing import Any
import math
import os
import threading
import time

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from results_analysis_app import resonance_checks, storage, sustained_sdpf, sustained_sdpf_heatmap
from results_analysis_app.common import (
    CancelFn,
    LogFn,
    check_cancel as _cancel,
    log_message as _log,
)
from results_analysis_app.envelope_chart import create_combined_envelope_plot, create_resonance_check_charts
from results_analysis_app.excel import autofit_workbook, excel_app
from results_analysis_app.exclusions import ExclusionMatcher, ExclusionRule
from results_analysis_app.models import (
    DEFAULT_ENVELOPE_CHART_HEIGHT,
    DEFAULT_ENVELOPE_CHART_WIDTH,
    DEFAULT_ENVELOPE_FALLBACK_FREQUENCY,
    DEFAULT_ENVELOPE_WORKERS,
    DEFAULT_ENVELOPE_TIME_END,
    DEFAULT_ENVELOPE_TIME_STEP,
    MAX_ENVELOPE_WORKERS,
    DEFAULT_HIGH_VOLTAGE_LIMIT_FACTOR,
    DEFAULT_NONCONV_CB_IIP_LIMIT,
    DEFAULT_NONCONV_CB_IIR_LIMIT,
    ScopeEntry,
    automatic_worker_count,
    normalize_positive_float,
)
from results_analysis_app.project_config import (
    ProjectTiming,
    input_data_workbook,
    load_project_timing,
    load_voltage_configs,
    normalize_voltage,
)
from pscad_plotter_app_v3.services.waveform_io import (
    InfDescriptor,
    hash_inf_file,
    load_out_columns,
    out_file_for_pgb,
    parse_inf_descriptors,
)
from pscad_plotter_app_v3.services.project_conventions import case_run_from_inf_path


FrequencyFallbackFn = Callable[[str], None]
OutFileData = tuple[np.ndarray, dict[int, np.ndarray]]
SUSTAINED_ENTRY = "__Sustained_SDPF__"

FREQUENCY_FALLBACK_MESSAGE_PREFIX = "FREQUENCY_FALLBACK|"
HIGH_VOLTAGE_LIMIT_FACTOR = DEFAULT_HIGH_VOLTAGE_LIMIT_FACTOR
NONCONV_CB_IIP_LIMIT = DEFAULT_NONCONV_CB_IIP_LIMIT
NONCONV_CB_IIR_LIMIT = DEFAULT_NONCONV_CB_IIR_LIMIT
TIME_STEP = DEFAULT_ENVELOPE_TIME_STEP
TIME_END = DEFAULT_ENVELOPE_TIME_END
MAX_WORKERS_ENV = "RESULTS_ANALYSIS_ENVELOPE_WORKERS"
# A stale envelope entry is an optimization failure, not a source-data
# failure.  The entry is kept with the project's other stage fingerprints and
# contains no processed waveform data.
LEGACY_ENVELOPE_MANIFEST_FILENAME = ".envelope_manifest.json"
ENVELOPE_MANIFEST_VERSION = 3
ENVELOPE_CALCULATION_VERSION = 1
ENVELOPE_PRESENTATION_VERSION = 1

ENVELOPE_HEADERS = [
    "Time (s)",
    "Max_A", "Max_B", "Max_C",
    "Case_A", "Case_B", "Case_C",
    "Run_A", "Run_B", "Run_C",
    "Fault_type_A", "Fault_type_B", "Fault_type_C",
    "MM_name_A", "MM_name_B", "MM_name_C",
    "Max_all", "Case_all", "Run_all", "Fault_type_all", "MM_name_all",
]
NONCONV_COLUMNS = ["Case", "Run", "Fault_type", "Source", "Signal", "Reason", "Value", "File"]
HIGH_VOLTAGE_COLUMNS = ["Case", "Run", "Fault_type", "MM_name", "Measurement", "Signal", "File", "Excluded_values", "Max_abs", "Limit"]


def _sustained_cache_versions(
    settings: sustained_sdpf.SustainedSDPFSettings | None,
) -> dict[str, int]:
    if settings is None or not settings.enabled:
        return {}
    return {
        "sustained_sdpf_result_version": sustained_sdpf.RESULT_VERSION,
        "sustained_sdpf_summary_version": sustained_sdpf.SUMMARY_WORKBOOK_VERSION,
    }


def _cache_json_value(value: Any) -> Any:
    if isinstance(value, pd.DataFrame):
        return _cache_json_value(value.to_dict("records"))
    if is_dataclass(value) and not isinstance(value, type):
        return {
            item.name: _cache_json_value(getattr(value, item.name))
            for item in fields(value)
        }
    if isinstance(value, Path):
        return value.as_posix()
    if isinstance(value, dict):
        return {str(key): _cache_json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_cache_json_value(item) for item in value]
    if isinstance(value, np.generic):
        return _cache_json_value(value.item())
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _cache_signature(value: Any) -> str:
    encoded = json.dumps(
        _cache_json_value(value),
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
    ).encode("utf-8")
    return hashlib.blake2b(encoded, digest_size=16).hexdigest()


def _relative_project_path(project_root: Path, path: Path) -> str:
    try:
        return path.resolve().relative_to(project_root.resolve()).as_posix()
    except (OSError, ValueError):
        return str(path.resolve())


def _file_manifest_entry(project_root: Path, path: Path) -> dict[str, Any]:
    relative = _relative_project_path(project_root, path)
    try:
        stat = path.stat()
    except OSError:
        return {"path": relative, "missing": True}
    return {
        "path": relative,
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


@dataclass(slots=True)
class _SourceFileInventory:
    """One directory listing shared by the envelope cache fingerprints."""

    records: dict[str, dict[str, Any]]
    paths_by_directory: dict[str, tuple[Path, ...]]

    @staticmethod
    def _key(path: Path) -> str:
        try:
            return str(path.resolve(strict=False)).casefold()
        except (OSError, RuntimeError):
            return str(path).casefold()

    def record(self, path: Path, project_root: Path) -> dict[str, Any]:
        record = self.records.get(self._key(path))
        if record is not None:
            return dict(record)
        return _file_manifest_entry(project_root, path)

    def directory_paths(self, directory: Path) -> tuple[Path, ...]:
        return self.paths_by_directory.get(self._key(directory), ())


def _source_file_inventory(
    project_root: Path,
    inf_paths: Iterable[Path],
    extra_paths: Iterable[Path] = (),
) -> _SourceFileInventory:
    """Collect source file metadata with one scan per input directory.

    The envelope and Sustained manifests need the same small size/mtime
    records.  Building the directory index once avoids a separate glob/stat
    pass for every `.inf` file and every selected scope.
    """
    root = Path(project_root)
    root_resolved = root.resolve()
    directories = list(dict.fromkeys(Path(path).parent for path in inf_paths))
    records: dict[str, dict[str, Any]] = {}
    paths_by_directory: dict[str, tuple[Path, ...]] = {}

    def relative_path(path: Path) -> str:
        try:
            return path.relative_to(root_resolved).as_posix()
        except (OSError, ValueError):
            return str(path.resolve(strict=False))

    def add_path(path: Path, stat_result: os.stat_result | None = None) -> None:
        key = _SourceFileInventory._key(path)
        if key in records:
            return
        if stat_result is None:
            try:
                stat_result = path.stat()
            except OSError:
                records[key] = {"path": relative_path(path), "missing": True}
                return
        records[key] = {
            "path": relative_path(path),
            "size": int(stat_result.st_size),
            "mtime_ns": int(stat_result.st_mtime_ns),
        }

    for directory in directories:
        found: list[Path] = []
        try:
            with os.scandir(directory) as entries:
                for entry in entries:
                    suffix = Path(entry.name).suffix.casefold()
                    if suffix not in {".inf", ".out"}:
                        continue
                    path = directory / entry.name
                    found.append(path)
                    try:
                        add_path(path, entry.stat())
                    except OSError:
                        add_path(path)
        except OSError:
            pass
        paths_by_directory[_SourceFileInventory._key(directory)] = tuple(
            sorted(found, key=lambda path: path.name.casefold())
        )

    for path in extra_paths:
        add_path(Path(path))
    return _SourceFileInventory(records, paths_by_directory)


def _remove_legacy_envelope_cache(project_root: Path) -> None:
    output_root = project_root / "Voltage_envelope"
    try:
        (output_root / LEGACY_ENVELOPE_MANIFEST_FILENAME).unlink(missing_ok=True)
    except OSError:
        pass
    shutil.rmtree(output_root / ".run_cache", ignore_errors=True)


def _run_source_manifest(
    project_root: Path,
    inf_path: Path,
    source_inventory: _SourceFileInventory | None = None,
) -> list[dict[str, Any]]:
    candidates = [inf_path]
    if source_inventory is None:
        try:
            candidates.extend(sorted(inf_path.parent.glob(f"{inf_path.stem}*.out")))
        except OSError:
            pass
        return [_file_manifest_entry(project_root, path) for path in dict.fromkeys(candidates)]

    stem = inf_path.stem.casefold()
    candidates.extend(
        path
        for path in source_inventory.directory_paths(inf_path.parent)
        if path.suffix.casefold() == ".out" and path.name.casefold().startswith(stem)
    )
    return [
        source_inventory.record(path, project_root)
        for path in dict.fromkeys(candidates)
    ]


def _run_source_manifests(
    project_root: Path,
    inf_paths: Iterable[Path],
    source_inventory: _SourceFileInventory | None = None,
) -> dict[Path, list[dict[str, Any]]]:
    return {
        path: _run_source_manifest(project_root, path, source_inventory)
        for path in sorted(set(inf_paths))
    }


def _validated_envelope_manifest(
    project_root: Path,
    calculation_signature: str,
) -> tuple[dict[str, Any], dict[str, Path]] | None:
    cache = storage.load_project_analysis_cache(project_root)
    payload = cache.get("envelope")
    if not isinstance(payload, dict) or payload.get("version") != ENVELOPE_MANIFEST_VERSION:
        return None
    stored_signature = payload.get("calculation_signature", payload.get("signature"))
    if stored_signature != calculation_signature:
        return None
    output_records = payload.get("outputs")
    calculation_records = payload.get("calculation_outputs", output_records)
    if not isinstance(output_records, list) or not isinstance(calculation_records, list):
        return None
    outputs_by_record: dict[str, Path] = {}
    for record in calculation_records:
        if not isinstance(record, dict) or not isinstance(record.get("path"), str):
            return None
        path = project_root / Path(record["path"])
        if _file_manifest_entry(project_root, path) != record:
            return None
        outputs_by_record[record["path"]] = path
    for record in output_records:
        if not isinstance(record, dict) or not isinstance(record.get("path"), str):
            return None
        if record["path"] in outputs_by_record:
            continue
        path = project_root / Path(record["path"])
        if path.is_file() and _file_manifest_entry(project_root, path) == record:
            outputs_by_record[record["path"]] = path
    return payload, outputs_by_record


def _envelope_manifest_payload(
    project_root: Path,
    calculation_signature: str,
) -> dict[str, Any] | None:
    validated = _validated_envelope_manifest(project_root, calculation_signature)
    return validated[0] if validated is not None else None


def _envelope_manifest_matches(
    project_root: Path,
    signature: str,
) -> list[Path] | None:
    validated = _validated_envelope_manifest(project_root, signature)
    if validated is None:
        return None
    payload, outputs_by_record = validated

    return_records = payload.get("return_outputs", payload.get("outputs", []))
    if not isinstance(return_records, list):
        return None
    outputs: list[Path] = []
    for record in return_records:
        if not isinstance(record, dict) or not isinstance(record.get("path"), str):
            return None
        path = outputs_by_record.get(record["path"])
        if path is None or record != _file_manifest_entry(project_root, path):
            return None
        outputs.append(path)
    return outputs


def _write_envelope_manifest(
    project_root: Path,
    signature: str,
    outputs: Iterable[Path],
    return_outputs: Iterable[Path] | None = None,
    *,
    presentation_signature: str | None = None,
    artifacts: dict[str, Any] | None = None,
    calculation_outputs: Iterable[Path] | None = None,
    presentation_outputs: Iterable[Path] | None = None,
) -> None:
    output_paths = list(dict.fromkeys(Path(output) for output in outputs))
    return_paths = list(
        dict.fromkeys(
            Path(output)
            for output in (return_outputs if return_outputs is not None else output_paths)
        )
    )
    records = [
        _file_manifest_entry(project_root, path)
        for path in output_paths
        if path.is_file()
    ]
    return_records = [
        _file_manifest_entry(project_root, path)
        for path in return_paths
        if path.is_file()
    ]
    calculation_paths = list(
        dict.fromkeys(
            Path(output)
            for output in (calculation_outputs if calculation_outputs is not None else output_paths)
        )
    )
    calculation_records = [
        _file_manifest_entry(project_root, path)
        for path in calculation_paths
        if path.is_file()
    ]
    presentation_paths = list(
        dict.fromkeys(
            Path(output)
            for output in (presentation_outputs if presentation_outputs is not None else ())
        )
    )
    presentation_records = [
        _file_manifest_entry(project_root, path)
        for path in presentation_paths
        if path.is_file()
    ]
    cache = storage.load_project_analysis_cache(project_root)
    envelope_cache: dict[str, Any] = {
        "version": ENVELOPE_MANIFEST_VERSION,
        "signature": signature,
        "calculation_signature": signature,
        "outputs": sorted(records, key=lambda item: str(item["path"]).casefold()),
        "calculation_outputs": sorted(
            calculation_records,
            key=lambda item: str(item["path"]).casefold(),
        ),
        "presentation_outputs": sorted(
            presentation_records,
            key=lambda item: str(item["path"]).casefold(),
        ),
        "return_outputs": sorted(
            return_records,
            key=lambda item: str(item["path"]).casefold(),
        ),
    }
    if presentation_signature is not None:
        envelope_cache["presentation_signature"] = presentation_signature
    if artifacts is not None:
        envelope_cache["artifacts"] = artifacts
    cache["envelope"] = envelope_cache
    storage.save_project_analysis_cache(project_root, cache)


def _refresh_cached_presentation(
    project_root: Path,
    payload: dict[str, Any] | None,
    voltage_keys: list[str],
    chart_axis_limits: dict[str, float | None],
    chart_axis_limits_by_voltage: dict[str, dict[str, float | None]] | None,
    event_times: dict[str, float] | None,
    show_sa_label: bool,
    chart_top_left_cell: str | None,
    chart_size: dict[str, float],
    build_charts: bool,
    log: LogFn | None,
    check_cancel: CancelFn | None,
    calculation_signature: str,
    presentation_signature: str,
) -> list[Path] | None:
    """Refresh chart/display artifacts while retaining valid calculations."""
    if not isinstance(payload, dict):
        return None
    artifacts = payload.get("artifacts")
    if not isinstance(artifacts, dict):
        return None

    def artifact_paths(key: str) -> list[Path] | None:
        raw_paths = artifacts.get(key, [])
        if not isinstance(raw_paths, list):
            return None
        paths: list[Path] = []
        for raw_path in raw_paths:
            if not isinstance(raw_path, str):
                return None
            path = project_root / Path(raw_path)
            if not path.is_file():
                return None
            paths.append(path)
        return paths

    result_outputs = artifact_paths("result_outputs")
    data_outputs = artifact_paths("data_outputs")
    resonance_workbooks = artifact_paths("resonance_workbooks")
    sustained_summary_outputs = artifact_paths("sustained_summary_outputs")
    raw_chart_inputs = artifacts.get("chart_inputs", [])
    if None in (result_outputs, data_outputs, resonance_workbooks, sustained_summary_outputs):
        return None
    if not isinstance(raw_chart_inputs, list):
        return None
    chart_inputs: list[tuple[str, Path, Path]] = []
    for item in raw_chart_inputs:
        if not isinstance(item, dict):
            return None
        voltage = str(item.get("voltage", ""))
        envelope_raw = item.get("envelope")
        combined_raw = item.get("combined")
        if not isinstance(envelope_raw, str) or not isinstance(combined_raw, str):
            return None
        envelope_path = project_root / Path(envelope_raw)
        combined_path = project_root / Path(combined_raw)
        if not envelope_path.is_file():
            return None
        chart_inputs.append((voltage, envelope_path, combined_path))

    outputs = list(result_outputs or [])
    chart_outputs: list[Path] = []
    if build_charts:
        with ExitStack() as stack:
            excel = stack.enter_context(excel_app()) if data_outputs or chart_inputs or resonance_workbooks else None
            for workbook_path in data_outputs or []:
                _cancel(check_cancel)
                autofit_workbook(excel, workbook_path)
            outputs.extend(sustained_summary_outputs or [])
            for voltage_key, envelope_path, combined_path in chart_inputs:
                if voltage_key not in voltage_keys:
                    continue
                _cancel(check_cancel)
                chart_outputs.append(
                    create_combined_envelope_plot(
                        envelope_path,
                        combined_path,
                        excel,
                        axis_limits_override=chart_axis_limits,
                        axis_limits_by_voltage=chart_axis_limits_by_voltage,
                        event_times=event_times,
                        show_sa_label=show_sa_label,
                        chart_top_left_cell=chart_top_left_cell,
                        chart_size=chart_size,
                    )
                )
            for workbook_path in resonance_workbooks or []:
                _cancel(check_cancel)
                create_resonance_check_charts(
                    workbook_path,
                    excel,
                    x_max=chart_axis_limits["x_max"],
                    x_major=chart_axis_limits["x_major"],
                )
        outputs.extend(chart_outputs)
    else:
        if data_outputs:
            with excel_app() as excel:
                for workbook_path in data_outputs:
                    _cancel(check_cancel)
                    autofit_workbook(excel, workbook_path)
        outputs.extend(data_outputs or [])
        outputs.extend(resonance_workbooks or [])

    manifest_outputs = [
        *outputs,
        *(data_outputs or []),
        *(resonance_workbooks or []),
        *(sustained_summary_outputs or []),
        *chart_outputs,
    ]
    artifacts = {
        "result_outputs": [
            _relative_project_path(project_root, path) for path in result_outputs or []
        ],
        "data_outputs": [
            _relative_project_path(project_root, path) for path in data_outputs or []
        ],
        "sustained_summary_outputs": [
            _relative_project_path(project_root, path)
            for path in sustained_summary_outputs or []
        ],
        "resonance_workbooks": [
            _relative_project_path(project_root, path)
            for path in resonance_workbooks or []
        ],
        "chart_inputs": [
            {
                "voltage": voltage,
                "envelope": _relative_project_path(project_root, envelope_path),
                "combined": _relative_project_path(project_root, combined_path),
            }
            for voltage, envelope_path, combined_path in chart_inputs
        ],
    }
    _write_envelope_manifest(
        project_root,
        calculation_signature,
        manifest_outputs,
        return_outputs=outputs,
        presentation_signature=presentation_signature,
        artifacts=artifacts,
        calculation_outputs=[
            *(result_outputs or []),
            *(data_outputs or []),
            *(resonance_workbooks or []),
        ],
        presentation_outputs=chart_outputs,
    )
    _log(log, f"Reused envelope calculations; refreshed presentation: {project_root.name}")
    return outputs


@dataclass
class _VoltageBuildResult:
    chart_inputs: list[tuple[str, Path, Path]]
    resonance_results: list[resonance_checks.ResonanceResult]
    data_outputs: list[Path]
    sustained_results: dict[tuple[str, str], sustained_sdpf.SustainedSDPFResult | None] = field(default_factory=dict)
    sustained_observations: dict[tuple[str, str], list[sustained_sdpf.SustainedSDPFResult]] = field(default_factory=dict)
    sustained_signature_inputs: dict[tuple[str, str], dict[str, Any]] = field(default_factory=dict)


@dataclass
class _BusWaveforms:
    times: list[np.ndarray]
    values: list[np.ndarray]
    absolute_values: list[np.ndarray]
    signals: list[str]
    context: str


def build_voltage_envelopes(
    project_root: Path,
    scopes: Iterable[ScopeEntry],
    voltages: Iterable[str],
    log: LogFn | None = None,
    check_cancel: CancelFn | None = None,
    envelope_workers: int | None = None,
    exclusions: list[ExclusionRule] | None = None,
    high_voltage_proposals: list[dict[str, object]] | None = None,
    high_voltage_include_overrides: list[tuple[str, str, int, str]] | None = None,
    envelope_time_step: float | None = None,
    envelope_time_end: float | None = None,
    envelope_fallback_frequency: float | None = None,
    envelope_chart_x_max: float | None = None,
    envelope_chart_x_major: float | None = None,
    envelope_chart_y_limits_by_voltage: dict[str, dict[str, float | None]] | None = None,
    envelope_chart_show_sa_label: bool = False,
    envelope_chart_top_left_cell: str | None = None,
    envelope_chart_width: float | None = None,
    envelope_chart_height: float | None = None,
    high_voltage_limit_factor: float | None = None,
    nonconv_cb_iip_limit: float | None = None,
    nonconv_cb_iir_limit: float | None = None,
    event_times: dict[str, float] | None = None,
    project_timing: ProjectTiming | None = None,
    voltage_um_overrides: dict[str, float] | None = None,
    resonance_settings: dict[str, Any] | None = None,
    build_charts: bool = True,
    sustained_sdpf_settings: dict[str, Any] | None = None,
    sustained_sdpf_limit_overrides: dict[str, dict[str, float]] | None = None,
    voltage_configs: dict[str, Any] | None = None,
    sustained_sdpf_limits_by_voltage: dict[str, sustained_sdpf.SDPFVoltageLimits] | None = None,
    nonconv_cases: Iterable[Any] | None = None,
) -> list[Path]:
    started = time.perf_counter()
    project_root = Path(project_root)
    _remove_legacy_envelope_cache(project_root)
    case_root = project_root / "Case_folder"
    if not case_root.is_dir():
        raise FileNotFoundError(f"Case_folder not found: {case_root}")

    total_workers = _configured_worker_count(envelope_workers)
    _log(log, f"Reading statistic and convergence data: {project_root.name}")
    stat_started = time.perf_counter()
    df_stat = _read_stat_files(project_root)
    cb_iip_limit = normalize_positive_float(nonconv_cb_iip_limit, NONCONV_CB_IIP_LIMIT)
    cb_iir_limit = normalize_positive_float(nonconv_cb_iir_limit, NONCONV_CB_IIR_LIMIT)
    nonconv_warnings: list[str] = []
    if nonconv_cases is None:
        df_nonconv = _find_non_convergent_cases(
            case_root,
            df_stat,
            cb_iip_limit,
            cb_iir_limit,
            warnings=nonconv_warnings,
            worker_count=total_workers,
            check_cancel=check_cancel,
        )
    else:
        df_nonconv = _nonconv_dataframe(nonconv_cases)
    for warning in nonconv_warnings:
        _log(log, warning)
    _log(log, f"Envelope setup complete: {project_root.name} | {_elapsed(stat_started)}")
    detected_nonconv_keys = _nonconv_keys(df_nonconv)
    exclusion_matcher = ExclusionMatcher(exclusions or [])
    include_override_keys = {
        key
        for voltage, case, run, bus in high_voltage_include_overrides or []
        if (key := _high_voltage_key(voltage, case, run, bus)) is not None
    }
    if detected_nonconv_keys:
        _log(log, f"Detected non-convergent proposals: {len(detected_nonconv_keys)}")
    if exclusion_matcher.rules:
        _log(log, f"Applied exclusion rules: {len(exclusion_matcher.rules)}")
    chart_axis_limits = _chart_axis_limits(envelope_chart_x_max, envelope_chart_x_major)
    limit_factor = normalize_positive_float(high_voltage_limit_factor, HIGH_VOLTAGE_LIMIT_FACTOR)
    time_step = normalize_positive_float(envelope_time_step, TIME_STEP)
    project_timing = project_timing or load_project_timing(project_root)
    automatic_time_end = envelope_time_end is None
    requested_time_end = (
        project_timing.final_duration or TIME_END
        if automatic_time_end
        else normalize_positive_float(envelope_time_end, TIME_END)
    )
    time_end = min(requested_time_end, project_timing.final_duration or requested_time_end)
    if project_timing.final_duration is not None:
        source = "Use project duration" if automatic_time_end else f"Settings limit {requested_time_end:g} s"
        _log(
            log,
            f"Envelope time end: {time_end:g} s "
            f"({source}; Input_Data Final duration {project_timing.final_duration:g} s)",
        )
    else:
        source = "automatic fallback" if automatic_time_end else "Settings"
        _log(log, f"Envelope time end from {source}: {time_end:g} s")
    resonance_settings_obj = resonance_checks.ResonanceSettings.from_mapping(resonance_settings)
    sustained_settings_obj = sustained_sdpf.SustainedSDPFSettings.from_mapping(sustained_sdpf_settings)
    sustained_worker_settings = sustained_settings_obj
    if resonance_settings_obj.enabled and not resonance_settings_obj.auto_release:
        if not (0.0 <= resonance_settings_obj.manual_analysis_start < time_end):
            raise ValueError("Manual analysis start time must be within [0, envelope time end).")
        if resonance_checks.NO_SETTLE_GROWTH in resonance_settings_obj.enabled_checks:
            _log(log, "No-settle growth skipped: manual analysis start time is enabled, so physical release detection is bypassed.")
    settings_fallback_frequency = normalize_positive_float(envelope_fallback_frequency, DEFAULT_ENVELOPE_FALLBACK_FREQUENCY)
    project_frequency = project_timing.frequency
    fallback_frequency = project_frequency or settings_fallback_frequency
    if project_frequency is not None:
        _log(log, f"Fallback frequency from Input_Data!B16: {fallback_frequency:g} Hz")
    else:
        _log(log, f"Fallback frequency from Settings: {fallback_frequency:g} Hz")
    frequency_warning_lock = threading.Lock()
    frequency_warning_sent = False

    def notify_frequency_fallback(context: str) -> None:
        nonlocal frequency_warning_sent
        with frequency_warning_lock:
            first_warning = not frequency_warning_sent
            frequency_warning_sent = True
        if project_frequency is not None:
            _log(log, f"Frequency auto-detection failed; using Input_Data!B16 {fallback_frequency:g} Hz: {context}")
            return
        if first_warning:
            _log(
                log,
                f"{FREQUENCY_FALLBACK_MESSAGE_PREFIX}{project_root.name}|{fallback_frequency:g}|{context}",
            )
        else:
            _log(log, f"Frequency auto-detection failed; using fallback {fallback_frequency:g} Hz: {context}")

    chart_size = {
        "width": normalize_positive_float(envelope_chart_width, DEFAULT_ENVELOPE_CHART_WIDTH),
        "height": normalize_positive_float(envelope_chart_height, DEFAULT_ENVELOPE_CHART_HEIGHT),
    }
    if voltage_configs is None or voltage_um_overrides:
        voltage_configs = load_voltage_configs(project_root, um_overrides=voltage_um_overrides)
    else:
        voltage_configs = dict(voltage_configs)
    sustained_limits_by_voltage: dict[str, sustained_sdpf.SDPFVoltageLimits] = {}
    sustained_limit_warnings: list[str] = []
    if sustained_settings_obj.enabled:
        if sustained_sdpf_limits_by_voltage is None:
            sustained_limits_by_voltage, sustained_limit_warnings = sustained_sdpf.resolve_project_limits(
                project_root,
                workbook_path=input_data_workbook(project_root),
                overrides_path=project_root / "Plots" / ".plottool_v3" / "limits.json",
                manual_overrides=sustained_sdpf_limit_overrides,
            )
        else:
            sustained_limits_by_voltage = sustained_sdpf.apply_limit_overrides(
                sustained_sdpf_limits_by_voltage,
                sustained_sdpf_limit_overrides,
            )
        for warning in sustained_limit_warnings:
            _log(log, f"Sustained SDPF: {warning}")
    outputs: list[Path] = []
    sustained_summary_outputs: list[Path] = []

    selected_scopes = list(scopes)
    inf_inventory = _inf_inventory(case_root, exclusion_matcher)
    scope_inf_paths = {
        scope.folder: _selected_inf_paths(inf_inventory, scope)
        for scope in selected_scopes
    }
    for scope in selected_scopes:
        if not scope_inf_paths[scope.folder]:
            _log(log, f"No PSCAD run files match scope: {project_root.name} | {scope.folder}")

    voltage_keys = [str(voltage).strip() for voltage in voltages if str(voltage).strip()]
    all_inf_paths = sorted({path for paths in scope_inf_paths.values() for path in paths})
    static_source_paths = [
        *sorted(project_root.glob("Input_Data_PSCAD*.xlsx")),
        project_root / "PSCAD_log.txt",
        project_root / "Plots" / ".plottool_v3" / "limits.json",
    ]
    source_inventory = _source_file_inventory(
        project_root,
        all_inf_paths,
        static_source_paths,
    )
    run_source_manifests = _run_source_manifests(
        project_root,
        all_inf_paths,
        source_inventory,
    )
    inf_descriptor_cache = _read_inf_descriptor_cache(all_inf_paths, total_workers, check_cancel, log)
    sustained_source_files_by_scope = {}
    if sustained_settings_obj.enabled:
        static_source_files = _sustained_source_manifest(
            project_root,
            [],
            source_inventory=source_inventory,
        )
        sustained_source_files_by_scope = {
            scope.folder: _sustained_source_manifest(
                project_root,
                scope_inf_paths[scope.folder],
                static_files=static_source_files,
                source_inventory=source_inventory,
            )
            for scope in selected_scopes
        }

    calculation_settings = {
        "project_timing": project_timing,
        "time_step": time_step,
        "time_end": time_end,
        "fallback_frequency": fallback_frequency,
        "high_voltage_limit_factor": limit_factor,
        "nonconv_cb_iip_limit": cb_iip_limit,
        "nonconv_cb_iir_limit": cb_iir_limit,
        "voltage_configs": voltage_configs,
        "exclusions": [rule.to_dict() for rule in exclusion_matcher.rules],
        "high_voltage_proposals": high_voltage_proposals or [],
        "high_voltage_include_overrides": sorted(include_override_keys),
        "stat": df_stat,
        "nonconv": df_nonconv,
        "resonance": resonance_settings_obj,
        "event_times": event_times or {},
        "sustained": sustained_settings_obj,
        "sustained_limits": sustained_limits_by_voltage,
        "sustained_limit_warnings": sustained_limit_warnings,
        "sustained_sources": sustained_source_files_by_scope,
    }
    calculation_settings.update(_sustained_cache_versions(sustained_settings_obj))
    presentation_settings = {
        "build_charts": build_charts,
        "chart_axis_limits": chart_axis_limits,
        "chart_y_limits_by_voltage": envelope_chart_y_limits_by_voltage,
        "chart_show_sa_label": envelope_chart_show_sa_label,
        "chart_top_left_cell": envelope_chart_top_left_cell,
        "chart_size": chart_size,
    }

    manifest_payload = {
        "version": ENVELOPE_MANIFEST_VERSION,
        "calculation_version": ENVELOPE_CALCULATION_VERSION,
        "presentation_version": ENVELOPE_PRESENTATION_VERSION,
        "scopes": [
            {
                "name": scope.name,
                "mode": scope.mode,
                "tokens": list(scope.tokens),
            }
            for scope in selected_scopes
        ],
        "scope_inf_paths": {
            scope.folder: [
                _relative_project_path(project_root, path)
                for path in scope_inf_paths[scope.folder]
            ]
            for scope in selected_scopes
        },
        "voltages": voltage_keys,
        "sources": {
            _relative_project_path(project_root, path): source_manifest
            for path, source_manifest in run_source_manifests.items()
        },
        "calculation": calculation_settings,
        "presentation": presentation_settings,
    }
    calculation_payload = {
        "version": ENVELOPE_CALCULATION_VERSION,
        "scopes": manifest_payload["scopes"],
        "scope_inf_paths": manifest_payload["scope_inf_paths"],
        "voltages": voltage_keys,
        "sources": manifest_payload["sources"],
        "settings": calculation_settings,
    }
    calculation_signature = _cache_signature(calculation_payload)
    presentation_signature = _cache_signature(
        {
            "version": ENVELOPE_PRESENTATION_VERSION,
            "calculation_signature": calculation_signature,
            "settings": presentation_settings,
        }
    )
    cached_manifest = _envelope_manifest_payload(project_root, calculation_signature)
    if cached_manifest is not None:
        if cached_manifest.get("presentation_signature") == presentation_signature:
            cached_outputs = _envelope_manifest_matches(project_root, calculation_signature)
            if cached_outputs is not None:
                _log(log, f"Skipping unchanged envelope build: {project_root.name}")
                return cached_outputs
        else:
            refreshed_outputs = _refresh_cached_presentation(
                project_root,
                cached_manifest,
                voltage_keys,
                chart_axis_limits,
                envelope_chart_y_limits_by_voltage,
                event_times,
                envelope_chart_show_sa_label,
                envelope_chart_top_left_cell,
                chart_size,
                build_charts,
                log,
                check_cancel,
                calculation_signature,
                presentation_signature,
            )
            if refreshed_outputs is not None:
                return refreshed_outputs

    chart_inputs: list[tuple[str, Path, Path]] = []
    data_outputs: list[Path] = []
    result_outputs: list[Path] = []
    resonance_results: list[resonance_checks.ResonanceResult] = []
    sustained_results: dict[tuple[str, str], sustained_sdpf.SustainedSDPFResult | None] = {}
    sustained_observations: dict[tuple[str, str], list[sustained_sdpf.SustainedSDPFResult]] = {}
    sustained_signature_inputs: dict[tuple[str, str], dict[str, Any]] = {}
    _log(log, f"Envelope worker plan: shared run-read pool={total_workers}; voltage reads concurrent")

    def build_voltage(voltage_key: str) -> _VoltageBuildResult:
        return _build_voltage_workbooks(
            project_root,
            selected_scopes,
            scope_inf_paths,
            all_inf_paths,
            inf_descriptor_cache,
            voltage_key,
            voltage_configs,
            exclusion_matcher,
            high_voltage_proposals or [],
            include_override_keys,
            limit_factor,
            time_step,
            time_end,
            fallback_frequency,
            df_stat,
            df_nonconv,
            total_workers,
            log,
            check_cancel,
            notify_frequency_fallback,
            resonance_settings_obj,
            event_times,
            executor,
            sustained_worker_settings,
            sustained_limits_by_voltage.get(voltage_key),
            project_frequency or fallback_frequency,
            sustained_source_files_by_scope,
        )

    with ProcessPoolExecutor(max_workers=total_workers) as executor:
        if len(voltage_keys) > 1:
            with ThreadPoolExecutor(max_workers=min(len(voltage_keys), 3)) as voltage_executor:
                voltage_futures = {
                    voltage_key: voltage_executor.submit(build_voltage, voltage_key)
                    for voltage_key in voltage_keys
                }
                voltage_results = [voltage_futures[voltage_key].result() for voltage_key in voltage_keys]
        else:
            voltage_results = [build_voltage(voltage_key) for voltage_key in voltage_keys]

        for result in voltage_results:
            chart_inputs.extend(result.chart_inputs)
            data_outputs.extend(result.data_outputs)
            resonance_results.extend(result.resonance_results)
            sustained_results.update(result.sustained_results)
            sustained_observations.update(result.sustained_observations)
            sustained_signature_inputs.update(result.sustained_signature_inputs)

    if sustained_settings_obj.enabled:
        for scope in selected_scopes:
            scope_results = {
                voltage: sustained_results.get((scope.folder, voltage))
                for voltage in voltage_keys
                if (scope.folder, voltage) in sustained_results
            }
            scope_observations = {
                voltage: sustained_observations.get((scope.folder, voltage), [])
                for voltage in voltage_keys
                if (scope.folder, voltage) in sustained_observations
            }
            scope_signature_inputs = {
                voltage: sustained_signature_inputs.get((scope.folder, voltage), {})
                for voltage in scope_results
            }
            result_path = sustained_sdpf.save_results(
                project_root,
                scope.folder,
                sustained_settings_obj,
                scope_results,
                sustained_limit_warnings,
                scope_signature_inputs,
                observations_by_voltage=scope_observations,
            )
            outputs.append(result_path)
            result_outputs.append(result_path)
            _log(log, f"Sustained SDPF result saved: {result_path}")
            summary_path = sustained_sdpf.write_summary_workbook(
                project_root,
                scope.folder,
                scope_observations,
                sustained_limits_by_voltage,
                settings=sustained_settings_obj,
                frequency_hz=project_frequency or fallback_frequency,
            )
            data_outputs.append(summary_path)
            sustained_summary_outputs.append(summary_path)
            _log(log, f"Sustained SDPF ranked summary saved: {summary_path}")
    else:
        for scope in selected_scopes:
            stale_result = sustained_sdpf.result_path(project_root, scope.folder)
            try:
                stale_result.unlink(missing_ok=True)
            except OSError:
                _log(log, f"Could not remove obsolete Sustained SDPF result: {stale_result}")
            stale_summary = sustained_sdpf.summary_path(project_root, scope.folder)
            try:
                stale_summary.unlink(missing_ok=True)
            except OSError:
                _log(log, f"Could not remove obsolete Sustained SDPF summary: {stale_summary}")
            sustained_sdpf_heatmap.clear_heatmaps(project_root, scope.folder)

    resonance_workbooks: list[Path] = []
    if resonance_settings_obj.enabled:
        for workbook_path in resonance_checks.write_workbooks(
            project_root,
            selected_scopes,
            resonance_results,
            resonance_settings_obj,
            event_times,
        ):
            resonance_workbooks.append(workbook_path)
            _log(log, f"Resonance checks workbook finished: {workbook_path}")

    if build_charts:
        chart_inputs.sort(key=lambda item: (voltage_keys.index(item[0]) if item[0] in voltage_keys else 999, str(item[1])))
        with ExitStack() as stack:
            excel = stack.enter_context(excel_app()) if data_outputs or chart_inputs or resonance_workbooks else None
            for workbook_path in data_outputs:
                autofit_workbook(excel, workbook_path)
                _log(log, f"Envelope workbook columns autofitted: {workbook_path.name}")
            outputs.extend(sustained_summary_outputs)
            for voltage_key, envelope_path, combined_path in chart_inputs:
                chart_started = time.perf_counter()
                create_combined_envelope_plot(
                    envelope_path,
                    combined_path,
                    excel,
                    axis_limits_override=chart_axis_limits,
                    axis_limits_by_voltage=envelope_chart_y_limits_by_voltage,
                    event_times=event_times,
                    show_sa_label=envelope_chart_show_sa_label,
                    chart_top_left_cell=envelope_chart_top_left_cell,
                    chart_size=chart_size,
                )
                outputs.append(combined_path)
                _log(log, f"Envelope chart finished: {voltage_key} kV | {combined_path.name} | {_elapsed(chart_started)}")
            for workbook_path in resonance_workbooks:
                chart_started = time.perf_counter()
                if create_resonance_check_charts(
                    workbook_path,
                    excel,
                    x_max=chart_axis_limits["x_max"],
                    x_major=chart_axis_limits["x_major"],
                ):
                    _log(log, f"Analysis charts finished: {workbook_path.name} | {_elapsed(chart_started)}")
    else:
        if data_outputs:
            with excel_app() as excel:
                for workbook_path in data_outputs:
                    autofit_workbook(excel, workbook_path)
                    _log(log, f"Envelope workbook columns autofitted: {workbook_path.name}")
        outputs.extend(data_outputs)
        outputs.extend(resonance_workbooks)

    manifest_outputs = [
        *outputs,
        *data_outputs,
        *resonance_workbooks,
        *(combined_path for _voltage, _envelope_path, combined_path in chart_inputs),
    ]
    artifacts = {
        "result_outputs": [
            _relative_project_path(project_root, path) for path in result_outputs
        ],
        "data_outputs": [
            _relative_project_path(project_root, path) for path in data_outputs
        ],
        "sustained_summary_outputs": [
            _relative_project_path(project_root, path)
            for path in sustained_summary_outputs
        ],
        "resonance_workbooks": [
            _relative_project_path(project_root, path)
            for path in resonance_workbooks
        ],
        "chart_inputs": [
            {
                "voltage": voltage,
                "envelope": _relative_project_path(project_root, envelope_path),
                "combined": _relative_project_path(project_root, combined_path),
            }
            for voltage, envelope_path, combined_path in chart_inputs
        ],
    }
    _write_envelope_manifest(
        project_root,
        calculation_signature,
        manifest_outputs,
        return_outputs=outputs,
        presentation_signature=presentation_signature,
        artifacts=artifacts,
        calculation_outputs=[
            *result_outputs,
            *data_outputs,
            *resonance_workbooks,
        ],
        presentation_outputs=[combined_path for _v, _e, combined_path in chart_inputs],
    )
    _log(log, f"Envelope build total: {project_root.name} | {_elapsed(started)}")
    return outputs


def _elapsed(started: float) -> str:
    return f"{time.perf_counter() - started:.1f}s"


def _configured_worker_count(envelope_workers: int | None = None) -> int:
    raw_value = envelope_workers
    if raw_value is None:
        raw_value = os.environ.get(MAX_WORKERS_ENV, str(DEFAULT_ENVELOPE_WORKERS))
    try:
        configured = int(raw_value)
    except (TypeError, ValueError):
        configured = 0
    return min(MAX_ENVELOPE_WORKERS, configured) if configured > 0 else _automatic_worker_count()


def _automatic_worker_count() -> int:
    return automatic_worker_count()


def _nominal_voltage_from_key(voltage_key: str, default: float) -> float:
    try:
        number = float(str(voltage_key).strip())
    except (TypeError, ValueError):
        return default
    return number if math.isfinite(number) and number > 0 else default


def _high_voltage_key(
    voltage: object,
    case: object,
    run: object,
    bus: object,
) -> tuple[str, str, int, str] | None:
    try:
        run_number = int(run)
    except (TypeError, ValueError):
        return None
    voltage_key = normalize_voltage(voltage)
    case_key = str(case).strip().casefold()
    bus_key = str(bus).strip().casefold()
    if not voltage_key or not case_key or not bus_key:
        return None
    return voltage_key, case_key, run_number, bus_key


def _chart_axis_limits(x_max: float | None, x_major: float | None) -> dict[str, float | None]:
    return {
        "x_max": float(x_max) if x_max is not None and x_max > 0 else None,
        "x_major": float(x_major) if x_major is not None and x_major > 0 else None,
    }


def _build_voltage_workbooks(
    project_root: Path,
    selected_scopes: list[ScopeEntry],
    scope_inf_paths: dict[str, list[Path]],
    all_inf_paths: list[Path],
    inf_descriptor_cache: dict[Path, list[InfDescriptor]],
    voltage_key: str,
    voltage_configs,
    exclusion_matcher: ExclusionMatcher,
    high_voltage_proposals: list[dict[str, object]],
    high_voltage_include_overrides: set[tuple[str, str, int, str]],
    high_voltage_limit_factor: float,
    time_step: float,
    time_end: float,
    fallback_frequency: float,
    df_stat: pd.DataFrame,
    df_nonconv: pd.DataFrame,
    worker_count: int,
    log: LogFn | None,
    check_cancel: CancelFn | None,
    frequency_fallback: FrequencyFallbackFn | None = None,
    resonance_settings: resonance_checks.ResonanceSettings | None = None,
    event_times: dict[str, float] | None = None,
    run_executor: Executor | None = None,
    sustained_settings: sustained_sdpf.SustainedSDPFSettings | None = None,
    sustained_limits: sustained_sdpf.SDPFVoltageLimits | None = None,
    sustained_frequency: float | None = None,
    sustained_source_files_by_scope: dict[str, list[dict[str, Any]]] | None = None,
) -> _VoltageBuildResult:
    started = time.perf_counter()
    _cancel(check_cancel)
    config = voltage_configs.get(voltage_key)
    if config is None:
        _log(log, f"Envelope voltage skipped; no MM_blocks Um config: {voltage_key} kV")
        return _VoltageBuildResult([], [], [])

    if not all_inf_paths:
        return _VoltageBuildResult([], [], [])

    bus_prefix, bus_um = config.bus_prefix, config.um
    nominal_voltage = _nominal_voltage_from_key(voltage_key, bus_um)

    _log(log, f"Envelope source read started: {voltage_key} kV | {len(all_inf_paths)} unique runs")
    read_started = time.perf_counter()
    run_data = _read_voltage_runs(
        all_inf_paths,
        voltage_key,
        bus_prefix,
        bus_um,
        exclusion_matcher,
        worker_count,
        log,
        check_cancel,
        high_voltage_include_overrides=high_voltage_include_overrides,
        high_voltage_limit_factor=high_voltage_limit_factor,
        time_step=time_step,
        time_end=time_end,
        fallback_frequency=fallback_frequency,
        frequency_fallback=frequency_fallback,
        inf_descriptor_cache=inf_descriptor_cache,
        executor=run_executor,
        process_pool=isinstance(run_executor, ProcessPoolExecutor),
        sustained_settings=sustained_settings,
        sustained_limits=sustained_limits,
        sustained_frequency=sustained_frequency,
    )
    _log(log, f"Envelope source read finished: {voltage_key} kV | {_elapsed(read_started)}")

    chart_inputs: list[tuple[str, Path, Path]] = []
    data_outputs: list[Path] = []
    resonance_results: list[resonance_checks.ResonanceResult] = []
    sustained_results: dict[tuple[str, str], sustained_sdpf.SustainedSDPFResult | None] = {}
    sustained_observations: dict[tuple[str, str], list[sustained_sdpf.SustainedSDPFResult]] = {}
    sustained_signature_inputs: dict[tuple[str, str], dict[str, Any]] = {}
    resonance_settings = resonance_settings or resonance_checks.ResonanceSettings()
    for scope in selected_scopes:
        inf_paths = scope_inf_paths[scope.folder]
        if not inf_paths:
            continue

        _cancel(check_cancel)
        scope_started = time.perf_counter()
        _log(log, f"Envelope merge started: {scope.folder} | {voltage_key} kV | {len(inf_paths)} runs")
        entries, high_voltage, sustained_rows = _entries_for_inf_paths(inf_paths, run_data)
        if resonance_settings.enabled:
            records = []
            for voltage_type in resonance_checks.VOLTAGE_TYPES:
                records.extend(
                    resonance_checks.records_from_entries(
                        scope.folder,
                        voltage_key,
                        voltage_type,
                        nominal_voltage,
                        entries,
                    )
                )
            scope_results = resonance_checks.analyze_records(records, resonance_settings, event_times)
            resonance_results.extend(scope_results)
            _log(
                log,
                f"Resonance checks finished: {scope.folder} | {voltage_key} kV | selected={len(scope_results)}",
            )
        if sustained_settings is not None and sustained_settings.enabled and sustained_limits is not None:
            assessment_frequency = sustained_frequency or fallback_frequency
            signature, signature_inputs = _sustained_signature(
                project_root,
                scope.folder,
                voltage_key,
                inf_paths,
                sustained_settings,
                sustained_limits,
                sustained_frequency,
                exclusion_matcher,
                high_voltage_include_overrides,
                source_files=(sustained_source_files_by_scope or {}).get(scope.folder),
            )
            sustained_candidates = _sustained_results_for_scope(
                sustained_rows,
                scope.folder,
                voltage_key,
                sustained_settings.effective_duration(sustained_frequency or fallback_frequency),
                df_stat,
                sustained_limits,
                signature,
                sustained_settings.required_cycles(assessment_frequency),
            )
            sustained_result = sustained_sdpf.select_governing(sustained_candidates)
            key = (scope.folder, voltage_key)
            sustained_results[key] = sustained_result
            sustained_observations[key] = sustained_candidates
            sustained_signature_inputs[key] = signature_inputs
        lg_df = _final_envelope([df for measurement, df in entries if measurement == "LGp"], df_stat, time_step)
        ll_df = _final_envelope([df for measurement, df in entries if measurement == "LLp"], df_stat, time_step)
        if lg_df.empty or ll_df.empty:
            _log(log, f"No waveform envelope data: {scope.folder} | {voltage_key} kV")
            continue
        high_voltage_df = _high_voltage_exclusion_dataframe(
            high_voltage,
            high_voltage_proposals,
            voltage_key,
            inf_paths,
            df_stat,
            available_keys={
                key
                for inf_path in inf_paths
                if inf_path in inf_descriptor_cache
                for descriptor in inf_descriptor_cache[inf_path]
                if descriptor.Group.startswith(bus_prefix)
                and "p_" in descriptor.Description
                for case, run in [case_run_from_inf_path(inf_path)]
                if (
                    key := _high_voltage_key(
                        voltage_key,
                        case,
                        run,
                        descriptor.Group,
                    )
                )
                is not None
            },
        )
        high_voltage_df = _filter_dataframe_by_scope(high_voltage_df, scope)
        nonconv_df = _filter_dataframe_by_scope(df_nonconv, scope)

        write_started = time.perf_counter()
        envelope_path = project_root / "Voltage_envelope" / scope.folder / f"MM_{voltage_key}.xlsx"
        combined_path = envelope_path.with_name(f"{envelope_path.stem}_with_combined_plot.xlsx")
        _write_workbook(envelope_path, lg_df, ll_df, nonconv_df, high_voltage_df)
        data_outputs.append(envelope_path)
        chart_inputs.append((voltage_key, envelope_path, combined_path))
        _log(
            log,
            f"Envelope workbook finished: {scope.folder} | {voltage_key} kV | "
            f"LG rows={len(lg_df)}, LL rows={len(ll_df)} | merge+write={_elapsed(scope_started)}, write={_elapsed(write_started)}",
        )

    _log(log, f"Envelope voltage finished: {voltage_key} kV | {_elapsed(started)}")
    return _VoltageBuildResult(
        chart_inputs=chart_inputs,
        resonance_results=resonance_results,
        data_outputs=data_outputs,
        sustained_results=sustained_results,
        sustained_observations=sustained_observations,
        sustained_signature_inputs=sustained_signature_inputs,
    )


def _sustained_results_for_scope(
    rows: list[dict[str, Any]],
    scope_folder: str,
    voltage: str,
    duration_s: float,
    df_stat: pd.DataFrame,
    limits: sustained_sdpf.SDPFVoltageLimits,
    signature: str,
    required_cycles: int,
) -> list[sustained_sdpf.SustainedSDPFResult]:
    fault_types: dict[tuple[str, int], str] = {}
    if not df_stat.empty and {"Case", "Run#", "Fault_type"} <= set(df_stat.columns):
        for _, row in df_stat[["Case", "Run#", "Fault_type"]].iterrows():
            try:
                case_key = str(row["Case"])
                run_key = int(float(row["Run#"]))
                fault_types[(case_key, run_key)] = str(row["Fault_type"] or "")
            except (TypeError, ValueError):
                continue

    grouped: dict[tuple[str, int, str], list[sustained_sdpf.PhaseStressResult]] = {}
    results: list[sustained_sdpf.SustainedSDPFResult] = []
    for row in rows:
        raw_result = row.get("result") if isinstance(row, dict) else None
        if isinstance(raw_result, dict):
            try:
                worker_result = sustained_sdpf.SustainedSDPFResult.from_dict(raw_result)
            except (KeyError, TypeError, ValueError):
                continue
            results.append(
                replace(
                    worker_result,
                    scope_folder=scope_folder,
                    voltage=str(voltage),
                    fault_type=fault_types.get((worker_result.case, worker_result.run), ""),
                    signature=signature,
                )
            )
            continue
        try:
            phase_result = sustained_sdpf.PhaseStressResult.from_dict(row)
            key = (str(row["case"]), int(row["run"]), str(row["mm_name"]))
        except (KeyError, TypeError, ValueError):
            continue
        # The worker used the same per-voltage limit object; reject malformed rows
        # rather than silently replacing a missing source value.
        expected_rms = limits.rms(phase_result.measurement)
        if abs(phase_result.sdpf_rms_kv - expected_rms) > 1e-9:
            continue
        grouped.setdefault(key, []).append(phase_result)

    for (case, run, mm_name), phase_rows in grouped.items():
        result = sustained_sdpf.classify_rows(
            scope_folder,
            voltage,
            case,
            run,
            mm_name,
            phase_rows,
            duration_s,
            fault_types.get((case, run), ""),
            signature,
            cycle_coverage=required_cycles,
        )
        if result is not None:
            results.append(result)
    return results


def _sustained_source_manifest(
    project_root: Path,
    inf_paths: list[Path],
    static_files: list[dict[str, Any]] | None = None,
    source_inventory: _SourceFileInventory | None = None,
) -> list[dict[str, Any]]:
    """Collect Sustained SDPF source metadata once per scope.

    The same run inventory is used for every voltage level.  Keeping this
    walk outside the voltage loop avoids repeated directory scans and makes
    the persisted manifest portable within the project folder.
    """
    candidates: list[Path] = []
    if static_files is None:
        candidates.extend(
            [
                *sorted(project_root.glob("Input_Data_PSCAD*.xlsx")),
                project_root / "PSCAD_log.txt",
                project_root / "Plots" / ".plottool_v3" / "limits.json",
            ]
        )
    candidates.extend(inf_paths)

    # The source directory is the smallest stable scope we can persist.  The
    # compact result stores these directories plus one fingerprint instead of
    # retaining every raw-file metadata triple.
    if source_inventory is not None:
        for directory in dict.fromkeys(path.parent for path in inf_paths):
            candidates.extend(source_inventory.directory_paths(directory))
    else:
        for directory in dict.fromkeys(path.parent for path in inf_paths):
            try:
                outputs = sorted(directory.glob("*.inf"))
                outputs.extend(sorted(directory.glob("*.out")))
            except OSError:
                outputs = []
            candidates.extend(outputs)

    root_path = Path(project_root)
    files: list[dict[str, Any]] = [dict(entry) for entry in (static_files or [])]
    for path in dict.fromkeys(candidates):
        files.append(
            source_inventory.record(path, root_path)
            if source_inventory is not None
            else _file_manifest_entry(root_path, path)
        )
    return files


def _sustained_signature(
    project_root: Path,
    scope_folder: str,
    voltage: str,
    inf_paths: list[Path],
    settings: sustained_sdpf.SustainedSDPFSettings,
    limits: sustained_sdpf.SDPFVoltageLimits,
    frequency: float | None,
    exclusion_matcher: ExclusionMatcher,
    high_voltage_include_overrides: set[tuple[str, str, int, str]],
    source_files: list[dict[str, Any]] | None = None,
) -> tuple[str, dict[str, Any]]:
    case_names: list[str] = []
    for inf_path in inf_paths:
        try:
            case_name, _run_number = case_run_from_inf_path(inf_path)
            if case_name not in case_names:
                case_names.append(case_name)
        except ValueError:
            pass
    files = (
        source_files
        if source_files is not None
        else _sustained_source_manifest(project_root, inf_paths)
    )
    rules = [rule.to_dict() if hasattr(rule, "to_dict") else repr(rule) for rule in exclusion_matcher.rules]
    inputs = {
        "version": sustained_sdpf.RESULT_VERSION,
        "scope": scope_folder,
        "voltage": str(voltage),
        "settings": settings.to_mapping(),
        "frequency": frequency,
        "limits": limits.to_dict(),
        "heatmap_case_names": case_names,
        "source_roots": sorted(
            {
                _relative_project_path(project_root, inf_path.parent)
                for inf_path in inf_paths
            }
        ),
        "exclusions": rules,
        "high_voltage_include_overrides": sorted(high_voltage_include_overrides),
        "files": files,
    }
    # File freshness is checked directly from the compact manifest.  Do not
    # hash the entire source list again for the per-result diagnostic signature.
    signature_payload = {key: value for key, value in inputs.items() if key != "files"}
    return sustained_sdpf.make_signature(signature_payload), inputs


def _scope_text(path: Path) -> str:
    try:
        case, run = case_run_from_inf_path(path)
    except ValueError:
        case, run = path.stem, ""
    return f"{case} {run} {path.parent.name} {path}".casefold()


def _path_matches_scope(path: Path, scope: ScopeEntry) -> bool:
    if scope.mode == "full":
        return True
    tokens = [token.casefold() for token in scope.tokens]
    matched = any(token in _scope_text(path) for token in tokens)
    return matched if scope.mode == "include" else not matched


def _inf_inventory(
    case_root: Path,
    exclusion_matcher: ExclusionMatcher | None = None,
) -> list[Path]:
    paths: list[Path] = []
    matcher = exclusion_matcher or ExclusionMatcher()
    for path in sorted(case_root.rglob("*.inf")):
        try:
            case, run = case_run_from_inf_path(path)
        except ValueError:
            continue
        if matcher.excludes_run("", case, run):
            continue
        paths.append(path)
    return paths


def _selected_inf_paths(inf_inventory: Iterable[Path], scope: ScopeEntry) -> list[Path]:
    return [path for path in inf_inventory if _path_matches_scope(path, scope)]


def _worker_count(item_count: int, configured: int) -> int:
    if item_count <= 0:
        return 1
    return max(1, min(item_count, os.cpu_count() or 1, configured))


def _read_inf_descriptor_cache(
    inf_paths: list[Path],
    worker_count: int,
    check_cancel: CancelFn | None,
    log: LogFn | None = None,
) -> dict[Path, list[InfDescriptor]]:
    if not inf_paths:
        return {}

    workers = _worker_count(len(inf_paths), worker_count)
    with ThreadPoolExecutor(max_workers=workers) as executor:
        hash_futures = {executor.submit(hash_inf_file, inf_path): inf_path for inf_path in inf_paths}
        paths_by_digest: dict[str, list[Path]] = {}
        for future in as_completed(hash_futures):
            _cancel(check_cancel)
            inf_path = hash_futures[future]
            paths_by_digest.setdefault(future.result(), []).append(inf_path)

        descriptor_cache: dict[Path, list[InfDescriptor]] = {}
        parse_futures = {
            executor.submit(parse_inf_descriptors, paths[0]): (digest, paths)
            for digest, paths in paths_by_digest.items()
        }
        for future in as_completed(parse_futures):
            _cancel(check_cancel)
            _digest, matching_paths = parse_futures[future]
            try:
                descriptors = future.result()
            except ValueError as exc:
                names = ", ".join(path.name for path in matching_paths[:3])
                extra = f" and {len(matching_paths) - 3} more" if len(matching_paths) > 3 else ""
                _log(log, f"Skipped invalid .inf layout: {names}{extra} | {exc}")
                continue
            for inf_path in matching_paths:
                descriptor_cache[inf_path] = descriptors
    return descriptor_cache


def _read_voltage_runs(
    inf_paths: list[Path],
    voltage_key: str,
    bus_prefix: str,
    bus_um: float,
    exclusion_matcher: ExclusionMatcher,
    worker_count: int,
    log: LogFn | None,
    check_cancel: CancelFn | None,
    high_voltage_include_overrides: set[tuple[str, str, int, str]] | None = None,
    high_voltage_limit_factor: float = HIGH_VOLTAGE_LIMIT_FACTOR,
    time_step: float = TIME_STEP,
    time_end: float = TIME_END,
    fallback_frequency: float = DEFAULT_ENVELOPE_FALLBACK_FREQUENCY,
    frequency_fallback: FrequencyFallbackFn | None = None,
    inf_descriptor_cache: dict[Path, list[InfDescriptor]] | None = None,
    executor: Executor | None = None,
    process_pool: bool = False,
    sustained_settings: sustained_sdpf.SustainedSDPFSettings | None = None,
    sustained_limits: sustained_sdpf.SDPFVoltageLimits | None = None,
    sustained_frequency: float | None = None,
) -> dict[Path, tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]]]]:
    limit = high_voltage_limit_factor * bus_um * math.sqrt(2)
    run_data: dict[Path, tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]]]] = {}

    workers = _worker_count(len(inf_paths), worker_count)
    completed = 0
    owns_executor = executor is None
    if executor is None:
        executor = ThreadPoolExecutor(max_workers=workers)
    try:
        futures = {}
        for inf_path in inf_paths:
            _cancel(check_cancel)
            case_name, run = case_run_from_inf_path(inf_path)
            if exclusion_matcher.excludes_run(voltage_key, case_name, run):
                run_data[inf_path] = ([], [])
                continue
            if process_pool:
                descriptors = (inf_descriptor_cache or {}).get(inf_path)
                if descriptors is not None:
                    descriptors = [
                        descriptor
                        for descriptor in descriptors
                        if descriptor.Group.startswith(bus_prefix)
                        and "p_" in descriptor.Description
                    ]
                future = executor.submit(
                    _read_run_entries_process,
                    inf_path,
                    limit,
                    voltage_key,
                    bus_prefix,
                    exclusion_matcher,
                    time_step,
                    time_end,
                    fallback_frequency,
                    descriptors,
                    high_voltage_include_overrides,
                    sustained_settings,
                    sustained_limits,
                    sustained_frequency,
                )
            else:
                future = executor.submit(
                    _read_run_entries,
                    inf_path,
                    limit,
                    voltage_key,
                    bus_prefix,
                    exclusion_matcher,
                    time_step,
                    time_end,
                    fallback_frequency,
                    frequency_fallback,
                    check_cancel,
                    inf_descriptor_cache,
                    high_voltage_include_overrides=high_voltage_include_overrides,
                    sustained_settings=sustained_settings,
                    sustained_limits=sustained_limits,
                    sustained_frequency=sustained_frequency,
                )
            futures[future] = inf_path
        completed = len(run_data)
        if completed and not futures:
            _log(log, f"Envelope source files read: {bus_prefix} | {completed}/{len(inf_paths)}")
        for future in as_completed(futures):
            _cancel(check_cancel)
            inf_path = futures[future]
            try:
                result = future.result()
                if process_pool:
                    run_entries, run_exclusions, fallback_contexts = result
                    if frequency_fallback is not None:
                        for context in fallback_contexts:
                            frequency_fallback(context)
                else:
                    run_entries, run_exclusions = result
                    fallback_contexts = []
            except ValueError as exc:
                if "no results" not in str(exc).casefold():
                    raise
                run_entries, run_exclusions, fallback_contexts = [], [], []
            run_data[inf_path] = (run_entries, run_exclusions)
            completed += 1
            if completed == 1 or completed == len(inf_paths) or completed % 10 == 0:
                _log(log, f"Envelope source files read: {bus_prefix} | {completed}/{len(inf_paths)}")
    except Exception:
        executor.shutdown(wait=True, cancel_futures=True)
        raise
    else:
        if owns_executor:
            executor.shutdown(wait=True)

    return run_data


def _read_run_entries_process(
    inf_path: Path,
    limit: float,
    voltage_key: str,
    bus_prefix: str,
    exclusion_matcher: ExclusionMatcher,
    time_step: float,
    time_end: float,
    fallback_frequency: float,
    inf_descriptors: list[InfDescriptor] | None,
    high_voltage_include_overrides: set[tuple[str, str, int, str]] | None,
    sustained_settings: sustained_sdpf.SustainedSDPFSettings | None = None,
    sustained_limits: sustained_sdpf.SDPFVoltageLimits | None = None,
    sustained_frequency: float | None = None,
) -> tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]], list[str]]:
    fallback_contexts: list[str] = []
    entries, exclusions = _read_run_entries(
        inf_path,
        limit,
        voltage_key,
        bus_prefix,
        exclusion_matcher,
        time_step=time_step,
        time_end=time_end,
        fallback_frequency=fallback_frequency,
        frequency_fallback=fallback_contexts.append,
        inf_descriptor_cache=(
            {inf_path: inf_descriptors}
            if inf_descriptors is not None
            else None
        ),
        high_voltage_include_overrides=high_voltage_include_overrides,
        sustained_settings=sustained_settings,
        sustained_limits=sustained_limits,
        sustained_frequency=sustained_frequency,
    )
    return entries, exclusions, fallback_contexts


def _entries_for_inf_paths(
    inf_paths: list[Path],
    run_data: dict[Path, tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]]]],
) -> tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]], list[dict[str, Any]]]:
    entries: list[tuple[str, pd.DataFrame]] = []
    high_voltage: list[dict[str, Any]] = []
    sustained: list[dict[str, Any]] = []
    for inf_path in inf_paths:
        run_entries, run_exclusions = run_data.get(inf_path, ([], []))
        for measurement, value in run_entries:
            if measurement == SUSTAINED_ENTRY:
                sustained.extend(value)
            else:
                entries.append((measurement, value))
        high_voltage.extend(run_exclusions)
    return entries, high_voltage, sustained


def _read_run_entries(
    inf_path: Path,
    limit: float,
    voltage_key: str,
    bus_prefix: str,
    exclusion_matcher: ExclusionMatcher,
    time_step: float = TIME_STEP,
    time_end: float = TIME_END,
    fallback_frequency: float = DEFAULT_ENVELOPE_FALLBACK_FREQUENCY,
    frequency_fallback: FrequencyFallbackFn | None = None,
    check_cancel: CancelFn | None = None,
    inf_descriptor_cache: dict[Path, list[InfDescriptor]] | None = None,
    high_voltage_include_overrides: set[tuple[str, str, int, str]] | None = None,
    sustained_settings: sustained_sdpf.SustainedSDPFSettings | None = None,
    sustained_limits: sustained_sdpf.SDPFVoltageLimits | None = None,
    sustained_frequency: float | None = None,
) -> tuple[list[tuple[str, pd.DataFrame]], list[dict[str, Any]]]:
    _cancel(check_cancel)
    case_name, run = case_run_from_inf_path(inf_path)
    if inf_descriptor_cache is None or inf_path not in inf_descriptor_cache:
        df_desc = _parse_inf(
            inf_path,
            bus_prefix,
            voltage_key,
            case_name,
            run,
            exclusion_matcher,
        )
    else:
        df_desc = _filter_inf_descriptors(
            inf_descriptor_cache[inf_path],
            bus_prefix,
            voltage_key,
            case_name,
            run,
            exclusion_matcher,
        )
    _cancel(check_cancel)
    out_cache = _read_out_files(df_desc, inf_path, check_cancel)
    candidates: list[tuple[str, str, _BusWaveforms]] = []
    high_voltage: list[dict[str, Any]] = []

    for measurement in ("LGp", "LLp"):
        _cancel(check_cancel)
        df_meas = df_desc[df_desc["Description"].str.contains(measurement, na=False)]
        for bus in df_meas["Group"].unique():
            _cancel(check_cancel)
            df_bus = df_meas[df_meas["Group"] == bus]
            waveforms, exclusions = _read_bus_waveforms(
                df_bus,
                bus,
                out_cache,
                inf_path,
                limit,
                case_name,
                run,
                measurement,
                check_cancel,
            )
            if waveforms is not None:
                candidates.append((measurement, bus, waveforms))
            key = _high_voltage_key(voltage_key, case_name, run, bus)
            if key not in (high_voltage_include_overrides or set()):
                high_voltage.extend(exclusions)

    entries: list[tuple[str, pd.DataFrame]] = []
    high_voltage_buses = {str(record["MM_name"]) for record in high_voltage}
    sustained_rows: list[dict[str, Any]] = []
    for measurement, bus, waveforms in candidates:
        _cancel(check_cancel)
        if bus in high_voltage_buses:
            continue
        envelope_df = _apply_bus_waveforms(
            waveforms,
            time_step,
            time_end,
            fallback_frequency,
            frequency_fallback,
            check_cancel,
        )
        if envelope_df is None:
            continue
        envelope_df["MM_name"] = bus
        envelope_df["Case"] = f"{case_name}-{run}"
        entries.append((measurement, envelope_df))

        if sustained_settings is not None and sustained_settings.enabled and sustained_limits is not None:
            assessment_frequency = sustained_frequency or fallback_frequency
            phase_results = sustained_sdpf.analyze_bus_phase_results(
                measurement,
                waveforms.times,
                waveforms.values,
                waveforms.signals,
                sustained_limits,
                sustained_settings.effective_duration(assessment_frequency),
                assessment_frequency,
            )
            bus_result = sustained_sdpf.classify_rows(
                "",
                voltage_key,
                case_name,
                run,
                bus,
                phase_results,
                sustained_settings.effective_duration(sustained_frequency or fallback_frequency),
                cycle_coverage=sustained_settings.required_cycles(assessment_frequency),
            )
            if bus_result is not None:
                # Select the worst fixed phase in the worker so only compact
                # per-bus metadata crosses the process boundary.
                sustained_rows.append({"result": bus_result.to_dict()})

    if sustained_rows:
        entries.append((SUSTAINED_ENTRY, sustained_rows))

    return entries, high_voltage


def _parse_inf(
    inf_path: Path,
    bus_prefix: str,
    voltage_key: str = "",
    case_name: str = "",
    run: int = 0,
    exclusion_matcher: ExclusionMatcher | None = None,
) -> pd.DataFrame:
    try:
        descriptors = parse_inf_descriptors(inf_path)
    except ValueError as exc:
        raise ValueError("There are no results in your data") from exc
    return _filter_inf_descriptors(
        descriptors,
        bus_prefix,
        voltage_key,
        case_name,
        run,
        exclusion_matcher,
    )


def _filter_inf_descriptors(
    descriptors: list[InfDescriptor],
    bus_prefix: str,
    voltage_key: str = "",
    case_name: str = "",
    run: int = 0,
    exclusion_matcher: ExclusionMatcher | None = None,
) -> pd.DataFrame:
    matcher = exclusion_matcher or ExclusionMatcher()
    records = [
        {"PGB": descriptor.PGB, "Description": descriptor.Description, "Group": descriptor.Group}
        for descriptor in descriptors
        if descriptor.Group.startswith(bus_prefix)
        and "p_" in descriptor.Description
        and not matcher.excludes(voltage_key, case_name, run, descriptor.Group)
    ]
    df_desc = pd.DataFrame(
        records,
        index=[int(record["PGB"]) - 1 for record in records],
        columns=["PGB", "Description", "Group"],
    )
    if df_desc.empty:
        raise ValueError("There are no results in your data")
    return df_desc


def _read_out_files(
    df_desc: pd.DataFrame,
    inf_path: Path,
    check_cancel: CancelFn | None = None,
) -> dict[Path, OutFileData]:
    file_col_map: dict[Path, set[int]] = {}
    for phase_index in df_desc.index:
        _cancel(check_cancel)
        out_path, col_number = out_file_for_pgb(inf_path, int(phase_index) + 1)
        file_col_map.setdefault(out_path, set()).add(col_number)

    out_cache: dict[Path, OutFileData] = {}
    for out_path, cols in file_col_map.items():
        _cancel(check_cancel)
        read_cols = sorted([0, *cols])
        columns = load_out_columns(out_path, read_cols, check_cancel)
        time_values = np.round(columns[0], 6)
        out_cache[out_path] = (
            time_values,
            {
                column: columns[column]
                for column in read_cols[1:]
            },
        )
        _cancel(check_cancel)
    return out_cache


def _read_bus_waveforms(
    df_bus: pd.DataFrame,
    bus: str,
    out_cache: dict[Path, OutFileData],
    inf_path: Path,
    limit: float,
    case_name: str,
    run: int,
    measurement: str,
    check_cancel: CancelFn | None = None,
) -> tuple[_BusWaveforms | None, list[dict[str, Any]]]:
    times: list[np.ndarray] = []
    values: list[np.ndarray] = []
    absolute_values: list[np.ndarray] = []
    signals: list[str] = []
    exclusions: list[dict[str, Any]] = []
    for phase_index in df_bus.index:
        _cancel(check_cancel)
        signal = bus + "-" + str(df_bus.loc[phase_index, "Description"]).replace("MM_", "")
        out_path, col_number = out_file_for_pgb(inf_path, int(phase_index) + 1)
        time_values, columns = out_cache[out_path]
        phase_values = columns[col_number]
        phase_absolute_values = np.abs(phase_values)
        times.append(time_values)
        values.append(phase_values)
        absolute_values.append(phase_absolute_values)
        signals.append(signal)

        over_limit = phase_absolute_values > limit
        if over_limit.any():
            exclusions.append(
                {
                    "Case": case_name,
                    "Run": int(run),
                    "MM_name": bus,
                    "Measurement": measurement,
                    "Signal": signal,
                    "File": out_path.name,
                    "Excluded_values": int(over_limit.sum()),
                    "Max_abs": float(phase_absolute_values[over_limit].max()),
                    "Limit": limit,
                }
            )

    if not values:
        return None, exclusions
    context = f"{case_name} run {run} | {bus} | {measurement}"
    return _BusWaveforms(times, values, absolute_values, signals, context), exclusions


def _apply_bus_waveforms(
    waveforms: _BusWaveforms,
    time_step: float,
    time_end: float,
    fallback_frequency: float,
    frequency_fallback: FrequencyFallbackFn | None = None,
    check_cancel: CancelFn | None = None,
) -> pd.DataFrame | None:
    _cancel(check_cancel)
    if _arrays_share_timebase(waveforms.times):
        envelope = _apply_envelope_arrays(
            waveforms.times[0],
            waveforms.values,
            time_step,
            time_end,
            fallback_frequency,
            waveforms.context,
            frequency_fallback,
            absolute_phase_values=waveforms.absolute_values,
        )
        _cancel(check_cancel)
        return envelope
    envelope = _apply_envelope(
        _phase_dataframe_from_arrays(
            waveforms.times,
            waveforms.values,
            waveforms.signals,
        ),
        time_step,
        time_end,
        fallback_frequency,
        waveforms.context,
        frequency_fallback,
    )
    _cancel(check_cancel)
    return envelope


def _arrays_share_timebase(times: list[np.ndarray]) -> bool:
    first = times[0]
    return all(time_values.shape == first.shape and np.array_equal(time_values, first) for time_values in times[1:])


def _phase_dataframe_from_arrays(
    times: list[np.ndarray],
    values: list[np.ndarray],
    signals: list[str],
) -> pd.DataFrame:
    output = pd.DataFrame({"Time (s)": times[0], signals[0]: values[0]})
    for time_values, phase_values, signal in zip(times[1:], values[1:], signals[1:]):
        phase_df = pd.DataFrame({"Time (s)": time_values, signal: phase_values})
        output = pd.merge(output, phase_df, on="Time (s)", how="outer")
    return output


def _apply_envelope_arrays(
    time_values: np.ndarray,
    phase_values: list[np.ndarray],
    time_step: float,
    time_end: float,
    fallback_frequency: float = DEFAULT_ENVELOPE_FALLBACK_FREQUENCY,
    frequency_context: str = "",
    frequency_fallback: FrequencyFallbackFn | None = None,
    absolute_phase_values: list[np.ndarray] | None = None,
) -> pd.DataFrame | None:
    dt = round(float(np.median(np.diff(time_values))), 6)
    if not np.isfinite(dt) or dt <= 0:
        return None

    detected_freq = _detect_frequency_from_arrays(
        time_values,
        phase_values,
        fallback_frequency,
        frequency_context,
        frequency_fallback,
    )

    window_size = max(1, int((1 / (2 * detected_freq)) / dt))
    new_time = _bounded_time_grid([time_values], time_step, time_end)
    interpolated = {}
    for index, values in enumerate(phase_values):
        abs_values = (
            absolute_phase_values[index]
            if absolute_phase_values is not None
            else np.abs(values)
        )
        if np.nanmax(abs_values) < 1e-4:
            continue
        rolling_values = (
            pd.Series(abs_values, index=time_values)
            .rolling(window=window_size, min_periods=1, center=True)
            .max()
            .to_numpy(dtype=float)
        )
        interpolated["Max_" + chr(65 + index)] = np.interp(
            new_time,
            time_values,
            rolling_values,
            left=rolling_values[0],
            right=np.nan,
        )
    if not interpolated:
        return None
    output = pd.DataFrame(interpolated, index=new_time)
    output.index.name = "Time (s)"
    return output


def _apply_envelope(
    df: pd.DataFrame,
    time_step: float = TIME_STEP,
    time_end: float = TIME_END,
    fallback_frequency: float = DEFAULT_ENVELOPE_FALLBACK_FREQUENCY,
    frequency_context: str = "",
    frequency_fallback: FrequencyFallbackFn | None = None,
) -> pd.DataFrame | None:
    df = df.copy()
    df.set_index("Time (s)", inplace=True)
    dt = round(df.index.to_series().diff().median(), 6)
    if not np.isfinite(dt) or dt <= 0:
        return None

    phase_cols = list(df.columns)
    max_phase_cols = {col: "Max_" + chr(65 + i) for i, col in enumerate(phase_cols)}
    detected_freq = _detect_frequency_from_arrays(
        df.index.to_numpy(dtype=float),
        [df[col].to_numpy(dtype=float) for col in phase_cols],
        fallback_frequency,
        frequency_context,
        frequency_fallback,
    )

    window_size = max(1, int((1 / (2 * detected_freq)) / dt))
    drop_cols = []
    for col in phase_cols:
        maximum = df[col].abs().max()
        if not np.isfinite(maximum) or maximum < 1e-4:
            drop_cols.append(col)
            continue
        df.loc[:, col] = df[col].abs().rolling(window=window_size, min_periods=1, center=True).max()
    df.drop(columns=drop_cols, inplace=True)
    if df.empty:
        return None

    valid_times = [
        df.index.to_numpy(dtype=float)[np.isfinite(df[col].to_numpy(dtype=float))]
        for col in df.columns
    ]
    new_time = _bounded_time_grid(valid_times, time_step, time_end)
    interpolated = {}
    for col in df.columns:
        x = df.index.to_numpy(dtype=float)
        y = df[col].to_numpy(dtype=float)
        finite = np.isfinite(x) & np.isfinite(y)
        valid_x = x[finite]
        valid_y = y[finite]
        interpolated[max_phase_cols[col]] = np.interp(
            new_time,
            valid_x,
            valid_y,
            left=valid_y[0],
            right=np.nan,
        )
    output = pd.DataFrame(interpolated, index=new_time)
    output.index.name = "Time (s)"
    return output


def _bounded_time_grid(
    time_arrays: Iterable[np.ndarray],
    time_step: float,
    requested_end: float,
) -> np.ndarray:
    available_ends = [
        float(np.max(finite))
        for values in time_arrays
        if len(finite := np.asarray(values, dtype=float)[np.isfinite(values)])
    ]
    if not available_ends:
        return np.array([], dtype=float)
    effective_end = min(float(requested_end), min(available_ends))
    return np.round(np.arange(0.0, max(0.0, effective_end), time_step), 6)


def _detect_frequency_from_arrays(
    time_values: np.ndarray,
    phase_values: list[np.ndarray],
    fallback_frequency: float,
    frequency_context: str = "",
    frequency_fallback: FrequencyFallbackFn | None = None,
) -> float:
    for values in phase_values:
        if np.nanmax(np.abs(values)) < 1e-4:
            continue
        zero_crossings = np.where(np.diff(np.sign(values)))[0]
        if len(zero_crossings) < 4:
            continue
        detected = round(1 / (2 * np.diff(time_values[zero_crossings[:4]]).mean()), 0)
        if np.isfinite(detected) and detected > 0:
            return float(detected)
    if frequency_fallback is not None:
        frequency_fallback(frequency_context)
    return fallback_frequency


def _final_envelope(
    source_dfs: list[pd.DataFrame],
    df_stat: pd.DataFrame,
    time_step: float = TIME_STEP,
) -> pd.DataFrame:
    if not source_dfs:
        return pd.DataFrame(columns=ENVELOPE_HEADERS)
    merged = _reduce_merge(source_dfs, time_step)
    merged = _map_fault_types_phase(merged, df_stat)
    return _organize_phase_columns(merged)


def _reduce_merge(source_dfs: list[pd.DataFrame], time_step: float = TIME_STEP) -> pd.DataFrame:
    phase_suffixes = list(
        dict.fromkeys(
            suffix
            for source_df in source_dfs
            for suffix in _phase_suffixes(source_df)
        )
    )
    row_count = max(len(source_df) for source_df in source_dfs)
    output = pd.DataFrame({"Time (s)": np.round(np.arange(row_count) * time_step, 6)})
    run_columns: dict[str, pd.Series] = {}

    for suffix in phase_suffixes:
        max_col = f"Max_{suffix}"
        max_arrays: list[np.ndarray] = []
        mm_arrays: list[np.ndarray] = []
        case_arrays: list[np.ndarray] = []
        for source_df in source_dfs:
            if max_col not in source_df.columns:
                continue
            values = source_df[max_col].to_numpy(dtype=float)
            order = np.argsort(np.where(np.isnan(values), -np.inf, values), kind="stable")[::-1]
            max_arrays.append(values[order])
            mm_arrays.append(source_df["MM_name"].to_numpy(dtype=object)[order])
            case_arrays.append(source_df["Case"].to_numpy(dtype=object)[order])

        max_matrix = np.full((row_count, len(max_arrays)), np.nan, dtype=float)
        mm_matrix = np.full((row_count, len(mm_arrays)), np.nan, dtype=object)
        case_matrix = np.full((row_count, len(case_arrays)), np.nan, dtype=object)
        for index, values in enumerate(max_arrays):
            max_matrix[: len(values), index] = values
            mm_matrix[: len(values), index] = mm_arrays[index]
            case_matrix[: len(values), index] = case_arrays[index]

        comparable = np.where(np.isnan(max_matrix), -np.inf, max_matrix)
        source_indices = comparable.argmax(axis=1)
        row_indices = np.arange(row_count)
        max_values = comparable[row_indices, source_indices]
        max_values[max_values == -np.inf] = np.nan
        output[max_col] = np.round(max_values, 1)
        output[f"MM_name_{suffix}"] = mm_matrix[row_indices, source_indices]
        selected_cases = pd.Series(case_matrix[row_indices, source_indices], dtype=object)
        parts = selected_cases.astype(str).str.split("-", n=1, expand=True)
        output[f"Case_{suffix}"] = parts[0].replace("nan", np.nan)
        run_columns[f"Run_{suffix}"] = pd.to_numeric(parts[1] if parts.shape[1] > 1 else np.nan, errors="coerce")
    for column, values in run_columns.items():
        output[column] = values
    return output


def _phase_suffixes(df: pd.DataFrame) -> list[str]:
    return list(dict.fromkeys(col.replace("Max_", "") for col in df.columns if str(col).startswith("Max_")))


def _map_fault_types_phase(df: pd.DataFrame, df_stat: pd.DataFrame) -> pd.DataFrame:
    if df_stat.empty:
        for suffix in _phase_suffixes(df):
            df[f"Fault_type_{suffix}"] = ""
        return df
    for suffix in _phase_suffixes(df):
        case_col = f"Case_{suffix}"
        run_col = f"Run_{suffix}"
        fault_col = f"Fault_type_{suffix}"
        df_stat_phase = df_stat[["Case", "Run#", "Fault_type"]].copy()
        df_stat_phase.columns = [case_col, run_col, fault_col]
        df = df.merge(df_stat_phase, left_on=[case_col, run_col], right_on=[case_col, run_col], how="left")
    return df


def _organize_phase_columns(df: pd.DataFrame) -> pd.DataFrame:
    for col in ENVELOPE_HEADERS:
        if col not in df.columns:
            df[col] = np.nan
    max_cols = ["Max_A", "Max_B", "Max_C"]
    df["Max_all"] = df[max_cols].max(axis=1, skipna=True)
    max_idx = df[max_cols].fillna(-np.inf).to_numpy().argmax(axis=1)
    row_indices = np.arange(len(df))
    for target, sources in (
        ("Case_all", ["Case_A", "Case_B", "Case_C"]),
        ("Run_all", ["Run_A", "Run_B", "Run_C"]),
        ("Fault_type_all", ["Fault_type_A", "Fault_type_B", "Fault_type_C"]),
        ("MM_name_all", ["MM_name_A", "MM_name_B", "MM_name_C"]),
    ):
        df[target] = df[sources].to_numpy(dtype=object)[row_indices, max_idx]
    return df[ENVELOPE_HEADERS]


def _statistic_files(project_root: Path) -> list[Path]:
    return sorted((project_root / "Case_folder").rglob("Statistic*.out"))


def _read_stat_file(path: Path) -> pd.DataFrame:
    case_name = path.parent.name.split(".")[0]
    raw = pd.read_csv(path, skiprows=[0], header=None, sep=" +Output+ ", engine="python")
    raw.reset_index(drop=True, inplace=True)
    raw[0] = raw[0].str.replace("Run #", "Run#", regex=False)
    summary_rows = raw[raw[0].str.startswith("Statistical Summary", na=False)].index
    last_row = int(summary_rows[0]) - 1 if len(summary_rows) else len(raw) - 1
    split = raw.iloc[0:last_row + 1, 0].str.split(" +", expand=True)
    if split.empty:
        return pd.DataFrame(columns=["Case", "Run#", "Fault_type"])
    df_stat = pd.DataFrame(split.values[1:], columns=split.iloc[0])
    df_stat = df_stat.apply(pd.to_numeric, errors="coerce")
    df_stat.dropna(subset=["Run#"], inplace=True)
    df_stat.insert(0, "Case", case_name)
    return _map_faults(df_stat)


def _read_stat_files(project_root: Path) -> pd.DataFrame:
    all_dfs = [_read_stat_file(path) for path in _statistic_files(project_root)]
    if not all_dfs:
        return pd.DataFrame(columns=["Case", "Run#", "Fault_type"])
    return pd.concat(all_dfs, ignore_index=True)


def read_project_fault_types(project_root: str | Path) -> dict[int, str]:
    """Read the project-wide Run -> Fault_type map from one Statistic file.

    PSCAD projects in this application use the same run/fault schedule in
    every case.  The first valid statistic file is therefore sufficient for
    project metadata; envelope calculations retain their existing full-table
    reader where case-specific rows are required.
    """
    for path in _statistic_files(Path(project_root)):
        try:
            frame = _read_stat_file(path)
        except (OSError, UnicodeError, KeyError, ValueError, pd.errors.ParserError):
            continue
        if frame.empty or not {"Run#", "Fault_type"} <= set(frame.columns):
            continue
        result: dict[int, str] = {}
        for row in frame[["Run#", "Fault_type"]].itertuples(index=False):
            try:
                run = int(row[0])
            except (TypeError, ValueError):
                continue
            value = "" if row[1] is None else str(row[1]).strip()
            if value and value.casefold() != "nan":
                result[run] = value
        if result:
            return result
    return {}


def _map_faults(df: pd.DataFrame) -> pd.DataFrame:
    replacement_dict = {0.0: "None", 11.0: "ABC", 1.0: "AG", 4.0: "ABG", 7.0: "ABCG", 8.0: "AB"}
    df = df.copy()
    if "Fault_type" in df.columns:
        df["Fault_type"] = df["Fault_type"].replace(replacement_dict, regex=True)
    return df


def _nonconv_dataframe(records: Iterable[Any]) -> pd.DataFrame:
    """Convert the scanner's cached NonConv records without rereading files."""
    rows: list[dict[str, Any]] = []
    for record in records:
        if isinstance(record, dict):
            row = record
        else:
            row = {
                column.casefold(): getattr(record, column.casefold(), "")
                for column in NONCONV_COLUMNS
            }
        rows.append(
            {
                column: row.get(column, row.get(column.casefold(), ""))
                for column in NONCONV_COLUMNS
            }
        )
    if not rows:
        return pd.DataFrame(columns=NONCONV_COLUMNS)
    return (
        pd.DataFrame(rows, columns=NONCONV_COLUMNS)
        .drop_duplicates()
        .sort_values(by=["Case", "Run", "Source", "Signal"])
        .reset_index(drop=True)
    )


def _find_non_convergent_cases(
    case_root: Path,
    df_stat: pd.DataFrame,
    cb_iip_limit: float = NONCONV_CB_IIP_LIMIT,
    cb_iir_limit: float = NONCONV_CB_IIR_LIMIT,
    *,
    warnings: list[str] | None = None,
    worker_count: int | None = None,
    check_cancel: CancelFn | None = None,
    fault_types_by_run: dict[int, str] | None = None,
) -> pd.DataFrame:
    paths = sorted(case_root.rglob("CB_*.out"))
    records: list[dict[str, Any]] = []
    results: dict[Path, tuple[list[dict[str, Any]], str | None]] = {}
    if worker_count is not None and worker_count > 1 and len(paths) > 1:
        workers = _worker_count(len(paths), worker_count)
        with ProcessPoolExecutor(max_workers=workers) as executor:
            futures = {
                executor.submit(_read_non_convergent_file, path, cb_iip_limit, cb_iir_limit): path
                for path in paths
            }
            for future in as_completed(futures):
                _cancel(check_cancel)
                results[futures[future]] = future.result()
    else:
        for path in paths:
            _cancel(check_cancel)
            results[path] = _read_non_convergent_file(path, cb_iip_limit, cb_iir_limit)

    for path in paths:
        _cancel(check_cancel)
        file_records, warning = results[path]
        records.extend(file_records)
        if warning and warnings is not None:
            warnings.append(warning)
    if not records:
        return pd.DataFrame(columns=NONCONV_COLUMNS)
    df_nonconv = pd.DataFrame(records).drop_duplicates()
    if fault_types_by_run is None:
        df_nonconv = _map_exclusion_fault_types(df_nonconv, df_stat, NONCONV_COLUMNS)
    else:
        df_nonconv["Fault_type"] = (
            pd.to_numeric(df_nonconv["Run"], errors="coerce")
            .map(fault_types_by_run)
            .fillna("")
        )
        for column in NONCONV_COLUMNS:
            if column not in df_nonconv.columns:
                df_nonconv[column] = ""
    return df_nonconv[NONCONV_COLUMNS].sort_values(by=["Case", "Run", "Source", "Signal"]).reset_index(drop=True)


def find_non_convergent_cases_for_project(
    project_root: str | Path,
    cb_iip_limit: float | None = None,
    cb_iir_limit: float | None = None,
    fault_types_by_run: dict[int, str] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    root = Path(project_root)
    case_root = root / "Case_folder"
    if not case_root.is_dir():
        return [], []
    warnings: list[str] = []
    df_stat = (
        _read_stat_files(root)
        if fault_types_by_run is None
        else pd.DataFrame(columns=["Case", "Run#", "Fault_type"])
    )
    df_nonconv = _find_non_convergent_cases(
        case_root,
        df_stat,
        normalize_positive_float(cb_iip_limit, NONCONV_CB_IIP_LIMIT),
        normalize_positive_float(cb_iir_limit, NONCONV_CB_IIR_LIMIT),
        warnings=warnings,
        fault_types_by_run=fault_types_by_run,
    )
    records = df_nonconv.astype(object).where(pd.notna(df_nonconv), "").to_dict("records")
    return records, warnings


def _read_non_convergent_file(
    path: Path,
    cb_iip_limit: float = NONCONV_CB_IIP_LIMIT,
    cb_iir_limit: float = NONCONV_CB_IIR_LIMIT,
) -> tuple[list[dict[str, Any]], str | None]:
    case_name = path.parent.name.split(".")[0]
    source = path.stem.replace("_01_0001", "")
    records: list[dict[str, Any]] = []
    df_summary, warning = _read_cb_summary(path)
    if df_summary.empty or "Run#" not in df_summary.columns:
        return records, warning

    current_cols = [col for col in df_summary.columns if "II" in str(col)]
    for _, row in df_summary.iterrows():
        run = pd.to_numeric(row["Run#"], errors="coerce")
        if pd.isna(run):
            continue
        for signal in current_cols:
            value_raw = row[signal]
            value = pd.to_numeric(value_raw, errors="coerce")
            reason = None
            if pd.isna(value):
                reason = "Non-numeric or NaN current"
            elif not np.isfinite(value):
                reason = "Non-finite current"
            elif signal == "CB_IIp" and abs(value) > cb_iip_limit:
                reason = f"CB_IIp above {cb_iip_limit:g}"
            elif signal == "CB_IIr" and abs(value) > cb_iir_limit:
                reason = f"CB_IIr above {cb_iir_limit:g}"
            if reason:
                records.append({"Case": case_name, "Run": int(run), "Source": source, "Signal": signal, "Reason": reason, "Value": value_raw, "File": path.name})
    return records, warning


def _read_cb_summary(path: Path) -> tuple[pd.DataFrame, str | None]:
    failures: list[str] = []
    for skiprows in ([0], None):
        try:
            df = pd.read_csv(path, skiprows=skiprows, header=0, sep=r"\s+", engine="python", dtype=str)
        except (OSError, UnicodeError, pd.errors.ParserError, ValueError) as exc:
            failures.append(str(exc))
            continue
        if "Run#" in df.columns:
            return df, None
        failures.append("missing Run# column")
    reason = "; ".join(dict.fromkeys(failures)) or "unsupported file structure"
    return pd.DataFrame(), f"Skipped unreadable CB summary {path.name}: {reason}"


def _nonconv_keys(df_nonconv: pd.DataFrame) -> set[tuple[str, int]]:
    if df_nonconv.empty:
        return set()
    return set(zip(df_nonconv["Case"], df_nonconv["Run"].astype(int)))


def _high_voltage_exclusion_dataframe(
    detected_rows: list[dict[str, Any]],
    proposals: list[dict[str, object]],
    voltage_key: str,
    inf_paths: list[Path],
    df_stat: pd.DataFrame,
    available_keys: set[tuple[str, str, int, str]] | None = None,
) -> pd.DataFrame:
    """Build the original detailed HV sheet and retain applied log-only rows."""
    actual_rows = [dict(row) for row in detected_rows]
    detected_keys = {
        key
        for row in detected_rows
        if (
            key := _high_voltage_key(
                voltage_key,
                row.get("Case", ""),
                row.get("Run", ""),
                row.get("MM_name", ""),
            )
        )
        is not None
    }
    scope_runs = {
        (case.casefold(), run)
        for inf_path in inf_paths
        for case, run in [case_run_from_inf_path(inf_path)]
    }
    normalized_voltage = normalize_voltage(voltage_key)
    for proposal in proposals:
        key = _high_voltage_key(
            proposal.get("Voltage", ""),
            proposal.get("Case", ""),
            proposal.get("Run", ""),
            proposal.get("MM_name", ""),
        )
        if (
            key is None
            or key[0] != normalized_voltage
            or (key[1], key[2]) not in scope_runs
            or (available_keys is not None and key not in available_keys)
            or not bool(proposal.get("Applied", False))
            or key in detected_keys
        ):
            continue
        actual_rows.append(
            {
                column: proposal.get(column, "")
                for column in HIGH_VOLTAGE_COLUMNS
                if column != "Fault_type"
            }
        )
        detected_keys.add(key)
    return _map_exclusion_fault_types(
        pd.DataFrame(actual_rows),
        df_stat,
        HIGH_VOLTAGE_COLUMNS,
    )


def _map_exclusion_fault_types(df: pd.DataFrame, df_stat: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=columns)
    df = df.copy()
    df["Run"] = pd.to_numeric(df["Run"], errors="coerce")
    if not df_stat.empty:
        df = df.merge(df_stat[["Case", "Run#", "Fault_type"]], left_on=["Case", "Run"], right_on=["Case", "Run#"], how="left")
        df.drop(columns=["Run#"], inplace=True)
    elif "Fault_type" not in df.columns:
        df["Fault_type"] = ""
    df["Run"] = df["Run"].astype("Int64")
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    return df[columns]


def _filter_dataframe_by_scope(df: pd.DataFrame, scope: ScopeEntry) -> pd.DataFrame:
    if df.empty or scope.mode == "full":
        return df
    tokens = [token.casefold() for token in scope.tokens]
    row_text = df.astype(str).agg(" ".join, axis=1).str.casefold()
    matched = row_text.apply(lambda value: any(token in value for token in tokens))
    return df[matched] if scope.mode == "include" else df[~matched]


def _write_workbook(
    path: Path,
    lg_df: pd.DataFrame,
    ll_df: pd.DataFrame,
    nonconv_df: pd.DataFrame,
    high_voltage_df: pd.DataFrame,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        lg_df.to_excel(writer, index=False, sheet_name="LGp")
        ll_df.to_excel(writer, index=False, sheet_name="LLp")
        nonconv_df.to_excel(writer, index=False, sheet_name="NonConv cases")
        high_voltage_df.to_excel(writer, index=False, sheet_name="High voltage exclusions")
        _format_workbook(writer.book)


def _format_workbook(workbook: Any) -> None:
    for ws in workbook.worksheets:
        if ws.max_column:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(1, col_idx)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")
