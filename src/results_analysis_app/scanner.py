from __future__ import annotations

from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, replace
from fnmatch import fnmatchcase
from pathlib import Path
import math
import os
import posixpath
import re
from typing import Any, Callable, Iterable
from zipfile import BadZipFile, ZipFile
import xml.etree.ElementTree as ET

import numpy as np

from results_analysis_app import sustained_sdpf
from results_analysis_app.models import DashboardFigure, ScopeEntry
from results_analysis_app.project_config import (
    VoltageConfig,
    discover_voltage_data,
    input_data_workbook,
    load_project_timing,
    load_voltage_configs,
    normalize_voltage,
)
from pscad_plotter_app_v3.services.waveform_io import (
    load_out_columns,
    out_file_for_pgb,
    parse_inf_descriptors,
)
from pscad_plotter_app_v3.services.project_conventions import case_run_from_inf_path


LEGACY_PSCAD_LOG_HIGH_VOLTAGE_SOURCE = "PSCAD_log.txt + MM results.csv"
PSCAD_LOG_HIGH_VOLTAGE_SOURCE = "PSCAD_log.txt + raw waveforms"
PSCAD_LOG_HIGH_VOLTAGE_SOURCES = {
    LEGACY_PSCAD_LOG_HIGH_VOLTAGE_SOURCE,
    PSCAD_LOG_HIGH_VOLTAGE_SOURCE,
}
HIGH_VOLTAGE_SOURCE_ANALYSIS = "Analysis"
HIGH_VOLTAGE_SOURCE_PSCAD_LOG = "PSCAD log"
HIGH_VOLTAGE_SOURCE_BOTH = "Both"


OOXML_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "office": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


@dataclass(frozen=True)
class CaseInfo:
    name: str
    inf_path: Path

    @property
    def match_text(self) -> str:
        return f"{self.name} {self.inf_path}".casefold()


@dataclass(frozen=True)
class NonConvergentCase:
    case: str
    run: int
    fault_type: str = ""
    source: str = ""
    signal: str = ""
    reason: str = ""
    value: str = ""
    file: str = ""


@dataclass(frozen=True)
class HighVoltageExclusion:
    voltage: str
    case: str
    run: int
    bus: str
    fault_type: str = ""
    measurement: str = ""
    signal: str = ""
    file: str = ""
    excluded_values: str = ""
    max_abs: str = ""
    limit: str = ""
    source: str = ""
    excluded: bool = False


@dataclass(frozen=True)
class HighVoltageMeasurement:
    voltage: str
    case: str
    run: int
    bus: str
    measurement: str
    signal: str
    file: str
    max_abs: float


@dataclass
class ProjectScan:
    path: Path
    exists: bool
    chips: list[str] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)
    case_infos: list[CaseInfo] = field(default_factory=list)
    nonconv_cases: list[NonConvergentCase] = field(default_factory=list)
    high_voltage_exclusions: list[HighVoltageExclusion] = field(default_factory=list)
    high_voltage_log_measurements: list[HighVoltageMeasurement] = field(default_factory=list)
    fault_types_by_run: dict[int, str] = field(default_factory=dict)
    available_voltages: list[str] = field(default_factory=list)
    voltage_configs: dict[str, VoltageConfig] = field(default_factory=dict)
    sustained_sdpf_limits: dict[str, sustained_sdpf.SDPFVoltageLimits] = field(default_factory=dict)
    sustained_sdpf_limit_warnings: list[str] = field(default_factory=list)
    project_frequency: float | None = None
    final_duration: float | None = None
    has_dashboards: bool = False
    dashboard_figures: list[DashboardFigure] = field(default_factory=list)
    dashboard_changed: bool = False
    has_envelopes: bool = False
    has_plots: bool = False
    has_reports: bool = False


@dataclass(frozen=True)
class ScopePreview:
    scope_folder: str
    matched_count: int
    total_count: int
    examples: tuple[str, ...]
    warning: str = ""


def _non_temp_files(root: Path, pattern: str) -> list[Path]:
    if not root.is_dir():
        return []
    return sorted(path for path in root.rglob(pattern) if path.is_file() and not path.name.startswith("~$"))


def _has_non_temp_file(root: Path, pattern: str) -> bool:
    if not root.is_dir():
        return False
    return any(path.is_file() and not path.name.startswith("~$") for path in root.rglob(pattern))


def _has_non_temp_files(root: Path, patterns: tuple[str, ...]) -> bool:
    """Check several extensions with one directory walk."""
    if not root.is_dir():
        return False
    folded = tuple(pattern.casefold() for pattern in patterns)
    for directory, _subdirs, names in os.walk(root):
        for name in names:
            if name.startswith("~$"):
                continue
            if any(fnmatchcase(name.casefold(), pattern) for pattern in folded):
                return True
    return False


def _has_dashboard_files(root: Path) -> bool:
    return _has_non_temp_files(root, ("*.xlsx", "*.xlsm", "*.xlsb"))


def _has_plot_files(root: Path) -> bool:
    return _has_non_temp_files(root, ("*.png", "*.jpg", "*.jpeg", "*.emf"))


def _file_metadata(path: Path, root: Path) -> dict[str, int | str] | None:
    try:
        stat = path.stat()
        relative_path = path.relative_to(root).as_posix()
    except (OSError, ValueError):
        return None
    return {
        "path": relative_path,
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


def _tree_metadata(
    root: Path,
    scan_root: Path,
    include_name: Callable[[str], bool],
) -> list[dict[str, int | str]]:
    if not scan_root.is_dir():
        return []
    files: list[dict[str, int | str]] = []
    for directory, _subdirs, names in os.walk(scan_root):
        for name in names:
            if name.startswith("~$") or not include_name(name):
                continue
            metadata = _file_metadata(Path(directory) / name, root)
            if metadata is not None:
                files.append(metadata)
    return sorted(files, key=lambda item: str(item["path"]))


def project_output_manifest(project_root: str | Path) -> dict[str, Any]:
    """Return output/dashboard metadata without walking raw PSCAD run files."""
    root = Path(project_root).resolve()
    if not root.is_dir():
        return {
            "dashboard_files": [],
            "envelope_files": [],
            "output_state": {
                "has_results": False,
                "has_dashboards": False,
                "has_envelopes": False,
                "has_plots": False,
                "has_reports": False,
            },
        }
    dashboard_files = _tree_metadata(
        root,
        root / "Dashboards",
        lambda name: name.casefold().endswith((".xlsx", ".xlsm", ".xlsb")),
    )
    envelope_files = _tree_metadata(
        root,
        root / "Voltage_envelope",
        lambda name: name.casefold().endswith(".xlsx"),
    )
    return {
        "dashboard_files": dashboard_files,
        "envelope_files": envelope_files,
        "output_state": {
            "has_results": _has_non_temp_file(root / "Results", "*.csv"),
            "has_dashboards": bool(dashboard_files),
            "has_envelopes": bool(envelope_files),
            "has_plots": _has_plot_files(root / "Plots" / "Generated"),
            "has_reports": _has_non_temp_file(root / "Reports", "*.docx"),
        },
    }


def project_scan_manifest(project_root: str | Path) -> dict[str, Any]:
    """Return metadata for files that can change an opening project scan."""
    root = Path(project_root).resolve()
    manifest: dict[str, Any] = {
        "exists": root.is_dir(),
        "files": [],
        "high_voltage_files": [],
        **project_output_manifest(root),
    }
    if not root.is_dir():
        return manifest

    files: list[dict[str, int | str]] = []
    high_voltage_files: list[dict[str, int | str]] = []

    def add_file(path: Path, target: list[dict[str, int | str]] = files) -> None:
        metadata = _file_metadata(path, root)
        if metadata is not None:
            target.append(metadata)

    log_path = _pscad_log_path(root)
    log_cases: set[str] = set()
    if log_path is not None:
        add_file(log_path, high_voltage_files)
        try:
            log_cases = {case.casefold() for case, _bus in _pscad_log_high_voltage_case_buses(log_path)}
        except OSError:
            pass

    case_root = root / "Case_folder"
    if case_root.is_dir():
        # `project_scan_manifest` is called for every cached project opening.
        # DirEntry.stat() avoids constructing a Path and issuing a second path
        # lookup for each raw file while retaining the same inclusion rules as
        # the former os.walk implementation.
        pending_directories = [case_root]
        while pending_directories:
            directory_path = pending_directories.pop()
            try:
                with os.scandir(directory_path) as entries:
                    directory_entries = []
                    for entry in entries:
                        try:
                            if entry.is_dir(follow_symlinks=False):
                                pending_directories.append(Path(entry.path))
                                continue
                            # os.walk places symlinked directories in its
                            # directory list even when it does not follow them.
                            if entry.is_symlink() and entry.is_dir():
                                continue
                        except OSError:
                            continue
                        directory_entries.append(entry)
            except OSError:
                continue

            inf_entries = [
                entry
                for entry in directory_entries
                if not entry.name.startswith("~$")
                and entry.name.casefold().endswith(".inf")
            ]
            matching_stems: list[str] = []

            def add_entry(entry, target=files) -> None:
                try:
                    stat = entry.stat()
                    relative_path = Path(entry.path).relative_to(root).as_posix()
                except (OSError, ValueError):
                    return
                target.append(
                    {
                        "path": relative_path,
                        "size": int(stat.st_size),
                        "mtime_ns": int(stat.st_mtime_ns),
                    }
                )

            for entry in inf_entries:
                add_entry(entry)
                path = Path(entry.path)
                if log_cases and _case_name_from_inf(path).casefold() in log_cases:
                    matching_stems.append(path.stem.casefold())

            for entry in directory_entries:
                name = entry.name
                if name.startswith("~$") or not name.casefold().endswith(".out"):
                    continue
                folded = name.casefold()
                if folded.startswith("statistic") or folded.startswith("cb_"):
                    add_entry(entry)
                elif any(folded.startswith(f"{stem}_") for stem in matching_stems):
                    add_entry(entry, high_voltage_files)
    for path in root.glob("Input_Data_PSCAD*.xlsx"):
        if not path.name.startswith("~$"):
            add_file(path)
    add_file(root / "Plots" / ".plottool_v3" / "limits.json")

    manifest["files"] = sorted(files, key=lambda item: str(item["path"]))
    manifest["high_voltage_files"] = sorted(
        high_voltage_files,
        key=lambda item: str(item["path"]),
    )
    return manifest


def _case_name_from_inf(path: Path) -> str:
    try:
        return case_run_from_inf_path(path)[0]
    except ValueError:
        return path.stem


def _case_infos_from_inf_paths(inf_paths: list[Path]) -> list[CaseInfo]:
    cases: dict[str, CaseInfo] = {}
    for inf_path in inf_paths:
        name = _case_name_from_inf(inf_path)
        cases.setdefault(name.casefold(), CaseInfo(name=name, inf_path=inf_path))
    return sorted(cases.values(), key=lambda item: item.name.casefold())


def _collect_case_infos_and_inf_paths(project_root: Path) -> tuple[list[CaseInfo], list[Path]]:
    case_root = project_root / "Case_folder"
    inf_paths = _non_temp_files(case_root, "*.inf")
    return _case_infos_from_inf_paths(inf_paths), inf_paths


def _as_text(value: Any) -> str:
    return "" if value is None else str(value)


def _as_int(value: Any) -> int | None:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def _pscad_log_path(project_root: Path) -> Path | None:
    preferred = project_root / "PSCAD_log.txt"
    if preferred.is_file():
        return preferred
    matches = sorted(path for path in project_root.glob("*log*.txt") if path.is_file())
    return matches[0] if matches else None


def _pscad_log_high_voltage_case_buses(log_path: Path) -> set[tuple[str, str]]:
    pattern = re.compile(
        r"WARNING:\s+(.*?)\s+-\s+(MM_\S+)\s+has very high values",
        flags=re.IGNORECASE,
    )
    case_buses: set[tuple[str, str]] = set()
    for line in log_path.read_text(encoding="utf-8", errors="ignore").splitlines():
        match = pattern.search(line)
        if match:
            case_buses.add((match.group(1).strip(), match.group(2).strip()))
    return case_buses


def _voltage_from_mm_bus(bus: str) -> str:
    match = re.match(r"MM_(\d+(?:\.\d+)?)(?:_|$)", bus.strip(), flags=re.IGNORECASE)
    return normalize_voltage(match.group(1)) if match else ""


def _raw_high_voltage_measurements_for_inf(
    inf_path: Path,
    buses: set[str],
    voltage_by_bus: dict[str, str],
    check_cancel: Callable[[], None] | None = None,
) -> tuple[list[HighVoltageMeasurement], list[str]]:
    if check_cancel is not None:
        check_cancel()
    try:
        case, run = case_run_from_inf_path(inf_path)
    except ValueError:
        return [], []
    channels_by_file: dict[Path, list[tuple[int, str, str, str]]] = defaultdict(list)
    try:
        descriptors = parse_inf_descriptors(inf_path)
    except OSError as exc:
        return [], [f"Could not read INF file: {inf_path.name} | {exc}"]
    except ValueError as exc:
        return [], [f"Could not parse INF file: {inf_path.name} | {exc}"]

    for descriptor in descriptors:
        if check_cancel is not None:
            check_cancel()
        description = descriptor.Description.strip()
        bus = descriptor.Group.strip()
        if bus not in buses:
            continue
        if "LGp" in description:
            measurement = "LGp"
        elif "LLp" in description:
            measurement = "LLp"
        else:
            continue
        out_path, column = out_file_for_pgb(inf_path, descriptor.PGB)
        signal = f"{bus}-{description.replace('MM_', '')}"
        channels_by_file[out_path].append((column, bus, measurement, signal))

    records: list[HighVoltageMeasurement] = []
    warnings: list[str] = []
    for out_path, channels in channels_by_file.items():
        if check_cancel is not None:
            check_cancel()
        columns = sorted({column for column, _bus, _measurement, _signal in channels})
        try:
            values_by_column = load_out_columns(out_path, columns, check_cancel)
        except OSError as exc:
            warnings.append(f"Could not read output file: {out_path.name} | {exc}")
            continue
        except ValueError as exc:
            warnings.append(f"Could not parse output file: {out_path.name} | {exc}")
            continue

        for column, bus, measurement, signal in channels:
            if check_cancel is not None:
                check_cancel()
            series = values_by_column.get(column)
            if series is None:
                continue
            abs_values = np.abs(series[np.isfinite(series)])
            if abs_values.size == 0:
                continue
            records.append(
                HighVoltageMeasurement(
                    voltage=voltage_by_bus[bus],
                    case=case,
                    run=int(run),
                    bus=bus,
                    measurement=measurement,
                    signal=signal,
                    file=out_path.name,
                    max_abs=float(abs_values.max()),
                )
            )
    maxima: dict[str, HighVoltageMeasurement] = {}
    for item in records:
        current = maxima.get(item.bus)
        if current is None or item.max_abs > current.max_abs:
            maxima[item.bus] = item
    return [maxima[bus] for bus in sorted(maxima)], warnings


def _collect_nonconv_cases(
    project_root: Path,
    nonconv_cb_iip_limit: float | None = None,
    nonconv_cb_iir_limit: float | None = None,
    fault_types_future=None,
) -> tuple[list[NonConvergentCase], list[str]]:
    from results_analysis_app.voltage_envelope import find_non_convergent_cases_for_project

    try:
        fault_types_by_run = fault_types_future.result() if fault_types_future is not None else None
    except Exception:
        fault_types_by_run = None
    rows = []
    records, warnings = find_non_convergent_cases_for_project(
        project_root,
        nonconv_cb_iip_limit,
        nonconv_cb_iir_limit,
        fault_types_by_run=fault_types_by_run,
    )
    for record in records:
        run = _as_int(record.get("Run"))
        case = _as_text(record.get("Case")).strip()
        if not case or run is None:
            continue
        rows.append(
            NonConvergentCase(
                case=case,
                run=run,
                fault_type=_as_text(record.get("Fault_type")),
                source=_as_text(record.get("Source")),
                signal=_as_text(record.get("Signal")),
                reason=_as_text(record.get("Reason")),
                value=_as_text(record.get("Value")),
                file=_as_text(record.get("File")),
            )
        )
    return rows, warnings


def _collect_fault_types_by_run(project_root: Path) -> dict[int, str]:
    from results_analysis_app.voltage_envelope import read_project_fault_types

    return read_project_fault_types(project_root)


def _read_xlsx_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        sheet = workbook[sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            return []
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        rows: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = {
                header: value
                for header, value in zip(headers, row)
                if header and value is not None
            }
            if record:
                rows.append(record)
        return rows
    finally:
        workbook.close()


def _collect_high_voltage_exclusions(project_root: Path) -> tuple[list[HighVoltageExclusion], list[str]]:
    envelope_root = project_root / "Voltage_envelope"
    rows: list[HighVoltageExclusion] = []
    warnings: list[str] = []
    seen: set[tuple[str, str, int, str, str]] = set()
    workbooks: dict[tuple[Path, str], dict[str, Path]] = {}
    for path in _non_temp_files(envelope_root, "MM_*.xlsx"):
        base_match = re.fullmatch(r"MM_(\d+(?:\.\d+)?)", path.stem, flags=re.IGNORECASE)
        combined_match = re.fullmatch(
            r"MM_(\d+(?:\.\d+)?)_with_combined_plot",
            path.stem,
            flags=re.IGNORECASE,
        )
        match = base_match or combined_match
        if match is None:
            continue
        kind = "base" if base_match is not None else "combined"
        workbooks.setdefault((path.parent, match.group(1)), {})[kind] = path

    selected_workbooks = [
        (voltage, choices.get("base") or choices["combined"])
        for (_parent, voltage), choices in sorted(
            workbooks.items(),
            key=lambda item: (str(item[0][0]).casefold(), float(item[0][1])),
        )
    ]
    for voltage, workbook_path in selected_workbooks:
        try:
            records = _read_xlsx_sheet_rows(workbook_path, "High voltage exclusions")
        except (OSError, BadZipFile, KeyError, ValueError) as exc:
            warnings.append(f"Could not read high-voltage proposals from {workbook_path.name}: {exc}")
            continue
        for record in records:
            run = _as_int(record.get("Run"))
            case = _as_text(record.get("Case")).strip()
            bus = _as_text(record.get("MM_name")).strip()
            signal = _as_text(record.get("Signal"))
            if not case or run is None or not bus:
                continue
            key = (voltage, case, run, bus, signal)
            if key in seen:
                continue
            seen.add(key)
            rows.append(
                HighVoltageExclusion(
                    voltage=voltage,
                    case=case,
                    run=run,
                    bus=bus,
                    fault_type=_as_text(record.get("Fault_type")),
                    measurement=_as_text(record.get("Measurement")),
                    signal=signal,
                    file=_as_text(record.get("File")),
                    excluded_values=_as_text(record.get("Excluded_values")),
                    max_abs=_as_text(record.get("Max_abs")),
                    limit=_as_text(record.get("Limit")),
                    source=(
                        HIGH_VOLTAGE_SOURCE_ANALYSIS
                        if _as_text(record.get("Source")).strip() in {"", "Full scan"}
                        else _as_text(record.get("Source")).strip()
                    ),
                    excluded=True,
                )
            )
    return rows, warnings


def _scan_high_voltage_measurements_from_pscad_log(
    project_root: str | Path,
    *,
    inf_paths: Iterable[Path] | None = None,
    check_cancel: Callable[[], None] | None = None,
) -> tuple[list[HighVoltageMeasurement], list[str]]:
    if check_cancel is not None:
        check_cancel()
    root = Path(project_root).resolve()
    log_path = _pscad_log_path(root)
    if log_path is None:
        return [], [f"PSCAD log not found: {root}"]

    try:
        case_buses = _pscad_log_high_voltage_case_buses(log_path)
    except OSError as exc:
        return [], [f"Could not read PSCAD log: {log_path.name} | {exc}"]
    if not case_buses:
        return [], [f"No treated high-voltage warnings found in {log_path.name}"]

    case_buses_by_case: dict[str, set[str]] = defaultdict(set)
    voltage_by_bus: dict[str, str] = {}
    for case, bus in sorted(case_buses):
        voltage = _voltage_from_mm_bus(bus)
        if voltage:
            case_buses_by_case[case].add(bus)
            voltage_by_bus[bus] = voltage

    case_root = root / "Case_folder"
    if not case_root.is_dir():
        return [], [f"Case_folder not found: {case_root}"]

    matched_cases: set[str] = set()
    matched_inf_paths: list[tuple[Path, set[str]]] = []
    inventory = sorted(inf_paths) if inf_paths is not None else sorted(case_root.rglob("*.inf"))
    for inf_path in inventory:
        if check_cancel is not None:
            check_cancel()
        try:
            case, _run = case_run_from_inf_path(inf_path)
        except ValueError:
            continue
        buses = case_buses_by_case.get(case)
        if buses:
            matched_cases.add(case)
            matched_inf_paths.append((inf_path, buses))

    measurements: list[HighVoltageMeasurement] = []
    warnings: list[str] = []
    if matched_inf_paths:
        worker_count = min(2, len(matched_inf_paths))
        with ThreadPoolExecutor(max_workers=worker_count) as executor:
            futures = {
                executor.submit(
                    _raw_high_voltage_measurements_for_inf,
                    inf_path,
                    buses,
                    voltage_by_bus,
                    check_cancel,
                ): inf_path
                for inf_path, buses in matched_inf_paths
            }
            for future in as_completed(futures):
                if check_cancel is not None:
                    check_cancel()
                records, record_warnings = future.result()
                measurements.extend(records)
                warnings.extend(record_warnings)

    missing_cases = sorted({case for case, _bus in case_buses} - matched_cases)
    if missing_cases:
        warnings.append(f"PSCAD log warning cases not found in Case_folder: {', '.join(missing_cases)}")
    if not matched_inf_paths:
        warnings.append("PSCAD log warning cases were not found in raw PSCAD run files.")
    return measurements, warnings


def high_voltage_exclusions_from_measurements(
    project_root: str | Path,
    measurements: list[HighVoltageMeasurement],
    high_voltage_limit_factor: float,
    um_overrides: dict[str, float] | None = None,
) -> tuple[list[HighVoltageExclusion], list[str]]:
    configs = load_voltage_configs(project_root, um_overrides=um_overrides)
    limit_factor = high_voltage_limit_factor if high_voltage_limit_factor > 0 else 1.0
    missing_um: set[str] = set()
    grouped: dict[tuple[str, str, int, str], dict[str, Any]] = {}
    for item in measurements:
        config = configs.get(item.voltage)
        if config is None:
            missing_um.add(item.voltage)
            continue
        limit = limit_factor * config.um * math.sqrt(2.0)
        if item.max_abs <= limit:
            continue
        key = (item.voltage, item.case, item.run, item.bus)
        value = grouped.setdefault(
            key,
            {
                "measurements": set(),
                "signals": set(),
                "max_abs": 0.0,
                "limit": limit,
            },
        )
        value["measurements"].add(item.measurement)
        value["signals"].add(item.signal)
        value["max_abs"] = max(float(value["max_abs"]), item.max_abs)

    rows = [
        HighVoltageExclusion(
            voltage=voltage,
            case=case,
            run=run,
            bus=bus,
            measurement=", ".join(sorted(value["measurements"])),
            signal=", ".join(sorted(value["signals"])),
            file=PSCAD_LOG_HIGH_VOLTAGE_SOURCE,
            max_abs=f"{float(value['max_abs']):g}",
            limit=f"{float(value['limit']):g}",
            source=HIGH_VOLTAGE_SOURCE_PSCAD_LOG,
        )
        for (voltage, case, run, bus), value in sorted(
            grouped.items(),
            key=lambda item: (float(item[0][0]), item[0][1], item[0][2], item[0][3]),
        )
    ]
    warnings = []
    if missing_um:
        warnings.append(f"Missing Um for voltage(s): {', '.join(sorted(missing_um))}")
    return rows, warnings


def refresh_high_voltage_log_exclusions(
    scan: ProjectScan,
    high_voltage_limit_factor: float,
    um_overrides: dict[str, float] | None = None,
) -> list[str]:
    log_rows, warnings = high_voltage_exclusions_from_measurements(
        scan.path,
        scan.high_voltage_log_measurements,
        high_voltage_limit_factor,
        um_overrides,
    )
    combined = []
    positions: dict[tuple[str, str, int, str], list[int]] = defaultdict(list)
    for row in scan.high_voltage_exclusions:
        if not row.excluded and (
            row.source == HIGH_VOLTAGE_SOURCE_PSCAD_LOG
            or row.file in PSCAD_LOG_HIGH_VOLTAGE_SOURCES
        ):
            continue
        source = row.source.strip()
        if source in {"", "Full scan"}:
            source = HIGH_VOLTAGE_SOURCE_ANALYSIS
        clean_row = replace(row, source=source)
        key = (clean_row.voltage, clean_row.case, clean_row.run, clean_row.bus)
        positions[key].append(len(combined))
        combined.append(clean_row)
    for row in log_rows:
        key = (row.voltage, row.case, row.run, row.bus)
        if key in positions:
            for position in positions[key]:
                existing = combined[position]
                source = (
                    HIGH_VOLTAGE_SOURCE_PSCAD_LOG
                    if existing.source == HIGH_VOLTAGE_SOURCE_PSCAD_LOG
                    else HIGH_VOLTAGE_SOURCE_BOTH
                )
                combined[position] = replace(existing, source=source)
        else:
            positions[key].append(len(combined))
            combined.append(row)
    scan.high_voltage_exclusions = combined
    _refresh_scan_chips(scan)
    return warnings


def refresh_high_voltage_log_scan(
    scan: ProjectScan,
    high_voltage_limit_factor: float,
    um_overrides: dict[str, float] | None = None,
    *,
    check_cancel: Callable[[], None] | None = None,
) -> list[str]:
    warnings: list[str] = []
    if _pscad_log_path(scan.path) is None:
        scan.high_voltage_log_measurements = []
    else:
        scan.high_voltage_log_measurements, warnings = (
            _scan_high_voltage_measurements_from_pscad_log(
                scan.path,
                check_cancel=check_cancel,
            )
        )
    warnings.extend(
        refresh_high_voltage_log_exclusions(
            scan,
            high_voltage_limit_factor,
            um_overrides,
        )
    )
    return warnings


def scan_high_voltage_from_pscad_log(
    project_root: str | Path,
    high_voltage_limit_factor: float,
    um_overrides: dict[str, float] | None = None,
    *,
    check_cancel: Callable[[], None] | None = None,
) -> tuple[list[HighVoltageExclusion], list[str]]:
    root = Path(project_root).resolve()
    measurements, warnings = _scan_high_voltage_measurements_from_pscad_log(
        root,
        check_cancel=check_cancel,
    )
    rows, config_warnings = high_voltage_exclusions_from_measurements(
        root,
        measurements,
        high_voltage_limit_factor,
        um_overrides,
    )
    warnings.extend(config_warnings)
    if measurements and not rows:
        limit_factor = high_voltage_limit_factor if high_voltage_limit_factor > 0 else 1.0
        warnings.append(
            "PSCAD log warnings matched raw run files, but no waveform maxima exceeded "
            f"{limit_factor:g} x Um x sqrt(2)."
        )
    return rows, warnings


def _collect_voltage_data(
    project_root: Path,
    inf_paths: list[Path],
) -> tuple[list[str], dict[str, VoltageConfig]]:
    return discover_voltage_data(project_root, inf_paths=inf_paths)


def _collect_sustained_sdpf_limits(
    project_root: Path,
) -> tuple[dict[str, sustained_sdpf.SDPFVoltageLimits], list[str]]:
    workbook_path = input_data_workbook(project_root)
    overrides_path = project_root / "Plots" / ".plottool_v3" / "limits.json"
    return sustained_sdpf.resolve_project_limits(
        project_root,
        workbook_path=workbook_path,
        overrides_path=overrides_path,
    )


def _refresh_scan_chips(scan: ProjectScan, has_result_files: bool | None = None) -> None:
    if not scan.exists:
        scan.chips = ["Missing Project"]
        return

    if has_result_files is None:
        base_chip = next(
            (chip for chip in scan.chips if chip in {"Ready", "Missing Results"}),
            "Ready" if scan.case_infos else "Missing Results",
        )
    else:
        base_chip = "Ready" if has_result_files or scan.case_infos else "Missing Results"

    scan.chips = [base_chip]
    if not scan.has_dashboards:
        scan.chips.append("No dashboards")
    if scan.dashboard_changed:
        scan.chips.append("Dashboards changed")
    if not scan.has_envelopes:
        scan.chips.append("No envelopes")
    if scan.has_plots:
        scan.chips.append("Plots exist")
    if scan.has_reports:
        scan.chips.append("Reports exist")
    if scan.nonconv_cases:
        scan.chips.append(f"NonConv proposals: {len(scan.nonconv_cases)}")
    if scan.high_voltage_exclusions:
        proposal_count = len(
            {
                (row.voltage, row.case, row.run, row.bus)
                for row in scan.high_voltage_exclusions
            }
        )
        scan.chips.append(f"High voltage proposals: {proposal_count}")


def mark_dashboard_changed(scan: ProjectScan) -> None:
    scan.dashboard_changed = True
    _refresh_scan_chips(scan)


def apply_project_output_state(scan: ProjectScan, state: Any) -> None:
    if not isinstance(state, dict):
        _refresh_scan_chips(scan)
        return
    scan.has_dashboards = bool(state.get("has_dashboards", scan.has_dashboards))
    scan.has_envelopes = bool(state.get("has_envelopes", scan.has_envelopes))
    scan.has_plots = bool(state.get("has_plots", scan.has_plots))
    scan.has_reports = bool(state.get("has_reports", scan.has_reports))
    _refresh_scan_chips(scan, has_result_files=bool(state.get("has_results", False)))


def refresh_project_scan_outputs(
    scan: ProjectScan,
    *,
    refresh_dashboards: bool = False,
    refresh_envelopes: bool = False,
    refresh_plots: bool = False,
    refresh_reports: bool = False,
    refresh_high_voltage_exclusions: bool = False,
    high_voltage_limit_factor: float | None = None,
    voltage_um_overrides: dict[str, float] | None = None,
) -> ProjectScan:
    """Refresh output flags after a targeted workflow action."""
    project_root = scan.path
    if refresh_dashboards:
        scan.has_dashboards = _has_dashboard_files(project_root / "Dashboards")
        scan.dashboard_changed = False
    if refresh_envelopes:
        scan.has_envelopes = _has_non_temp_file(project_root / "Voltage_envelope", "*.xlsx")
    if refresh_plots:
        scan.has_plots = _has_plot_files(project_root / "Plots" / "Generated")
    if refresh_reports:
        scan.has_reports = _has_non_temp_file(project_root / "Reports", "*.docx")
    if refresh_high_voltage_exclusions:
        scan.high_voltage_exclusions, warnings = _collect_high_voltage_exclusions(project_root)
        scan.messages.extend(warnings)
        if high_voltage_limit_factor is not None:
            scan.messages.extend(
                refresh_high_voltage_log_exclusions(
                    scan,
                    high_voltage_limit_factor,
                    voltage_um_overrides,
                )
            )
    _refresh_scan_chips(scan)
    return scan


def scan_project(
    path: str | Path,
    nonconv_cb_iip_limit: float | None = None,
    nonconv_cb_iir_limit: float | None = None,
    high_voltage_limit_factor: float = 1.0,
    voltage_um_overrides: dict[str, float] | None = None,
    check_cancel: Callable[[], None] | None = None,
) -> ProjectScan:
    project_root = Path(path).resolve()
    scan = ProjectScan(path=project_root, exists=project_root.is_dir())
    if not scan.exists:
        scan.chips.append("Missing Project")
        scan.messages.append(f"Project folder does not exist: {project_root}")
        return scan

    scan.case_infos, inf_paths = _collect_case_infos_and_inf_paths(project_root)
    dashboard_root = project_root / "Dashboards"
    envelope_root = project_root / "Voltage_envelope"
    generated_root = project_root / "Plots" / "Generated"
    reports_root = project_root / "Reports"
    results_dir = project_root / "Results"

    with ThreadPoolExecutor(max_workers=8) as executor:
        fault_types_future = executor.submit(
            _collect_fault_types_by_run,
            project_root,
        )
        futures = {
            "has_result_files": executor.submit(_has_non_temp_file, results_dir, "*.csv"),
            "nonconv_cases": executor.submit(
                _collect_nonconv_cases,
                project_root,
                nonconv_cb_iip_limit,
                nonconv_cb_iir_limit,
                fault_types_future,
            ),
            "fault_types_by_run": fault_types_future,
            "has_dashboards": executor.submit(_has_dashboard_files, dashboard_root),
            "dashboard_figures": executor.submit(scan_dashboard_figures, project_root),
            "has_envelopes": executor.submit(_has_non_temp_file, envelope_root, "*.xlsx"),
            "has_plots": executor.submit(_has_plot_files, generated_root),
            "has_reports": executor.submit(_has_non_temp_file, reports_root, "*.docx"),
            "voltage_data": executor.submit(_collect_voltage_data, project_root, inf_paths),
            "sustained_sdpf_limits": executor.submit(_collect_sustained_sdpf_limits, project_root),
            "high_voltage_exclusions": executor.submit(_collect_high_voltage_exclusions, project_root),
            "project_timing": executor.submit(load_project_timing, project_root),
        }
        if _pscad_log_path(project_root) is not None:
            futures["high_voltage_log"] = executor.submit(
                _scan_high_voltage_measurements_from_pscad_log,
                project_root,
                inf_paths=inf_paths,
                check_cancel=check_cancel,
            )
        has_result_files = bool(futures["has_result_files"].result())
        try:
            scan.nonconv_cases, warnings = futures["nonconv_cases"].result()
            scan.messages.extend(warnings)
        except Exception as exc:
            scan.messages.append(f"Could not scan NonConv proposals: {exc}")
        try:
            scan.fault_types_by_run = futures["fault_types_by_run"].result()
        except Exception as exc:
            scan.messages.append(f"Could not scan fault types: {exc}")
        scan.has_dashboards = futures["has_dashboards"].result()
        try:
            scan.dashboard_figures, dashboard_warnings = futures["dashboard_figures"].result()
            scan.messages.extend(
                f"Dashboard: {warning}" for warning in dashboard_warnings
            )
        except Exception as exc:
            scan.messages.append(f"Could not scan dashboard figures: {exc}")
        scan.has_envelopes = futures["has_envelopes"].result()
        scan.has_plots = futures["has_plots"].result()
        scan.has_reports = futures["has_reports"].result()
        scan.available_voltages, scan.voltage_configs = futures["voltage_data"].result()
        try:
            (
                scan.sustained_sdpf_limits,
                scan.sustained_sdpf_limit_warnings,
            ) = futures["sustained_sdpf_limits"].result()
            scan.messages.extend(
                f"Sustained SDPF: {warning}"
                for warning in scan.sustained_sdpf_limit_warnings
            )
        except Exception as exc:
            scan.messages.append(f"Could not scan Sustained SDPF limits: {exc}")
        scan.high_voltage_exclusions, warnings = futures["high_voltage_exclusions"].result()
        scan.messages.extend(warnings)
        timing = futures["project_timing"].result()
        scan.project_frequency = timing.frequency
        scan.final_duration = timing.final_duration
        if "high_voltage_log" in futures:
            scan.high_voltage_log_measurements, warnings = futures["high_voltage_log"].result()
            scan.messages.extend(warnings)
            scan.messages.extend(
                refresh_high_voltage_log_exclusions(
                    scan,
                    high_voltage_limit_factor,
                    voltage_um_overrides,
                )
            )

    has_results = bool(has_result_files or scan.case_infos)

    _refresh_scan_chips(scan, has_result_files=has_results)

    return scan


def scope_matches_case(scope: ScopeEntry, case_info: CaseInfo) -> bool:
    if scope.mode == "full":
        return True
    tokens = [token.casefold() for token in scope.tokens]
    matched = any(token in case_info.match_text for token in tokens)
    if scope.mode == "include":
        return matched
    if scope.mode == "exclude":
        return not matched
    return True


def preview_scope(scan: ProjectScan, scope: ScopeEntry) -> ScopePreview:
    if not scan.case_infos:
        return ScopePreview(
            scope_folder=scope.folder,
            matched_count=0,
            total_count=0,
            examples=(),
            warning="no cases found",
        )
    matched = [case for case in scan.case_infos if scope_matches_case(scope, case)]
    warning = "no matched cases" if not matched else ""
    return ScopePreview(
        scope_folder=scope.folder,
        matched_count=len(matched),
        total_count=len(scan.case_infos),
        examples=tuple(case.name for case in matched[:4]),
        warning=warning,
    )


def _figure_id(workbook: str, sheet: str, chart_index: int, title: str) -> str:
    raw = f"{workbook}|{sheet}|{chart_index}|{title}"
    return re.sub(r"\s+", " ", raw).strip()


def _rels_path(part_name: str) -> str:
    folder, name = posixpath.split(part_name)
    return posixpath.join(folder, "_rels", f"{name}.rels")


def _resolve_part(base_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _read_rels(workbook: ZipFile, part_name: str) -> dict[str, dict[str, str]]:
    rels_name = _rels_path(part_name)
    if rels_name not in workbook.namelist():
        return {}
    root = ET.fromstring(workbook.read(rels_name))
    return {
        str(rel.attrib.get("Id")): dict(rel.attrib)
        for rel in root.findall("rel:Relationship", OOXML_NS)
        if rel.attrib.get("Id")
    }


def _chart_title_from_xml(workbook: ZipFile, chart_part: str) -> str:
    root = ET.fromstring(workbook.read(chart_part))
    title = root.find(".//c:title", OOXML_NS)
    if title is None:
        return ""
    return "".join(text.text or "" for text in title.findall(".//a:t", OOXML_NS)).strip()


def _scan_dashboard_workbook(workbook_path: Path) -> list[DashboardFigure]:
    figures: list[DashboardFigure] = []
    with ZipFile(workbook_path) as workbook:
        workbook_part = "xl/workbook.xml"
        root = ET.fromstring(workbook.read(workbook_part))
        workbook_rels = _read_rels(workbook, workbook_part)
        for sheet in root.findall("main:sheets/main:sheet", OOXML_NS):
            sheet_name = str(sheet.attrib.get("name", ""))
            sheet_id = sheet.attrib.get(f"{{{OOXML_NS['office']}}}id")
            sheet_rel = workbook_rels.get(str(sheet_id))
            if not sheet_rel:
                continue
            sheet_part = _resolve_part(workbook_part, str(sheet_rel.get("Target", "")))
            sheet_rels = _read_rels(workbook, sheet_part)
            if not sheet_rels:
                continue

            sheet_root = ET.fromstring(workbook.read(sheet_part))
            chart_index = 0
            for drawing in sheet_root.findall("main:drawing", OOXML_NS):
                drawing_id = drawing.attrib.get(f"{{{OOXML_NS['office']}}}id")
                drawing_rel = sheet_rels.get(str(drawing_id))
                if not drawing_rel:
                    continue
                drawing_part = _resolve_part(sheet_part, str(drawing_rel.get("Target", "")))
                drawing_rels = _read_rels(workbook, drawing_part)
                drawing_root = ET.fromstring(workbook.read(drawing_part))
                for graphic_data in drawing_root.findall(".//a:graphicData", OOXML_NS):
                    chart = graphic_data.find("c:chart", OOXML_NS)
                    if chart is None:
                        continue
                    chart_id = chart.attrib.get(f"{{{OOXML_NS['office']}}}id")
                    chart_rel = drawing_rels.get(str(chart_id))
                    if not chart_rel:
                        continue
                    chart_index += 1
                    chart_part = _resolve_part(drawing_part, str(chart_rel.get("Target", "")))
                    title = _chart_title_from_xml(workbook, chart_part)
                    figures.append(
                        DashboardFigure(
                            id=_figure_id(workbook_path.name, sheet_name, chart_index, title),
                            workbook=workbook_path.name,
                            sheet=sheet_name,
                            chart_index=chart_index,
                            title=title,
                        )
                    )
    return figures


def scan_dashboard_figures(project_root: str | Path) -> tuple[list[DashboardFigure], list[str]]:
    """Read dashboard workbook chart metadata without changing the workbooks."""
    root = Path(project_root).resolve() / "Dashboards"
    if not root.is_dir():
        return [], [f"Dashboard folder not found: {root}"]

    figures: list[DashboardFigure] = []
    warnings: list[str] = []
    workbooks = [
        path
        for pattern in ("*.xlsx", "*.xlsm")
        for path in sorted(root.glob(pattern))
        if not path.name.startswith("~$")
    ]

    for workbook_path in workbooks:
        try:
            figures.extend(_scan_dashboard_workbook(workbook_path))
        except (BadZipFile, ET.ParseError, KeyError, OSError) as exc:
            warnings.append(f"Could not read charts from {workbook_path.name}: {exc}")

    return figures, warnings
