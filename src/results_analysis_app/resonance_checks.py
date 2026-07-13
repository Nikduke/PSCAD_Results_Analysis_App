from __future__ import annotations

from dataclasses import dataclass, field, fields as dataclass_fields, replace
from pathlib import Path
from typing import Any
import math
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from results_analysis_app.models import (
    DEFAULT_RESONANCE_AUTO_RELEASE,
    DEFAULT_RESONANCE_CHECKS,
    DEFAULT_RESONANCE_GROWTH_WINDOW_FRACTION,
    DEFAULT_RESONANCE_LIMIT_MULTIPLIER,
    DEFAULT_RESONANCE_LOG_FLOOR_ABSOLUTE,
    DEFAULT_RESONANCE_LOG_FLOOR_VLIM_FACTOR,
    DEFAULT_RESONANCE_MANUAL_START,
    DEFAULT_RESONANCE_MIN_GROWTH_DELTA_FACTOR,
    DEFAULT_RESONANCE_MIN_GROWTH_RATIO,
    DEFAULT_RESONANCE_MIN_LEVEL_OVER_VLIM,
    DEFAULT_RESONANCE_MIN_POSITIVE_FRACTION,
    DEFAULT_RESONANCE_PEAK_SEARCH_FRACTION,
    DEFAULT_RESONANCE_RELEASE_DECAY_RATIO,
    DEFAULT_RESONANCE_RELEASE_HOLD_TIME,
    DEFAULT_RESONANCE_RELEASE_REBOUND_RATIO,
    DEFAULT_RESONANCE_ROLLING_MIN_SAMPLES,
    DEFAULT_RESONANCE_ROLLING_P95_WINDOW,
    DEFAULT_RESONANCE_TOP_N,
    normalize_fraction,
    normalize_nonnegative_float,
    normalize_positive_float,
    normalize_positive_int,
)


WORKBOOK_NAME = "Resonance_Checks.xlsx"
VOLTAGE_TYPES = ("LGp", "LLp")
POST_EVENT_STRESS = "Post_Event_Stress"
LATE_GROWTH = "Late_Growth"
NO_SETTLE_GROWTH = "No_Settle_Growth"

CHECK_DEFINITIONS = {
    POST_EVENT_STRESS: {
        "label": "Post-event stress",
        "folder": "Post_Event_Stress",
        "sheet": "Post_Stress",
        "abbr": "PES",
    },
    LATE_GROWTH: {
        "label": "Late growth",
        "folder": "Late_Growth",
        "sheet": "Late_Growth",
        "abbr": "LG",
    },
    NO_SETTLE_GROWTH: {
        "label": "No-settle growth",
        "folder": "No_Settle_Growth",
        "sheet": "No_Settle",
        "abbr": "NS",
    },
}

RESONANCE_SESSION_FIELDS = {
    "top_n": "resonance_top_n",
    "limit_multiplier": "resonance_limit_multiplier",
    "auto_release": "resonance_auto_release",
    "manual_analysis_start": "resonance_manual_analysis_start",
    "peak_search_fraction": "resonance_peak_search_fraction",
    "release_decay_ratio": "resonance_release_decay_ratio",
    "release_rebound_ratio": "resonance_release_rebound_ratio",
    "release_hold_time": "resonance_release_hold_time",
    "rolling_p95_window": "resonance_rolling_p95_window",
    "rolling_min_samples": "resonance_rolling_min_samples",
    "log_floor_vlim_factor": "resonance_log_floor_vlim_factor",
    "log_floor_absolute": "resonance_log_floor_absolute",
    "growth_window_fraction": "resonance_growth_window_fraction",
    "min_positive_fraction": "resonance_min_positive_fraction",
    "min_growth_ratio": "resonance_min_growth_ratio",
    "min_level_over_vlim": "resonance_min_level_over_vlim",
    "min_growth_delta_factor": "resonance_min_growth_delta_factor",
}


@dataclass(frozen=True)
class ResonanceSettings:
    enabled_checks: tuple[str, ...] = ()
    top_n: int = DEFAULT_RESONANCE_TOP_N
    limit_multiplier: float = DEFAULT_RESONANCE_LIMIT_MULTIPLIER
    auto_release: bool = DEFAULT_RESONANCE_AUTO_RELEASE
    manual_analysis_start: float = DEFAULT_RESONANCE_MANUAL_START
    peak_search_fraction: float = DEFAULT_RESONANCE_PEAK_SEARCH_FRACTION
    release_decay_ratio: float = DEFAULT_RESONANCE_RELEASE_DECAY_RATIO
    release_rebound_ratio: float = DEFAULT_RESONANCE_RELEASE_REBOUND_RATIO
    release_hold_time: float = DEFAULT_RESONANCE_RELEASE_HOLD_TIME
    rolling_p95_window: float = DEFAULT_RESONANCE_ROLLING_P95_WINDOW
    rolling_min_samples: int = DEFAULT_RESONANCE_ROLLING_MIN_SAMPLES
    log_floor_vlim_factor: float = DEFAULT_RESONANCE_LOG_FLOOR_VLIM_FACTOR
    log_floor_absolute: float = DEFAULT_RESONANCE_LOG_FLOOR_ABSOLUTE
    growth_window_fraction: float = DEFAULT_RESONANCE_GROWTH_WINDOW_FRACTION
    min_positive_fraction: float = DEFAULT_RESONANCE_MIN_POSITIVE_FRACTION
    min_growth_ratio: float = DEFAULT_RESONANCE_MIN_GROWTH_RATIO
    min_level_over_vlim: float = DEFAULT_RESONANCE_MIN_LEVEL_OVER_VLIM
    min_growth_delta_factor: float = DEFAULT_RESONANCE_MIN_GROWTH_DELTA_FACTOR

    @classmethod
    def from_mapping(cls, data: dict[str, Any] | None) -> "ResonanceSettings":
        if not data:
            return cls()
        return cls(
            enabled_checks=tuple(
                str(value)
                for value in data.get("enabled_checks", ())
                if str(value) in DEFAULT_RESONANCE_CHECKS
            ),
            top_n=normalize_positive_int(data.get("top_n", DEFAULT_RESONANCE_TOP_N), DEFAULT_RESONANCE_TOP_N),
            limit_multiplier=normalize_positive_float(
                data.get("limit_multiplier", DEFAULT_RESONANCE_LIMIT_MULTIPLIER),
                DEFAULT_RESONANCE_LIMIT_MULTIPLIER,
            ),
            auto_release=bool(data.get("auto_release", DEFAULT_RESONANCE_AUTO_RELEASE)),
            manual_analysis_start=normalize_nonnegative_float(
                data.get("manual_analysis_start", DEFAULT_RESONANCE_MANUAL_START),
                DEFAULT_RESONANCE_MANUAL_START,
            ),
            peak_search_fraction=normalize_fraction(
                data.get("peak_search_fraction", DEFAULT_RESONANCE_PEAK_SEARCH_FRACTION),
                DEFAULT_RESONANCE_PEAK_SEARCH_FRACTION,
            ),
            release_decay_ratio=normalize_fraction(
                data.get("release_decay_ratio", DEFAULT_RESONANCE_RELEASE_DECAY_RATIO),
                DEFAULT_RESONANCE_RELEASE_DECAY_RATIO,
            ),
            release_rebound_ratio=normalize_fraction(
                data.get("release_rebound_ratio", DEFAULT_RESONANCE_RELEASE_REBOUND_RATIO),
                DEFAULT_RESONANCE_RELEASE_REBOUND_RATIO,
            ),
            release_hold_time=normalize_positive_float(
                data.get("release_hold_time", DEFAULT_RESONANCE_RELEASE_HOLD_TIME),
                DEFAULT_RESONANCE_RELEASE_HOLD_TIME,
            ),
            rolling_p95_window=normalize_positive_float(
                data.get("rolling_p95_window", DEFAULT_RESONANCE_ROLLING_P95_WINDOW),
                DEFAULT_RESONANCE_ROLLING_P95_WINDOW,
            ),
            rolling_min_samples=normalize_positive_int(
                data.get("rolling_min_samples", DEFAULT_RESONANCE_ROLLING_MIN_SAMPLES),
                DEFAULT_RESONANCE_ROLLING_MIN_SAMPLES,
            ),
            log_floor_vlim_factor=normalize_positive_float(
                data.get("log_floor_vlim_factor", DEFAULT_RESONANCE_LOG_FLOOR_VLIM_FACTOR),
                DEFAULT_RESONANCE_LOG_FLOOR_VLIM_FACTOR,
            ),
            log_floor_absolute=normalize_positive_float(
                data.get("log_floor_absolute", DEFAULT_RESONANCE_LOG_FLOOR_ABSOLUTE),
                DEFAULT_RESONANCE_LOG_FLOOR_ABSOLUTE,
            ),
            growth_window_fraction=normalize_fraction(
                data.get("growth_window_fraction", DEFAULT_RESONANCE_GROWTH_WINDOW_FRACTION),
                DEFAULT_RESONANCE_GROWTH_WINDOW_FRACTION,
            ),
            min_positive_fraction=normalize_fraction(
                data.get("min_positive_fraction", DEFAULT_RESONANCE_MIN_POSITIVE_FRACTION),
                DEFAULT_RESONANCE_MIN_POSITIVE_FRACTION,
            ),
            min_growth_ratio=normalize_positive_float(
                data.get("min_growth_ratio", DEFAULT_RESONANCE_MIN_GROWTH_RATIO),
                DEFAULT_RESONANCE_MIN_GROWTH_RATIO,
            ),
            min_level_over_vlim=normalize_positive_float(
                data.get("min_level_over_vlim", DEFAULT_RESONANCE_MIN_LEVEL_OVER_VLIM),
                DEFAULT_RESONANCE_MIN_LEVEL_OVER_VLIM,
            ),
            min_growth_delta_factor=normalize_positive_float(
                data.get("min_growth_delta_factor", DEFAULT_RESONANCE_MIN_GROWTH_DELTA_FACTOR),
                DEFAULT_RESONANCE_MIN_GROWTH_DELTA_FACTOR,
            ),
        )

    @classmethod
    def from_session(cls, session: Any) -> "ResonanceSettings":
        defaults = cls()
        data = {
            "enabled_checks": getattr(session, "resonance_enabled_checks", ()),
            **{
                key: getattr(session, session_attr, getattr(defaults, key))
                for key, session_attr in RESONANCE_SESSION_FIELDS.items()
            },
        }
        return cls.from_mapping(data)

    @property
    def enabled(self) -> bool:
        return bool(self.enabled_checks)

    @property
    def effective_enabled_checks(self) -> tuple[str, ...]:
        return tuple(
            check
            for check in self.enabled_checks
            if check != NO_SETTLE_GROWTH or self.auto_release
        )

    def to_mapping(self) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for item in dataclass_fields(self):
            value = getattr(self, item.name)
            result[item.name] = list(value) if item.name == "enabled_checks" else value
        return result


@dataclass(frozen=True)
class ChronologicalEnvelope:
    scope_folder: str
    voltage: str
    voltage_type: str
    case: str
    run: int
    mm_name: str
    vnom: float
    time: np.ndarray
    envelope: np.ndarray


@dataclass
class ResonanceResult:
    check: str
    scope_folder: str
    voltage: str
    voltage_type: str
    case: str
    run: int
    mm_name: str
    vnom: float
    vlim: float
    t_guard: float
    t_start: float
    main_metric: float
    metrics: dict[str, float | int | str]
    rank_key: tuple[float, ...]
    time: np.ndarray = field(repr=False)
    envelope: np.ndarray = field(repr=False)
    smooth: np.ndarray = field(repr=False)
    rank: int = 0


def selected_plot_events(settings: ResonanceSettings | dict[str, Any] | None) -> list[str]:
    parsed = settings if isinstance(settings, ResonanceSettings) else ResonanceSettings.from_mapping(settings)
    events: list[str] = []
    for check in parsed.effective_enabled_checks:
        for voltage_type in VOLTAGE_TYPES:
            events.append(plot_event_name(check, voltage_type))
    return events


def plot_event_name(check: str, voltage_type: str) -> str:
    return f"{CHECK_DEFINITIONS[check]['folder']}__{voltage_type}"


def output_dir_for_event(project_root: Path, scope_folder: str, event_name: str) -> Path | None:
    for check, definition in CHECK_DEFINITIONS.items():
        prefix = f"{definition['folder']}__"
        if event_name.startswith(prefix):
            voltage_type = event_name.removeprefix(prefix)
            if voltage_type in VOLTAGE_TYPES:
                return project_root / "Plots" / "Generated" / scope_folder / definition["folder"] / voltage_type
    return None


def report_folder(check: str, voltage_type: str) -> tuple[str, str]:
    return CHECK_DEFINITIONS[check]["folder"], voltage_type


def sheet_name(check: str, voltage_type: str) -> str:
    return f"{CHECK_DEFINITIONS[check]['sheet']}_{voltage_type}"


def records_from_entries(
    scope_folder: str,
    voltage: str,
    voltage_type: str,
    vnom: float,
    entries: list[tuple[str, pd.DataFrame]],
) -> list[ChronologicalEnvelope]:
    records: list[ChronologicalEnvelope] = []
    for measurement, df in entries:
        if measurement != voltage_type or df.empty or "Case" not in df.columns or "MM_name" not in df.columns:
            continue
        max_cols = [col for col in df.columns if str(col).startswith("Max_")]
        if not max_cols:
            continue
        time = df.index.to_numpy(dtype=float)
        envelope = df[max_cols].max(axis=1, skipna=True).to_numpy(dtype=float)
        case_text = str(df["Case"].iloc[0])
        case, run = _split_case_run(case_text)
        records.append(
            ChronologicalEnvelope(
                scope_folder=scope_folder,
                voltage=str(voltage),
                voltage_type=voltage_type,
                case=case,
                run=run,
                mm_name=str(df["MM_name"].iloc[0]),
                vnom=float(vnom),
                time=time,
                envelope=envelope,
            )
        )
    return records


def analyze_records(
    records: list[ChronologicalEnvelope],
    settings: ResonanceSettings,
    event_times: dict[str, float] | None,
) -> list[ResonanceResult]:
    if not records or not settings.enabled:
        return []
    if not settings.auto_release and NO_SETTLE_GROWTH in settings.enabled_checks:
        settings = replace(settings, enabled_checks=settings.effective_enabled_checks)
    if not settings.enabled_checks:
        return []

    if not settings.auto_release:
        t_end_min = min(float(record.time[-1]) for record in records if len(record.time))
        if not (0.0 <= settings.manual_analysis_start < t_end_min):
            raise ValueError(
                "Manual analysis start time must be within [0, t_end) for resonance checks."
            )

    t_guard = analysis_guard_time(event_times)
    by_group: dict[tuple[str, str, str], list[ResonanceResult]] = {}
    for record in records:
        if len(record.time) < 3:
            continue
        vlim = record.vnom * settings.limit_multiplier
        smooth = _rolling_p95(record.time, record.envelope, settings)
        release = _release_time(record.time, smooth, t_guard, settings)
        start = settings.manual_analysis_start if not settings.auto_release else release

        if POST_EVENT_STRESS in settings.enabled_checks and start is not None:
            result = _post_event_result(record, settings, smooth, vlim, t_guard, float(start))
            if result is not None:
                by_group.setdefault((POST_EVENT_STRESS, record.voltage, record.voltage_type), []).append(result)

        if LATE_GROWTH in settings.enabled_checks and start is not None:
            result = _late_growth_result(record, settings, smooth, vlim, t_guard, float(start))
            if result is not None:
                by_group.setdefault((LATE_GROWTH, record.voltage, record.voltage_type), []).append(result)

        if (
            settings.auto_release
            and NO_SETTLE_GROWTH in settings.enabled_checks
            and release is None
        ):
            result = _no_settle_result(record, settings, smooth, vlim, t_guard)
            if result is not None:
                by_group.setdefault((NO_SETTLE_GROWTH, record.voltage, record.voltage_type), []).append(result)

    selected: list[ResonanceResult] = []
    for results in by_group.values():
        for rank, result in enumerate(sorted(results, key=lambda item: item.rank_key, reverse=True)[: settings.top_n], start=1):
            result.rank = rank
            selected.append(result)
    return selected


def write_workbooks(
    project_root: Path,
    scopes: list[Any],
    results: list[ResonanceResult],
    settings: ResonanceSettings,
    event_times: dict[str, float] | None,
) -> list[Path]:
    if not settings.enabled:
        return []
    written: list[Path] = []
    for scope in scopes:
        scope_results = [result for result in results if result.scope_folder == scope.folder]
        path = project_root / "Voltage_envelope" / scope.folder / WORKBOOK_NAME
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.active.title = "Settings"
        _write_settings_sheet(wb["Settings"], settings, event_times)
        for check in settings.enabled_checks:
            if check not in settings.effective_enabled_checks:
                continue
            for voltage_type in VOLTAGE_TYPES:
                rows = [
                    result
                    for result in scope_results
                    if result.check == check and result.voltage_type == voltage_type
                ]
                _write_result_sheet(wb.create_sheet(sheet_name(check, voltage_type)), rows)
        for result in scope_results:
            _write_chart_data_sheet(wb, result)
        _format_workbook(wb)
        wb.save(path)
        wb.close()
        written.append(path)
    return written


def create_plot_batch_rows(
    project_root: Path,
    scope_folder: str,
    checks: tuple[str, ...],
) -> dict[str, list[dict[str, Any]]]:
    path = project_root / "Voltage_envelope" / scope_folder / WORKBOOK_NAME
    if not path.is_file() or not checks:
        return {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        output: dict[str, list[dict[str, Any]]] = {}
        for check in checks:
            if check not in CHECK_DEFINITIONS:
                continue
            for voltage_type in VOLTAGE_TYPES:
                name = sheet_name(check, voltage_type)
                if name not in wb.sheetnames:
                    continue
                rows = _rows_from_sheet(wb[name])
                if rows:
                    output[plot_event_name(check, voltage_type)] = rows
        return output
    finally:
        wb.close()


def _rows_from_sheet(ws) -> list[dict[str, Any]]:
    headers = {
        str(ws.cell(1, column).value).strip(): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(1, column).value is not None
    }
    rows: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        case = ws.cell(row, headers.get("Case", 0)).value if "Case" in headers else None
        run = ws.cell(row, headers.get("Run", 0)).value if "Run" in headers else None
        mm_name = ws.cell(row, headers.get("MM", 0)).value if "MM" in headers else None
        if not case or not run or not mm_name:
            continue
        rows.append(
            {
                "case": case,
                "run": run,
                "element": mm_name,
                "trace": "Both",
                "overview": True,
                "tov_windows": True,
                "tov_window_count": 4,
                "limits": True,
                "excel_export": True,
            }
        )
    return rows


def _post_event_result(
    record: ChronologicalEnvelope,
    settings: ResonanceSettings,
    smooth: np.ndarray,
    vlim: float,
    t_guard: float,
    t_start: float,
) -> ResonanceResult | None:
    mask = record.time >= t_start
    if mask.sum() < 2:
        return None
    area = _area_above(record.time[mask], record.envelope[mask], vlim)
    if area <= 0:
        return None
    above = record.envelope[mask] > vlim
    metrics = {
        "A_post": area,
        "T_above_post": _time_above(record.time[mask], above),
        "E_peak_post": float(np.nanmax(record.envelope[mask])),
    }
    return _result(record, POST_EVENT_STRESS, settings, smooth, vlim, t_guard, t_start, area, metrics, (area,))


def _late_growth_result(
    record: ChronologicalEnvelope,
    settings: ResonanceSettings,
    smooth: np.ndarray,
    vlim: float,
    t_guard: float,
    t_start: float,
) -> ResonanceResult | None:
    windows = _growth_windows(record.time, smooth, t_start, settings)
    if windows is None:
        return None
    previous, tail = windows
    previous_p95 = _p95(smooth[previous])
    tail_p95 = _p95(smooth[tail])
    growth_ratio = tail_p95 / previous_p95 if previous_p95 > 0 else math.inf
    growth_delta = tail_p95 - previous_p95
    positive_fraction = _positive_fraction(smooth[tail])
    sigma = _log_slope(record.time[tail], smooth[tail], vlim, settings)
    if sigma <= 0:
        return None
    if not _passes_relevance(settings, positive_fraction, tail_p95 / vlim, growth_ratio, growth_delta, vlim):
        return None
    metrics = {
        "sigma_tail": sigma,
        "growth_ratio_tail": growth_ratio,
        "growth_delta_tail": growth_delta,
        "positive_fraction_tail": positive_fraction,
        "E_tail_p95_over_Vlim": tail_p95 / vlim,
    }
    rank_key = (sigma, growth_ratio, positive_fraction, tail_p95 / vlim)
    return _result(record, LATE_GROWTH, settings, smooth, vlim, t_guard, t_start, sigma, metrics, rank_key)


def _no_settle_result(
    record: ChronologicalEnvelope,
    settings: ResonanceSettings,
    smooth: np.ndarray,
    vlim: float,
    t_guard: float,
) -> ResonanceResult | None:
    mask = record.time >= t_guard
    if mask.sum() < 4:
        return None
    duration = record.time[-1] - t_guard
    if duration <= 0:
        return None
    window_s = max(duration * settings.growth_window_fraction, _median_dt(record.time) * settings.rolling_min_samples)
    start_mask = (record.time >= t_guard) & (record.time <= t_guard + window_s)
    end_mask = record.time >= record.time[-1] - window_s
    if start_mask.sum() < 2 or end_mask.sum() < 2:
        return None
    start_p95 = _p95(smooth[start_mask])
    end_p95 = _p95(smooth[end_mask])
    growth_ratio = end_p95 / start_p95 if start_p95 > 0 else math.inf
    growth_delta = end_p95 - start_p95
    positive_fraction = _positive_fraction(smooth[mask])
    sigma = _log_slope(record.time[mask], smooth[mask], vlim, settings)
    if sigma <= 0:
        return None
    if not _passes_relevance(settings, positive_fraction, end_p95 / vlim, growth_ratio, growth_delta, vlim):
        return None
    metrics = {
        "sigma_post": sigma,
        "growth_ratio_post": growth_ratio,
        "growth_delta_post": growth_delta,
        "positive_fraction_post": positive_fraction,
        "positive_window_count": _positive_window_count(smooth[mask]),
        "E_end_p95_over_Vlim": end_p95 / vlim,
        "A_from_start": _area_above(record.time[mask], record.envelope[mask], vlim),
    }
    rank_key = (
        sigma,
        positive_fraction,
        float(metrics["positive_window_count"]),
        growth_ratio,
        end_p95 / vlim,
        float(metrics["A_from_start"]),
    )
    return _result(record, NO_SETTLE_GROWTH, settings, smooth, vlim, t_guard, t_guard, sigma, metrics, rank_key)


def _result(
    record: ChronologicalEnvelope,
    check: str,
    settings: ResonanceSettings,
    smooth: np.ndarray,
    vlim: float,
    t_guard: float,
    t_start: float,
    main_metric: float,
    metrics: dict[str, float | int | str],
    rank_key: tuple[float, ...],
) -> ResonanceResult:
    return ResonanceResult(
        check=check,
        scope_folder=record.scope_folder,
        voltage=record.voltage,
        voltage_type=record.voltage_type,
        case=record.case,
        run=record.run,
        mm_name=record.mm_name,
        vnom=record.vnom,
        vlim=vlim,
        t_guard=t_guard,
        t_start=t_start,
        main_metric=main_metric,
        metrics=metrics,
        rank_key=rank_key,
        time=record.time.copy(),
        envelope=record.envelope.copy(),
        smooth=smooth.copy(),
    )


def _rolling_p95(time: np.ndarray, values: np.ndarray, settings: ResonanceSettings) -> np.ndarray:
    dt = _median_dt(time)
    window = max(settings.rolling_min_samples, int(math.ceil(settings.rolling_p95_window / dt)))
    return (
        pd.Series(values)
        .rolling(window=window, min_periods=1)
        .quantile(0.95)
        .to_numpy(dtype=float)
    )


def _release_time(
    time: np.ndarray,
    smooth: np.ndarray,
    t_guard: float,
    settings: ResonanceSettings,
) -> float | None:
    if t_guard >= float(time[-1]):
        return None
    search_end = t_guard + (float(time[-1]) - t_guard) * settings.peak_search_fraction
    peak_indices = _peak_candidates(time, smooth, t_guard, search_end)
    if not peak_indices:
        return None
    hold = max(1, int(math.ceil(settings.release_hold_time / _median_dt(time))))
    for peak_index in peak_indices:
        release = _release_after_peak(time, smooth, peak_index, hold, settings)
        if release is not None:
            return release
    return None


def _peak_candidates(time: np.ndarray, smooth: np.ndarray, t_guard: float, search_end: float) -> list[int]:
    post_guard = np.flatnonzero((time >= t_guard) & np.isfinite(smooth))
    if len(post_guard) < 2:
        return []
    preferred = post_guard[time[post_guard] <= search_end]
    candidates: list[int] = []
    for search in (preferred, post_guard):
        if len(search) < 2:
            continue
        peak_index = int(search[int(np.nanargmax(smooth[search]))])
        if peak_index not in candidates:
            candidates.append(peak_index)
    return candidates


def _release_after_peak(
    time: np.ndarray,
    smooth: np.ndarray,
    peak_index: int,
    hold: int,
    settings: ResonanceSettings,
) -> float | None:
    if peak_index >= len(time) - 2:
        return None
    peak = float(smooth[peak_index])
    if not math.isfinite(peak) or peak <= 0:
        return None
    last = len(time) - hold + 1
    for index in range(peak_index + 1, max(peak_index + 1, last)):
        window = smooth[index : index + hold]
        if len(window) < hold:
            break
        if smooth[index] <= peak * settings.release_decay_ratio and np.nanmax(window) <= peak * settings.release_rebound_ratio:
            return float(time[index])
    return None


def _growth_windows(
    time: np.ndarray,
    smooth: np.ndarray,
    t_start: float,
    settings: ResonanceSettings,
) -> tuple[np.ndarray, np.ndarray] | None:
    t_end = float(time[-1])
    duration = t_end - t_start
    if duration <= 0:
        return None
    window_s = max(duration * settings.growth_window_fraction, _median_dt(time) * settings.rolling_min_samples)
    tail_start = t_end - window_s
    previous_start = tail_start - window_s
    if previous_start < t_start:
        return None
    previous = (time >= previous_start) & (time < tail_start) & np.isfinite(smooth)
    tail = (time >= tail_start) & np.isfinite(smooth)
    if previous.sum() < 2 or tail.sum() < 2:
        return None
    return previous, tail


def _passes_relevance(
    settings: ResonanceSettings,
    positive_fraction: float,
    level_over_vlim: float,
    growth_ratio: float,
    growth_delta: float,
    vlim: float,
) -> bool:
    return positive_fraction >= settings.min_positive_fraction and (
        level_over_vlim >= settings.min_level_over_vlim
        or (
            growth_ratio >= settings.min_growth_ratio
            and growth_delta >= settings.min_growth_delta_factor * vlim
        )
    )


def _log_slope(time: np.ndarray, values: np.ndarray, vlim: float, settings: ResonanceSettings) -> float:
    finite = np.isfinite(time) & np.isfinite(values)
    if finite.sum() < 2:
        return 0.0
    t = time[finite]
    y = values[finite]
    floor = max(settings.log_floor_absolute, settings.log_floor_vlim_factor * vlim)
    if float(t[-1]) == float(t[0]):
        return 0.0
    return float(np.polyfit(t - t[0], np.log(np.maximum(y, floor)), 1)[0])


def _positive_fraction(values: np.ndarray) -> float:
    diffs = np.diff(values[np.isfinite(values)])
    return float(np.mean(diffs > 0)) if len(diffs) else 0.0


def _positive_window_count(values: np.ndarray) -> int:
    best = current = 0
    for grows in np.diff(values[np.isfinite(values)]) > 0:
        current = current + 1 if grows else 0
        best = max(best, current)
    return best


def _area_above(time: np.ndarray, values: np.ndarray, vlim: float) -> float:
    if len(time) < 2:
        return 0.0
    return float(np.trapezoid(np.maximum(0.0, values - vlim), time))


def _time_above(time: np.ndarray, above: np.ndarray) -> float:
    if len(time) < 2 or not np.any(above):
        return 0.0
    dt = _median_dt(time)
    return float(np.count_nonzero(above) * dt)


def _p95(values: np.ndarray) -> float:
    finite = values[np.isfinite(values)]
    return float(np.percentile(finite, 95)) if len(finite) else 0.0


def _median_dt(time: np.ndarray) -> float:
    diffs = np.diff(time)
    finite = diffs[np.isfinite(diffs) & (diffs > 0)]
    return float(np.median(finite)) if len(finite) else 1.0


def _float(value: Any, default: float) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return number if math.isfinite(number) else default


def analysis_guard_time(event_times: dict[str, float] | None) -> float:
    event_time = _float((event_times or {}).get("SFO"), 0.004)
    tov_time = _float((event_times or {}).get("TOV"), 0.03)
    return max(0.0, event_time + tov_time)


def _split_case_run(text: str) -> tuple[str, int]:
    case, _, run_text = text.rpartition("-")
    if not case:
        return text, 0
    try:
        return case, int(float(run_text))
    except ValueError:
        return text, 0


def _write_settings_sheet(ws, settings: ResonanceSettings, event_times: dict[str, float] | None) -> None:
    ws.append(["Setting", "Value"])
    values = {
        "enabled_checks": ", ".join(settings.effective_enabled_checks),
        "top_n": settings.top_n,
        "limit_multiplier": settings.limit_multiplier,
        "auto_detect_release": settings.auto_release,
        "manual_analysis_start_s": settings.manual_analysis_start,
        "SFO_time_s": (event_times or {}).get("SFO", ""),
        "TOV_time_s": (event_times or {}).get("TOV", ""),
        "analysis_guard_s": analysis_guard_time(event_times),
        "peak_search_fraction": settings.peak_search_fraction,
        "release_decay_ratio": settings.release_decay_ratio,
        "release_rebound_ratio": settings.release_rebound_ratio,
        "release_hold_time_s": settings.release_hold_time,
        "rolling_p95_window_s": settings.rolling_p95_window,
        "rolling_min_samples": settings.rolling_min_samples,
        "log_floor_vlim_factor": settings.log_floor_vlim_factor,
        "log_floor_absolute": settings.log_floor_absolute,
        "growth_window_fraction": settings.growth_window_fraction,
        "min_positive_fraction": settings.min_positive_fraction,
        "min_growth_ratio": settings.min_growth_ratio,
        "min_level_over_vlim": settings.min_level_over_vlim,
        "min_growth_delta_factor": settings.min_growth_delta_factor,
    }
    for key, value in values.items():
        ws.append([key, value])


def _write_result_sheet(ws, rows: list[ResonanceResult]) -> None:
    metric_keys = []
    for row in rows:
        for key in row.metrics:
            if key not in metric_keys:
                metric_keys.append(key)
    headers = ["Rank", "kV", "Case", "Run", "MM", "Vnom", "Vlim", "t_guard", "t_start", *metric_keys, "Waveform folder"]
    ws.append(headers)
    for result in sorted(rows, key=lambda item: (_voltage_sort_key(item.voltage), item.rank)):
        folder, voltage_type = report_folder(result.check, result.voltage_type)
        ws.append(
            [
                result.rank,
                result.voltage,
                result.case,
                result.run,
                result.mm_name,
                result.vnom,
                result.vlim,
                result.t_guard,
                result.t_start,
                *[result.metrics.get(key, "") for key in metric_keys],
                f"Plots/Generated/{result.scope_folder}/{folder}/{voltage_type}",
            ]
        )


def _write_chart_data_sheet(wb: Workbook, result: ResonanceResult) -> None:
    base = f"{CHECK_DEFINITIONS[result.check]['abbr']}_{result.voltage_type}_{result.voltage}_{result.rank}"
    title = _unique_sheet_name(wb, base[:31])
    ws = wb.create_sheet(title)
    ws.append(["Time (s)", "E(t)", "E_s(t)", "Vlim"])
    for time_s, value, smooth_value in zip(result.time, result.envelope, result.smooth):
        ws.append([float(time_s), float(value), float(smooth_value), float(result.vlim)])
    ws["F1"] = "Case"
    ws["G1"] = result.case
    ws["F2"] = "Run"
    ws["G2"] = result.run
    ws["F3"] = "MM"
    ws["G3"] = result.mm_name
    ws["F4"] = "Start"
    ws["G4"] = result.t_start
    ws["F5"] = "Main metric"
    ws["G5"] = result.main_metric
    ws["F6"] = "Vlim"
    ws["G6"] = result.vlim



def _unique_sheet_name(wb: Workbook, base: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_]+", "_", base).strip("_") or "Chart"
    if name not in wb.sheetnames:
        return name
    for index in range(2, 100):
        suffix = f"_{index}"
        candidate = f"{name[:31 - len(suffix)]}{suffix}"
        if candidate not in wb.sheetnames:
            return candidate
    return name[:28] + "_99"


def _voltage_sort_key(value: str) -> tuple[int, float | str]:
    try:
        return 0, float(value)
    except ValueError:
        return 1, value


def _format_workbook(wb: Workbook) -> None:
    for ws in wb.worksheets:
        if ws.max_column:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"
        for col_idx in range(1, ws.max_column + 1):
            header = ws.cell(1, col_idx).value
            cell = ws.cell(1, col_idx)
            cell.font = Font(name="Arial", bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(32, len(str(header or "")) + 2))
