from __future__ import annotations


def test_project_timing_reads_input_data_labels(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.project_config import load_project_timing

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Input_Data"
    sheet.append(["Parameter", "Value"])
    sheet.append(["Frequency", 60])
    sheet.append(["Final duration", 0.5])
    workbook.save(tmp_path / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    timing = load_project_timing(tmp_path)

    assert timing.frequency == 60.0
    assert timing.final_duration == 0.5


def test_project_scan_cache_keeps_project_timing(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app import project_scan_cache, scanner

    (tmp_path / "Case_folder").mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Input_Data"
    sheet.append(["Frequency", 60])
    sheet.append(["Final duration", 0.75])
    workbook.save(tmp_path / "Input_Data_PSCAD_test.xlsx")
    workbook.close()
    project_path = str(tmp_path.resolve())
    cache_path = tmp_path / "cache.json"

    scan = scanner.scan_project(tmp_path)
    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        (450.0, 250.0),
        cache_path,
    )
    cached = project_scan_cache.cached_scan(
        project_scan_cache.load(cache_path),
        project_path,
        (450.0, 250.0),
    )

    assert cached is not None
    assert cached.project_frequency == 60.0
    assert cached.final_duration == 0.75


def test_envelope_descriptor_cache_deduplicates_identical_inf(tmp_path, monkeypatch) -> None:
    from results_analysis_app import voltage_envelope

    inf_content = 'PGB(1) Desc="MM_66_LGp_A" Group="MM_66_BUS1" Units="kV"\n'
    inf_one = tmp_path / "C1_r00001.inf"
    inf_two = tmp_path / "C2_r00001.inf"
    inf_one.write_text(inf_content, encoding="utf-8")
    inf_two.write_text(inf_content, encoding="utf-8")

    original_parser = voltage_envelope.parse_inf_descriptors
    parse_count = 0

    def counting_parser(path):
        nonlocal parse_count
        parse_count += 1
        return original_parser(path)

    monkeypatch.setattr(voltage_envelope, "parse_inf_descriptors", counting_parser)
    cache = voltage_envelope._read_inf_descriptor_cache([inf_one, inf_two], 2, None)

    assert parse_count == 1
    assert cache[inf_one] is cache[inf_two]


def test_output_tree_does_not_precreate_generated_plot_folders(tmp_path) -> None:
    from results_analysis_app.actions import ensure_output_tree
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    ensure_output_tree(project, [ScopeEntry.full()])

    assert (project / "Voltage_envelope" / "Full").is_dir()
    assert (project / "Reports" / "Full").is_dir()
    assert (project / "Plots" / "Plot_batch").is_dir()
    assert not (project / "Plots" / "Generated").exists()


def test_run_analysis_pipeline_passes_project_dashboard_mapping_to_reports(
    tmp_path, monkeypatch
) -> None:
    from results_analysis_app import actions
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    project_key = str(project.resolve())
    captured: dict[str, object] = {}

    monkeypatch.setattr(actions, "build_voltage_envelopes", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(actions, "create_plot_batches", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(actions, "render_plot_batches", lambda *_args, **_kwargs: None)

    def capture_reports(*_args, **kwargs):
        captured["dashboard_figure_ids_by_project"] = kwargs[
            "dashboard_figure_ids_by_project"
        ]
        return []

    monkeypatch.setattr(actions, "build_reports_from_existing_plots", capture_reports)

    actions.run_analysis_pipeline(
        [project],
        [ScopeEntry.full()],
        ["66"],
        [],
        dashboard_figure_ids_by_project={project_key: ["dashboard.xlsx|Graphs|1|Figure"]},
    )

    assert captured["dashboard_figure_ids_by_project"] == {
        project_key: ["dashboard.xlsx|Graphs|1|Figure"]
    }


def test_run_analysis_pipeline_reuses_sustained_cache_validation(
    tmp_path, monkeypatch
) -> None:
    from results_analysis_app import actions, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    payload = {"version": sustained_sdpf.RESULT_VERSION}
    validation = sustained_sdpf.SustainedSDPFCacheValidation(
        valid=True,
        shared_manifest_current=True,
    )
    validation_calls = 0
    captured: dict[str, dict[str, object]] = {}

    monkeypatch.setattr(actions, "build_voltage_envelopes", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(
        actions.sustained_sdpf,
        "load_results",
        lambda *_args, **_kwargs: payload,
    )

    def validate(*_args, **_kwargs):
        nonlocal validation_calls
        validation_calls += 1
        return validation

    monkeypatch.setattr(actions.sustained_sdpf, "validate_result_cache", validate)
    def capture_create(*_args, **kwargs):
        captured["create"] = kwargs
        return []

    def capture_render(*_args, **kwargs):
        captured["render"] = kwargs

    def capture_report(*_args, **kwargs):
        captured["report"] = kwargs
        return []

    monkeypatch.setattr(actions, "create_plot_batches", capture_create)
    monkeypatch.setattr(actions, "render_plot_batches", capture_render)
    monkeypatch.setattr(actions, "build_reports_from_existing_plots", capture_report)

    actions.run_analysis_pipeline(
        [project],
        [ScopeEntry.full()],
        ["66"],
        [],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 30.0},
    )

    assert validation_calls == 1
    for stage in ("create", "render"):
        assert captured[stage]["sustained_cache_validations_by_scope"] == {
            "Full": validation
        }
    assert captured["report"]["sustained_cache_validations_by_project_scope"] == {
        str(project.resolve()): {"Full": validation}
    }


def test_rebuild_envelope_charts_uses_existing_workbooks(tmp_path, monkeypatch) -> None:
    from results_analysis_app import actions
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    input_path = envelope_dir / "MM_66.xlsx"
    input_path.write_bytes(b"placeholder")

    class DummyExcelApp:
        def __enter__(self):
            return object()

        def __exit__(self, _exc_type, _exc, _tb):
            return False

    calls = []

    def fake_create_combined(input_file, output_file, excel, **kwargs):
        calls.append((input_file, output_file, excel, kwargs))
        output_file.write_bytes(b"chart")
        return output_file

    monkeypatch.setattr(actions, "excel_app", DummyExcelApp)
    monkeypatch.setattr(actions, "create_combined_envelope_plot", fake_create_combined)
    project_key = str(project.resolve())

    outputs = actions.rebuild_envelope_charts(
        [project],
        [ScopeEntry.full()],
        ["66", "230"],
        event_times={"SFO": 0.005},
        envelope_chart_x_max_by_project={project_key: 0.7},
        envelope_chart_x_major_by_project={project_key: 0.1},
        envelope_chart_y_limits_by_voltage={"66": {"y_min": 40.0}},
        envelope_chart_show_sa_label=True,
        envelope_chart_top_left_cell="J2",
        envelope_chart_width=800.0,
        envelope_chart_height=400.0,
    )

    assert outputs == [envelope_dir / "MM_66_with_combined_plot.xlsx"]
    assert len(calls) == 1
    assert calls[0][0] == input_path
    assert calls[0][1] == envelope_dir / "MM_66_with_combined_plot.xlsx"
    assert calls[0][3]["axis_limits_override"] == {"x_max": 0.7, "x_major": 0.1}
    assert calls[0][3]["axis_limits_by_voltage"] == {"66": {"y_min": 40.0}}
    assert calls[0][3]["event_times"] == {"SFO": 0.005}
    assert calls[0][3]["show_sa_label"] is True
    assert calls[0][3]["chart_top_left_cell"] == "J2"
    assert calls[0][3]["chart_size"] == {"width": 800.0, "height": 400.0}


def test_corrupt_autosave_falls_back_to_default(tmp_path, monkeypatch) -> None:
    from results_analysis_app import storage

    autosave_path = tmp_path / "last_session.json"
    autosave_path.write_text("{not valid json", encoding="utf-8")
    monkeypatch.setattr(storage, "AUTOSAVE_PATH", autosave_path)

    session = storage.load_autosave()

    assert session.projects == []
    assert session.scopes[0].mode == "full"


def test_semantically_malformed_autosave_falls_back_to_valid_defaults(tmp_path, monkeypatch) -> None:
    import json

    from results_analysis_app import storage

    autosave_path = tmp_path / "last_session.json"
    autosave_path.write_text(
        json.dumps(
            {
                "projects": None,
                "scopes": None,
                "events": None,
                "voltages": None,
                "status_cache": None,
                "dashboard_figure_selection": None,
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(storage, "AUTOSAVE_PATH", autosave_path)

    session = storage.load_autosave()

    assert session.projects == []
    assert session.scopes[0].mode == "full"
    assert session.events == ["SFO", "TOV", "SA"]


def test_combined_envelope_chart_failure_keeps_previous_output(tmp_path, monkeypatch) -> None:
    import pytest

    from results_analysis_app import envelope_chart

    source = tmp_path / "MM_66.xlsx"
    output = tmp_path / "MM_66_with_combined_plot.xlsx"
    source.write_bytes(b"new base")
    output.write_bytes(b"previous good chart")

    def fail(*_args, **_kwargs):
        raise RuntimeError("Excel failed")

    monkeypatch.setattr(envelope_chart, "_create_combined_envelope_plot_with_excel", fail)

    with pytest.raises(RuntimeError, match="Excel failed"):
        envelope_chart.create_combined_envelope_plot(source, output, excel=object())

    assert output.read_bytes() == b"previous good chart"
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_fast_array_envelope_matches_dataframe_path() -> None:
    import numpy as np
    import pandas as pd

    from results_analysis_app.voltage_envelope import _apply_envelope, _apply_envelope_arrays

    time = np.round(np.arange(0.0, 0.1, 0.0002), 6)
    phases = [
        np.sin(2 * np.pi * 50 * time) * 100,
        np.sin(2 * np.pi * 50 * time + 2.1) * 80,
        np.zeros_like(time),
    ]
    source = pd.DataFrame(
        {
            "Time (s)": time,
            "A": phases[0],
            "B": phases[1],
            "C": phases[2],
        }
    )

    slow = _apply_envelope(source, time_step=0.002, time_end=0.05)
    fast = _apply_envelope_arrays(time, phases, time_step=0.002, time_end=0.05)
    fast_reused = _apply_envelope_arrays(
        time,
        phases,
        time_step=0.002,
        time_end=0.05,
        absolute_phase_values=[np.abs(phase) for phase in phases],
    )

    pd.testing.assert_frame_equal(fast, slow)
    pd.testing.assert_frame_equal(fast_reused, fast)


def test_envelope_processing_stops_at_available_waveform_end() -> None:
    import numpy as np
    import pandas as pd

    from results_analysis_app.voltage_envelope import _apply_envelope, _apply_envelope_arrays

    time = np.round(np.arange(0.0, 0.101, 0.001), 6)
    phases = [np.full_like(time, 10.0), np.full_like(time, 8.0)]
    source = pd.DataFrame({"Time (s)": time, "A": phases[0], "B": phases[1]})

    slow = _apply_envelope(source, time_step=0.01, time_end=0.2)
    fast = _apply_envelope_arrays(time, phases, time_step=0.01, time_end=0.2)

    pd.testing.assert_frame_equal(fast, slow)
    assert fast.index[-1] == 0.09
    assert fast.notna().all().all()
    assert (fast["Max_A"] == 10.0).all()


def test_chart_x_max_is_bounded_by_available_data() -> None:
    from results_analysis_app.envelope_chart import _bounded_x_axis

    assert _bounded_x_axis({"x_max": 1.0}, 0.5)["x_max"] == 0.5
    assert _bounded_x_axis({"x_max": 0.2}, 0.5)["x_max"] == 0.2
    assert _bounded_x_axis({"x_max": None}, 0.5)["x_max"] == 0.5
    assert _bounded_x_axis({"x_max": None, "x_major": None}, 0.5)["x_major"] == 0.05


def test_automatic_time_major_uses_readable_ten_interval_step() -> None:
    from results_analysis_app.project_config import automatic_time_major

    assert automatic_time_major(0.5) == 0.05
    assert automatic_time_major(0.3) == 0.025
    assert automatic_time_major(2.0) == 0.2
    assert automatic_time_major(None) is None


def test_envelope_inf_inventory_is_filtered_before_scope_selection(tmp_path) -> None:
    from results_analysis_app.exclusions import ExclusionMatcher, ExclusionRule
    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.voltage_envelope import _inf_inventory, _selected_inf_paths

    case_root = tmp_path / "Case_folder"
    case_root.mkdir()
    first = case_root / "C1_r00001.inf"
    second = case_root / "C2_r00001.inf"
    first.write_text("", encoding="utf-8")
    second.write_text("", encoding="utf-8")

    inventory = _inf_inventory(
        case_root,
        ExclusionMatcher([ExclusionRule(case="C2", run=1)]),
    )

    assert inventory == [first]
    assert _selected_inf_paths(inventory, ScopeEntry.full()) == [first]


def test_envelope_workbook_is_formatted_during_write(tmp_path) -> None:
    from openpyxl import load_workbook
    import pandas as pd

    from results_analysis_app.voltage_envelope import HIGH_VOLTAGE_COLUMNS, NONCONV_COLUMNS, _write_workbook

    path = tmp_path / "MM_66.xlsx"
    data = pd.DataFrame({"Time (s)": [0.0], "Max_A": [1.0]})
    _write_workbook(
        path,
        data,
        data,
        pd.DataFrame(columns=NONCONV_COLUMNS),
        pd.DataFrame(columns=HIGH_VOLTAGE_COLUMNS),
    )

    workbook = load_workbook(path, read_only=False)
    try:
        assert workbook["LGp"].freeze_panes == "A2"
        assert workbook["LGp"].auto_filter.ref == "A1:B1"
        assert [cell.value for cell in workbook["High voltage exclusions"][1]] == HIGH_VOLTAGE_COLUMNS
    finally:
        workbook.close()


def test_autofit_workbook_uses_excel_native_autofit(tmp_path) -> None:
    from results_analysis_app.excel import autofit_workbook

    class DummyColumns:
        def __init__(self) -> None:
            self.calls = 0

        def AutoFit(self) -> None:
            self.calls += 1

    class DummyWorksheet:
        def __init__(self) -> None:
            self.columns = DummyColumns()
            self.UsedRange = type("UsedRange", (), {"Columns": self.columns})()

    class DummyWorkbook:
        def __init__(self) -> None:
            self.Worksheets = [DummyWorksheet(), DummyWorksheet()]
            self.saved = False
            self.closed = False

        def Save(self) -> None:
            self.saved = True

        def Close(self, SaveChanges: bool) -> None:
            self.closed = SaveChanges

    class DummyWorkbooks:
        def __init__(self, workbook) -> None:
            self.workbook = workbook
            self.open_args = None

        def Open(self, *args, **kwargs):
            self.open_args = (args, kwargs)
            return self.workbook

    workbook = DummyWorkbook()
    excel = type("Excel", (), {})()
    excel.Workbooks = DummyWorkbooks(workbook)
    path = tmp_path / "MM_66.xlsx"

    autofit_workbook(excel, path)

    assert excel.Workbooks.open_args[0] == (str(path.resolve()),)
    assert excel.Workbooks.open_args[1] == {"UpdateLinks": 0, "ReadOnly": False}
    assert [sheet.columns.calls for sheet in workbook.Worksheets] == [1, 1]
    assert workbook.saved is True
    assert workbook.closed is True


def test_analysis_high_voltage_violation_excludes_the_bus_from_both_envelopes(
    tmp_path,
    monkeypatch,
) -> None:
    import pandas as pd

    from results_analysis_app import voltage_envelope
    from results_analysis_app.exclusions import ExclusionMatcher

    inf_path = tmp_path / "C1_r00001.inf"
    inf_path.write_text("", encoding="utf-8")
    descriptors = pd.DataFrame(
        {
            "PGB": [1, 2],
            "Description": ["MM_LGp_a", "MM_LLp_a"],
            "Group": ["MM_66_A", "MM_66_A"],
        }
    )
    monkeypatch.setattr(voltage_envelope, "_parse_inf", lambda *_args: descriptors)
    monkeypatch.setattr(voltage_envelope, "_read_out_files", lambda *_args: {})

    def build_bus(_df, bus, *_args):
        measurement = _args[5]
        exclusions = []
        if measurement == "LGp":
            exclusions.append(
                {
                    "Case": "C1",
                    "Run": 1,
                    "MM_name": bus,
                    "Measurement": measurement,
                    "Signal": "Va",
                    "File": "C1_01.out",
                    "Excluded_values": 1,
                    "Max_abs": 600.0,
                    "Limit": 500.0,
                    "Source": "Analysis",
                }
            )
        return object(), exclusions

    monkeypatch.setattr(voltage_envelope, "_read_bus_waveforms", build_bus)
    monkeypatch.setattr(
        voltage_envelope,
        "_apply_bus_waveforms",
        lambda *_args: pd.DataFrame({"Time (s)": [0.0], "Max_A": [1.0]}),
    )

    entries, exclusions = voltage_envelope._read_run_entries(
        inf_path,
        500.0,
        "66",
        "MM_66",
        ExclusionMatcher(),
    )

    assert entries == []
    assert len(exclusions) == 1

    entries, exclusions = voltage_envelope._read_run_entries(
        inf_path,
        500.0,
        "66",
        "MM_66",
        ExclusionMatcher(),
        high_voltage_include_overrides={
            voltage_envelope._high_voltage_key("66", "C1", 1, "MM_66_A")
        },
    )

    assert [measurement for measurement, _df in entries] == ["LGp", "LLp"]
    assert exclusions == []


def test_high_voltage_sheet_preserves_per_signal_analysis_rows() -> None:
    import pandas as pd

    from results_analysis_app.voltage_envelope import HIGH_VOLTAGE_COLUMNS, _map_exclusion_fault_types

    detected = [
        {
            "Case": "C1",
            "Run": 1,
            "MM_name": "MM_66_A",
            "Measurement": "LGp",
            "Signal": "Va",
            "File": "C1_01.out",
            "Excluded_values": 2,
            "Max_abs": 600.0,
            "Limit": 500.0,
            "Source": "Analysis",
        },
        {
            "Case": "C1",
            "Run": 1,
            "MM_name": "MM_66_A",
            "Measurement": "LLp",
            "Signal": "Vb",
            "File": "C1_02.out",
            "Excluded_values": 3,
            "Max_abs": 650.0,
            "Limit": 500.0,
        },
    ]

    result = _map_exclusion_fault_types(
        pd.DataFrame(detected),
        pd.DataFrame(),
        HIGH_VOLTAGE_COLUMNS,
    )

    assert list(result.columns) == HIGH_VOLTAGE_COLUMNS
    assert result["Case"].tolist() == ["C1", "C1"]
    assert result["Measurement"].tolist() == ["LGp", "LLp"]
    assert result["Signal"].tolist() == ["Va", "Vb"]
    assert result["Excluded_values"].tolist() == [2, 3]
    assert result["Max_abs"].tolist() == [600.0, 650.0]


def test_frequency_detection_tries_later_phases_before_fallback() -> None:
    import numpy as np

    from results_analysis_app.voltage_envelope import _detect_frequency_from_arrays

    time = np.round(np.arange(0.0, 0.1, 0.0002), 6)
    no_crossings = np.ones_like(time)
    valid_60hz = np.sin(2 * np.pi * 60 * time)
    fallback_events: list[str] = []

    detected = _detect_frequency_from_arrays(
        time,
        [no_crossings, valid_60hz],
        50.0,
        "case run bus",
        fallback_events.append,
    )

    assert detected == 60.0
    assert fallback_events == []

    detected = _detect_frequency_from_arrays(
        time,
        [no_crossings],
        50.0,
        "case run bus",
        fallback_events.append,
    )

    assert detected == 50.0
    assert fallback_events == ["case run bus"]


def test_voltage_run_reads_cancel_queued_reads(tmp_path, monkeypatch) -> None:
    import pytest

    from results_analysis_app import voltage_envelope
    from results_analysis_app.exclusions import ExclusionMatcher

    paths = [tmp_path / f"C1_r{i:05d}.inf" for i in range(20)]
    state = {"cancelled": False, "calls": 0}

    def check_cancel() -> None:
        if state["cancelled"]:
            raise RuntimeError("Operation stopped by user.")

    def fake_read_run_entries(*args, **_kwargs):
        state["calls"] += 1
        state["cancelled"] = True
        next(arg for arg in reversed(args) if callable(arg))()
        return [], []

    monkeypatch.setattr(voltage_envelope, "_read_run_entries", fake_read_run_entries)

    with pytest.raises(RuntimeError, match="Operation stopped by user"):
        voltage_envelope._read_voltage_runs(
            paths,
            "66",
            "MM_66",
            72.5,
            ExclusionMatcher(),
            1,
            None,
            check_cancel,
        )

    assert state["calls"] == 1


def test_voltage_run_process_worker_returns_envelope_data(tmp_path) -> None:
    from concurrent.futures import ProcessPoolExecutor

    import numpy as np

    from results_analysis_app import voltage_envelope
    from results_analysis_app.exclusions import ExclusionMatcher
    from pscad_plotter_app_v3.services.waveform_io import parse_inf_descriptors

    inf_path = tmp_path / "C1_r00001.inf"
    inf_path.write_text(
        '\n'.join(
            [
                'PGB(1) Output Desc="MM_LGp_a" Group="MM_66_A" Units="kV"',
                'PGB(2) Output Desc="MM_LLp_a" Group="MM_66_A" Units="kV"',
            ]
        ),
        encoding="utf-8",
    )
    time_values = np.arange(0.0, 0.1, 0.0002)
    out_rows = [
        " ".join(f"{value:.12g}" for value in (time, np.sin(2 * np.pi * 50 * time), np.cos(2 * np.pi * 50 * time)))
        for time in time_values
    ]
    out_path = inf_path.with_name("C1_r00001_01.out")
    out_path.write_text("header\n" + "\n".join(out_rows), encoding="utf-8")
    descriptors = parse_inf_descriptors(inf_path)

    with ProcessPoolExecutor(max_workers=1) as executor:
        cache = voltage_envelope._read_voltage_runs(
            [inf_path],
            "66",
            "MM_66",
            72.5,
            ExclusionMatcher(),
            1,
            None,
            None,
            time_end=0.1,
            inf_descriptor_cache={inf_path: descriptors},
            executor=executor,
            process_pool=True,
        )

    entries, exclusions = cache[inf_path]
    assert [measurement for measurement, _df in entries] == ["LGp", "LLp"]
    assert exclusions == []


def test_envelope_manifest_validates_artifacts_and_return_outputs(tmp_path) -> None:
    from results_analysis_app import voltage_envelope

    project_root = tmp_path / "Project"
    output_dir = project_root / "Voltage_envelope" / "Full"
    output_dir.mkdir(parents=True)
    returned = output_dir / "MM_66_with_combined_plot.xlsx"
    auxiliary = output_dir / "MM_66.xlsx"
    returned.write_bytes(b"chart")
    auxiliary.write_bytes(b"base")

    voltage_envelope._write_envelope_manifest(
        project_root,
        "signature",
        [returned, auxiliary],
        return_outputs=[returned],
    )

    assert voltage_envelope._envelope_manifest_matches(project_root, "signature") == [returned]
    auxiliary.unlink()
    assert voltage_envelope._envelope_manifest_matches(project_root, "signature") is None


def test_envelope_source_inventory_matches_per_file_manifest(tmp_path) -> None:
    from results_analysis_app import voltage_envelope

    project_root = tmp_path / "Project"
    source_dir = project_root / "Case_folder" / "C1.1"
    source_dir.mkdir(parents=True)
    inf_path = source_dir / "C1_r00001.inf"
    inf_path.write_text("descriptor", encoding="utf-8")
    matching_out = source_dir / "C1_r00001_01.out"
    matching_out.write_text("waveform", encoding="utf-8")
    (source_dir / "C1_r00002_01.out").write_text("other", encoding="utf-8")

    inventory = voltage_envelope._source_file_inventory(project_root, [inf_path])

    assert voltage_envelope._run_source_manifest(project_root, inf_path) == (
        voltage_envelope._run_source_manifest(project_root, inf_path, inventory)
    )


def test_envelope_source_inventory_indexes_case_prefix_outputs(tmp_path) -> None:
    from results_analysis_app import voltage_envelope

    project_root = tmp_path / "Project"
    source_dir = project_root / "Case_folder" / "C1.1"
    source_dir.mkdir(parents=True)
    inf_path = source_dir / "C1_r00001.inf"
    inf_path.write_text("descriptor", encoding="utf-8")
    expected = {
        "C1_r00001.out",
        "C1_r00001_01.out",
        "C1_r000010_extra.out",
    }
    for name in expected | {"C1_r00002_01.out"}:
        (source_dir / name).write_text("waveform", encoding="utf-8")

    inventory = voltage_envelope._source_file_inventory(project_root, [inf_path])

    assert {
        path.name for path in inventory.output_paths(source_dir, inf_path.stem)
    } == expected


def test_sustained_cache_signature_tracks_summary_workbook_version() -> None:
    from results_analysis_app import sustained_sdpf, voltage_envelope

    settings = sustained_sdpf.SustainedSDPFSettings(enabled=True)
    versions = voltage_envelope._sustained_cache_versions(settings)

    assert versions == {
        "sustained_sdpf_result_version": sustained_sdpf.RESULT_VERSION,
        "sustained_sdpf_summary_version": sustained_sdpf.SUMMARY_WORKBOOK_VERSION,
    }
    original_signature = voltage_envelope._cache_signature({"settings": versions})
    versions["sustained_sdpf_summary_version"] += 1
    assert voltage_envelope._cache_signature({"settings": versions}) != original_signature
    assert voltage_envelope._sustained_cache_versions(
        sustained_sdpf.SustainedSDPFSettings(enabled=False)
    ) == {}


def test_nonconvergent_scan_process_workers_preserve_results(tmp_path) -> None:
    import pandas as pd

    from results_analysis_app import voltage_envelope

    case_root = tmp_path / "Case_folder" / "C1.1"
    case_root.mkdir(parents=True)
    (case_root / "CB_01_0001.out").write_text(
        "metadata\nRun# CB_IIp CB_IIr\n1 600 300\n2 10 20\n",
        encoding="utf-8",
    )

    result = voltage_envelope._find_non_convergent_cases(
        tmp_path / "Case_folder",
        pd.DataFrame(columns=["Case", "Run#", "Fault_type"]),
        500.0,
        250.0,
        worker_count=2,
    )

    assert result[["Case", "Run", "Signal", "Value"]].to_dict("records") == [
        {"Case": "C1", "Run": 1, "Signal": "CB_IIp", "Value": "600"},
        {"Case": "C1", "Run": 1, "Signal": "CB_IIr", "Value": "300"},
    ]


def test_automatic_envelope_workers_use_eighty_percent_with_safe_cap(monkeypatch) -> None:
    from results_analysis_app import voltage_envelope

    monkeypatch.setattr(voltage_envelope.os, "process_cpu_count", lambda: 32, raising=False)
    assert voltage_envelope._configured_worker_count(0) == 26

    monkeypatch.setattr(voltage_envelope.os, "process_cpu_count", lambda: 128, raising=False)
    assert voltage_envelope._configured_worker_count(0) == 60

    monkeypatch.setattr(voltage_envelope.os, "process_cpu_count", lambda: None, raising=False)
    monkeypatch.setattr(voltage_envelope.os, "cpu_count", lambda: 8)
    assert voltage_envelope._configured_worker_count(0) == 7

    monkeypatch.setattr(voltage_envelope.os, "cpu_count", lambda: 6)
    assert voltage_envelope._configured_worker_count(0) == 5

    monkeypatch.setattr(voltage_envelope.os, "process_cpu_count", lambda: 1, raising=False)
    assert voltage_envelope._configured_worker_count(0) == 1
    assert voltage_envelope._configured_worker_count(4) == 4
    assert voltage_envelope._configured_worker_count(100) == 60


def test_envelope_build_reuses_shared_run_executor(tmp_path, monkeypatch) -> None:
    import pandas as pd

    from results_analysis_app import voltage_envelope
    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.project_config import ProjectTiming

    project = tmp_path / "Project"
    (project / "Case_folder").mkdir(parents=True)
    inf_path = project / "Case_folder" / "C1_r00001.inf"
    inf_path.write_text("", encoding="utf-8")

    monkeypatch.setattr(
        voltage_envelope,
        "_read_stat_files",
        lambda _root, **_kwargs: pd.DataFrame(),
    )
    monkeypatch.setattr(voltage_envelope, "_find_non_convergent_cases", lambda *_args, **_kwargs: pd.DataFrame())
    monkeypatch.setattr(voltage_envelope, "_selected_inf_paths", lambda *_args: [inf_path])
    monkeypatch.setattr(voltage_envelope, "_read_inf_descriptor_cache", lambda *_args: {})
    monkeypatch.setattr(
        voltage_envelope,
        "load_project_timing",
        lambda _root: ProjectTiming(final_duration=2.5),
    )
    monkeypatch.setattr(voltage_envelope, "load_voltage_configs", lambda *_args, **_kwargs: {})

    calls = []

    def fake_build(*args, **kwargs):
        calls.append((args[5], args[12], args[16], args[22]))
        return voltage_envelope._VoltageBuildResult([], [], [])

    monkeypatch.setattr(voltage_envelope, "_build_voltage_workbooks", fake_build)

    logs = []
    voltage_envelope.build_voltage_envelopes(
        project,
        [ScopeEntry.full()],
        ["66", "161"],
        envelope_workers=4,
        build_charts=False,
        log=logs.append,
    )

    assert any(
        "logical CPUs=" in message and "shared run-read pool=4" in message
        for message in logs
    )
    assert sorted(call[0] for call in calls) == ["161", "66"]
    assert [call[1] for call in calls] == [2.5, 2.5]
    assert [call[2] for call in calls] == [4, 4]
    assert len({id(call[3]) for call in calls}) == 1

    calls.clear()
    voltage_envelope.build_voltage_envelopes(
        project,
        [ScopeEntry.full()],
        ["66"],
        envelope_workers=4,
        envelope_time_end=0.75,
        build_charts=False,
    )

    assert [call[1] for call in calls] == [0.75]


def test_envelope_merge_ignores_shorter_runs_after_their_end() -> None:
    import pandas as pd

    from results_analysis_app.voltage_envelope import _reduce_merge

    short = pd.DataFrame(
        {
            "Max_A": [10.0, 9.0],
            "MM_name": ["MM_short"] * 2,
            "Case": ["C1-1"] * 2,
        }
    )
    long = pd.DataFrame(
        {
            "Max_A": [8.0, 7.0, 6.0],
            "MM_name": ["MM_long"] * 3,
            "Case": ["C2-1"] * 3,
        }
    )

    result = _reduce_merge([short, long], time_step=0.5)

    assert result["Max_A"].tolist() == [10.0, 9.0, 6.0]
    assert not (result["Max_A"] == 0.0).any()


def test_envelope_reduction_preserves_ranked_source_provenance() -> None:
    import numpy as np
    import pandas as pd

    from results_analysis_app.voltage_envelope import _reduce_merge

    first = pd.DataFrame(
        {
            "Max_A": [10.0, 1.0, np.nan],
            "Max_B": [9.0, 5.0, np.nan],
            "MM_name": ["MM_1"] * 3,
            "Case": ["C1-1"] * 3,
        }
    )
    second = pd.DataFrame(
        {
            "Max_A": [9.0, 8.0, 7.0],
            "Max_B": [9.0, 7.0, np.nan],
            "MM_name": ["MM_2"] * 3,
            "Case": ["C2-2"] * 3,
        }
    )

    result = _reduce_merge([first, second], time_step=0.01)

    assert result["Max_A"].tolist() == [10.0, 8.0, 7.0]
    assert result["Case_A"].tolist() == ["C1", "C2", "C2"]
    assert result["Run_A"].tolist() == [1, 2, 2]
    assert result.loc[0, "Case_B"] == "C1"
    assert result.loc[1, "Case_B"] == "C2"
    assert np.isnan(result.loc[2, "Max_B"])
