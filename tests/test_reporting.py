from __future__ import annotations


def test_dashboard_voltage_slicer_mapping() -> None:
    from results_analysis_app.reporting import (
        _dashboard_figure_uses_saved_voltage_filter,
        _dashboard_slicer_value,
    )

    assert _dashboard_slicer_value("22") == "23"
    assert _dashboard_slicer_value("66") == "66"
    assert _dashboard_slicer_value(" 230 ") == "230"
    assert _dashboard_figure_uses_saved_voltage_filter("Initial voltages") is True
    assert _dashboard_figure_uses_saved_voltage_filter("Initial Active Power") is True
    assert _dashboard_figure_uses_saved_voltage_filter("Initial Reactive Power") is True
    assert _dashboard_figure_uses_saved_voltage_filter("LG instantaneous overvoltages") is False


def test_report_image_index_reuses_directory_scan(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import reporting

    image_dir = tmp_path / "Generated" / "Full" / "SFO"
    image_dir.mkdir(parents=True)
    (image_dir / "C1_MM_66_001.png").write_bytes(b"png")
    (image_dir / "C1_MM_230_001.png").write_bytes(b"png")

    original_rglob = Path.rglob
    calls = []

    def counting_rglob(path, pattern):
        calls.append((path, pattern))
        return original_rglob(path, pattern)

    monkeypatch.setattr(Path, "rglob", counting_rglob)
    cache = {}

    assert [path.name for path in reporting._find_voltage_images(image_dir, "66", cache)] == [
        "C1_MM_66_001.png"
    ]
    assert [path.name for path in reporting._find_voltage_images(image_dir, "230", cache)] == [
        "C1_MM_230_001.png"
    ]
    assert len(calls) == 3


def test_plot_report_heading_parses_generated_mm_filename() -> None:
    from pathlib import Path

    from results_analysis_app.reporting import _plot_heading_from_image_path

    heading = _plot_heading_from_image_path(
        Path("C1_S1_66OFT2_MM_66_StA_AG_004_LGp_LLp_3Ph.png")
    )

    assert heading == "Case: C1_S1_66OFT2 | Run: 4 | Element: MM_66_StA | Fault: AG | Trace: LGp & LLp"

    no_fault_heading = _plot_heading_from_image_path(
        Path("O1_V3_S3_MM_161_TPC1_002_LGp_LLp_3Ph.png")
    )

    assert no_fault_heading == "Case: O1_V3_S3 | Run: 2 | Element: MM_161_TPC1 | Trace: LGp & LLp"


def test_report_figure_fields_and_heading_numbering() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from results_analysis_app.reporting import (
        _FigureRegistry,
        _add_figure_reference_sentence,
        _add_report_heading,
        _configure_report_header_footer,
        _configure_report_styles,
        _create_heading_numbering,
    )

    document = Document()
    _configure_report_styles(document)
    _configure_report_header_footer(document, "PSCAD Results Analysis - Test project - Full")
    numbering_id = _create_heading_numbering(document)
    _add_report_heading(document, "161 kV", level=1)
    registry = _FigureRegistry()
    reference = registry.allocate()
    _add_figure_reference_sentence(document, "", reference, " shows the result.")
    registry.add_caption(document, reference, "Representative envelope at 161 kV")

    assert document.paragraphs[0]._p.pPr.numPr is None
    heading_num = next(
        item
        for item in document.part.numbering_part.element.findall(qn("w:num"))
        if item.get(qn("w:numId")) == str(numbering_id)
    )
    heading_abstract_id = heading_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    heading_abstract = next(
        item
        for item in document.part.numbering_part.element.findall(qn("w:abstractNum"))
        if item.get(qn("w:abstractNumId")) == heading_abstract_id
    )
    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        assert style._element.pPr.numPr.numId.val == numbering_id
        if level == 1:
            assert style._element.pPr.numPr.ilvl is None
        else:
            assert style._element.pPr.numPr.ilvl.val == level - 1
        heading_level = heading_abstract.findall(qn("w:lvl"))[level - 1]
        assert heading_level.find(qn("w:pStyle")).get(qn("w:val")) == f"Heading{level}"
    assert "STYLEREF 1 \\s" in document._element.xml
    assert "SEQ Figure \\* ARABIC \\s 1" in document._element.xml
    assert "REF ReportFigure1 \\h" in document._element.xml
    assert "updateFields" not in document.settings._element.xml
    assert document.paragraphs[-1].text == "Figure 1-1 \u2013 Representative envelope at 161 kV"
    assert document.paragraphs[1].style.name == "Body Text"
    assert all(run.bold is True for run in document.paragraphs[1].runs[1:-1])
    assert all(run.bold is True for run in document.paragraphs[-1].runs[:-1])
    assert document.paragraphs[-1].runs[-1].bold is not True
    assert document.styles["Body Text"].font.name == "Calibri Light"
    assert document.styles["Body Text"].paragraph_format.alignment == 3
    assert document.styles["Heading 1"].font.color.rgb == (112, 155, 50)
    assert document.styles["Heading 1"].font.bold is None
    assert document.styles["Heading 2"].font.bold is None
    assert document.styles["Caption"].font.italic is True
    assert document.styles["Caption"].font.bold is None
    assert document.styles["Caption"].font.color.rgb == (112, 155, 50)
    assert document.styles["Plot Heading"].font.color.rgb == (112, 155, 50)
    assert document.styles["Plot Heading"]._element.pPr.outlineLvl.get(qn("w:val")) == "2"
    assert document.sections[0].header.paragraphs[0].text == (
        "PSCAD Results Analysis - Test project - Full"
    )
    assert "PAGE \\* MERGEFORMAT" in document.sections[0].footer._element.xml


def test_report_builder_embeds_styles_without_template(tmp_path) -> None:
    from docx import Document

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.reporting import build_reports_from_existing_plots

    project = tmp_path / "Project"
    project.mkdir()

    outputs = build_reports_from_existing_plots(
        [project],
        [ScopeEntry.full()],
        ["161"],
        [],
    )

    assert len(outputs) == 1
    document = Document(outputs[0])
    assert document.sections[0].header.paragraphs[0].text.startswith(
        "PSCAD Results Analysis - Project -"
    )
    assert document.sections[0].footer.paragraphs[0].text == "Page 1"
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].text.endswith(" - Full")
    assert document.paragraphs[1].style.name == "Body Text"


def test_report_plot_heading_precedes_cross_reference(tmp_path, monkeypatch) -> None:
    from docx import Document

    from results_analysis_app import reporting
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    image_path = tmp_path / "O1_V3_S1_MM_161_TPC1_001_LGp_LLp_3Ph.png"
    image_path.write_bytes(b"placeholder")

    monkeypatch.setattr(reporting, "_find_event_images", lambda *_args: [image_path])
    monkeypatch.setattr(reporting, "_find_resonance_images", lambda *_args: [])
    monkeypatch.setattr(reporting, "_is_report_image", lambda _path: True)
    monkeypatch.setattr(reporting, "_add_centered_report_image", lambda *_args: None)

    outputs = reporting.build_reports_from_existing_plots(
        [project],
        [ScopeEntry.full()],
        ["161"],
        ["SFO"],
    )

    document = Document(outputs[0])
    texts = [paragraph.text for paragraph in document.paragraphs]
    heading_index = texts.index("Case: O1_V3_S1 | Run: 1 | Element: MM_161_TPC1 | Trace: LGp & LLp")
    reference_index = next(index for index, text in enumerate(texts) if text.startswith("Figure 1-1 shows"))
    assert heading_index < reference_index
    assert document.paragraphs[heading_index].style.name == "Plot Heading"


def test_report_dashboard_heading_precedes_cross_reference(tmp_path, monkeypatch) -> None:
    from docx import Document

    from results_analysis_app import reporting
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    image_path = tmp_path / "dashboard.png"
    image_path.write_bytes(b"placeholder")

    monkeypatch.setattr(
        reporting,
        "_export_dashboard_figures",
        lambda *_args, **_kwargs: [("Initial Reactive Power", image_path)],
    )
    monkeypatch.setattr(reporting, "_add_centered_report_image", lambda *_args: None)

    outputs = reporting.build_reports_from_existing_plots(
        [project],
        [ScopeEntry.full()],
        ["161"],
        [],
        dashboard_figure_ids=["dashboard.xlsx|Graphs|1|Initial Reactive Power"],
    )

    document = Document(outputs[0])
    texts = [paragraph.text for paragraph in document.paragraphs]
    heading_index = texts.index("Initial Reactive Power")
    reference_index = next(index for index, text in enumerate(texts) if text.startswith("The initial reactive power"))
    assert heading_index < reference_index
    assert document.paragraphs[heading_index].style.name == "Plot Heading"


def test_report_orders_initial_dashboard_figures(tmp_path, monkeypatch) -> None:
    from docx import Document

    from results_analysis_app import reporting
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    image_path = tmp_path / "dashboard.png"
    image_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        reporting,
        "_export_dashboard_figures",
        lambda *_args, **_kwargs: [
            ("Initial Active Power", image_path),
            ("Other dashboard figure", image_path),
            ("Initial voltages", image_path),
            ("Initial Reactive Power", image_path),
        ],
    )
    monkeypatch.setattr(reporting, "_add_centered_report_image", lambda *_args: None)

    outputs = reporting.build_reports_from_existing_plots(
        [project],
        [ScopeEntry.full()],
        ["161"],
        [],
        dashboard_figure_ids=["selected"],
    )

    texts = [paragraph.text for paragraph in Document(outputs[0]).paragraphs]
    positions = {
        title: texts.index(title)
        for title in (
            "Initial voltages",
            "Initial Reactive Power",
            "Initial Active Power",
            "Other dashboard figure",
        )
    }
    assert positions["Initial voltages"] < positions["Initial Reactive Power"]
    assert positions["Initial Reactive Power"] < positions["Initial Active Power"]
    assert positions["Initial Active Power"] < positions["Other dashboard figure"]


def test_report_places_envelope_section_before_dashboard_section(tmp_path, monkeypatch) -> None:
    from docx import Document

    from results_analysis_app import reporting
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    project.mkdir()
    image_path = tmp_path / "figure.png"
    image_path.write_bytes(b"placeholder")
    monkeypatch.setattr(reporting, "_export_envelope_chart", lambda *_args, **_kwargs: image_path)
    monkeypatch.setattr(
        reporting,
        "_export_dashboard_figures",
        lambda *_args, **_kwargs: [("Initial voltages", image_path)],
    )
    monkeypatch.setattr(reporting, "_add_centered_report_image", lambda *_args: None)

    outputs = reporting.build_reports_from_existing_plots(
        [project],
        [ScopeEntry.full()],
        ["161"],
        [],
        dashboard_figure_ids=["selected"],
    )

    document = Document(outputs[0])
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 2"
    ]
    captions = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption"
    ]
    assert headings == [
        "161 kV - Envelope - Full",
        "161 kV - Dashboard Figures - Full",
    ]
    assert captions[0].startswith("Figure 1-1")
    assert captions[1].startswith("Figure 1-2")


def test_report_dashboard_wording_uses_voltage_rise_for_rms_overvoltage() -> None:
    from results_analysis_app.reporting import (
        FigureReference,
        _add_dashboard_figure_sentence,
        _dashboard_figure_caption,
        _report_dashboard_title,
    )
    from docx import Document

    assert _report_dashboard_title("LG RMS voltage drop x switching time") == "LG RMS voltage drop x switching time"
    assert _report_dashboard_title("LG RMS overvoltages x switching time") == "LG RMS voltage rise x switching time"
    assert _dashboard_figure_caption("Initial voltages", "161") == "Initial voltage across voltage levels in the OWF"
    assert _dashboard_figure_caption("Initial Reactive Power", "161") == "Initial reactive power at the POC"
    assert _dashboard_figure_caption("LG instantaneous overvoltages", "161") == "Overvoltage results at the 161 kV side"

    document = Document()
    _add_dashboard_figure_sentence(
        document,
        "Initial Active Power",
        "161",
        FigureReference("ReportFigure1", "Figure 1-1"),
    )
    assert document.paragraphs[0].text == (
        "The initial active power conditions of all the cases are shown in the Figure 1-1."
    )


def test_envelope_summary_rows_read_sfo_tov_from_llp(tmp_path) -> None:
    import pytest
    from openpyxl import Workbook

    from results_analysis_app.reporting import _envelope_summary_rows

    workbook = Workbook()
    ws = workbook.active
    ws.title = "LGp"
    ws.append(["Time (s)", "Max_all"])
    ws.append([0.004, 999.0])
    ws.append([0.03, 888.0])
    ws.append([0.1, 300.0])
    ll = workbook.create_sheet("LLp")
    ll.append(["Time (s)", "Max_all"])
    ll.append([0.004, 141.421356])
    ll.append([0.03, 200.0])
    ll.append([0.1, 111.0])
    path = tmp_path / "MM_66.xlsx"
    workbook.save(path)
    workbook.close()

    event_times = {"SFO": 0.004, "TOV": 0.03, "SA": 0.1}
    rows = _envelope_summary_rows(path, event_times)

    assert [(row.event, row.measurement) for row in rows] == [("SFO", "LLp"), ("TOV", "LLp")]
    assert rows[0].rms_kv == pytest.approx(100.0)

    rows = _envelope_summary_rows(path, event_times, events=["SFO", "TOV", "SA"])

    assert [(row.event, row.measurement) for row in rows] == [
        ("SFO", "LLp"),
        ("TOV", "LLp"),
        ("SA", "LGp"),
    ]
    assert rows[2].peak_kv == pytest.approx(300.0)
    assert rows[2].rms_kv == pytest.approx(300.0 / 2**0.5)

    line_ground_tov = _envelope_summary_rows(
        path,
        event_times,
        events=["TOV"],
        measurement_overrides={"TOV": "LGp"},
    )
    assert line_ground_tov[0].peak_kv == pytest.approx(888.0)


def test_envelope_summary_nearest_row_tie_keeps_earlier_source_row(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.reporting import _envelope_summary_rows

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LLp"
    sheet.append(["Time (s)", "Max_all"])
    sheet.append([0.0, 100.0])
    sheet.append([0.002, 200.0])
    path = tmp_path / "MM_66.xlsx"
    workbook.save(path)
    workbook.close()

    rows = _envelope_summary_rows(path, {"SFO": 0.001}, events=["SFO"])

    assert rows[0].peak_kv == 100.0


def test_stale_combined_envelope_chart_is_not_exported(tmp_path) -> None:
    import os

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.reporting import _export_envelope_chart

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    base = envelope_dir / "MM_66.xlsx"
    combined = envelope_dir / "MM_66_with_combined_plot.xlsx"
    base.write_bytes(b"new")
    combined.write_bytes(b"old")
    old = combined.stat().st_mtime_ns
    os.utime(base, ns=(old + 1_000_000, old + 1_000_000))
    logs = []

    result = _export_envelope_chart(
        project,
        ScopeEntry.full(),
        "66",
        tmp_path / "exports",
        logs.append,
        lambda: (_ for _ in ()).throw(AssertionError("Excel must not open")),
    )

    assert result is None
    assert any("older than its data" in message for message in logs)


def test_report_cancellation_cleans_temporary_exports(tmp_path) -> None:
    import pytest

    from results_analysis_app.background import OperationCancelled
    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.reporting import build_reports_from_existing_plots

    report_dir = tmp_path / "Reports" / "Full"
    temporary = report_dir / "_envelope_exports"
    temporary.mkdir(parents=True)
    (temporary / "stale.png").write_bytes(b"old")
    calls = 0

    def cancel() -> None:
        nonlocal calls
        calls += 1
        if calls == 4:
            raise OperationCancelled()

    with pytest.raises(OperationCancelled):
        build_reports_from_existing_plots(
            [tmp_path],
            [ScopeEntry.full()],
            ["66"],
            [],
            check_cancel=cancel,
        )

    assert not temporary.exists()


def test_sa_report_text_uses_line_ground_tov_rms_value() -> None:
    from docx import Document

    from results_analysis_app.reporting import (
        EnvelopeSummaryRow,
        FigureReference,
        _add_sa_figure_sentence,
        _time_domain_figure_caption,
    )

    document = Document()
    _add_sa_figure_sentence(
        document,
        "161",
        FigureReference("ReportFigure7", "Figure 4-7"),
        EnvelopeSummaryRow("TOV", "LGp", 0.03, 188.0, 133.0),
    )

    assert document.paragraphs[0].text == (
        "Figure 4-7 shows the highest TOV line-ground value found at 161 kV, "
        "which will be used for surge arrester selection."
    )
    assert document.paragraphs[1].text == "This TOV value is 133 kVRMS and was assumed to last for 300 ms."
    assert document.paragraphs[1].runs[3].font.subscript is True
    assert _time_domain_figure_caption("SA", "161") == (
        "Time-domain plot of the highest TOV at 161 kV for surge arrester selection"
    )


def test_envelope_summary_list_uses_subscripted_units() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from results_analysis_app.reporting import (
        EnvelopeSummaryRow,
        _add_envelope_summary_list,
        _configure_report_styles,
        _create_heading_numbering,
    )

    doc = Document()
    _configure_report_styles(doc)
    _create_heading_numbering(doc)
    _add_envelope_summary_list(
        doc,
        [EnvelopeSummaryRow("TOV", "LLp", 0.03, 200.0, 141.4)],
    )

    bullet = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name == "List Bullet")
    bullet_num_id = bullet._p.pPr.numPr.numId.val
    assert bullet_num_id == 2
    number = next(
        number
        for number in doc.part.numbering_part.element.findall(qn("w:num"))
        if int(number.get(qn("w:numId"))) == bullet_num_id
    )
    abstract_id = number.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(
        abstract
        for abstract in doc.part.numbering_part.element.findall(qn("w:abstractNum"))
        if abstract.get(qn("w:abstractNumId")) == abstract_id
    )
    level = abstract.find(qn("w:lvl"))
    assert level.find(qn("w:numFmt")).get(qn("w:val")) == "bullet"
    assert level.find(qn("w:pStyle")).get(qn("w:val")) == "ListBullet"
    assert level.find(qn("w:lvlText")).get(qn("w:val")) == "\u25e6"
    run_properties = level.find(qn("w:rPr"))
    assert run_properties.find(qn("w:rFonts")).get(qn("w:ascii")) == "Calibri"
    assert run_properties.find(qn("w:color")).get(qn("w:val")) == "709B32"
    assert run_properties.find(qn("w:position")).get(qn("w:val")) == "-4"
    assert run_properties.find(qn("w:sz")).get(qn("w:val")) == "36"
    assert doc.styles["List Bullet"]._element.pPr.numPr.numId.val == 2
    assert [run.text for run in bullet.runs] == ["TOV: 200", " kV", "peak", " (141.4", " kV", "RMS", ")"]
    assert bullet.runs[2].font.subscript is True
    assert bullet.runs[5].font.subscript is True


def test_envelope_bullets_do_not_change_heading_numbering() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from results_analysis_app.reporting import (
        EnvelopeSummaryRow,
        _add_envelope_summary_list,
        _add_report_heading,
        _configure_report_styles,
        _create_heading_numbering,
    )

    doc = Document()
    _configure_report_styles(doc)
    heading_num_id = _create_heading_numbering(doc)
    _add_report_heading(doc, "230 kV - Envelope - Full", level=2)
    _add_envelope_summary_list(doc, [EnvelopeSummaryRow("SFO", "LLp", 0, 1, 1)])

    heading = doc.paragraphs[0]
    bullet = doc.paragraphs[1]
    assert heading._p.pPr.numPr is None
    assert doc.styles["Heading 2"]._element.pPr.numPr.numId.val == heading_num_id
    assert heading_num_id == 1
    assert bullet._p.pPr.numPr.numId.val == 2
    assert bullet.style.name == "List Bullet"

    numbering_children = list(doc.part.numbering_part.element)
    last_abstract_index = max(
        index for index, child in enumerate(numbering_children) if child.tag == qn("w:abstractNum")
    )
    first_number_index = min(
        index for index, child in enumerate(numbering_children) if child.tag == qn("w:num")
    )
    assert last_abstract_index < first_number_index

    heading_number = next(
        number
        for number in doc.part.numbering_part.element.findall(qn("w:num"))
        if int(number.get(qn("w:numId"))) == heading_num_id
    )
    heading_abstract_id = heading_number.find(qn("w:abstractNumId")).get(qn("w:val"))
    heading_abstract = next(
        abstract
        for abstract in doc.part.numbering_part.element.findall(qn("w:abstractNum"))
        if abstract.get(qn("w:abstractNumId")) == heading_abstract_id
    )
    level_2 = next(
        level
        for level in heading_abstract.findall(qn("w:lvl"))
        if level.get(qn("w:ilvl")) == "1"
    )
    assert level_2.find(qn("w:numFmt")).get(qn("w:val")) == "decimal"


def test_invalid_inf_and_cb_inputs_are_reported(tmp_path) -> None:
    from results_analysis_app import voltage_envelope

    invalid_inf = tmp_path / "bad.inf"
    invalid_inf.write_text("not a descriptor", encoding="utf-8")
    messages: list[str] = []
    assert voltage_envelope._read_inf_descriptor_cache([invalid_inf], 1, None, messages.append) == {}
    assert messages and "Skipped invalid .inf layout" in messages[0]

    invalid_cb = tmp_path / "CB_bad.out"
    invalid_cb.write_bytes(b"\x00\x00")
    records, warning = voltage_envelope._read_non_convergent_file(invalid_cb)
    assert records == []
    assert warning and "Skipped unreadable CB summary" in warning
