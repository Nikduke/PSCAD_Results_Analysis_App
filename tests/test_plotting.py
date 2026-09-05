from __future__ import annotations


def _sustained_cache_payload(canonical, observations, duration_ms=40.0):
    from results_analysis_app import sustained_sdpf

    representatives = sustained_sdpf.select_representatives(observations)
    selections, extras = sustained_sdpf._build_representative_cache(
        canonical,
        representatives,
    )
    return {
        "version": sustained_sdpf.RESULT_VERSION,
        "scope_folder": "Full",
        "settings": {"enabled": True, "duration_ms": duration_ms},
        "results": {"66": canonical.to_dict() if canonical is not None else None},
        "representative_results": {"66": extras},
        "selections": {"66": selections},
    }


def test_embedded_plotter_core_imports() -> None:
    import pscad_plotter_app_v3
    from pscad_plotter_app_v3.models import PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    assert getattr(pscad_plotter_app_v3, "__embedded_source__", "")
    assert PlotMode.MM.value == "MM Overvoltage"
    assert list(PlotMode) == [PlotMode.MM]
    assert "MM" in BatchExcelService.SHEET_DEFINITIONS
    assert MatplotlibRenderer.FIGSIZE


def test_plot_execution_combines_png_and_requested_excel_export(tmp_path) -> None:
    from pathlib import Path

    from pscad_plotter_app_v3.models import PlotJob, PlotMode
    from pscad_plotter_app_v3.services.plot_execution import execute_plot_job

    class FakeRenderer:
        def render(self, job):
            output = Path(job.output_dir) / "plot.png"
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(b"png")
            return output

    class FakeExporter:
        def __init__(self) -> None:
            self.calls = 0

        def export(self, job):
            self.calls += 1
            output = Path(job.output_dir) / "Excel" / "plot.xlsx"
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(b"xlsx")
            return output

    job = PlotJob(
        mode=PlotMode.MM,
        case_name="C1",
        run_number=1,
        group_label="MM_66_A",
        output_dir=str(tmp_path),
        excel_export=True,
    )
    exporter = FakeExporter()

    output = execute_plot_job(FakeRenderer(), exporter, job)

    assert output.png_path.read_bytes() == b"png"
    assert output.excel_path is not None
    assert output.excel_path.read_bytes() == b"xlsx"
    assert exporter.calls == 1


def test_automatic_plot_worker_count_uses_threshold_and_cap(monkeypatch) -> None:
    from pscad_plotter_app_v3.services import plot_execution

    monkeypatch.setattr(plot_execution.os, "cpu_count", lambda: 16)
    assert plot_execution.automatic_plot_worker_count(2) == 1
    assert plot_execution.automatic_plot_worker_count(3) == 3
    assert plot_execution.automatic_plot_worker_count(10) == 4
    assert plot_execution.automatic_plot_worker_count(100) == 4

    monkeypatch.setattr(plot_execution.os, "cpu_count", lambda: 2)
    assert plot_execution.automatic_plot_worker_count(100) == 1


def test_cancelled_parallel_executor_terminates_active_workers() -> None:
    from results_analysis_app.analysis_engine import _terminate_plot_executor

    class FakeProcess:
        def __init__(self) -> None:
            self.terminated = False
            self.joined = False

        def terminate(self) -> None:
            self.terminated = True

        def join(self, timeout: float) -> None:
            assert timeout == 1.0
            self.joined = True

    class FakeExecutor:
        def __init__(self, processes) -> None:
            self._processes = {index: process for index, process in enumerate(processes)}
            self.shutdown_args = None

        def shutdown(self, **kwargs) -> None:
            self.shutdown_args = kwargs

    processes = [FakeProcess(), FakeProcess()]
    executor = FakeExecutor(processes)

    _terminate_plot_executor(executor)

    assert all(process.terminated and process.joined for process in processes)
    assert executor.shutdown_args == {"wait": False, "cancel_futures": True}


def test_shared_out_readers_check_cancellation_during_file_read(tmp_path) -> None:
    import numpy as np
    import pytest

    from pscad_plotter_app_v3.services.waveform_io import (
        load_out_columns,
        load_out_frame,
    )

    out_path = tmp_path / "C1_r00001_01.out"
    values = np.column_stack(
        [
            np.arange(5000, dtype=float),
            np.arange(5000, dtype=float) * 2.0,
        ]
    )
    np.savetxt(out_path, values, header="time signal", comments="")

    for reader in (
        lambda callback: load_out_columns(out_path, [0, 1], callback),
        lambda callback: load_out_frame(out_path, callback),
    ):
        calls = 0

        def cancel() -> None:
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("cancelled")

        with pytest.raises(RuntimeError, match="cancelled"):
            reader(cancel)
        assert calls == 2


def test_create_plot_batches_reads_base_envelope_workbook(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "LLp"
    ws.append(["Time (s)", "Case_all", "Run_all", "MM_name_all", "Fault_type_all"])
    ws.append([0.004, "C1", 4, "MM_66_StA", "AG"])
    ws.append([0.03, "C2", 5, "MM_66_StB", "AB"])
    lg = workbook.create_sheet("LGp")
    lg.append(["Time (s)", "Case_all", "Run_all", "MM_name_all", "Fault_type_all"])
    lg.append([0.1, "C3", 6, "MM_66_StC", "ABC"])
    workbook.save(envelope_dir / "MM_66.xlsx")
    workbook.close()

    (envelope_dir / "MM_66_with_combined_plot.xlsx").write_bytes(b"not an xlsx")

    written = []

    def fake_create_batch(path, rows):
        written.append((path.name, rows))
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("ok", encoding="utf-8")

    monkeypatch.setattr(analysis_engine, "_create_batch_workbook", fake_create_batch)

    outputs = analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        ["66"],
        ["SFO", "TOV", "SA"],
        {"SFO": 0.004, "TOV": 0.03, "SA": 0.1},
    )

    assert [path.name for path in outputs] == [
        "batch_paste_Full_SFO.xlsx",
        "batch_paste_Full_TOV.xlsx",
        "batch_paste_Full_SA.xlsx",
    ]
    rows_by_name = {name: rows for name, rows in written}
    assert rows_by_name["batch_paste_Full_SFO.xlsx"][0]["case"] == "C1"
    assert rows_by_name["batch_paste_Full_TOV.xlsx"][0]["case"] == "C2"
    assert rows_by_name["batch_paste_Full_SA.xlsx"][0]["case"] == "C3"
    assert all(
        row["excel_export"]
        for rows in rows_by_name.values()
        for row in rows
    )


def test_create_plot_batches_can_disable_automatic_excel_exports(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    (envelope_dir / "MM_66.xlsx").write_bytes(b"placeholder")
    written = []

    monkeypatch.setattr(
        analysis_engine,
        "_read_batch_points",
        lambda *_args, **_kwargs: {
            "SFO": {"case": "C1", "run": 1, "element": "MM_66_A"},
        },
    )
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, rows: written.append((path, rows)),
    )

    analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        ["66"],
        ["SFO"],
        {"SFO": 0.004},
        excel_waveform_exports_enabled=False,
    )

    assert written[0][1][0]["excel_export"] is False


def test_event_only_batch_creation_preserves_sustained_batch(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    stale_batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_Sustained_SDPF.xlsx"
    stale_batch.parent.mkdir(parents=True)
    stale_batch.write_bytes(b"keep")
    monkeypatch.setattr(
        analysis_engine.sustained_sdpf_heatmap,
        "clear_heatmaps",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("event-only creation must not clear Sustained heatmaps")
        ),
    )

    analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        [],
        ["SFO"],
        {"SFO": 0.004},
    )

    assert stale_batch.read_bytes() == b"keep"


def test_event_only_render_does_not_rebuild_or_clear_sustained_heatmaps(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    monkeypatch.setattr(analysis_engine, "_render_plot_batches_direct", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        analysis_engine,
        "render_sustained_heatmaps",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("event-only rendering must not touch Sustained heatmaps")
        ),
    )

    analysis_engine.render_plot_batches(
        tmp_path,
        [ScopeEntry.full()],
        ["SFO"],
    )


def test_create_plot_batches_marks_sustained_plot_with_standard_tov_window(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    phase = sustained_sdpf.PhaseStressResult(
        "LGp", "A-G", 1.0, 2**0.5, 0.85, 0.85 * 2**0.5,
        0.04, 0.04, 1.5, 1.0, 1.06,
        "SDPF exceeded for at least 30 ms of peak-envelope persistence (peak)",
        margin_exceeded=True,
        sdpf_exceeded=True,
    )
    result = sustained_sdpf.SustainedSDPFResult(
        "Full", "66", "C1", 1, "MM_66_A", "AG", 0.04, phase, (phase,), cycle_coverage=2
    )
    payload = _sustained_cache_payload(result, [result])
    validation = sustained_sdpf.SustainedSDPFCacheValidation(
        valid=True,
        shared_manifest_current=True,
    )
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "load_results", lambda *_args: payload)
    monkeypatch.setattr(
        analysis_engine.sustained_sdpf,
        "validate_result_cache",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("the connected workflow should reuse validation")
        ),
    )
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "result_inputs_current", lambda *_args, **_kwargs: True)
    written = []
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, rows: written.append((path, rows)),
    )

    outputs = analysis_engine.create_plot_batches(
        tmp_path,
        [ScopeEntry.full()],
        ["66"],
        [],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 40.0},
        sustained_cache_validations_by_scope={"Full": validation},
    )

    assert outputs[-1].name == "batch_paste_Full_Sustained_SDPF.xlsx"
    row = written[-1][1][0]
    assert row["tov_windows"] is True
    assert "tov_window_s" not in row
    assert row["tov_window_count"] == 4
    assert "time_start_s" not in row
    assert "time_end_s" not in row


def test_create_plot_batches_includes_distinct_sustained_selections(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    def result(case: str) -> sustained_sdpf.SustainedSDPFResult:
        phase = sustained_sdpf.PhaseStressResult(
            "LGp", "A-G", 1.0, 2**0.5, 0.85, 0.85 * 2**0.5,
            0.04, 0.04, 1.5, 1.0, 1.06,
            "SDPF exceeded",
            margin_exceeded=True,
            sdpf_exceeded=True,
            sdpf_excess_area_norm_ms=5.0 if case == "Area" else 2.0,
            sdpf_longest_continuous_s=0.01 if case == "Area" else 0.025,
        )
        return sustained_sdpf.SustainedSDPFResult(
            "Full", "66", case, 1, f"MM_66_{case}", "AG", 0.04, phase, (phase,)
        )

    area = result("Area")
    duration = result("Duration")
    payload = _sustained_cache_payload(area, [area, duration])
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "load_results", lambda *_args: payload)
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "source_manifest_current", lambda *_args: True)
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "result_inputs_current", lambda *_args, **_kwargs: True)
    written = []
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, rows: written.append((path, rows)),
    )

    analysis_engine.create_plot_batches(
        tmp_path,
        [ScopeEntry.full()],
        ["66"],
        [],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 40.0},
    )

    sustained_rows = written[-1][1]
    assert {(row["case"], row["element"]) for row in sustained_rows} == {
        ("Area", "MM_66_Area"),
        ("Duration", "MM_66_Duration"),
    }


def test_sustained_render_uses_batch_voltage_keys(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    from pscad_plotter_app_v3.models import DEFAULT_TOV_WINDOW_S, PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from results_analysis_app import analysis_engine, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch_path = project / "Plots" / "Plot_batch" / "batch_paste_Full_Sustained_SDPF.xlsx"
    batch_path.parent.mkdir(parents=True)
    batch_path.write_bytes(b"batch")
    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=True,
        case_name="C1",
        run_numbers=[1],
        elements=["MM_66_A"],
        trace_type="Both",
        show_three_phase_overview=True,
        show_tov_windows=True,
        tov_window_s=DEFAULT_TOV_WINDOW_S,
        tov_window_count=4,
        excel_export=False,
        time_start_s=None,
        time_end_s=None,
    )
    loaded_voltage_keys = []

    monkeypatch.setattr(
        BatchExcelService,
        "load_requests",
        lambda *_args, **_kwargs: SimpleNamespace(requests=[request], errors=[]),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (
            SimpleNamespace(),
            {},
            SimpleNamespace(run_index={}),
            object(),
        ),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_expected_sustained_plot_results",
        lambda *args, **_kwargs: (
            loaded_voltage_keys.extend(args[5])
            or {("66", "C1", 1, "MM_66_A"): object()}
        ),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_sustained_request_matches",
        lambda *_args, **_kwargs: True,
    )
    monkeypatch.setattr(
        analysis_engine,
        "_build_mm_plot_jobs",
        lambda *_args, **_kwargs: [object()],
    )
    monkeypatch.setattr(analysis_engine, "_plot_manifest_matches", lambda *_args: False)
    monkeypatch.setattr(analysis_engine, "_render_plot_plans", lambda *_args: None)

    analysis_engine._render_plot_batches_direct(
        project,
        [ScopeEntry.full()],
        [sustained_sdpf.SUSTAINED_SDPF],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 30.0},
    )

    assert loaded_voltage_keys == [66.0]


def test_sustained_batch_current_requires_exact_expected_selection_set(monkeypatch) -> None:
    from types import SimpleNamespace

    from results_analysis_app import analysis_engine

    request = SimpleNamespace(
        voltage_kv=66.0,
        case_name="C1",
        run_numbers=[1],
        elements=["MM_66_A"],
    )
    expected = {("66", "C1", 1, "MM_66_A"): object()}
    monkeypatch.setattr(analysis_engine, "_sustained_request_matches", lambda *_args: True)

    assert not analysis_engine._sustained_batch_is_current([], expected)
    assert analysis_engine._sustained_batch_is_current([request], expected)
    assert not analysis_engine._sustained_batch_is_current(
        [request],
        {**expected, ("66", "C2", 1, "MM_66_B"): object()},
    )


def test_sustained_render_repairs_empty_stale_batch_from_current_results(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    from pscad_plotter_app_v3.models import DEFAULT_TOV_WINDOW_S, PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from results_analysis_app import analysis_engine, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch_path = project / "Plots" / "Plot_batch" / "batch_paste_Full_Sustained_SDPF.xlsx"
    batch_path.parent.mkdir(parents=True)
    batch_path.write_bytes(b"empty-or-stale")
    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=True,
        case_name="C1",
        run_numbers=[1],
        elements=["MM_66_A"],
        trace_type="Both",
        show_three_phase_overview=True,
        show_tov_windows=True,
        tov_window_s=DEFAULT_TOV_WINDOW_S,
        tov_window_count=4,
        excel_export=False,
        time_start_s=None,
        time_end_s=None,
    )
    import_calls = []

    def load_requests(*_args, **_kwargs):
        import_calls.append(True)
        return SimpleNamespace(
            requests=[] if len(import_calls) == 1 else [request],
            errors=[],
        )

    monkeypatch.setattr(BatchExcelService, "load_requests", load_requests)
    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (SimpleNamespace(), {}, object(), object()),
    )
    result = SimpleNamespace(case="C1", run=1, mm_name="MM_66_A")
    expected = {("66", "C1", 1, "MM_66_A"): result}
    monkeypatch.setattr(
        analysis_engine,
        "_expected_sustained_plot_results",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        analysis_engine,
        "_sustained_request_matches",
        lambda *_args, **_kwargs: True,
    )
    cleared = []
    monkeypatch.setattr(
        analysis_engine,
        "_clear_sustained_outputs",
        lambda *_args, **_kwargs: cleared.append(True),
    )
    rebuilt = []
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, rows: rebuilt.append((path, rows)),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_build_mm_plot_jobs",
        lambda *_args, **_kwargs: [object()],
    )
    monkeypatch.setattr(analysis_engine, "_plot_manifest_matches", lambda *_args: False)
    rendered = []
    monkeypatch.setattr(
        analysis_engine,
        "_render_plot_plans",
        lambda plans, *_args: rendered.append(plans),
    )

    analysis_engine._render_plot_batches_direct(
        project,
        [ScopeEntry.full()],
        [sustained_sdpf.SUSTAINED_SDPF],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 30.0},
        sustained_voltage_keys=["66"],
        excel_waveform_exports_enabled=False,
    )

    assert len(import_calls) == 2
    assert rebuilt[0][0] == batch_path
    assert rebuilt[0][1] == [
        {
            "case": "C1",
            "run": 1,
            "element": "MM_66_A",
            "trace": "Both",
            "overview": True,
            "tov_windows": True,
            "tov_window_count": 4,
            "limits": True,
            "excel_export": False,
        }
    ]
    assert cleared == [True]
    assert rendered


def test_create_plot_batches_does_not_write_empty_batch_for_invalid_cache(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine, sustained_sdpf
    from results_analysis_app.models import ScopeEntry

    batch_path = tmp_path / "Plots" / "Plot_batch" / "batch_paste_Full_Sustained_SDPF.xlsx"
    batch_path.parent.mkdir(parents=True)
    batch_path.write_bytes(b"stale")
    payload = {
        "version": sustained_sdpf.RESULT_VERSION,
        "settings": {"enabled": True, "duration_ms": 30.0},
        "results": {"66": None},
    }
    monkeypatch.setattr(analysis_engine.sustained_sdpf, "load_results", lambda *_args: payload)
    written = []
    logs = []
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, rows: written.append((path, rows)),
    )

    outputs = analysis_engine.create_plot_batches(
        tmp_path,
        [ScopeEntry.full()],
        ["66"],
        [],
        sustained_sdpf_settings={"enabled": True, "duration_ms": 30.0},
        log=logs.append,
    )

    assert outputs == []
    assert written == []
    assert not batch_path.exists()
    assert any("source files or their fingerprint" in message for message in logs)


def test_sustained_plot_limit_overrides_preserve_siwl_values() -> None:
    from pscad_plotter_app_v3.models import VoltageLimitSet
    from results_analysis_app.analysis_engine import _apply_sustained_sdpf_plot_limit_overrides

    limits = {
        "66": VoltageLimitSet(
            voltage_kv=66.0,
            sdpf_lg=100.0,
            sdpf_ll=120.0,
            siwl_lg=140.0,
            siwl_ll=160.0,
        )
    }

    adjusted = _apply_sustained_sdpf_plot_limit_overrides(
        limits,
        {"66.0": {"LGp": 90.0}},
    )

    assert adjusted is not limits
    assert adjusted["66"].sdpf_lg == 90.0
    assert adjusted["66"].sdpf_ll == 120.0
    assert adjusted["66"].siwl_lg == 140.0
    assert adjusted["66"].siwl_ll == 160.0
    assert adjusted["66"].source == "manual"


def test_nearest_envelope_row_tie_keeps_earlier_source_row() -> None:
    from openpyxl import Workbook

    from results_analysis_app.analysis_engine import _batch_points_from_sheet

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LLp"
    sheet.append(["Time (s)", "Case_all", "Run_all", "MM_name_all"])
    sheet.append([0.0, "Earlier", 1, "MM_66_A"])
    sheet.append([0.002, "Later", 2, "MM_66_B"])

    points = _batch_points_from_sheet(workbook, "LLp", {"SFO": 0.001})

    assert points["SFO"]["case"] == "Earlier"


def test_empty_batch_removes_stale_generated_event_directory(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_SFO.xlsx"
    analysis_engine._create_batch_workbook(batch, [])
    generated = project / "Plots" / "Generated" / "Full" / "SFO"
    generated.mkdir(parents=True)
    (generated / "stale.png").write_bytes(b"old")

    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (SimpleNamespace(), {}, object(), object()),
    )

    analysis_engine.render_plot_batches(project, [ScopeEntry.full()], ["SFO"])

    assert not generated.exists()


def test_plot_render_failure_keeps_previous_event_output(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    import pytest

    from pscad_plotter_app_v3.models import PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services import batching
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_SFO.xlsx"
    batch.parent.mkdir(parents=True)
    batch.write_bytes(b"batch")
    generated = project / "Plots" / "Generated" / "Full" / "SFO"
    generated.mkdir(parents=True)
    previous = generated / "previous.png"
    previous.write_bytes(b"good")
    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=False,
        case_name="C1",
        elements=["MM_66_A"],
    )
    monkeypatch.setattr(
        BatchExcelService,
        "load_requests",
        lambda *_args, **_kwargs: SimpleNamespace(requests=[request], errors=[]),
    )

    def fake_jobs(plot_request, _limit):
        return [
            SimpleNamespace(
                case_name="C1",
                group_label="MM_66_A",
                run_number=1,
                excel_export=False,
                output_dir=plot_request.output_dir,
            )
        ]

    monkeypatch.setattr(batching, "build_mm_jobs", fake_jobs)

    class FailingRenderer:
        def render(self, _job):
            raise RuntimeError("render failed")

    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (SimpleNamespace(), {}, FailingRenderer(), object()),
    )

    with pytest.raises(RuntimeError, match="Plot rendering failed"):
        analysis_engine.render_plot_batches(project, [ScopeEntry.full()], ["SFO"])

    assert previous.read_bytes() == b"good"
    assert not list(generated.parent.glob(".SFO.render-*"))


def test_large_plot_batches_share_one_process_pool_and_commit_each_batch(tmp_path, monkeypatch) -> None:
    from concurrent.futures import Future
    import json
    from pathlib import Path
    from types import SimpleNamespace

    from pscad_plotter_app_v3.models import PlotJob, PlotMode
    from pscad_plotter_app_v3.services import plot_execution
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services import batching
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch_dir = project / "Plots" / "Plot_batch"
    batch_dir.mkdir(parents=True)
    for event_name in ("SFO", "TOV"):
        (batch_dir / f"batch_paste_Full_{event_name}.xlsx").write_bytes(b"batch")

    for event_name in ("SFO", "TOV"):
        generated = project / "Plots" / "Generated" / "Full" / event_name
        generated.mkdir(parents=True)
        (generated / "stale.png").write_bytes(b"old")

    def load_requests(_service, path, _catalog):
        event_name = "TOV" if "_TOV" in path.name else "SFO"
        request = SimpleNamespace(
            mode=PlotMode.MM,
            output_dir="",
            voltage_kv=66.0,
            show_limits=False,
            case_name=event_name,
            elements=["MM_66_A"],
        )
        return SimpleNamespace(requests=[request], errors=[])

    monkeypatch.setattr(BatchExcelService, "load_requests", load_requests)
    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (
            SimpleNamespace(),
            {},
            SimpleNamespace(run_index={"C": {1: Path("C:/C_r00001.inf")}}),
            object(),
        ),
    )

    def fake_jobs(request, _limit):
        return [
            PlotJob(
                mode=PlotMode.MM,
                case_name=request.case_name,
                run_number=1,
                group_label=f"MM_66_{index}",
                output_dir=request.output_dir,
                show_limits=False,
                excel_export=True,
            )
            for index in range(5)
        ]

    monkeypatch.setattr(batching, "build_mm_jobs", fake_jobs)

    def fake_execute(job):
        output_dir = Path(job.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        png_path = output_dir / f"{job.case_name}_{job.group_label}.png"
        png_path.write_bytes(b"png")
        excel_dir = output_dir / "Excel"
        excel_dir.mkdir(exist_ok=True)
        excel_path = excel_dir / f"{job.case_name}_{job.group_label}.xlsx"
        excel_path.write_bytes(b"xlsx")
        return plot_execution.PlotExecutionOutput(png_path, excel_path)

    monkeypatch.setattr(plot_execution, "initialize_plot_process", lambda _run_index: None)
    monkeypatch.setattr(plot_execution, "execute_plot_job_in_process", fake_execute)
    monkeypatch.setattr(plot_execution, "automatic_plot_worker_count", lambda _job_count: 2)

    class ImmediateProcessExecutor:
        instances = []

        def __init__(self, **kwargs):
            self.kwargs = kwargs
            self.submitted = []
            self.instances.append(self)

        def submit(self, function, job):
            self.submitted.append(job)
            future = Future()
            try:
                future.set_result(function(job))
            except Exception as exc:
                future.set_exception(exc)
            return future

        def shutdown(self, **_kwargs):
            return None

    monkeypatch.setattr(analysis_engine, "ProcessPoolExecutor", ImmediateProcessExecutor)

    analysis_engine._render_plot_batches_direct(
        project,
        [ScopeEntry.full()],
        ["SFO", "TOV"],
    )

    assert len(ImmediateProcessExecutor.instances) == 1
    executor = ImmediateProcessExecutor.instances[0]
    assert executor.kwargs["max_workers"] == 2
    assert len(executor.submitted) == 2
    assert sorted(len(group) for group in executor.submitted) == [5, 5]
    for event_name in ("SFO", "TOV"):
        output_dir = project / "Plots" / "Generated" / "Full" / event_name
        assert len(list(output_dir.glob("*.png"))) == 5
        assert len(list((output_dir / "Excel").glob("*.xlsx"))) == 5
        cache_path = project / ".state" / "analysis_cache.json"
        assert cache_path.is_file()
        cache = json.loads(cache_path.read_text(encoding="utf-8"))
        assert output_dir.relative_to(project).as_posix() in cache["plots"]
        assert not (output_dir / analysis_engine.LEGACY_PLOT_MANIFEST_FILENAME).exists()
        assert not list(output_dir.parent.glob(f".{event_name}.render-*"))

    analysis_engine._render_plot_batches_direct(
        project,
        [ScopeEntry.full()],
        ["SFO", "TOV"],
    )

    assert len(ImmediateProcessExecutor.instances) == 1


def test_broken_plot_process_pool_retries_sequentially(tmp_path, monkeypatch) -> None:
    from concurrent.futures import Future
    from concurrent.futures.process import BrokenProcessPool
    from pathlib import Path
    from types import SimpleNamespace

    from pscad_plotter_app_v3.models import PlotJob, PlotMode
    from pscad_plotter_app_v3.services import plot_execution
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services import batching
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_SFO.xlsx"
    batch.parent.mkdir(parents=True)
    batch.write_bytes(b"batch")
    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=False,
        case_name="C1",
        elements=["MM_66_A"],
    )
    monkeypatch.setattr(
        BatchExcelService,
        "load_requests",
        lambda *_args, **_kwargs: SimpleNamespace(requests=[request], errors=[]),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (
            SimpleNamespace(),
            {},
            SimpleNamespace(run_index={"C1": {1: Path("C:/C1_r00001.inf")}}),
            object(),
        ),
    )
    monkeypatch.setattr(
        batching,
        "build_mm_jobs",
        lambda plot_request, _limit: [
            PlotJob(
                mode=PlotMode.MM,
                case_name="C1",
                run_number=index + 1,
                group_label=f"MM_66_{index}",
                output_dir=plot_request.output_dir,
                show_limits=False,
            )
            for index in range(3)
        ],
    )

    def sequential_execute(_renderer, _exporter, job):
        output_dir = Path(job.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / f"{job.group_label}.png"
        output.write_bytes(b"png")
        return plot_execution.PlotExecutionOutput(output)

    monkeypatch.setattr(plot_execution, "execute_plot_job", sequential_execute)
    monkeypatch.setattr(plot_execution, "initialize_plot_process", lambda _run_index: None)
    monkeypatch.setattr(plot_execution, "automatic_plot_worker_count", lambda _job_count: 2)

    class BrokenProcessExecutor:
        def __init__(self, **_kwargs):
            pass

        def submit(self, _function, _job):
            future = Future()
            future.set_exception(BrokenProcessPool("child process terminated"))
            return future

        def shutdown(self, **_kwargs):
            return None

    monkeypatch.setattr(analysis_engine, "ProcessPoolExecutor", BrokenProcessExecutor)
    logs = []

    analysis_engine._render_plot_batches_direct(
        project,
        [ScopeEntry.full()],
        ["SFO"],
        log=logs.append,
    )

    output_dir = project / "Plots" / "Generated" / "Full" / "SFO"
    assert len(list(output_dir.glob("*.png"))) == 3
    assert "retrying sequentially" in "\n".join(logs)
    assert not list(output_dir.parent.glob(".SFO.render-*"))


def test_renderer_waveform_cache_is_bounded_by_bytes(tmp_path) -> None:
    import numpy as np

    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer
    from pscad_plotter_app_v3.services.waveform_io import WaveformFrame

    renderer = MatplotlibRenderer({})
    frame = WaveformFrame(np.zeros((2, 2), dtype=float))
    oversized = WaveformFrame(np.zeros((3, 3), dtype=float))
    frame_bytes = frame.values.nbytes
    renderer.MAX_OUT_CACHE_BYTES = frame_bytes * 2

    renderer._remember_out_frame(tmp_path / "first.out", frame)
    renderer._remember_out_frame(tmp_path / "second.out", frame)
    renderer._remember_out_frame(tmp_path / "third.out", frame)

    assert list(renderer._out_file_cache) == [tmp_path / "second.out", tmp_path / "third.out"]
    assert renderer._out_cache_bytes == frame_bytes * 2

    renderer._remember_out_frame(tmp_path / "oversized.out", oversized)

    assert tmp_path / "oversized.out" not in renderer._out_file_cache
    assert renderer._out_cache_bytes == frame_bytes * 2


def test_parallel_plot_failure_preserves_previous_output_and_cleans_stage(tmp_path, monkeypatch) -> None:
    from concurrent.futures import Future
    from pathlib import Path
    from types import SimpleNamespace

    import pytest

    from pscad_plotter_app_v3.models import PlotJob, PlotMode
    from pscad_plotter_app_v3.services import plot_execution
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services import batching
    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_SFO.xlsx"
    batch.parent.mkdir(parents=True)
    batch.write_bytes(b"batch")
    generated = project / "Plots" / "Generated" / "Full" / "SFO"
    generated.mkdir(parents=True)
    previous = generated / "previous.png"
    previous.write_bytes(b"good")

    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=False,
        case_name="C1",
        elements=["MM_66_A"],
    )
    monkeypatch.setattr(
        BatchExcelService,
        "load_requests",
        lambda *_args, **_kwargs: SimpleNamespace(requests=[request], errors=[]),
    )
    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (
            SimpleNamespace(),
            {},
            SimpleNamespace(run_index={"C": {1: Path("C:/C_r00001.inf")}}),
            object(),
        ),
    )
    monkeypatch.setattr(
        batching,
        "build_mm_jobs",
        lambda plot_request, _limit: [
            PlotJob(
                mode=PlotMode.MM,
                case_name="C1",
                run_number=1,
                group_label=f"MM_66_{index}",
                output_dir=plot_request.output_dir,
                show_limits=False,
            )
            for index in range(10)
        ],
    )

    def failing_execute(job):
        if job.group_label == "MM_66_5":
            raise RuntimeError("render failed")
        output_dir = Path(job.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / f"{job.group_label}.png"
        output.write_bytes(b"partial")
        return plot_execution.PlotExecutionOutput(output)

    monkeypatch.setattr(plot_execution, "initialize_plot_process", lambda _run_index: None)
    monkeypatch.setattr(plot_execution, "execute_plot_job_in_process", failing_execute)
    monkeypatch.setattr(plot_execution, "automatic_plot_worker_count", lambda _job_count: 2)

    class ImmediateProcessExecutor:
        def __init__(self, **_kwargs):
            pass

        def submit(self, function, job):
            future = Future()
            try:
                future.set_result(function(job))
            except Exception as exc:
                future.set_exception(exc)
            return future

        def shutdown(self, **_kwargs):
            return None

    monkeypatch.setattr(analysis_engine, "ProcessPoolExecutor", ImmediateProcessExecutor)

    with pytest.raises(RuntimeError, match="Plot rendering failed"):
        analysis_engine._render_plot_batches_direct(
            project,
            [ScopeEntry.full()],
            ["SFO"],
        )

    assert previous.read_bytes() == b"good"
    assert not list(generated.parent.glob(".SFO.render-*"))


def test_successful_render_directory_commit_removes_obsolete_files(tmp_path) -> None:
    from results_analysis_app.analysis_engine import _replace_generated_directory

    output = tmp_path / "SFO"
    output.mkdir()
    (output / "obsolete.png").write_bytes(b"old")
    stage = tmp_path / ".SFO.render-test"
    stage.mkdir()
    (stage / "current.png").write_bytes(b"new")

    _replace_generated_directory(stage, output)

    assert sorted(path.name for path in output.iterdir()) == ["current.png"]
    assert not stage.exists()


def test_successful_render_directory_commit_retries_transient_lock(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import analysis_engine

    output = tmp_path / "SFO"
    output.mkdir()
    (output / "obsolete.png").write_bytes(b"old")
    stage = tmp_path / ".SFO.render-test"
    stage.mkdir()
    (stage / "current.png").write_bytes(b"new")

    original_replace = Path.replace
    attempts = 0

    def flaky_replace(source, target):
        nonlocal attempts
        if source == stage and attempts == 0:
            attempts += 1
            raise PermissionError("transient Windows lock")
        return original_replace(source, target)

    monkeypatch.setattr(Path, "replace", flaky_replace)

    analysis_engine._replace_generated_directory(stage, output)

    assert attempts == 1
    assert sorted(path.name for path in output.iterdir()) == ["current.png"]
    assert not stage.exists()


def test_plot_render_cancellation_is_not_wrapped_as_failure(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    import pytest

    from pscad_plotter_app_v3.models import PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services import batching
    from results_analysis_app import analysis_engine
    from results_analysis_app.background import OperationCancelled
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch = project / "Plots" / "Plot_batch" / "batch_paste_Full_SFO.xlsx"
    batch.parent.mkdir(parents=True)
    batch.write_bytes(b"batch")
    request = SimpleNamespace(
        mode=PlotMode.MM,
        output_dir="",
        voltage_kv=66.0,
        show_limits=False,
        case_name="C1",
        elements=["MM_66_A"],
    )
    monkeypatch.setattr(
        BatchExcelService,
        "load_requests",
        lambda *_args, **_kwargs: SimpleNamespace(requests=[request], errors=[]),
    )
    monkeypatch.setattr(
        batching,
        "build_mm_jobs",
        lambda plot_request, _limit: [
            SimpleNamespace(
                case_name="C1",
                group_label="MM_66_A",
                run_number=1,
                excel_export=False,
                output_dir=plot_request.output_dir,
            )
        ],
    )

    class CancelledRenderer:
        def render(self, _job):
            raise OperationCancelled()

    monkeypatch.setattr(
        analysis_engine,
        "_load_embedded_plotter_session",
        lambda *_args, **_kwargs: (SimpleNamespace(), {}, CancelledRenderer(), object()),
    )

    with pytest.raises(OperationCancelled):
        analysis_engine.render_plot_batches(project, [ScopeEntry.full()], ["SFO"])


def test_resonance_batches_are_authoritative_for_current_settings(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine, resonance_checks
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch_dir = project / "Plots" / "Plot_batch"
    batch_dir.mkdir(parents=True)
    disabled_event = resonance_checks.plot_event_name("Late_Growth", "LGp")
    stale = batch_dir / f"batch_paste_Full_{disabled_event}.xlsx"
    stale.write_bytes(b"stale")
    written = []

    def fake_write(path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"current")
        written.append((path.name, rows))

    monkeypatch.setattr(analysis_engine, "_create_batch_workbook", fake_write)

    analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        [],
        [],
        resonance_settings={"enabled_checks": ["Post_Event_Stress"]},
    )

    expected = {
        f"batch_paste_Full_{resonance_checks.plot_event_name('Post_Event_Stress', voltage_type)}.xlsx"
        for voltage_type in resonance_checks.VOLTAGE_TYPES
    }
    assert {name for name, _rows in written} == expected
    assert all(rows == [] for _name, rows in written)
    assert not stale.exists()


def test_event_batch_creation_does_not_touch_resonance_batches(tmp_path, monkeypatch) -> None:
    from results_analysis_app import analysis_engine, resonance_checks
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    batch_dir = project / "Plots" / "Plot_batch"
    batch_dir.mkdir(parents=True)
    event_name = resonance_checks.plot_event_name("Late_Growth", "LGp")
    resonance_batch = batch_dir / f"batch_paste_Full_{event_name}.xlsx"
    resonance_batch.write_bytes(b"existing analysis batch")
    monkeypatch.setattr(
        analysis_engine,
        "_create_batch_workbook",
        lambda path, _rows: path.write_bytes(b"event batch"),
    )

    analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        [],
        ["SFO"],
        resonance_settings=None,
    )

    assert resonance_batch.read_bytes() == b"existing analysis batch"


def test_inf_case_run_parser_accepts_uppercase_run_marker(tmp_path) -> None:
    from pscad_plotter_app_v3.services.project_conventions import case_run_from_inf_path

    path = tmp_path / "Case_A_R00012.inf"

    assert case_run_from_inf_path(path) == ("Case_A", 12)


def test_mm_batch_workbook_contains_only_supported_sheet(tmp_path) -> None:
    from openpyxl import load_workbook

    from results_analysis_app.analysis_engine import _create_batch_workbook

    path = tmp_path / "batch.xlsx"
    _create_batch_workbook(path, [{"case": "C1", "run": 1, "element": "MM_66_A"}])
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ["MM"]
    finally:
        workbook.close()


def test_embedded_plotter_session_loads_minimal_mm_catalog(tmp_path) -> None:
    from results_analysis_app.analysis_engine import _load_embedded_plotter_session

    project = tmp_path / "Project"
    case_root = project / "Case_folder"
    results = project / "Results"
    case_root.mkdir(parents=True)
    results.mkdir()
    inf_path = case_root / "C1_r00001.inf"
    inf_path.write_text(
        'PGB(1) Output Desc="MM_LGp_a" Group="MM_66_A" Units="kV"',
        encoding="utf-8",
    )
    (results / "MM results.csv").write_text(
        "Case name,Run#,Bus voltage [kV],Bus name,Fault_type\n"
        "C1,1,66,MM_66_A,AG\n",
        encoding="utf-8",
    )

    check_cancel = lambda: None
    catalog, limits, renderer, exporter = _load_embedded_plotter_session(
        project,
        check_cancel=check_cancel,
    )

    assert catalog.runs_by_case == {"C1": [1]}
    assert [(record.element_name, record.available_runs) for record in catalog.mm_elements] == [
        ("MM_66_A", [1])
    ]
    assert limits == {}
    assert renderer.resolve_inf_path("C1", 1) == inf_path
    assert renderer._check_cancel is check_cancel
    assert exporter.renderer is renderer


def test_project_discovery_finds_inf_files_case_insensitively(tmp_path) -> None:
    from pscad_plotter_app_v3.services.project import ProjectDiscoveryService

    case_root = tmp_path / "Case_folder"
    case_root.mkdir()
    expected = case_root / "C1_R00001.INF"
    expected.write_text("", encoding="utf-8")
    (case_root / "not_an_inf.txt").write_text("", encoding="utf-8")

    context = ProjectDiscoveryService().discover(tmp_path)

    assert context.inf_paths_by_dir[case_root] == [expected]


def test_embedded_mm_renderer_and_excel_export(tmp_path) -> None:
    import numpy as np
    from openpyxl import load_workbook

    from pscad_plotter_app_v3.models import PlotJob, PlotMode
    from pscad_plotter_app_v3.services.exporter import ExcelExporter
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    case_root = tmp_path / "C1.if18"
    case_root.mkdir()
    inf_path = case_root / "C1_r00001.inf"
    descriptions = ["MM_LGp_a", "MM_LGp_b", "MM_LGp_c", "MM_LLp_a", "MM_LLp_b", "MM_LLp_c"]
    inf_path.write_text(
        "\n".join(
            f'PGB({index}) Output Desc="{description}" Group="MM_66_A" Units="kV"'
            for index, description in enumerate(descriptions, start=1)
        ),
        encoding="utf-8",
    )
    time = np.arange(0.0, 0.101, 0.001)
    phases = [120.0 * np.sin(2 * np.pi * 50 * time + phase) for phase in (0.0, 2.1, 4.2)]
    values = np.column_stack([time, *phases, *(1.2 * phase for phase in phases)])
    np.savetxt(case_root / "C1_r00001_01.out", values, header="time a b c ab bc ca", comments="")

    renderer = MatplotlibRenderer({"C1": {1: inf_path}})
    job = PlotJob(
        mode=PlotMode.MM,
        case_name="C1",
        run_number=1,
        group_label="MM_66_A",
        output_dir=str(tmp_path / "plots"),
        trace_type="Both",
        show_limits=False,
        excel_export=True,
        voltage_kv=66.0,
    )

    image_path = renderer.render(job)
    workbook_path = ExcelExporter(renderer).export(job)

    assert image_path.is_file() and image_path.stat().st_size > 1000
    workbook = load_workbook(workbook_path, read_only=True)
    try:
        assert workbook.sheetnames == ["LGp", "LLp"]
    finally:
        workbook.close()


def test_renderer_reads_selected_columns_once_and_reuses_group_frame(tmp_path, monkeypatch) -> None:
    import numpy as np

    from pscad_plotter_app_v3.services import renderer as renderer_module
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    case_root = tmp_path / "C1.if18"
    case_root.mkdir()
    inf_path = case_root / "C1_r00001.inf"
    inf_path.write_text(
        "\n".join(
            f'PGB({index}) Output Desc="MM_LGp_{phase}" Group="MM_66_A" Units="kV"'
            for index, phase in enumerate(("a", "b", "c"), start=1)
        ),
        encoding="utf-8",
    )
    out_path = case_root / "C1_r00001_01.out"
    time = np.arange(0.0, 0.005, 0.001)
    values = np.column_stack([time, time + 1.0, time + 2.0, time + 3.0, time + 4.0])
    np.savetxt(out_path, values, header="time a b c unused", comments="")

    calls = []
    original_loader = renderer_module.load_out_columns

    def load_columns(path, columns, check_cancel=None):
        calls.append((path, tuple(columns)))
        return original_loader(path, columns, check_cancel)

    monkeypatch.setattr(renderer_module, "load_out_columns", load_columns)

    renderer = MatplotlibRenderer({"C1": {1: inf_path}})
    descriptors = renderer.load_inf_descriptors(inf_path)
    first = renderer.load_standardized_group_frame(inf_path, descriptors, "MM_66_A", "LGp")
    second = renderer.load_standardized_group_frame(inf_path, descriptors, "MM_66_A", "LGp")

    assert first is second
    assert first.columns == ["Time (s)", "V_a", "V_b", "V_c"]
    assert calls == [(out_path, (0, 1, 2, 3))]


def test_renderer_parses_each_statistic_file_once_for_multiple_runs(tmp_path, monkeypatch) -> None:
    from pscad_plotter_app_v3.services import renderer as renderer_module
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    stat_path = tmp_path / "Statistic_01.out"
    stat_path.write_text("placeholder", encoding="utf-8")
    find_calls = 0
    parse_calls = 0

    def fake_find(directory):
        nonlocal find_calls
        find_calls += 1
        assert directory == tmp_path
        return stat_path

    def fake_parse(path):
        nonlocal parse_calls
        parse_calls += 1
        assert path == stat_path
        return [
            {"run_number": 1, "fault_label": "AG"},
            {"run_number": 2, "fault_label": "AB"},
        ]

    monkeypatch.setattr(renderer_module, "find_stat_file", fake_find)
    monkeypatch.setattr(renderer_module, "parse_stat_rows", fake_parse)
    renderer = MatplotlibRenderer({})

    assert renderer._find_stat_file(tmp_path) == stat_path
    assert renderer._find_stat_file(tmp_path) == stat_path
    assert renderer._load_event_info(stat_path, 1)["fault_label"] == "AG"
    assert renderer._load_event_info(stat_path, 2)["fault_label"] == "AB"
    assert renderer._load_event_info(stat_path, 1)["fault_label"] == "AG"
    assert find_calls == 1
    assert parse_calls == 1
