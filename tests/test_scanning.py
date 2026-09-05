from __future__ import annotations


def test_voltage_configs_read_mm_blocks(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.project_config import load_voltage_configs

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MM_blocks"
    sheet.append(["Group", "Un", "Um"])
    sheet.append(["MM_161_A", 161, 170.0])
    sheet.append(["MM_330_B", 330, 362.0])
    workbook.save(tmp_path / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    configs = load_voltage_configs(tmp_path)

    assert configs["161"].bus_prefix == "MM_161"
    assert configs["161"].um == 170.0
    assert configs["330"].bus_prefix == "MM_330"
    assert configs["330"].um == 362.0


def test_project_frequency_reads_input_data_b16(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.project_config import load_project_frequency

    workbook = Workbook()
    workbook.active.title = "Input_Data"
    workbook["Input_Data"]["B16"] = "60 Hz"
    workbook.save(tmp_path / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    assert load_project_frequency(tmp_path) == 60.0


def test_shared_float_parser_rejects_nonfinite_values() -> None:
    from results_analysis_app.common import as_float

    assert as_float(" 12.5 ") == 12.5
    assert as_float("") is None
    assert as_float("nan") is None
    assert as_float("inf") is None
    assert as_float("inf", finite=False) == float("inf")


def test_voltage_configs_missing_workbook_has_no_legacy_fallback(tmp_path) -> None:
    from results_analysis_app.project_config import load_voltage_configs

    assert load_voltage_configs(tmp_path) == {}


def test_voltage_configs_use_inf_prefix_with_um_override(tmp_path) -> None:
    from results_analysis_app.project_config import discover_voltage_prefixes, load_voltage_configs

    case_dir = tmp_path / "Case_folder" / "C1.if15_x86"
    case_dir.mkdir(parents=True)
    (case_dir / "C1_r00001.inf").write_text(
        '1 Desc="MM_330_LGp_A" Group="MM_330_BUS1"\n',
        encoding="utf-8",
    )

    assert discover_voltage_prefixes(tmp_path) == {"330": "MM_330"}
    configs = load_voltage_configs(tmp_path, um_overrides={"330": 362.0})

    assert configs["330"].bus_prefix == "MM_330"
    assert configs["330"].um == 362.0


def test_project_scan_voltage_sources_ignore_results_csv_and_outputs(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app import scanner

    project = tmp_path / "Project"
    case_dir = project / "Case_folder" / "C1.if18"
    case_dir.mkdir(parents=True)
    (case_dir / "C1_r00001.inf").write_text(
        '1 Desc="MM_66_LGp_A" Group="MM_66_BUS1"\n',
        encoding="utf-8",
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MM_blocks"
    sheet.append(["Group", "Un", "Um"])
    sheet.append(["MM_161_A", 161, 170.0])
    workbook.save(project / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    results = project / "Results"
    results.mkdir()
    (results / "MM results.csv").write_text(
        "Case name,Bus name,Bus voltage [kV],Run#\nC1,MM_999_A,999,1\n",
        encoding="utf-8",
    )
    general_plots = results / "General Plots"
    general_plots.mkdir()
    (general_plots / "MM_888_dummy.png").write_bytes(b"not used")

    scan = scanner.scan_project(project)

    assert scan.available_voltages == ["66", "161"]


def test_project_scan_caches_existing_um_and_sdpf_limits(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app import scanner

    project = tmp_path / "Project"
    (project / "Case_folder").mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MM_blocks"
    sheet.append(["Group", "Un", "Um", "SDPF_LG", "SDPF_LL"])
    sheet.append(["MM_161_A", 161, 170.0, 325.0, 325.0])
    workbook.save(project / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    scan = scanner.scan_project(project)

    assert scan.voltage_configs["161"].um == 170.0
    assert scan.sustained_sdpf_limits["161"].rms("LGp") == 325.0
    assert scan.sustained_sdpf_limits["161"].peak("LLp") == 325.0 * 2**0.5


def test_project_scan_ignores_corrupt_envelope_workbook(tmp_path) -> None:
    from results_analysis_app import scanner

    project = tmp_path / "Project"
    (project / "Case_folder").mkdir(parents=True)
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    (envelope_dir / "MM_66_with_combined_plot.xlsx").write_bytes(b"not an xlsx")

    scan = scanner.scan_project(project)

    assert scan.exists is True
    assert scan.has_envelopes is True
    assert any("Could not read high-voltage proposals" in message for message in scan.messages)


def test_cold_project_scan_populates_dashboard_figure_catalog(tmp_path, monkeypatch) -> None:
    from results_analysis_app import scanner

    figure = scanner.DashboardFigure(
        id="Dashboard.xlsx|Graphs|1|Initial voltages",
        workbook="Dashboard.xlsx",
        sheet="Graphs",
        chart_index=1,
        title="Initial voltages",
    )
    monkeypatch.setattr(scanner, "scan_dashboard_figures", lambda _path: ([figure], []))

    project = tmp_path / "Project"
    project.mkdir()
    scan = scanner.scan_project(project)

    assert scan.dashboard_figures == [figure]


def test_high_voltage_proposals_prefer_base_envelope_workbook(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.scanner import _collect_high_voltage_exclusions

    envelope_dir = tmp_path / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    for name, case in (
        ("MM_66.xlsx", "BaseCase"),
        ("MM_66_with_combined_plot.xlsx", "StaleCombinedCase"),
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "High voltage exclusions"
        sheet.append(["Case", "Run", "MM_name", "Signal"])
        sheet.append([case, 1, "MM_66_A", "V_a"])
        workbook.save(envelope_dir / name)
        workbook.close()

    rows, warnings = _collect_high_voltage_exclusions(tmp_path)

    assert warnings == []
    assert [row.case for row in rows] == ["BaseCase"]
    assert rows[0].source == "Analysis"
    assert rows[0].excluded is True


def test_high_voltage_proposals_mark_log_and_analysis_sources(tmp_path, monkeypatch) -> None:
    from results_analysis_app import scanner

    scan = scanner.ProjectScan(
        path=tmp_path,
        exists=True,
        high_voltage_exclusions=[
            scanner.HighVoltageExclusion(
                voltage="66",
                case="C1",
                run=1,
                bus="MM_66_A",
                signal="Va",
                source=scanner.HIGH_VOLTAGE_SOURCE_ANALYSIS,
            ),
            scanner.HighVoltageExclusion(
                voltage="66",
                case="C1",
                run=1,
                bus="MM_66_A",
                signal="Vb",
                source=scanner.HIGH_VOLTAGE_SOURCE_ANALYSIS,
            ),
        ],
    )
    log_rows = [
        scanner.HighVoltageExclusion(
            voltage="66",
            case="C1",
            run=1,
            bus="MM_66_A",
            source=scanner.HIGH_VOLTAGE_SOURCE_PSCAD_LOG,
        ),
        scanner.HighVoltageExclusion(
            voltage="230",
            case="C2",
            run=2,
            bus="MM_230_B",
            source=scanner.HIGH_VOLTAGE_SOURCE_PSCAD_LOG,
        ),
    ]
    monkeypatch.setattr(
        scanner,
        "high_voltage_exclusions_from_measurements",
        lambda *_args, **_kwargs: (log_rows, []),
    )

    warnings = scanner.refresh_high_voltage_log_exclusions(scan, 5.0)

    assert warnings == []
    assert [row.source for row in scan.high_voltage_exclusions] == [
        "Both",
        "Both",
        "PSCAD log",
    ]
    assert "High voltage proposals: 2" in scan.chips


def test_scan_cache_deserializer_tolerates_null_collections(tmp_path) -> None:
    from results_analysis_app.project_scan_cache import _deserialize_scan

    scan = _deserialize_scan(
        str(tmp_path),
        {
            "case_infos": None,
            "nonconv_cases": None,
            "high_voltage_exclusions": None,
            "high_voltage_log_measurements": None,
            "chips": None,
            "messages": None,
            "available_voltages": None,
        },
    )

    assert scan is not None
    assert scan.case_infos == []
    assert scan.messages == []


def test_project_fault_types_read_one_statistic_file(tmp_path, monkeypatch) -> None:
    from results_analysis_app import voltage_envelope

    project = tmp_path / "Project"
    case_root = project / "Case_folder" / "C1.if18"
    case_root.mkdir(parents=True)
    statistic = case_root / "Statistic_0001.out"
    statistic.write_text(
        " Multiple Run Output File\n"
        "  Run #               T_sw         FLT_type          Tswitch       Fault_type       Fault_time\n"
        "    1          10.20000000         1             10.20000000         1             10.10000000\n"
        "    2          10.20083300         4             10.20083300         4             10.10083300\n",
        encoding="utf-8",
    )
    calls = []
    original_read = voltage_envelope._read_stat_file

    def counted_read(path):
        calls.append(path)
        return original_read(path)

    monkeypatch.setattr(voltage_envelope, "_read_stat_file", counted_read)

    assert voltage_envelope.read_project_fault_types(project) == {1: "AG", 2: "ABG"}
    assert calls == [statistic]


def test_project_scan_cache_reuses_unchanged_scan_and_invalidates_changed_inputs(tmp_path) -> None:
    import json

    from results_analysis_app import project_scan_cache, scanner
    from results_analysis_app.project_config import VoltageConfig
    from results_analysis_app.sustained_sdpf import SDPFVoltageLimits

    project = tmp_path / "Project"
    case_root = project / "Case_folder"
    case_root.mkdir(parents=True)
    inf_path = case_root / "C1_r00001.inf"
    inf_path.write_text('1 Desc="MM_66_LGp_A" Group="MM_66_BUS1"\n', encoding="utf-8")
    project_path = str(project.resolve())
    scan = scanner.ProjectScan(
        path=project.resolve(),
        exists=True,
        chips=["Ready"],
        case_infos=[scanner.CaseInfo(name="C1", inf_path=inf_path)],
        fault_types_by_run={1: "AG"},
        high_voltage_exclusions=[
            scanner.HighVoltageExclusion(
                voltage="66",
                case="C1",
                run=1,
                bus="MM_66_BUS1",
                source=scanner.HIGH_VOLTAGE_SOURCE_BOTH,
                excluded=True,
            )
        ],
        available_voltages=["66"],
        voltage_configs={
            "66": VoltageConfig("66", "MM_66", 72.0),
        },
        sustained_sdpf_limits={
            "66": SDPFVoltageLimits(66.0, 140.0, 140.0),
        },
        sustained_sdpf_limit_warnings=["example warning"],
    )
    cache_path = tmp_path / "project_scan_cache.json"

    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        (450.0, 250.0),
        cache_path,
    )
    cache_text = cache_path.read_text(encoding="utf-8")
    assert cache_text == json.dumps(
        json.loads(cache_text),
        separators=(",", ":"),
    ) + "\n"
    cached = project_scan_cache.cached_scan(
        project_scan_cache.load(cache_path),
        project_path,
        (450.0, 250.0),
    )

    assert cached is not None
    assert cached.available_voltages == ["66"]
    assert cached.fault_types_by_run == {1: "AG"}
    assert cached.voltage_configs["66"].um == 72.0
    assert cached.sustained_sdpf_limits["66"].rms("LGp") == 140.0
    assert cached.sustained_sdpf_limit_warnings == ["example warning"]
    assert cached.high_voltage_exclusions[0].source == "Both"
    assert cached.high_voltage_exclusions[0].excluded is True
    inf_path.write_text(inf_path.read_text(encoding="utf-8") + "changed\n", encoding="utf-8")

    assert (
        project_scan_cache.cached_scan(
            project_scan_cache.load(cache_path),
            project_path,
            (450.0, 250.0),
        )
        is None
    )


def test_project_scan_cache_round_trips_dashboard_figure_catalog(tmp_path) -> None:
    from results_analysis_app import project_scan_cache, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    figure = scanner.DashboardFigure(
        id="Dashboard.xlsx|Graphs|1|Initial voltages",
        workbook="Dashboard.xlsx",
        sheet="Graphs",
        chart_index=1,
        title="Initial voltages",
    )
    scan = scanner.ProjectScan(
        path=project.resolve(),
        exists=True,
        chips=["Ready"],
        has_dashboards=True,
        dashboard_figures=[figure],
    )
    cache_path = tmp_path / "project_scan_cache.json"

    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        (450.0, 250.0),
        cache_path,
    )

    cached = project_scan_cache.cached_scan(
        project_scan_cache.load(cache_path),
        project_path,
        (450.0, 250.0),
    )

    assert cached is not None
    assert cached.dashboard_figures == [figure]


def test_project_scan_manifest_ignores_unrelated_waveform_files(tmp_path) -> None:
    from results_analysis_app import scanner

    project = tmp_path / "Project"
    run_dir = project / "Case_folder" / "CaseA"
    run_dir.mkdir(parents=True)
    tracked = {
        run_dir / "CaseA_r00001.inf",
        run_dir / "Statistic1.out",
        run_dir / "CB_1.out",
    }
    for path in tracked:
        path.write_text("tracked\n", encoding="utf-8")
    (run_dir / "CaseA_r00001_01.out").write_text("waveform\n", encoding="utf-8")
    (run_dir / "~$Statistic2.out").write_text("temporary\n", encoding="utf-8")

    manifest = scanner.project_scan_manifest(project)
    manifest_paths = {str(item["path"]) for item in manifest["files"]}

    assert manifest_paths == {
        path.relative_to(project).as_posix()
        for path in tracked
    }


def test_output_only_changes_update_cached_status_without_core_rescan(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import project_scan_runner, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    calls = 0

    def fake_scan(path, **_kwargs):
        nonlocal calls
        calls += 1
        return scanner.ProjectScan(path=Path(path).resolve(), exists=True, chips=["Missing Results"])

    monkeypatch.setattr(project_scan_runner.scanner, "scan_project", fake_scan)
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    )
    generated = project / "Plots" / "Generated" / "Full" / "SFO"
    generated.mkdir(parents=True)
    (generated / "plot.png").write_bytes(b"plot")

    refreshed = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    ).scans[project_path]

    assert calls == 1
    assert refreshed.has_plots is True
    assert "Plots exist" in refreshed.chips


def test_envelope_change_refreshes_hv_proposals_without_core_rescan(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from results_analysis_app import project_scan_runner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    envelope = project / "Voltage_envelope" / "Full" / "MM_66.xlsx"
    envelope.parent.mkdir(parents=True)

    def write_proposal(case):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "High voltage exclusions"
        sheet.append(["Case", "Run", "MM_name", "Signal"])
        sheet.append([case, 1, "MM_66_A", "V_a"])
        workbook.save(envelope)
        workbook.close()

    write_proposal("OldCase")
    first = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    ).scans[project_path]
    assert [row.case for row in first.high_voltage_exclusions] == ["OldCase"]
    write_proposal("NewCaseWithLongerName")
    monkeypatch.setattr(
        project_scan_runner.scanner,
        "scan_project",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("envelope-only changes must not rescan core project data")
        ),
    )

    refreshed = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    ).scans[project_path]

    assert [row.case for row in refreshed.high_voltage_exclusions] == ["NewCaseWithLongerName"]


def test_cache_miss_hot_start_and_force_each_build_manifest_once(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import project_scan_runner, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    manifest_calls = 0
    original_manifest = scanner.project_scan_manifest

    def count_manifest(path):
        nonlocal manifest_calls
        manifest_calls += 1
        return original_manifest(path)

    monkeypatch.setattr(project_scan_runner.scanner, "project_scan_manifest", count_manifest)
    monkeypatch.setattr(
        project_scan_runner.scanner,
        "scan_project",
        lambda path, **_kwargs: scanner.ProjectScan(
            path=Path(path).resolve(),
            exists=True,
            chips=["Ready"],
        ),
    )

    project_scan_runner.scan_projects_cached(
        [project_path], project_path, 1.0, 1.0, cache_path=cache_path
    )
    project_scan_runner.scan_projects_cached(
        [project_path], project_path, 1.0, 1.0, cache_path=cache_path
    )
    project_scan_runner.scan_projects_cached(
        [project_path], project_path, 1.0, 1.0, force=True, cache_path=cache_path
    )

    assert manifest_calls == 3


def test_post_action_cache_update_preserves_input_manifest(tmp_path, monkeypatch) -> None:
    from results_analysis_app import project_scan_cache, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    scan = scanner.ProjectScan(path=project.resolve(), exists=True, chips=["Ready"])
    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        (1.0, 1.0),
        cache_path,
    )
    generated = project / "Plots" / "Generated" / "Full" / "SFO"
    generated.mkdir(parents=True)
    (generated / "plot.png").write_bytes(b"plot")
    scan.has_plots = True

    monkeypatch.setattr(
        project_scan_cache.scanner,
        "project_scan_manifest",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("post-action cache writes must not rescan core inputs")
        ),
    )
    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        (1.0, 1.0),
        cache_path,
        preserve_input_manifest=True,
    )

    entry = project_scan_cache.load(cache_path)["projects"][project_path]
    assert entry["manifest"]["output_state"]["has_plots"] is True


def test_dashboard_change_refreshes_cached_figures_without_rescanning_project(
    tmp_path,
    monkeypatch,
) -> None:
    from results_analysis_app import project_scan_cache, project_scan_runner, scanner

    project = tmp_path / "Project"
    case_root = project / "Case_folder"
    dashboard_root = project / "Dashboards"
    case_root.mkdir(parents=True)
    dashboard_root.mkdir()
    inf_path = case_root / "C1_r00001.inf"
    inf_path.write_text("signals\n", encoding="utf-8")
    dashboard_path = dashboard_root / "Dashboard.xlsx"
    dashboard_path.write_bytes(b"first")
    project_path = str(project.resolve())
    limits = (450.0, 250.0)
    cache_path = tmp_path / "project_scan_cache.json"
    scan = scanner.ProjectScan(
        path=project.resolve(),
        exists=True,
        chips=["Ready"],
        case_infos=[scanner.CaseInfo(name="C1", inf_path=inf_path)],
        has_dashboards=True,
    )
    project_scan_cache.update_project_scans(
        {project_path: scan},
        [project_path],
        limits,
        cache_path,
    )
    dashboard_path.write_bytes(b"changed dashboard")
    scan_calls = 0
    logs: list[str] = []

    def fake_scan(*_args, **_kwargs):
        nonlocal scan_calls
        scan_calls += 1
        raise AssertionError("Dashboard-only changes must not rescan the project")

    monkeypatch.setattr(project_scan_runner.scanner, "scan_project", fake_scan)
    monkeypatch.setattr(
        project_scan_runner.scanner,
        "scan_dashboard_figures",
        lambda *_args, **_kwargs: ([], []),
    )
    result = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        *limits,
        log=logs.append,
        cache_path=cache_path,
    )
    changed_scan = result.scans[project_path]

    assert scan_calls == 0
    assert changed_scan.dashboard_changed is False
    assert "Dashboards changed" not in changed_scan.chips
    assert any("Refreshing dashboard figures" in message for message in logs)

    project_scan_cache.update_project_scans(
        result.scans,
        [project_path],
        limits,
        cache_path,
    )
    still_current = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        *limits,
        cache_path=cache_path,
    ).scans[project_path]
    assert still_current.dashboard_changed is False
    assert "Dashboards changed" not in still_current.chips


def test_cached_project_scan_runner_skips_hot_scan_and_supports_force(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import project_scan_runner, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "project_scan_cache.json"
    calls = 0
    cache_updates = 0

    def fake_scan(path, **_kwargs):
        nonlocal calls
        calls += 1
        return scanner.ProjectScan(path=Path(path).resolve(), exists=True, chips=["Ready"])

    monkeypatch.setattr(project_scan_runner.scanner, "scan_project", fake_scan)
    original_update = project_scan_runner.project_scan_cache.update_project_scans

    def counting_update(*args, **kwargs):
        nonlocal cache_updates
        cache_updates += 1
        return original_update(*args, **kwargs)

    monkeypatch.setattr(project_scan_runner.project_scan_cache, "update_project_scans", counting_update)
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    )
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        cache_path=cache_path,
    )
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        force=True,
        cache_path=cache_path,
    )

    assert calls == 2
    assert cache_updates == 2


def test_fresh_scan_messages_are_logged_once_but_not_replayed_from_cache(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import project_scan_runner, scanner

    project = tmp_path / "Project"
    project.mkdir()
    project_path = str(project.resolve())
    cache_path = tmp_path / "project_scan_cache.json"
    calls = 0

    def fake_scan(path, **_kwargs):
        nonlocal calls
        calls += 1
        return scanner.ProjectScan(
            path=Path(path).resolve(),
            exists=True,
            chips=["Ready"],
            messages=["Input workbook warning", "Input workbook warning"],
        )

    monkeypatch.setattr(project_scan_runner.scanner, "scan_project", fake_scan)
    first_logs = []
    second_logs = []

    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        log=first_logs.append,
        cache_path=cache_path,
    )
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        log=second_logs.append,
        cache_path=cache_path,
    )

    assert calls == 1
    assert sum("Input workbook warning" in message for message in first_logs) == 1
    assert not any("Input workbook warning" in message for message in second_logs)


def test_dashboard_figure_selection_uses_only_the_requested_project() -> None:
    from types import SimpleNamespace

    from results_analysis_app.main_window import MainWindow
    from results_analysis_app.models import AppSession, DashboardFigure, ProjectEntry

    selected_without_figures = r"C:\Project\O3"
    selected_with_figures = r"C:\Project\O1"
    unselected_with_figures = r"C:\Project\O2"
    session = AppSession.default()
    session.projects = [
        ProjectEntry(selected_without_figures, selected=True),
        ProjectEntry(selected_with_figures, selected=True),
        ProjectEntry(unselected_with_figures, selected=False),
    ]
    session.dashboard_figure_selection_by_project = {
        selected_with_figures: ["id-1"],
        unselected_with_figures: ["id-2"],
    }
    session.dashboard_figure_apply_to_all = False
    dashboard_figures = {
        selected_with_figures: [DashboardFigure("id-1", "A.xlsx", "Graphs", 1, "Chart")],
        unselected_with_figures: [DashboardFigure("id-2", "B.xlsx", "Graphs", 1, "Chart")],
    }
    window = SimpleNamespace(session=session, dashboard_figures=dashboard_figures)

    assert MainWindow._dashboard_figure_ids_for_project(window, selected_without_figures) == []
    assert MainWindow._dashboard_figure_ids_for_project(window, selected_with_figures) == ["id-1"]
    assert MainWindow._dashboard_figure_ids_for_project(window, unselected_with_figures) == ["id-2"]


def test_pscad_log_high_voltage_scan_uses_raw_waveform_runs(tmp_path) -> None:
    import pytest

    from results_analysis_app import scanner

    project = tmp_path / "Project"
    project.mkdir()
    (project / "PSCAD_log.txt").write_text(
        " WARNING: C1 - MM_66_A has very high values. Please check it. Values have been treated to lower\n",
        encoding="utf-8",
    )
    case_dir = project / "Case_folder" / "C1.if18"
    case_dir.mkdir(parents=True)
    for run, ll_peak in ((1, 260.0), (2, 150.0)):
        stem = f"C1_r{run:05d}"
        (case_dir / f"{stem}.inf").write_text(
            "\n".join(
                [
                    'PGB(1) Output Desc="MM_LGp_a" Group="MM_66_A" Max=2 Min=-2 Units="kV"',
                    'PGB(2) Output Desc="MM_LLp_a" Group="MM_66_A" Max=2 Min=-2 Units="kV"',
                    'PGB(3) Output Desc="MM_LLp_a" Group="MM_66_B" Max=2 Min=-2 Units="kV"',
                ]
            ),
            encoding="utf-8",
        )
        (case_dir / f"{stem}_01.out").write_text(
            "\n".join(
                [
                    "time ch1 ch2 ch3",
                    f"0.0 100.0 {ll_peak} 9999.0",
                    "0.1 120.0 130.0 9999.0",
                ]
            ),
            encoding="utf-8",
        )

    rows, warnings = scanner.scan_high_voltage_from_pscad_log(
        project,
        high_voltage_limit_factor=2.0,
        um_overrides={"66": 72.5},
    )

    assert warnings == []
    assert len(rows) == 1
    assert rows[0].case == "C1"
    assert rows[0].run == 1
    assert rows[0].bus == "MM_66_A"
    assert rows[0].signal == "MM_66_A-LLp_a"
    assert rows[0].file == scanner.PSCAD_LOG_HIGH_VOLTAGE_SOURCE
    assert float(rows[0].max_abs) == pytest.approx(260.0)
    assert float(rows[0].limit) == pytest.approx(2.0 * 72.5 * 2**0.5)


def test_project_open_scans_and_caches_hv_log_maxima_for_factor_changes(
    tmp_path,
    monkeypatch,
) -> None:
    import pytest

    from results_analysis_app import project_scan_runner

    project = tmp_path / "Project"
    case_dir = project / "Case_folder" / "C1.if18"
    case_dir.mkdir(parents=True)
    (project / "PSCAD_log.txt").write_text(
        " WARNING: C1 - MM_66_A has very high values. Please check it. Values have been treated to lower\n",
        encoding="utf-8",
    )
    inf_path = case_dir / "C1_r00001.inf"
    inf_path.write_text(
        'PGB(1) Output Desc="MM_LLp_a" Group="MM_66_A" Units="kV"\n',
        encoding="utf-8",
    )
    out_path = case_dir / "C1_r00001_01.out"
    out_path.write_text("time ch1\n0.0 260.0\n0.1 130.0\n", encoding="utf-8")
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    overrides = {project_path: {"66": 72.5}}

    first = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        3.0,
        overrides,
        cache_path=cache_path,
    ).scans[project_path]

    assert first.high_voltage_exclusions == []
    assert len(first.high_voltage_log_measurements) == 1
    assert first.high_voltage_log_measurements[0].max_abs == pytest.approx(260.0)

    monkeypatch.setattr(
        project_scan_runner.scanner,
        "scan_project",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("factor changes must not rescan the project")
        ),
    )
    monkeypatch.setattr(
        project_scan_runner.scanner,
        "_scan_high_voltage_measurements_from_pscad_log",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("factor changes must reuse cached maxima")
        ),
    )
    second = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        2.0,
        overrides,
        cache_path=cache_path,
    ).scans[project_path]

    assert len(second.high_voltage_exclusions) == 1
    assert float(second.high_voltage_exclusions[0].max_abs) == pytest.approx(260.0)
    assert float(second.high_voltage_exclusions[0].limit) == pytest.approx(2.0 * 72.5 * 2**0.5)


def test_raw_hv_file_change_refreshes_only_hv_cache(tmp_path, monkeypatch) -> None:
    import os

    from results_analysis_app import project_scan_runner

    project = tmp_path / "Project"
    case_dir = project / "Case_folder"
    case_dir.mkdir(parents=True)
    (project / "PSCAD_log.txt").write_text(
        "WARNING: C1 - MM_66_A has very high values\n",
        encoding="utf-8",
    )
    (case_dir / "C1_r00001.inf").write_text(
        'PGB(1) Output Desc="MM_LLp_a" Group="MM_66_A" Units="kV"\n',
        encoding="utf-8",
    )
    out_path = case_dir / "C1_r00001_01.out"
    out_path.write_text("time ch1\n0.0 220.0\n", encoding="utf-8")
    project_path = str(project.resolve())
    cache_path = tmp_path / "cache.json"
    overrides = {project_path: {"66": 72.5}}
    project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        2.0,
        overrides,
        cache_path=cache_path,
    )

    previous_stat = out_path.stat()
    out_path.write_text("time ch1\n0.0 280.0\n", encoding="utf-8")
    os.utime(
        out_path,
        ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns + 1_000_000),
    )
    monkeypatch.setattr(
        project_scan_runner.scanner,
        "scan_project",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("raw HV changes must not rescan core project data")
        ),
    )
    refreshed = project_scan_runner.scan_projects_cached(
        [project_path],
        project_path,
        450.0,
        250.0,
        2.0,
        overrides,
        cache_path=cache_path,
    ).scans[project_path]

    assert float(refreshed.high_voltage_exclusions[0].max_abs) == 280.0


def test_raw_high_voltage_scan_checks_cancellation_inside_descriptor_loop(tmp_path) -> None:
    import pytest

    from results_analysis_app import scanner

    inf_path = tmp_path / "C1_r00001.inf"
    inf_path.write_text(
        '\n'.join(
            f'PGB({index}) Output Desc="MM_LGp_a" Group="MM_66_A" Units="kV"'
            for index in range(1, 6)
        ),
        encoding="utf-8",
    )
    calls = 0

    def cancel() -> None:
        nonlocal calls
        calls += 1
        if calls >= 3:
            raise RuntimeError("cancelled")

    with pytest.raises(RuntimeError, match="cancelled"):
        scanner._raw_high_voltage_measurements_for_inf(
            inf_path,
            {"MM_66_A"},
            {"MM_66_A": "66"},
            cancel,
        )

    assert calls == 3


def test_cached_project_scan_updates_only_changed_entry(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app import project_scan_runner, scanner

    projects = [tmp_path / "A", tmp_path / "B"]
    for project in projects:
        case_root = project / "Case_folder"
        case_root.mkdir(parents=True)
        (case_root / f"{project.name}_r00001.inf").write_text("one\n", encoding="utf-8")
    paths = [str(project.resolve()) for project in projects]
    cache_path = tmp_path / "cache.json"
    scanned: list[str] = []

    def fake_scan(path, **_kwargs):
        resolved = str(Path(path).resolve())
        scanned.append(resolved)
        return scanner.ProjectScan(path=Path(resolved), exists=True, chips=["Ready"])

    monkeypatch.setattr(project_scan_runner.scanner, "scan_project", fake_scan)
    project_scan_runner.scan_projects_cached(paths, paths[0], 1.0, 1.0, cache_path=cache_path)
    scanned.clear()
    updated: list[list[str]] = []
    original_update = project_scan_runner.project_scan_cache.update_project_scans

    def capture_update(scans, project_paths, limits, path, **kwargs):
        selected = list(project_paths)
        updated.append(selected)
        return original_update(scans, selected, limits, path, **kwargs)

    monkeypatch.setattr(project_scan_runner.project_scan_cache, "update_project_scans", capture_update)
    changed_inf = projects[0] / "Case_folder" / "A_r00001.inf"
    changed_inf.write_text("changed content\n", encoding="utf-8")
    project_scan_runner.scan_projects_cached(paths, paths[0], 1.0, 1.0, cache_path=cache_path)

    assert scanned == [paths[0]]
    assert updated == [[paths[0]]]
