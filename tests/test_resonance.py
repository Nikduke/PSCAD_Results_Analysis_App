from __future__ import annotations


def test_resonance_nominal_voltage_uses_voltage_key_not_um() -> None:
    from results_analysis_app.voltage_envelope import _nominal_voltage_from_key

    assert _nominal_voltage_from_key("230", 245.0) == 230.0
    assert _nominal_voltage_from_key("66", 72.5) == 66.0


def test_rebuild_analysis_charts_updates_existing_resonance_workbook(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from results_analysis_app import actions
    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.resonance_checks import WORKBOOK_NAME

    project = tmp_path / "Project"
    workbook_path = project / "Voltage_envelope" / "Full" / WORKBOOK_NAME
    workbook_path.parent.mkdir(parents=True)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "PES_LGp_66_1"
    ws.append(["Time (s)", "E(t)", "E_s(t)", "Vlim"])
    ws.append([0.0, 1.0, 1.0, 1.5])
    ws.append([0.1, 2.0, 1.8, 1.5])
    ws["F1"] = "Case"
    ws["G1"] = "C1"
    ws["F2"] = "Run"
    ws["G2"] = 4
    ws["F3"] = "MM"
    ws["G3"] = "MM_66_StA"
    ws["F4"] = "Start"
    ws["G4"] = 0.034
    ws["F6"] = "Vlim"
    ws["G6"] = 1.5
    workbook.save(workbook_path)
    workbook.close()

    excel_obj = object()

    class DummyExcel:
        def __enter__(self):
            return excel_obj

        def __exit__(self, _exc_type, _exc, _tb):
            return False

    calls = []

    def fake_create_resonance_check_charts(path, excel, x_max=None, x_major=None):
        calls.append((path, excel, x_max, x_major))
        return True

    monkeypatch.setattr(actions, "excel_app", DummyExcel)
    monkeypatch.setattr(actions, "create_resonance_check_charts", fake_create_resonance_check_charts)
    project_key = str(project.resolve())

    assert actions.rebuild_analysis_charts(
        [project],
        [ScopeEntry.full()],
        envelope_chart_x_max_by_project={project_key: 0.5},
        envelope_chart_x_major_by_project={project_key: 0.05},
    ) == [workbook_path]
    assert calls == [(workbook_path, excel_obj, 0.5, 0.05)]


def test_resonance_post_event_uses_chronological_envelope() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        POST_EVENT_STRESS,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = np.where(time < 0.05, time * 200.0, np.where(time < 0.4, 2.0, 0.5))
    envelope[5] = 10.0
    record = ChronologicalEnvelope("Full", "66", "LGp", "C1", 4, "MM_66_StA", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(POST_EVENT_STRESS,),
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
        release_hold_time=0.02,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert len(results) == 1
    assert results[0].check == POST_EVENT_STRESS
    assert results[0].voltage_type == "LGp"
    assert results[0].case == "C1"
    assert results[0].run == 4
    assert results[0].main_metric > 0


def test_resonance_manual_mode_skips_no_settle() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        LATE_GROWTH,
        NO_SETTLE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = 0.4 + time
    record = ChronologicalEnvelope("Full", "66", "LLp", "C1", 4, "MM_66_StA", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(LATE_GROWTH, NO_SETTLE_GROWTH),
        auto_release=False,
        manual_analysis_start=0.1,
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert {result.check for result in results} == {LATE_GROWTH}
    assert results[0].voltage_type == "LLp"


def test_resonance_no_settle_flags_growing_case_without_release() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        NO_SETTLE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = 0.2 + time
    record = ChronologicalEnvelope("Full", "66", "LGp", "C2", 5, "MM_66_StB", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(NO_SETTLE_GROWTH,),
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert len(results) == 1
    assert results[0].check == NO_SETTLE_GROWTH
    assert results[0].main_metric > 0
    assert abs(results[0].t_start - 0.034) < 1e-9


def test_resonance_growth_requires_positive_sigma() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        LATE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = np.where(time < 0.1, 3.0, 2.0 - 0.2 * time)
    record = ChronologicalEnvelope("Full", "66", "LGp", "C3", 6, "MM_66_StC", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(LATE_GROWTH,),
        auto_release=False,
        manual_analysis_start=0.1,
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    assert analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03}) == []


def test_resonance_workbook_omits_empty_result_sheets_in_manual_mode(tmp_path) -> None:
    from openpyxl import load_workbook

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.resonance_checks import (
        NO_SETTLE_GROWTH,
        POST_EVENT_STRESS,
        ResonanceSettings,
        WORKBOOK_NAME,
        write_workbooks,
    )

    settings = ResonanceSettings(
        enabled_checks=(POST_EVENT_STRESS, NO_SETTLE_GROWTH),
        auto_release=False,
    )

    write_workbooks(tmp_path, [ScopeEntry.full()], [], settings, {"SFO": 0.004, "TOV": 0.03})
    workbook_path = tmp_path / "Voltage_envelope" / "Full" / WORKBOOK_NAME
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        setting_rows = dict(workbook["Settings"].iter_rows(min_row=2, values_only=True))
        assert setting_rows["enabled_checks"] == POST_EVENT_STRESS
        assert workbook.sheetnames == ["Settings"]
    finally:
        workbook.close()


def test_resonance_workbook_writes_only_populated_result_sheets(tmp_path) -> None:
    import numpy as np
    from openpyxl import load_workbook

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        POST_EVENT_STRESS,
        ResonanceSettings,
        WORKBOOK_NAME,
        analyze_records,
        write_workbooks,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = np.where(time < 0.05, time * 200.0, np.where(time < 0.4, 2.0, 0.5))
    envelope[5] = 10.0
    settings = ResonanceSettings(
        enabled_checks=(POST_EVENT_STRESS,),
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
        release_hold_time=0.02,
    )
    results = analyze_records(
        [ChronologicalEnvelope("Full", "66", "LGp", "C1", 4, "MM_66_StA", 1.0, time, envelope)],
        settings,
        {"SFO": 0.004, "TOV": 0.03},
    )

    write_workbooks(tmp_path, [ScopeEntry.full()], results, settings, {"SFO": 0.004, "TOV": 0.03})
    workbook_path = tmp_path / "Voltage_envelope" / "Full" / WORKBOOK_NAME
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        assert "Post_Stress_LGp" in workbook.sheetnames
        assert "Post_Stress_LLp" not in workbook.sheetnames
    finally:
        workbook.close()
