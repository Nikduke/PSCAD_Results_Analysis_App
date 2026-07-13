from __future__ import annotations


def test_embedded_plotter_core_imports() -> None:
    import pscad_plotter_app_v3
    from pscad_plotter_app_v3.models import PlotMode
    from pscad_plotter_app_v3.services.batch_excel import BatchExcelService
    from pscad_plotter_app_v3.services.renderer import MatplotlibRenderer

    assert getattr(pscad_plotter_app_v3, "__embedded_source__", "")
    assert PlotMode.MM.value == "MM Overvoltage"
    assert "MM" in BatchExcelService.SHEET_DEFINITIONS
    assert MatplotlibRenderer.FIGSIZE


def test_waveform_catalog_hashes_and_deduplicates_identical_inf(tmp_path, monkeypatch) -> None:
    from pscad_plotter_app_v3.models import ProjectCatalog
    from pscad_plotter_app_v3.services import project as project_service
    from pscad_plotter_app_v3.services.project import CatalogCache, ResultsCatalogService, WaveformCatalogService

    inf_content = 'PGB(1) Desc="MM_66_LGp_A" Group="MM_66_BUS1" Units="kV"\n'
    inf_one = tmp_path / "C1_r00001.inf"
    inf_two = tmp_path / "C2_r00001.inf"
    inf_one.write_text(inf_content, encoding="utf-8")
    inf_two.write_text(inf_content, encoding="utf-8")

    original_parser = project_service.parse_inf_descriptors
    parse_count = 0

    def counting_parser(path):
        nonlocal parse_count
        parse_count += 1
        return original_parser(path)

    monkeypatch.setattr(project_service, "parse_inf_descriptors", counting_parser)
    cache = CatalogCache(tmp_path / "state")
    try:
        run_index = {"C1": {1: inf_one}, "C2": {1: inf_two}}
        ResultsCatalogService().load_waveform_catalog(
            ProjectCatalog(),
            run_index,
            WaveformCatalogService(),
            cache,
        )

        assert parse_count == 1
        hashes = cache.load_inf_hashes([inf_one, inf_two])
        assert hashes[inf_one.resolve()] == hashes[inf_two.resolve()]
        inf_two.write_text(inf_content + "changed\n", encoding="utf-8")
        assert inf_two.resolve() not in cache.load_inf_hashes([inf_two])
    finally:
        cache.close()


def test_dashboard_voltage_slicer_mapping() -> None:
    from results_analysis_app.reporting import _dashboard_slicer_value

    assert _dashboard_slicer_value("22") == "23"
    assert _dashboard_slicer_value("66") == "66"
    assert _dashboard_slicer_value(" 230 ") == "230"


def test_plot_report_heading_parses_generated_mm_filename() -> None:
    from pathlib import Path

    from results_analysis_app.reporting import _plot_heading_from_image_path

    heading = _plot_heading_from_image_path(
        Path("C1_S1_66OFT2_MM_66_StA_AG_004_LGp_LLp_3Ph.png")
    )

    assert heading == "Case: C1_S1_66OFT2 | Run: 4 | Element: MM_66_StA | Fault: AG | Trace: LGp & LLp"


def test_envelope_summary_rows_read_sfo_tov_from_llp(tmp_path) -> None:
    import pytest
    from openpyxl import Workbook

    from results_analysis_app.reporting import _envelope_summary_rows

    workbook = Workbook()
    ws = workbook.active
    ws.title = "LGp"
    ws.append(["Time (s)", "Max_all"])
    ws.append([0.004, 999.0])
    ws.append([0.03, 888.0])
    ws.append([0.1, 300.0])
    ll = workbook.create_sheet("LLp")
    ll.append(["Time (s)", "Max_all"])
    ll.append([0.004, 141.421356])
    ll.append([0.03, 200.0])
    ll.append([0.1, 111.0])
    path = tmp_path / "MM_66.xlsx"
    workbook.save(path)
    workbook.close()

    event_times = {"SFO": 0.004, "TOV": 0.03, "SA": 0.1}
    rows = _envelope_summary_rows(path, event_times)

    assert [(row.event, row.measurement) for row in rows] == [("SFO", "LLp"), ("TOV", "LLp")]
    assert rows[0].rms_kv == pytest.approx(100.0)

    rows = _envelope_summary_rows(path, event_times, events=["SFO", "TOV", "SA"])

    assert [(row.event, row.measurement) for row in rows] == [
        ("SFO", "LLp"),
        ("TOV", "LLp"),
        ("SA", "LGp"),
    ]
    assert rows[2].peak_kv == pytest.approx(300.0)
    assert rows[2].rms_kv == pytest.approx(300.0 / 2**0.5)


def test_envelope_summary_list_uses_subscripted_units() -> None:
    from docx import Document

    from results_analysis_app.reporting import EnvelopeSummaryRow, _add_envelope_summary_list

    doc = Document()
    _add_envelope_summary_list(
        doc,
        [EnvelopeSummaryRow("TOV", "LLp", 0.03, 200.0, 141.4)],
    )

    bullet = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name == "List Bullet")
    assert [run.text for run in bullet.runs] == ["TOV: 200", " kV", "peak", " (141.4", " kV", "RMS", ")"]
    assert bullet.runs[2].font.subscript is True
    assert bullet.runs[5].font.subscript is True


def test_voltage_configs_read_mm_blocks(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.project_config import load_voltage_configs

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MM_blocks"
    sheet.append(["Group", "Un", "Um"])
    sheet.append(["MM_161_A", 161, 170.0])
    sheet.append(["MM_330_B", 330, 362.0])
    workbook.save(tmp_path / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    configs = load_voltage_configs(tmp_path)

    assert configs["161"].bus_prefix == "MM_161"
    assert configs["161"].um == 170.0
    assert configs["330"].bus_prefix == "MM_330"
    assert configs["330"].um == 362.0


def test_project_frequency_reads_input_data_b16(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.project_config import load_project_frequency

    workbook = Workbook()
    workbook.active.title = "Input_Data"
    workbook["Input_Data"]["B16"] = "60 Hz"
    workbook.save(tmp_path / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    assert load_project_frequency(tmp_path) == 60.0


def test_voltage_configs_missing_workbook_has_no_legacy_fallback(tmp_path) -> None:
    from results_analysis_app.project_config import load_voltage_configs

    assert load_voltage_configs(tmp_path) == {}


def test_voltage_configs_use_inf_prefix_with_um_override(tmp_path) -> None:
    from results_analysis_app.project_config import discover_voltage_prefixes, load_voltage_configs

    case_dir = tmp_path / "Case_folder" / "C1.if15_x86"
    case_dir.mkdir(parents=True)
    (case_dir / "C1_r00001.inf").write_text(
        '1 Desc="MM_330_LGp_A" Group="MM_330_BUS1"\n',
        encoding="utf-8",
    )

    assert discover_voltage_prefixes(tmp_path) == {"330": "MM_330"}
    configs = load_voltage_configs(tmp_path, um_overrides={"330": 362.0})

    assert configs["330"].bus_prefix == "MM_330"
    assert configs["330"].um == 362.0


def test_project_scan_voltage_sources_ignore_results_csv_and_outputs(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app import scanner

    project = tmp_path / "Project"
    case_dir = project / "Case_folder" / "C1.if18"
    case_dir.mkdir(parents=True)
    (case_dir / "C1_r00001.inf").write_text(
        '1 Desc="MM_66_LGp_A" Group="MM_66_BUS1"\n',
        encoding="utf-8",
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MM_blocks"
    sheet.append(["Group", "Un", "Um"])
    sheet.append(["MM_161_A", 161, 170.0])
    workbook.save(project / "Input_Data_PSCAD_Python_v11.xlsx")
    workbook.close()

    results = project / "Results"
    results.mkdir()
    (results / "MM results.csv").write_text(
        "Case name,Bus name,Bus voltage [kV],Run#\nC1,MM_999_A,999,1\n",
        encoding="utf-8",
    )
    general_plots = results / "General Plots"
    general_plots.mkdir()
    (general_plots / "MM_888_dummy.png").write_bytes(b"not used")

    scan = scanner.scan_project(project)

    assert scan.available_voltages == ["66", "161"]


def test_project_scan_ignores_corrupt_envelope_workbook(tmp_path) -> None:
    from results_analysis_app import scanner

    project = tmp_path / "Project"
    (project / "Case_folder").mkdir(parents=True)
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    (envelope_dir / "MM_66_with_combined_plot.xlsx").write_bytes(b"not an xlsx")

    scan = scanner.scan_project(project)

    assert scan.exists is True
    assert scan.has_envelopes is True


def test_resonance_nominal_voltage_uses_voltage_key_not_um() -> None:
    from results_analysis_app.voltage_envelope import _nominal_voltage_from_key

    assert _nominal_voltage_from_key("230", 245.0) == 230.0
    assert _nominal_voltage_from_key("66", 72.5) == 66.0


def test_rebuild_envelope_charts_uses_existing_workbooks(tmp_path, monkeypatch) -> None:
    from results_analysis_app import actions
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)
    input_path = envelope_dir / "MM_66.xlsx"
    input_path.write_bytes(b"placeholder")

    class DummyExcelApp:
        def __enter__(self):
            return object()

        def __exit__(self, _exc_type, _exc, _tb):
            return False

    calls = []

    def fake_create_combined(input_file, output_file, excel, **kwargs):
        calls.append((input_file, output_file, excel, kwargs))
        output_file.write_bytes(b"chart")
        return output_file

    monkeypatch.setattr(actions, "excel_app", DummyExcelApp)
    monkeypatch.setattr(actions, "create_combined_envelope_plot", fake_create_combined)

    outputs = actions.rebuild_envelope_charts(
        [project],
        [ScopeEntry.full()],
        ["66", "230"],
        event_times={"SFO": 0.005},
        envelope_chart_x_max=0.7,
        envelope_chart_x_major=0.1,
        envelope_chart_y_limits_by_voltage={"66": {"y_min": 40.0}},
        envelope_chart_show_sa_label=True,
        envelope_chart_top_left_cell="J2",
        envelope_chart_width=800.0,
        envelope_chart_height=400.0,
    )

    assert outputs == [envelope_dir / "MM_66_with_combined_plot.xlsx"]
    assert len(calls) == 1
    assert calls[0][0] == input_path
    assert calls[0][1] == envelope_dir / "MM_66_with_combined_plot.xlsx"
    assert calls[0][3]["axis_limits_override"] == {"x_max": 0.7, "x_major": 0.1}
    assert calls[0][3]["axis_limits_by_voltage"] == {"66": {"y_min": 40.0}}
    assert calls[0][3]["event_times"] == {"SFO": 0.005}
    assert calls[0][3]["show_sa_label"] is True
    assert calls[0][3]["chart_top_left_cell"] == "J2"
    assert calls[0][3]["chart_size"] == {"width": 800.0, "height": 400.0}


def test_dashboard_figure_list_falls_back_to_scanned_selected_projects() -> None:
    from types import SimpleNamespace

    from results_analysis_app.main_window import MainWindow
    from results_analysis_app.models import AppSession, DashboardFigure, ProjectEntry

    selected_without_figures = r"C:\Project\O3"
    selected_with_figures = r"C:\Project\O1"
    unselected_with_figures = r"C:\Project\O2"
    session = AppSession.default()
    session.projects = [
        ProjectEntry(selected_without_figures, selected=True),
        ProjectEntry(selected_with_figures, selected=True),
        ProjectEntry(unselected_with_figures, selected=False),
    ]
    dashboard_figures = {
        selected_with_figures: [DashboardFigure("id-1", "A.xlsx", "Graphs", 1, "Chart")],
        unselected_with_figures: [DashboardFigure("id-2", "B.xlsx", "Graphs", 1, "Chart")],
    }
    window = SimpleNamespace(session=session, dashboard_figures=dashboard_figures)

    assert MainWindow._dashboard_figure_source_paths(window, None) == [selected_with_figures]
    assert MainWindow._dashboard_figure_source_paths(window, selected_without_figures) == [selected_with_figures]
    assert MainWindow._dashboard_figure_source_paths(window, selected_with_figures) == [selected_with_figures]


def test_pscad_log_high_voltage_scan_uses_raw_waveform_runs(tmp_path) -> None:
    import pytest

    from results_analysis_app import scanner

    project = tmp_path / "Project"
    project.mkdir()
    (project / "PSCAD_log.txt").write_text(
        " WARNING: C1 - MM_66_A has very high values. Please check it. Values have been treated to lower\n",
        encoding="utf-8",
    )
    case_dir = project / "Case_folder" / "C1.if18"
    case_dir.mkdir(parents=True)
    for run, ll_peak in ((1, 260.0), (2, 150.0)):
        stem = f"C1_r{run:05d}"
        (case_dir / f"{stem}.inf").write_text(
            "\n".join(
                [
                    'PGB(1) Output Desc="MM_LGp_a" Group="MM_66_A" Max=2 Min=-2 Units="kV"',
                    'PGB(2) Output Desc="MM_LLp_a" Group="MM_66_A" Max=2 Min=-2 Units="kV"',
                    'PGB(3) Output Desc="MM_LLp_a" Group="MM_66_B" Max=2 Min=-2 Units="kV"',
                ]
            ),
            encoding="utf-8",
        )
        (case_dir / f"{stem}_01.out").write_text(
            "\n".join(
                [
                    "time ch1 ch2 ch3",
                    f"0.0 100.0 {ll_peak} 9999.0",
                    "0.1 120.0 130.0 9999.0",
                ]
            ),
            encoding="utf-8",
        )

    rows, warnings = scanner.scan_high_voltage_from_pscad_log(
        project,
        high_voltage_limit_factor=2.0,
        um_overrides={"66": 72.5},
    )

    assert warnings == []
    assert len(rows) == 1
    assert rows[0].case == "C1"
    assert rows[0].run == 1
    assert rows[0].bus == "MM_66_A"
    assert rows[0].signal == "MM_66_A-LLp_a"
    assert rows[0].file == scanner.PSCAD_LOG_HIGH_VOLTAGE_SOURCE
    assert float(rows[0].max_abs) == pytest.approx(260.0)
    assert float(rows[0].limit) == pytest.approx(2.0 * 72.5 * 2**0.5)


def test_create_plot_batches_reads_base_envelope_workbook(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from results_analysis_app import analysis_engine
    from results_analysis_app.models import ScopeEntry

    project = tmp_path / "Project"
    envelope_dir = project / "Voltage_envelope" / "Full"
    envelope_dir.mkdir(parents=True)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "LLp"
    ws.append(["Time (s)", "Case_all", "Run_all", "MM_name_all", "Fault_type_all"])
    ws.append([0.004, "C1", 4, "MM_66_StA", "AG"])
    ws.append([0.03, "C2", 5, "MM_66_StB", "AB"])
    lg = workbook.create_sheet("LGp")
    lg.append(["Time (s)", "Case_all", "Run_all", "MM_name_all", "Fault_type_all"])
    lg.append([0.1, "C3", 6, "MM_66_StC", "ABC"])
    workbook.save(envelope_dir / "MM_66.xlsx")
    workbook.close()

    (envelope_dir / "MM_66_with_combined_plot.xlsx").write_bytes(b"not an xlsx")

    written = []

    def fake_create_batch(path, rows):
        written.append((path.name, rows))
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("ok", encoding="utf-8")

    monkeypatch.setattr(analysis_engine, "_create_batch_workbook", fake_create_batch)

    outputs = analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        ["66"],
        ["SFO", "TOV", "SA"],
        {"SFO": 0.004, "TOV": 0.03, "SA": 0.1},
    )

    assert [path.name for path in outputs] == [
        "batch_paste_Full_SFO.xlsx",
        "batch_paste_Full_TOV.xlsx",
        "batch_paste_Full_SA.xlsx",
    ]
    rows_by_name = {name: rows for name, rows in written}
    assert rows_by_name["batch_paste_Full_SFO.xlsx"][0]["case"] == "C1"
    assert rows_by_name["batch_paste_Full_TOV.xlsx"][0]["case"] == "C2"
    assert rows_by_name["batch_paste_Full_SA.xlsx"][0]["case"] == "C3"


def test_rebuild_analysis_charts_updates_existing_resonance_workbook(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from results_analysis_app import actions
    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.resonance_checks import WORKBOOK_NAME

    project = tmp_path / "Project"
    workbook_path = project / "Voltage_envelope" / "Full" / WORKBOOK_NAME
    workbook_path.parent.mkdir(parents=True)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "PES_LGp_66_1"
    ws.append(["Time (s)", "E(t)", "E_s(t)", "Vlim"])
    ws.append([0.0, 1.0, 1.0, 1.5])
    ws.append([0.1, 2.0, 1.8, 1.5])
    ws["F1"] = "Case"
    ws["G1"] = "C1"
    ws["F2"] = "Run"
    ws["G2"] = 4
    ws["F3"] = "MM"
    ws["G3"] = "MM_66_StA"
    ws["F4"] = "Start"
    ws["G4"] = 0.034
    ws["F6"] = "Vlim"
    ws["G6"] = 1.5
    workbook.save(workbook_path)
    workbook.close()

    excel_obj = object()

    class DummyExcel:
        def __enter__(self):
            return excel_obj

        def __exit__(self, _exc_type, _exc, _tb):
            return False

    calls = []

    def fake_create_resonance_check_charts(path, excel):
        calls.append((path, excel))
        return True

    monkeypatch.setattr(actions, "excel_app", DummyExcel)
    monkeypatch.setattr(actions, "create_resonance_check_charts", fake_create_resonance_check_charts)

    assert actions.rebuild_analysis_charts([project], [ScopeEntry.full()]) == [workbook_path]
    assert calls == [(workbook_path, excel_obj)]


def test_session_event_times_round_trip() -> None:
    from results_analysis_app.models import AppSession

    session = AppSession.default()
    session.event_times["SFO"] = 0.005
    session.envelope_workers = 32
    session.envelope_time_step = 0.001
    session.envelope_time_end = 0.5
    session.envelope_fallback_frequency = 60.0
    session.envelope_chart_x_max = 1.0
    session.envelope_chart_x_major = 0.1
    session.envelope_chart_top_left_cell = "J2"
    session.envelope_chart_width = 800.0
    session.envelope_chart_height = 400.0
    session.envelope_chart_y_limits_by_voltage = {"330": {"y_min": 200.0, "y_max": 800.0, "y_major": 100.0}}
    session.envelope_chart_show_sa_label = True
    session.high_voltage_limit_factor = 4.5
    session.nonconv_cb_iip_limit = 450.0
    session.nonconv_cb_iir_limit = 250.0
    session.resonance_enabled_checks = ["Post_Event_Stress", "Late_Growth"]
    session.resonance_top_n = 2
    session.resonance_limit_multiplier = 1.8
    session.resonance_auto_release = False
    session.resonance_manual_analysis_start = 0.05
    session.resonance_peak_search_fraction = 0.3
    session.resonance_release_decay_ratio = 0.75
    session.resonance_release_rebound_ratio = 0.9
    session.resonance_release_hold_time = 0.04
    session.resonance_rolling_p95_window = 0.03
    session.resonance_rolling_min_samples = 5
    session.resonance_log_floor_vlim_factor = 1e-5
    session.resonance_log_floor_absolute = 1e-3
    session.resonance_growth_window_fraction = 0.25
    session.resonance_min_positive_fraction = 0.7
    session.resonance_min_growth_ratio = 1.1
    session.resonance_min_level_over_vlim = 0.6
    session.resonance_min_growth_delta_factor = 0.02
    project = r"C:\Project"
    session.bus_exclusions_by_project = {project: {"66": ["MM_66_StA", "MM_66_StB"]}}
    session.manual_case_run_exclusions_by_project = {project: [("C5_S1_66OFT2", 4)]}
    session.disabled_nonconv_by_project = {project: [("C7_S1_66OFT2", 8)]}
    session.high_voltage_exclusions_by_project = {project: [("230", "C1", 39, "MM_230_StA")]}
    session.voltage_um_overrides_by_project = {project: {"330": 362.0}}

    loaded = AppSession.from_dict(session.to_dict())

    assert loaded.event_times["TOV"] == 0.03
    assert loaded.event_times["SFO"] == 0.005
    assert loaded.event_times["SA"] == 0.1
    assert loaded.envelope_workers == 32
    assert loaded.envelope_time_step == 0.001
    assert loaded.envelope_time_end == 0.5
    assert loaded.envelope_fallback_frequency == 60.0
    assert loaded.envelope_chart_x_max == 1.0
    assert loaded.envelope_chart_x_major == 0.1
    assert loaded.envelope_chart_top_left_cell == "J2"
    assert loaded.envelope_chart_width == 800.0
    assert loaded.envelope_chart_height == 400.0
    assert loaded.envelope_chart_y_limits_by_voltage["330"] == {"y_min": 200.0, "y_max": 800.0, "y_major": 100.0}
    assert loaded.envelope_chart_show_sa_label is True
    assert loaded.high_voltage_limit_factor == 4.5
    assert loaded.nonconv_cb_iip_limit == 450.0
    assert loaded.nonconv_cb_iir_limit == 250.0
    assert loaded.resonance_enabled_checks == ["Post_Event_Stress", "Late_Growth"]
    assert loaded.resonance_top_n == 2
    assert loaded.resonance_limit_multiplier == 1.8
    assert loaded.resonance_auto_release is False
    assert loaded.resonance_manual_analysis_start == 0.05
    assert loaded.resonance_peak_search_fraction == 0.3
    assert loaded.resonance_release_decay_ratio == 0.75
    assert loaded.resonance_release_rebound_ratio == 0.9
    assert loaded.resonance_release_hold_time == 0.04
    assert loaded.resonance_rolling_p95_window == 0.03
    assert loaded.resonance_rolling_min_samples == 5
    assert loaded.resonance_log_floor_vlim_factor == 1e-5
    assert loaded.resonance_log_floor_absolute == 1e-3
    assert loaded.resonance_growth_window_fraction == 0.25
    assert loaded.resonance_min_positive_fraction == 0.7
    assert loaded.resonance_min_growth_ratio == 1.1
    assert loaded.resonance_min_level_over_vlim == 0.6
    assert loaded.resonance_min_growth_delta_factor == 0.02
    assert loaded.events == ["SFO", "TOV", "SA"]
    assert loaded.bus_exclusions_by_project == {project: {"66": ["MM_66_StA", "MM_66_StB"]}}
    assert loaded.manual_case_run_exclusions_by_project == {project: [("C5_S1_66OFT2", 4)]}
    assert loaded.disabled_nonconv_by_project == {project: [("C7_S1_66OFT2", 8)]}
    assert loaded.high_voltage_exclusions_by_project == {project: [("230", "C1", 39, "MM_230_StA")]}
    assert loaded.voltage_um_overrides_by_project == {project: {"330": 362.0}}


def test_corrupt_autosave_falls_back_to_default(tmp_path, monkeypatch) -> None:
    from results_analysis_app import storage

    autosave_path = tmp_path / "last_session.json"
    autosave_path.write_text("{not valid json", encoding="utf-8")
    monkeypatch.setattr(storage, "AUTOSAVE_PATH", autosave_path)

    session = storage.load_autosave()

    assert session.projects == []
    assert session.scopes[0].mode == "full"


def test_exclusion_normalizers_drop_invalid_values() -> None:
    from results_analysis_app.models import (
        normalize_bus_exclusions,
        normalize_case_run_exclusions,
        normalize_high_voltage_exclusions,
    )

    assert normalize_bus_exclusions(
        {"66.0": [" MM_66_StA ", "", "MM_66_StA"], "bad": ["MM_bad"]}
    ) == {"66": ["MM_66_StA"]}
    assert normalize_case_run_exclusions(
        [{"case": "C1", "run": "4.0"}, ["C1", 4], ["C2", "bad"], ["", 1]]
    ) == [("C1", 4)]
    assert normalize_high_voltage_exclusions(
        [{"voltage": "230.0", "case": "C1", "run": "39", "bus": "MM_230_StA"}, ["", "C2", 1, "MM"]]
    ) == [("230", "C1", 39, "MM_230_StA")]


def test_fast_array_envelope_matches_dataframe_path() -> None:
    import numpy as np
    import pandas as pd

    from results_analysis_app.voltage_envelope import _apply_envelope, _apply_envelope_arrays

    time = np.round(np.arange(0.0, 0.1, 0.0002), 6)
    phases = [
        np.sin(2 * np.pi * 50 * time) * 100,
        np.sin(2 * np.pi * 50 * time + 2.1) * 80,
        np.zeros_like(time),
    ]
    source = pd.DataFrame(
        {
            "Time (s)": time,
            "A": phases[0],
            "B": phases[1],
            "C": phases[2],
        }
    )

    slow = _apply_envelope(source, time_step=0.002, time_end=0.05)
    fast = _apply_envelope_arrays(time, phases, time_step=0.002, time_end=0.05)

    pd.testing.assert_frame_equal(fast, slow)


def test_frequency_detection_tries_later_phases_before_fallback() -> None:
    import numpy as np

    from results_analysis_app.voltage_envelope import _detect_frequency_from_arrays

    time = np.round(np.arange(0.0, 0.1, 0.0002), 6)
    no_crossings = np.ones_like(time)
    valid_60hz = np.sin(2 * np.pi * 60 * time)
    fallback_events: list[str] = []

    detected = _detect_frequency_from_arrays(
        time,
        [no_crossings, valid_60hz],
        50.0,
        "case run bus",
        fallback_events.append,
    )

    assert detected == 60.0
    assert fallback_events == []

    detected = _detect_frequency_from_arrays(
        time,
        [no_crossings],
        50.0,
        "case run bus",
        fallback_events.append,
    )

    assert detected == 50.0
    assert fallback_events == ["case run bus"]


def test_voltage_run_cache_cancels_queued_reads(tmp_path, monkeypatch) -> None:
    import pytest

    from results_analysis_app import voltage_envelope

    paths = [tmp_path / f"C1_r{i:05d}.inf" for i in range(20)]
    state = {"cancelled": False, "calls": 0}

    def check_cancel() -> None:
        if state["cancelled"]:
            raise RuntimeError("Operation stopped by user.")

    def fake_read_run_entries(*args):
        state["calls"] += 1
        state["cancelled"] = True
        args[-1]()
        return [], []

    monkeypatch.setattr(voltage_envelope, "_read_run_entries", fake_read_run_entries)

    with pytest.raises(RuntimeError, match="Operation stopped by user"):
        voltage_envelope._read_voltage_run_cache(
            paths,
            "MM_66",
            72.5,
            set(),
            set(),
            1,
            None,
            check_cancel,
        )

    assert state["calls"] == 1


def test_resonance_post_event_uses_chronological_envelope() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        POST_EVENT_STRESS,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = np.where(time < 0.05, time * 200.0, np.where(time < 0.4, 2.0, 0.5))
    envelope[5] = 10.0
    record = ChronologicalEnvelope("Full", "66", "LGp", "C1", 4, "MM_66_StA", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(POST_EVENT_STRESS,),
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
        release_hold_time=0.02,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert len(results) == 1
    assert results[0].check == POST_EVENT_STRESS
    assert results[0].voltage_type == "LGp"
    assert results[0].case == "C1"
    assert results[0].run == 4
    assert results[0].main_metric > 0


def test_resonance_manual_mode_skips_no_settle() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        LATE_GROWTH,
        NO_SETTLE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = 0.4 + time
    record = ChronologicalEnvelope("Full", "66", "LLp", "C1", 4, "MM_66_StA", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(LATE_GROWTH, NO_SETTLE_GROWTH),
        auto_release=False,
        manual_analysis_start=0.1,
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert {result.check for result in results} == {LATE_GROWTH}
    assert results[0].voltage_type == "LLp"


def test_resonance_no_settle_flags_growing_case_without_release() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        NO_SETTLE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = 0.2 + time
    record = ChronologicalEnvelope("Full", "66", "LGp", "C2", 5, "MM_66_StB", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(NO_SETTLE_GROWTH,),
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    results = analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03})

    assert len(results) == 1
    assert results[0].check == NO_SETTLE_GROWTH
    assert results[0].main_metric > 0
    assert abs(results[0].t_start - 0.034) < 1e-9


def test_resonance_growth_requires_positive_sigma() -> None:
    import numpy as np

    from results_analysis_app.resonance_checks import (
        ChronologicalEnvelope,
        LATE_GROWTH,
        ResonanceSettings,
        analyze_records,
    )

    time = np.round(np.arange(0.0, 1.0, 0.01), 6)
    envelope = np.where(time < 0.1, 3.0, 2.0 - 0.2 * time)
    record = ChronologicalEnvelope("Full", "66", "LGp", "C3", 6, "MM_66_StC", 1.0, time, envelope)
    settings = ResonanceSettings(
        enabled_checks=(LATE_GROWTH,),
        auto_release=False,
        manual_analysis_start=0.1,
        limit_multiplier=1.0,
        rolling_min_samples=1,
        rolling_p95_window=0.01,
    )

    assert analyze_records([record], settings, {"SFO": 0.004, "TOV": 0.03}) == []


def test_resonance_workbook_uses_effective_checks_in_manual_mode(tmp_path) -> None:
    from openpyxl import load_workbook

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app.resonance_checks import (
        NO_SETTLE_GROWTH,
        POST_EVENT_STRESS,
        ResonanceSettings,
        WORKBOOK_NAME,
        write_workbooks,
    )

    settings = ResonanceSettings(
        enabled_checks=(POST_EVENT_STRESS, NO_SETTLE_GROWTH),
        auto_release=False,
    )

    write_workbooks(tmp_path, [ScopeEntry.full()], [], settings, {"SFO": 0.004, "TOV": 0.03})
    workbook_path = tmp_path / "Voltage_envelope" / "Full" / WORKBOOK_NAME
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        setting_rows = dict(workbook["Settings"].iter_rows(min_row=2, values_only=True))
        assert setting_rows["enabled_checks"] == POST_EVENT_STRESS
        assert "No_Settle_LGp" not in workbook.sheetnames
        assert "Post_Stress_LGp" in workbook.sheetnames
    finally:
        workbook.close()
