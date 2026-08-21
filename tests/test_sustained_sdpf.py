from __future__ import annotations


def _cycle_spikes(peaks: list[float], frequency: float = 60.0):
    import numpy as np

    period = 1.0 / frequency
    time = np.arange(0.0, (len(peaks) + 2) * period, 0.0001)
    waveform = np.zeros_like(time)
    for cycle, peak in enumerate(peaks):
        sample_time = (cycle + 0.5) * period
        waveform[np.argmin(np.abs(time - sample_time))] = peak
    return time, waveform


def test_sustained_sdpf_phase_labels_ignore_bus_phase_tokens() -> None:
    from results_analysis_app.sustained_sdpf import phase_labels

    signals = [
        "MM_66_A-LGp_a",
        "MM_66_A-LGp_b",
        "MM_66_A-LGp_c",
    ]
    assert phase_labels("LGp", signals) == ["A-G", "B-G", "C-G"]
    assert phase_labels(
        "LLp",
        [signal.replace("LGp", "LLp") for signal in signals],
    ) == ["A-B", "B-C", "C-A"]


def test_sustained_sdpf_requires_peak_envelope_duration() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    import numpy as np

    time = np.arange(0.0, 0.2, 0.0001)
    amplitude = np.where((time >= 0.005) & (time < 0.1), 2.0, 0.0)
    result = analyze_phase_amplitude(time, amplitude, "LGp", "A-G", 1.0, 0.03, 60.0)

    assert result.longest_margin_s >= 0.03
    assert result.longest_sdpf_s >= 0.03
    assert result.sustained_peak_kv == 2.0
    assert result.sdpf_exceeded
    assert 0.0 < result.sdpf_event_start_s < result.sdpf_event_end_s
    assert result.sdpf_event_end_s - result.sdpf_event_start_s >= 0.03
    assert result.sdpf_excess_area_kv_ms > 0.0
    assert result.sdpf_excess_area_norm_ms > 0.0
    assert result.margin_excess_area_norm_ms > result.sdpf_excess_area_norm_ms


def test_sustained_sdpf_rejects_a_short_event_touching_two_cycle_bins() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    import numpy as np

    time = np.arange(0.0, 0.2, 0.0001)
    amplitude = np.where((time >= 0.011) & (time < 0.031), 2.0, 0.0)
    result = analyze_phase_amplitude(time, amplitude, "LGp", "A-G", 1.0, 0.03, 60.0)

    assert result.peak_kv == 2.0
    assert result.longest_sdpf_s == 0.0
    assert not result.sdpf_exceeded


def test_sustained_sdpf_uses_the_largest_polarity_peak_per_cycle() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    import numpy as np

    frequency = 60.0
    time = np.arange(0.0, 0.3, 0.0001)
    # Positive peaks reach 2.0 while negative excursions reach only -1.2.
    # The positive half-wave alone must qualify each complete cycle.
    waveform = 1.6 * np.sin(2.0 * np.pi * frequency * time) + 0.4
    result = analyze_phase_amplitude(time, waveform, "LGp", "A-G", 1.0, 0.03, frequency)

    assert result.sdpf_exceeded
    assert result.sustained_peak_kv > 1.9


def test_sustained_sdpf_keeps_a_later_event_separate_from_a_short_first_event() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    import numpy as np

    time = np.arange(0.0, 0.25, 0.0001)
    amplitude = np.zeros_like(time)
    amplitude[(time >= 0.005) & (time < 0.020)] = 2.0
    amplitude[(time >= 0.060) & (time < 0.140)] = 2.0
    result = analyze_phase_amplitude(time, amplitude, "LGp", "A-G", 1.0, 0.03, 60.0)

    assert result.sdpf_exceeded
    assert result.sdpf_event_start_s >= 0.05


def test_sustained_sdpf_rejects_an_isolated_peak() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    time, amplitude = _cycle_spikes([2.0, 0.0, 0.0, 0.0])
    result = analyze_phase_amplitude(time, amplitude, "LGp", "A-G", 1.0, 0.03, 60.0)

    assert result.peak_kv == 2.0
    assert result.longest_sdpf_s == 0.0
    assert not result.sdpf_exceeded
    assert result.sustained_peak_kv == 0.0


def test_sustained_sdpf_does_not_double_count_a_short_boundary_crossing() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    frequency = 60.0
    period = 1.0 / frequency
    step = period / 100.0
    time = np.arange(0.0, 5.0 * period + step / 2.0, step)
    waveform = np.zeros_like(time)
    boundary = 2.0 * period
    for sample_time in (boundary - step, boundary, boundary + step):
        waveform[np.argmin(np.abs(time - sample_time))] = 1.3

    result = analyze_phase_amplitude(time, waveform, "LLp", "B-C", 1.0, 0.03, frequency)

    assert result.peak_kv == 1.3
    assert result.longest_margin_s == 0.0
    assert not result.margin_exceeded


def test_sustained_sdpf_rms_diagnostic_does_not_qualify_without_peak() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    time = np.arange(0.0, 0.2, 0.0001)
    waveform = np.full_like(time, 1.2)
    result = analyze_phase_amplitude(time, waveform, "LGp", "A-G", 1.0, 0.03, 60.0)

    assert result.sdpf_rms_exceeded
    assert not result.margin_exceeded
    assert not result.sdpf_exceeded


def test_sustained_sdpf_keeps_peak_and_rms_stress_distinct() -> None:
    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude

    import numpy as np

    time = np.arange(0.0, 0.2, 0.0001)
    amplitude = np.where((time >= 0.005) & (time < 0.1), 1.3, 0.0)
    result = analyze_phase_amplitude(time, amplitude, "LLp", "B-C", 1.0, 0.03, 60.0)

    assert result.peak_kv == 1.3
    assert result.margin_exceeded
    assert not result.sdpf_exceeded
    assert result.sustained_peak_kv == 1.3
    assert result.sustained_rms_kv < result.sustained_peak_kv
    assert result.classification.startswith("SDPF safety margin")


def test_sustained_sdpf_does_not_stitch_nan_gaps_or_phases() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import (
        analyze_phase_amplitude,
        classify_rows,
    )

    time = np.arange(0.0, 0.2, 0.001)
    first = np.where((time >= 0.01) & (time < 0.03), 2.0, 0.0)
    second = np.where((time >= 0.05) & (time < 0.08), 2.0, 0.0)
    phase_a = analyze_phase_amplitude(time, first, "LGp", "A-G", 1.0, 0.2, 60.0)
    phase_b = analyze_phase_amplitude(time, second, "LGp", "B-G", 1.0, 0.2, 60.0)
    case = classify_rows("Full", "66", "C1", 1, "MM_66_A", [phase_a, phase_b], 0.2, cycle_coverage=12)
    assert case is not None
    assert case.governing.phase in {"A-G", "B-G"}
    assert phase_a.longest_margin_s < 0.2
    assert phase_b.longest_margin_s < 0.2


def test_sustained_sdpf_uses_separate_lg_and_ll_limits() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import (
        SDPFVoltageLimits,
        analyze_bus_waveforms,
    )

    time = np.arange(0.0, 0.3, 0.0002)
    values = [np.sin(2 * np.pi * 50.0 * time) * 2.0]
    limits = SDPFVoltageLimits(66.0, 1.0, 3.0)
    lg = analyze_bus_waveforms("LGp", [time], values, ["MM_LGp_a"], limits, "C1", 1, "MM_66_A", 0.03, 50.0)
    ll = analyze_bus_waveforms("LLp", [time], values, ["MM_LLp_a"], limits, "C1", 1, "MM_66_A", 0.03, 50.0)

    assert lg[0]["sdpf_rms_kv"] == 1.0
    assert ll[0]["sdpf_rms_kv"] == 3.0


def test_sustained_sdpf_uses_cigre_divisor_for_safety_threshold() -> None:
    from results_analysis_app.sustained_sdpf import SDPFVoltageLimits

    limits = SDPFVoltageLimits(66.0, 115.0, 230.0)

    assert abs(limits.margin_rms("LGp") - 100.0) < 1e-12
    assert abs(limits.margin_peak("LGp") - 100.0 * 2**0.5) < 1e-12


def test_sustained_sdpf_50_and_60_hz_tracks_are_frequency_aware() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_waveform

    for frequency in (50.0, 60.0):
        time = np.arange(0.0, 0.5, 0.0002)
        waveform = np.sin(2 * np.pi * frequency * time) * 2.0
        result = analyze_waveform(time, waveform, "LGp", "A-G", 1.0, 0.03, frequency)
        assert result.sustained_peak_kv > 1.9
        assert result.sustained_ratio > 1.3


def test_sustained_sdpf_does_not_join_nonconsecutive_violations() -> None:
    from results_analysis_app.sustained_sdpf import analyze_waveform

    time, waveform = _cycle_spikes([1.4, 1.2, 0.5, 0.4, 0.3])

    result = analyze_waveform(time, waveform, "LLp", "B-C", 1.0, 0.03, 60.0)

    assert abs(result.peak_kv - 1.4) < 0.02
    assert result.sustained_peak_kv == 0.0
    assert not result.margin_exceeded


def test_sustained_sdpf_registers_a_later_separate_qualifying_run() -> None:
    from results_analysis_app.sustained_sdpf import analyze_waveform

    time, waveform = _cycle_spikes([1.4, 0.0, 1.3, 1.3, 1.3, 1.3, 0.0])
    result = analyze_waveform(time, waveform, "LLp", "B-C", 1.0, 0.03, 60.0)

    assert result.margin_exceeded
    assert not result.sdpf_exceeded
    assert result.longest_margin_s >= 0.03
    assert abs(result.sustained_peak_kv - 1.3) < 0.02


def test_sustained_sdpf_settings_use_physical_duration() -> None:
    from results_analysis_app.sustained_sdpf import SustainedSDPFSettings

    settings = SustainedSDPFSettings(enabled=True)
    assert settings.duration_ms == 30.0
    assert settings.effective_duration(60.0) == 0.03
    assert SustainedSDPFSettings.from_mapping({"duration_ms": 1.5}).duration_ms == 1.5
    assert SustainedSDPFSettings.from_mapping({"duration_ms": 0}).duration_ms == 30.0
    assert settings.required_cycles(60.0) == 2
    assert settings.required_cycles(50.0) == 2


def test_sustained_sdpf_sustained_t_rejects_a_short_internal_spike() -> None:
    import numpy as np
    import pytest

    from results_analysis_app.sustained_sdpf import _maximum_sustained_envelope_level

    times = np.asarray([0.0, 0.01, 0.02, 0.03])
    values = np.asarray([1.0, 1.3, 1.0, 1.0])

    sustained = _maximum_sustained_envelope_level(times, values, 0.02)

    assert sustained == pytest.approx(1.0, abs=1e-9)


def test_sustained_sdpf_representatives_are_independent_by_population_and_metric() -> None:
    from results_analysis_app.sustained_sdpf import (
        ACTUAL_SDPF_POPULATION,
        CUMULATIVE_STRESS_SELECTION,
        CONTINUOUS_DURATION_SELECTION,
        HIGHEST_VOLTAGE_SUSTAINED_SELECTION,
        PhaseStressResult,
        SAFETY_MARGIN_ONLY_POPULATION,
        SustainedSDPFResult,
        select_representatives,
    )

    def result(case: str, *, actual: bool, t_ratio: float, area: float, duration: float):
        phase = PhaseStressResult(
            "LGp", "A-G", 1.0, 2**0.5, 1.0 / 1.15, 2**0.5 / 1.15,
            0.04, 0.04 if actual else 0.0, 1.5, 1.0, 1.06, "candidate",
            margin_exceeded=True,
            sdpf_exceeded=actual,
            sdpf_excess_area_norm_ms=area if actual else 0.0,
            margin_excess_area_norm_ms=area if not actual else area / 2.0,
            sdpf_longest_continuous_s=duration if actual else 0.0,
            margin_longest_continuous_s=duration if not actual else duration / 2.0,
            sdpf_sustained_t_peak_kv=t_ratio * 2**0.5,
            sdpf_sustained_t_rms_kv=t_ratio,
            sdpf_sustained_t_ratio=t_ratio if actual else 0.0,
            margin_sustained_t_peak_kv=t_ratio * 2**0.5,
            margin_sustained_t_rms_kv=t_ratio,
            margin_sustained_t_ratio=t_ratio if not actual else 0.0,
        )
        return SustainedSDPFResult(
            "Full", "66", case, 1, f"MM_66_{case}", "AG", 0.03, phase, (phase,)
        )

    actual_t = result("ActualT", actual=True, t_ratio=1.20, area=1.0, duration=0.01)
    actual_area = result("ActualArea", actual=True, t_ratio=1.05, area=3.0, duration=0.01)
    actual_duration = result("ActualDuration", actual=True, t_ratio=1.05, area=1.0, duration=0.03)
    margin_t = result("MarginT", actual=False, t_ratio=0.94, area=1.0, duration=0.01)
    margin_area = result("MarginArea", actual=False, t_ratio=0.90, area=3.0, duration=0.01)
    margin_duration = result("MarginDuration", actual=False, t_ratio=0.90, area=1.0, duration=0.03)

    selected = select_representatives(
        [actual_t, actual_area, actual_duration, margin_t, margin_area, margin_duration]
    )

    assert selected[ACTUAL_SDPF_POPULATION][HIGHEST_VOLTAGE_SUSTAINED_SELECTION].case == "ActualT"
    assert selected[ACTUAL_SDPF_POPULATION][CUMULATIVE_STRESS_SELECTION].case == "ActualArea"
    assert selected[ACTUAL_SDPF_POPULATION][CONTINUOUS_DURATION_SELECTION].case == "ActualDuration"
    assert selected[SAFETY_MARGIN_ONLY_POPULATION][HIGHEST_VOLTAGE_SUSTAINED_SELECTION].case == "MarginT"
    assert selected[SAFETY_MARGIN_ONLY_POPULATION][CUMULATIVE_STRESS_SELECTION].case == "MarginArea"
    assert selected[SAFETY_MARGIN_ONLY_POPULATION][CONTINUOUS_DURATION_SELECTION].case == "MarginDuration"


def test_sustained_sdpf_does_not_stitch_nan_gap() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_waveform

    time = np.arange(0.0, 0.12, 0.001)
    waveform = np.full_like(time, 2.0)
    waveform[(time >= 0.04) & (time < 0.06)] = np.nan
    result = analyze_waveform(time, waveform, "LGp", "A-G", 1.0, 0.2, 50.0)

    assert result.longest_margin_s < 0.07
    assert result.classification.startswith("No sustained SDPF")


def test_sustained_sdpf_requires_enough_complete_cycles() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_waveform

    time = np.arange(0.0, 0.41, 0.001)
    values = np.ones_like(time)
    result = analyze_waveform(time, values, "LGp", "A-G", 1.0, 0.6, 10.0)

    assert result.sustained_peak_kv == 0.0
    assert not result.margin_exceeded


def test_sustained_sdpf_retains_no_margin_case_for_reporting() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import analyze_phase_amplitude, classify_rows, select_governing

    time = np.arange(0.0, 0.1, 0.001)
    phase = analyze_phase_amplitude(time, np.full_like(time, 0.1), "LGp", "A-G", 1.0, 0.03)

    result = classify_rows("Full", "66", "C1", 1, "MM_66_A", [phase], 0.03)

    assert result is not None
    assert result.governing.classification.startswith("No sustained SDPF")
    assert select_governing([result]) is None


def test_sustained_sdpf_governing_selection_uses_normalized_excess_area() -> None:
    from results_analysis_app.sustained_sdpf import (
        PhaseStressResult,
        SustainedSDPFResult,
        select_governing,
    )

    def phase(area: float, peak: float) -> PhaseStressResult:
        return PhaseStressResult(
            "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
            0.1, 0.1, peak, peak / 2**0.5, peak / 1.414, "x",
            sdpf_excess_area_kv_ms=area * 1.414,
            sdpf_excess_area_norm_ms=area,
            margin_exceeded=True,
            sdpf_exceeded=True,
        )

    low_limit_high_kv = SustainedSDPFResult(
        "Full", "66", "C1", 1, "MM_66_A", "", 0.03, phase(1.0, 1.7), (phase(1.0, 1.7),)
    )
    high_limit_lower_kv = SustainedSDPFResult(
        "Full", "161", "C2", 1, "MM_161_A", "", 0.03, phase(3.0, 1.5), (phase(3.0, 1.5),)
    )

    assert select_governing([low_limit_high_kv, high_limit_lower_kv]) is high_limit_lower_kv


def test_sustained_sdpf_full_wave_ranking_keeps_lower_opposite_peak_in_envelope() -> None:
    import numpy as np

    from results_analysis_app.sustained_sdpf import (
        _CycleMetrics,
        _full_wave_envelope_metrics,
    )

    metrics = _CycleMetrics(
        starts=np.asarray([0.0, 0.02]),
        peaks=np.asarray([1.3, 1.3]),
        peak_times=np.asarray([0.005, 0.025]),
        positive_peaks=np.asarray([1.3, 1.3]),
        positive_peak_times=np.asarray([0.005, 0.025]),
        negative_peaks=np.asarray([0.8, 0.8]),
        negative_peak_times=np.asarray([0.015, 0.035]),
        rms=np.asarray([1.0, 1.0]),
    )

    area_kv_s, area_norm_s, longest_s = _full_wave_envelope_metrics(
        metrics,
        0,
        2,
        1.0,
    )

    assert area_kv_s == area_norm_s
    assert 0.002 < area_norm_s < 0.003
    assert 0.011 < longest_s < 0.013


def test_sustained_sdpf_has_independent_area_and_duration_selections() -> None:
    from results_analysis_app.sustained_sdpf import (
        PhaseStressResult,
        SustainedSDPFResult,
        select_representatives,
    )

    def result(case: str, area: float, duration_s: float) -> SustainedSDPFResult:
        phase = PhaseStressResult(
            "LGp",
            "A-G",
            1.0,
            1.414,
            0.85,
            1.202,
            0.04,
            0.04,
            1.5,
            1.0,
            1.06,
            "SDPF exceeded",
            margin_exceeded=True,
            sdpf_exceeded=True,
            sdpf_excess_area_norm_ms=area,
            sdpf_longest_continuous_s=duration_s,
        )
        return SustainedSDPFResult(
            "Full",
            "66",
            case,
            1,
            f"MM_66_{case}",
            "",
            0.03,
            phase,
            (phase,),
        )

    area_winner = result("Area", 5.0, 0.010)
    duration_winner = result("Duration", 2.0, 0.025)

    representatives = select_representatives([area_winner, duration_winner])
    assert representatives["actual_sdpf"]["cumulative_stress"].case == "Area"
    assert representatives["actual_sdpf"]["continuous_duration"].case == "Duration"


def test_sustained_sdpf_persists_only_the_extra_duration_selection(tmp_path) -> None:
    import json

    from results_analysis_app.sustained_sdpf import (
        CUMULATIVE_STRESS_SELECTION,
        CONTINUOUS_DURATION_SELECTION,
        PhaseStressResult,
        SustainedSDPFResult,
        SustainedSDPFSettings,
        current_representatives_for_voltage,
        save_results,
    )

    def result(case: str, area: float, duration_s: float) -> SustainedSDPFResult:
        phase = PhaseStressResult(
            "LGp",
            "A-G",
            1.0,
            1.414,
            0.85,
            1.202,
            0.04,
            0.04,
            1.5,
            1.0,
            1.06,
            "SDPF exceeded",
            margin_exceeded=True,
            sdpf_exceeded=True,
            sdpf_excess_area_norm_ms=area,
            sdpf_longest_continuous_s=duration_s,
        )
        return SustainedSDPFResult(
            "Full",
            "66",
            case,
            1,
            f"MM_66_{case}",
            "",
            0.03,
            phase,
            (phase,),
        )

    source = tmp_path / "run.out"
    source.write_text("source", encoding="utf-8")
    stat = source.stat()
    settings = SustainedSDPFSettings(enabled=True, duration_ms=30.0)
    save_results(
        tmp_path,
        "Full",
        settings,
        {"66": result("Area", 5.0, 0.010)},
        signature_inputs={
            "66": {
                "files": [
                    {"path": str(source), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}
                ]
            }
        },
        observations_by_voltage={
            "66": [result("Area", 5.0, 0.010), result("Duration", 2.0, 0.025)]
        },
    )

    payload = json.loads(
        (tmp_path / "Voltage_envelope" / "Full" / "Sustained_SDpf.json").read_text(
            encoding="utf-8"
        )
    )
    assert payload["results"]["66"]["case"] == "Area"
    assert payload["selections"]["66"]["actual_sdpf"][CONTINUOUS_DURATION_SELECTION]["case"] == "Duration"
    assert "result" not in payload["selections"]["66"]["actual_sdpf"][CONTINUOUS_DURATION_SELECTION]
    assert len(payload["representative_results"]["66"]) == 1
    selections, validation = current_representatives_for_voltage(
        payload,
        "66",
        tmp_path,
        settings,
    )

    assert validation.valid
    assert selections["actual_sdpf"][CUMULATIVE_STRESS_SELECTION].case == "Area"
    assert selections["actual_sdpf"][CONTINUOUS_DURATION_SELECTION].case == "Duration"


def test_sustained_sdpf_governing_selection_prioritizes_actual_crossing() -> None:
    from results_analysis_app.sustained_sdpf import (
        PhaseStressResult,
        SustainedSDPFResult,
        select_governing,
    )

    sustained_only = PhaseStressResult(
        "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
        0.04, 0.04, 1.8, 1.8 / 2**0.5, 1.27,
        "SDPF exceeded for at least 30 ms of peak-envelope persistence (peak)",
        margin_exceeded=True, sdpf_exceeded=True,
    )
    short_margin = PhaseStressResult(
        "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
        0.01, 0.0, 0.2, 0.2 / 2**0.5, 0.14, "No SDPF",
        peak_kv=1.3,
    )
    result = lambda case, phase: SustainedSDPFResult(
        "Full", "66", case, 1, "MM_66_A", "", 0.03, phase, (phase,)
    )

    assert select_governing([result("sustained", sustained_only), result("actual", short_margin)]).case == "sustained"


def test_sustained_sdpf_uses_the_worst_path_area_without_summing_paths() -> None:
    from dataclasses import replace

    from results_analysis_app.sustained_sdpf import PhaseStressResult, classify_rows

    phase_a = PhaseStressResult(
        "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
        0.05, 0.05, 1.6, 1.1, 1.13, "actual",
        sdpf_excess_area_kv_ms=1.414,
        sdpf_excess_area_norm_ms=1.0,
        margin_excess_area_kv_ms=2.0,
        margin_excess_area_norm_ms=1.7,
        margin_exceeded=True,
        sdpf_exceeded=True,
    )
    phase_b = replace(
        phase_a,
        phase="B-G",
        sdpf_excess_area_kv_ms=4.242,
        sdpf_excess_area_norm_ms=3.0,
    )

    result = classify_rows(
        "Full", "66", "C1", 1, "MM_66_A", [phase_a, phase_b], 0.03,
        cycle_coverage=2,
    )

    assert result is not None
    assert result.governing.phase == "B-G"
    assert result.governing.sdpf_excess_area_norm_ms == 3.0


def test_sustained_sdpf_cache_recomputes_governing_path_from_phase_population() -> None:
    from dataclasses import replace

    from results_analysis_app.sustained_sdpf import (
        PhaseStressResult,
        SustainedSDPFResult,
    )

    actual = PhaseStressResult(
        "LLp", "B-C", 1.0, 1.414, 0.85, 1.202,
        0.04, 0.04, 1.8, 1.2, 1.27, "actual",
        margin_exceeded=True,
        sdpf_exceeded=True,
        sdpf_excess_area_norm_ms=2.0,
    )
    margin = replace(actual, phase="A-G", sdpf_exceeded=False, sdpf_excess_area_norm_ms=0.0)
    result = SustainedSDPFResult(
        "Full", "66", "C1", 1, "MM_66_A", "AG", 0.03,
        margin,
        (margin, actual),
    )
    payload = result.to_dict()
    assert "governing" not in payload
    assert "signature" not in payload
    payload["governing"] = margin.to_dict()

    loaded = SustainedSDPFResult.from_dict(payload)

    assert loaded.governing.phase == "B-C"
    assert loaded.governing.sdpf_exceeded


def test_sustained_sdpf_bus_analysis_reuses_common_cycle_index(monkeypatch) -> None:
    import numpy as np

    from results_analysis_app import sustained_sdpf
    from results_analysis_app.sustained_sdpf import (
        SDPFVoltageLimits,
        analyze_bus_phase_results,
        analyze_phase_amplitude,
    )

    frequency = 60.0
    time = np.arange(0.0, 0.3, 0.0001)
    values = [
        1.6 * np.sin(2.0 * np.pi * frequency * time + phase)
        for phase in (0.0, 0.8, 1.6)
    ]
    expected = [
        analyze_phase_amplitude(
            time,
            waveform,
            "LGp",
            phase,
            1.0,
            0.03,
            frequency,
        )
        for waveform, phase in zip(values, ("A-G", "B-G", "C-G"))
    ]
    calls = 0
    original = sustained_sdpf._full_cycle_tracks

    def counted(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(sustained_sdpf, "_full_cycle_tracks", counted)
    actual = analyze_bus_phase_results(
        "LGp",
        [time, time, time],
        values,
        ["MM_LGp_a", "MM_LGp_b", "MM_LGp_c"],
        SDPFVoltageLimits(66.0, 1.0, 1.0),
        0.03,
        frequency,
    )

    assert calls == 1
    assert actual == expected


def test_sustained_sdpf_governing_selection_prioritizes_sdpf_candidate_over_margin_only() -> None:
    from results_analysis_app.sustained_sdpf import (
        PhaseStressResult,
        SustainedSDPFResult,
        select_governing,
    )

    margin_only = PhaseStressResult(
        "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
        0.08, 0.01, 2.0, 2.0 / 2**0.5, 1.41,
        "SDPF safety margin exceeded; SDPF not exceeded for at least 30 ms of peak-envelope persistence (peak)",
        margin_exceeded=True,
    )
    actual = PhaseStressResult(
        "LGp", "A-G", 1.0, 1.414, 0.85, 1.202,
        0.04, 0.04, 1.2, 1.2 / 2**0.5, 0.85,
        "SDPF exceeded for at least 30 ms of peak-envelope persistence (peak)",
        margin_exceeded=True,
        sdpf_exceeded=True,
    )

    def result(case: str, phase: PhaseStressResult) -> SustainedSDPFResult:
        return SustainedSDPFResult(
            "Full", "66", case, 1, "MM_66_A", "", 0.03, phase, (phase,)
        )

    selected = select_governing([result("margin", margin_only), result("actual", actual)])

    assert selected is not None
    assert selected.case == "actual"


def test_sustained_sdpf_summary_is_ranked_per_mm_and_keeps_phase_detail(tmp_path) -> None:
    from dataclasses import replace
    from openpyxl import load_workbook

    from results_analysis_app.sustained_sdpf import (
        SDPFVoltageLimits,
        PhaseStressResult,
        SustainedSDPFResult,
        write_summary_workbook,
    )

    def phase(measurement: str, phase_name: str, ratio: float, category: int) -> PhaseStressResult:
        margin = category >= 1
        actual = category >= 2
        return PhaseStressResult(
            measurement,
            phase_name,
            100.0,
            2**0.5 * 100.0,
            100.0 / 1.15,
            2**0.5 * 100.0 / 1.15,
            0.03 if margin else 0.0,
            0.03 if actual else 0.0,
            ratio * 2**0.5 * 100.0,
            ratio * 100.0,
            ratio,
            "SDPF exceeded" if actual else ("SDPF safety margin exceeded" if margin else "No SDPF"),
            peak_kv=ratio * 2**0.5 * 100.0,
            margin_exceeded=margin,
            sdpf_exceeded=actual,
        )

    def result(case: str, mm_name: str, row: PhaseStressResult) -> SustainedSDPFResult:
        return SustainedSDPFResult(
            "Full", "66", case, 1, mm_name, "AG", 0.03, row, (row,)
        )

    first_lg = result("C1", "MM_66_A", phase("LGp", "A-G", 0.90, 1))
    first_ll = replace(
        result("C1", "MM_66_A", phase("LLp", "B-C", 0.70, 0)),
        fault_type="AG",
    )
    second = result("C2", "MM_66_B", phase("LLp", "B-C", 1.20, 2))

    path = write_summary_workbook(
        tmp_path,
        "Full",
        {"66": [first_lg, first_ll, second]},
        {"66": SDPFVoltageLimits(66.0, 100.0, 100.0)},
    )
    workbook = load_workbook(path, read_only=False, data_only=True)
    ranked = list(workbook["Ranked cases"].values)
    header_index = next(
        index
        for index, row in enumerate(ranked)
        if row and row[0] == "Population rank"
    )
    headers = ranked[header_index]
    data_rows = ranked[header_index + 1 :]

    assert data_rows[0][headers.index("Case")] == "C2"
    assert len(data_rows) == 3
    assert any(row[headers.index("Measurement")] == "LLp" for row in data_rows)
    assert "Normalized excess area (pu*ms)" in headers
    assert "Longest continuous duration (ms)" in headers
    assert "Event start (s)" in headers
    assert "Event end (s)" in headers
    assert "Qualifying paths" in headers
    assert data_rows[0][headers.index("Qualifying paths")] == "LLp:B-C"
    assert set(workbook["Ranked cases"].tables) == {"RankedCases"}
    assert workbook["Ranked cases"].auto_filter.ref is None
    context = {
        row[0]: row[1]
        for row in ranked[1:header_index]
        if row and row[0]
    }
    assert context["Result version"] == 17
    assert context["Minimum persistence (ms)"] == 30.0

    representative = list(workbook["Representative selections"].values)
    representative_header_index = next(
        index
        for index, row in enumerate(representative)
        if row and row[0] == "Voltage (kV)"
    )
    representative_headers = representative[representative_header_index]
    representative_rows = representative[representative_header_index + 1 :]
    assert len(representative_rows) == 6
    assert "Selection criterion" in representative_headers
    assert set(workbook["Representative selections"].tables) == {
        "RepresentativeSelections"
    }
    workbook.close()


def test_sustained_sdpf_invalid_limit_rows_are_reported(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.sustained_sdpf import resolve_project_limits

    project = tmp_path / "Project"
    project.mkdir()
    workbook = Workbook()
    workbook.active.title = "MM_blocks"
    sheet = workbook.active
    sheet.append(["Un", "SDPF_LG", "SDPF_LL"])
    sheet.append([66, 1.0, 2.0])
    sheet.append([66, 1.1, 2.0])
    workbook.save(project / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    limits, warnings = resolve_project_limits(project)

    assert "66" not in limits
    assert any("Inconsistent" in warning for warning in warnings)


def test_sustained_sdpf_ignores_blank_mm_blocks_rows(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.sustained_sdpf import resolve_project_limits

    project = tmp_path / "Project"
    project.mkdir()
    workbook = Workbook()
    workbook.active.title = "MM_blocks"
    sheet = workbook.active
    sheet.append(["Un", "SDPF_LG", "SDPF_LL"])
    sheet.append([66, 1.0, 2.0])
    for _ in range(5):
        sheet.append([None, None, None])
    workbook.save(project / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    limits, warnings = resolve_project_limits(project)

    assert set(limits) == {"66"}
    assert warnings == []


def test_sustained_sdpf_invalid_duplicate_limit_skips_voltage(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.sustained_sdpf import resolve_project_limits

    project = tmp_path / "Project"
    project.mkdir()
    workbook = Workbook()
    workbook.active.title = "MM_blocks"
    sheet = workbook.active
    sheet.append(["Un", "SDPF_LG", "SDPF_LL"])
    sheet.append([66, 1.0, 2.0])
    sheet.append([66, "bad", 2.0])
    workbook.save(project / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    limits, warnings = resolve_project_limits(project)

    assert "66" not in limits
    assert any("Invalid SDPF values for 66 kV" in warning for warning in warnings)


def test_sustained_sdpf_manual_limit_overrides_replace_rms_values(tmp_path) -> None:
    from openpyxl import Workbook

    from results_analysis_app.sustained_sdpf import resolve_project_limits

    project = tmp_path / "Project"
    project.mkdir()
    workbook = Workbook()
    workbook.active.title = "MM_blocks"
    sheet = workbook.active
    sheet.append(["Un", "SDPF_LG", "SDPF_LL"])
    sheet.append([66, 140.0, 141.0])
    workbook.save(project / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    limits, warnings = resolve_project_limits(
        project,
        manual_overrides={"66": {"LGp": 150.0}},
    )

    assert warnings == []
    assert limits["66"].source == "manual"
    assert limits["66"].rms("LGp") == 150.0
    assert limits["66"].rms("LLp") == 141.0


def test_sustained_sdpf_result_manifest_marks_source_change_stale(tmp_path) -> None:
    from results_analysis_app.sustained_sdpf import (
        RESULT_VERSION,
        SOURCE_MANIFEST_VERSION,
        make_signature,
        result_inputs_current,
    )

    source = tmp_path / "run.bin"
    source.write_text("before", encoding="utf-8")
    stat = source.stat()
    source_entries = [["run.bin", stat.st_size, stat.st_mtime_ns]]
    payload = {
        "version": RESULT_VERSION,
        "source_manifest": {
            "version": SOURCE_MANIFEST_VERSION,
            "roots": [],
            "include_static": False,
            "extra_files": ["run.bin"],
            "fingerprint": make_signature({"files": source_entries}),
        },
        "signature_inputs": {
            "66": {},
        }
    }

    assert result_inputs_current(payload, "66", tmp_path)
    source.write_text("after", encoding="utf-8")
    assert not result_inputs_current(payload, "66", tmp_path)


def test_sustained_sdpf_cache_validation_reports_settings_mismatch(tmp_path) -> None:
    from results_analysis_app.sustained_sdpf import (
        RESULT_VERSION,
        SustainedSDPFSettings,
        validate_result_cache,
    )

    payload = {
        "version": RESULT_VERSION,
        "settings": {"enabled": True, "duration_ms": 30.0},
        "results": {},
    }

    status = validate_result_cache(
        payload,
        tmp_path,
        SustainedSDPFSettings(enabled=True, duration_ms=40.0),
    )

    assert not status.valid
    assert "settings" in status.reason


def test_sustained_sdpf_cache_validation_reports_stale_source(tmp_path) -> None:
    from results_analysis_app.sustained_sdpf import (
        SustainedSDPFSettings,
        save_results,
        validate_result_cache,
    )

    source = tmp_path / "run.out"
    source.write_text("before", encoding="utf-8")
    stat = source.stat()
    save_results(
        tmp_path,
        "Full",
        SustainedSDPFSettings(enabled=True, duration_ms=30.0),
        {"66": None},
        signature_inputs={
            "66": {
                "files": [
                    {"path": str(source), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}
                ]
            }
        },
    )
    source.write_text("after", encoding="utf-8")

    import json

    payload = json.loads(
        (tmp_path / "Voltage_envelope" / "Full" / "Sustained_SDpf.json").read_text(
            encoding="utf-8"
        )
    )
    status = validate_result_cache(
        payload,
        tmp_path,
        SustainedSDPFSettings(enabled=True, duration_ms=30.0),
    )

    assert not status.valid
    assert "stale" in status.reason


def test_sustained_sdpf_current_result_normalizes_voltage_key(tmp_path) -> None:
    from results_analysis_app import sustained_sdpf

    phase = sustained_sdpf.PhaseStressResult(
        "LLp",
        "A-B",
        100.0,
        100.0 * 2**0.5,
        100.0 / 1.15,
        100.0 * 2**0.5 / 1.15,
        0.04,
        0.0,
        160.0,
        113.0,
        1.6,
        "SDPF safety margin exceeded",
        margin_exceeded=True,
    )
    result = sustained_sdpf.SustainedSDPFResult(
        "Full",
        "66",
        "C1",
        1,
        "MM_66_A",
        "AG",
        0.04,
        phase,
        (phase,),
    )
    source = tmp_path / "run.out"
    source.write_text("source", encoding="utf-8")
    stat = source.stat()
    settings = sustained_sdpf.SustainedSDPFSettings(enabled=True, duration_ms=40.0)
    sustained_sdpf.save_results(
        tmp_path,
        "Full",
        settings,
        {"66": result},
        signature_inputs={
            "66": {
                "files": [
                    {"path": str(source), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}
                ]
            }
        },
    )

    import json

    payload = json.loads(
        (tmp_path / "Voltage_envelope" / "Full" / "Sustained_SDpf.json").read_text(
            encoding="utf-8"
        )
    )
    representatives, status = sustained_sdpf.current_representatives_for_voltage(
        payload,
        "66.0",
        tmp_path,
        settings,
    )

    assert status.valid
    loaded = representatives["safety_margin_only"][sustained_sdpf.CUMULATIVE_STRESS_SELECTION]
    assert loaded is not None
    assert loaded.voltage == "66"


def test_sustained_sdpf_rejects_pre_duration_qualification_cache(tmp_path) -> None:
    import json

    from results_analysis_app.sustained_sdpf import load_results, result_path

    path = result_path(tmp_path, "Full")
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps({"version": 3, "results": {}}), encoding="utf-8")

    assert load_results(tmp_path, "Full") == {}


def test_sustained_sdpf_invalidation_removes_ranked_summary(tmp_path) -> None:
    from results_analysis_app.sustained_sdpf import (
        invalidate_results,
        result_path,
        summary_path,
    )

    result = result_path(tmp_path, "Full")
    summary = summary_path(tmp_path, "Full")
    result.parent.mkdir(parents=True)
    result.write_text("{}", encoding="utf-8")
    summary.write_text("stale", encoding="utf-8")

    invalidate_results(tmp_path)

    assert not result.exists()
    assert not summary.exists()


def test_sustained_sdpf_shared_manifest_validation_can_be_reused(tmp_path, monkeypatch) -> None:
    from pathlib import Path

    from results_analysis_app.sustained_sdpf import (
        RESULT_VERSION,
        SOURCE_MANIFEST_VERSION,
        make_signature,
        result_inputs_current,
        source_manifest_current,
    )

    source = tmp_path / "run.bin"
    source.write_text("source", encoding="utf-8")
    stat = source.stat()
    manifest = [["run.bin", stat.st_size, stat.st_mtime_ns]]
    payload = {
        "version": RESULT_VERSION,
        "source_manifest": {
            "version": SOURCE_MANIFEST_VERSION,
            "roots": [],
            "include_static": False,
            "extra_files": ["run.bin"],
            "fingerprint": make_signature({"files": manifest}),
        },
        "signature_inputs": {"66": {}, "161": {}},
    }

    stat_paths: list[str] = []
    original_stat = Path.stat

    def counted_stat(path, *args, **kwargs):
        stat_paths.append(str(path))
        return original_stat(path, *args, **kwargs)

    monkeypatch.setattr(Path, "stat", counted_stat)
    shared_current = source_manifest_current(payload, tmp_path)

    assert shared_current
    assert result_inputs_current(
        payload,
        "66",
        tmp_path,
        shared_manifest_current=shared_current,
    )
    assert result_inputs_current(
        payload,
        "161",
        tmp_path,
        shared_manifest_current=shared_current,
    )
    assert stat_paths.count(str(source)) == 1


def test_sustained_source_manifest_is_relative_and_scope_limited(tmp_path) -> None:
    from results_analysis_app.voltage_envelope import _sustained_source_manifest

    selected_dir = tmp_path / "Case_folder" / "Selected.if18"
    other_dir = tmp_path / "Case_folder" / "Other.if18"
    selected_dir.mkdir(parents=True)
    other_dir.mkdir(parents=True)
    selected_inf = selected_dir / "Selected_r00001.inf"
    selected_inf.write_text("layout", encoding="utf-8")
    (selected_dir / "Selected_r00001_01.out").write_text("raw", encoding="utf-8")
    (selected_dir / "Statistic_0001.out").write_text("stat", encoding="utf-8")
    (selected_dir / "CB_0001.out").write_text("cb", encoding="utf-8")
    (other_dir / "Other_r00001_01.out").write_text("unselected", encoding="utf-8")

    entries = _sustained_source_manifest(tmp_path, [selected_inf])
    paths = {str(entry["path"]) for entry in entries}

    assert "Case_folder/Selected.if18/Selected_r00001.inf" in paths
    assert "Case_folder/Selected.if18/Selected_r00001_01.out" in paths
    assert "Case_folder/Selected.if18/Statistic_0001.out" in paths
    assert "Case_folder/Selected.if18/CB_0001.out" in paths
    assert "Case_folder/Other.if18/Other_r00001_01.out" not in paths
    assert all(not path.startswith(str(tmp_path)) for path in paths)


def test_sustained_sdpf_persists_one_shared_manifest_for_identical_voltages(tmp_path) -> None:
    import json

    from results_analysis_app.sustained_sdpf import (
        SustainedSDPFSettings,
        result_inputs_current,
        save_results,
    )

    source = tmp_path / "run.out"
    source.write_text("source", encoding="utf-8")
    stat = source.stat()
    manifest = [{"path": str(source), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}]

    save_results(
        tmp_path,
        "Full",
        SustainedSDPFSettings(enabled=True, duration_ms=30.0),
        {},
        signature_inputs={
            "66": {"files": manifest, "heatmap_case_names": ["C"]},
            "161": {"files": list(manifest), "heatmap_case_names": ["C"]},
        },
    )

    payload = json.loads((tmp_path / "Voltage_envelope" / "Full" / "Sustained_SDpf.json").read_text())
    assert "files" not in payload["signature_inputs"]["66"]
    assert payload["source_manifest"]["version"] == 1
    assert payload["source_manifest"]["roots"] == ["."]
    assert isinstance(payload["source_manifest"]["fingerprint"], str)
    assert result_inputs_current(payload, "66", tmp_path)
    assert result_inputs_current(payload, "161", tmp_path)


def test_sustained_sdpf_rejects_legacy_full_source_manifest(tmp_path) -> None:
    from results_analysis_app.sustained_sdpf import RESULT_VERSION, result_inputs_current

    source = tmp_path / "run.out"
    source.write_text("source", encoding="utf-8")
    stat = source.stat()
    payload = {
        "version": RESULT_VERSION,
        "source_manifest": {
            "files": [["run.out", stat.st_size, stat.st_mtime_ns]],
        },
        "signature_inputs": {"66": {}},
    }

    assert not result_inputs_current(payload, "66", tmp_path)


def test_sustained_sdpf_report_section_is_compact() -> None:
    from docx import Document

    from results_analysis_app.reporting import _FigureRegistry, _add_sustained_sdpf_section
    from results_analysis_app.sustained_sdpf import PhaseStressResult, SustainedSDPFResult

    phase = PhaseStressResult(
        "LGp", "A-G", 1.0, 2**0.5, 0.85, 0.85 * 2**0.5,
        0.04, 0.02, 1.5, 1.5 / 2**0.5, 1.5 / 2**0.5,
        "SDPF safety margin exceeded; SDPF not exceeded for at least 30 ms of peak-envelope persistence (peak)",
        margin_exceeded=True,
    )
    result = SustainedSDPFResult(
        "Full", "66", "C1", 1, "MM_66_A", "AG", 0.03, phase, (phase,)
    )
    document = Document()

    _add_sustained_sdpf_section(document, result, "66", None, _FigureRegistry())

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Sustained SDPF" in text
    assert "Candidates require peak-envelope persistence" in text
    assert "Sustained-TOV criterion" not in text
    assert "Maximum qualifying sustained-run" not in text
    assert "Highest qualifying sustained-run" not in text
    assert len(document.tables) == 1
    headers = [cell.text for cell in document.tables[0].rows[0].cells]
    assert headers == [
        "Selection criteria",
        "Case / Run",
        "Fault",
        "MM element",
        "Path",
        "Vₜ (kVₚₑₐₖ / %SDPF)",
        "Area (pu·ms)",
        "Duration (ms)",
    ]
    assert document.tables[0].rows[1].cells[0].text == (
        "Cumulative stress; Longest duration"
    )
    captions = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption"
    ]
    assert any("Table 1-1" in caption and "Safety-margin-only" in caption for caption in captions)


def test_sustained_sdpf_report_merges_criteria_for_one_case() -> None:
    from docx import Document

    from results_analysis_app.reporting import _FigureRegistry, _add_sustained_sdpf_section
    from results_analysis_app.sustained_sdpf import (
        CONTINUOUS_DURATION_SELECTION,
        CUMULATIVE_STRESS_SELECTION,
        HIGHEST_VOLTAGE_SUSTAINED_SELECTION,
        PhaseStressResult,
        SAFETY_MARGIN_ONLY_POPULATION,
        SustainedSDPFResult,
    )

    phase = PhaseStressResult(
        "LLp", "B-C", 1.0, 2**0.5, 0.85, 0.85 * 2**0.5,
        0.03, 0.0, 1.3, 1.3 / 2**0.5, 0.92,
        "SDPF safety margin exceeded; SDPF not exceeded for at least 30 ms of peak-envelope persistence (peak)",
        peak_kv=1.3,
        margin_exceeded=True,
        margin_excess_area_norm_ms=0.7,
        margin_longest_continuous_s=0.04,
        margin_sustained_t_peak_kv=1.3,
        margin_sustained_t_ratio=0.92,
    )
    result = SustainedSDPFResult(
        "Full", "66", "C1", 1, "MM_66_A", "AG", 0.03, phase, (phase,)
    )
    document = Document()

    _add_sustained_sdpf_section(
        document,
        result,
        "66",
        None,
        _FigureRegistry(),
        selection_results_by_population={
            SAFETY_MARGIN_ONLY_POPULATION: {
                HIGHEST_VOLTAGE_SUSTAINED_SELECTION: result,
                CUMULATIVE_STRESS_SELECTION: result,
                CONTINUOUS_DURATION_SELECTION: result,
            }
        },
    )

    table = document.tables[0]
    assert len(table.rows) == 2
    assert table.rows[1].cells[0].text == (
        "Highest sustained Vₜ; Cumulative stress; Longest duration"
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Maximum qualifying sustained-run" not in text
    assert "RMS diagnostic" not in text


def test_sustained_sdpf_report_lists_three_distinct_selected_cases() -> None:
    from dataclasses import replace

    from docx import Document

    from results_analysis_app.reporting import _FigureRegistry, _add_sustained_sdpf_section
    from results_analysis_app.sustained_sdpf import (
        CONTINUOUS_DURATION_SELECTION,
        CUMULATIVE_STRESS_SELECTION,
        HIGHEST_VOLTAGE_SUSTAINED_SELECTION,
        PhaseStressResult,
        SAFETY_MARGIN_ONLY_POPULATION,
        SustainedSDPFResult,
    )

    phase = PhaseStressResult(
        "LGp", "A-G", 1.0, 2**0.5, 0.85, 0.85 * 2**0.5,
        0.03, 0.0, 1.3, 1.3 / 2**0.5, 0.92,
        "candidate",
        margin_exceeded=True,
        margin_excess_area_norm_ms=0.7,
        margin_longest_continuous_s=0.04,
        margin_sustained_t_peak_kv=1.3,
        margin_sustained_t_ratio=0.92,
    )
    result = SustainedSDPFResult(
        "Full", "66", "CaseT", 1, "MM_66_A", "AG", 0.03, phase, (phase,)
    )
    selected = {
        HIGHEST_VOLTAGE_SUSTAINED_SELECTION: replace(result, case="CaseT"),
        CUMULATIVE_STRESS_SELECTION: replace(result, case="CaseArea"),
        CONTINUOUS_DURATION_SELECTION: replace(result, case="CaseDuration"),
    }
    document = Document()

    _add_sustained_sdpf_section(
        document,
        result,
        "66",
        None,
        _FigureRegistry(),
        selection_results_by_population={SAFETY_MARGIN_ONLY_POPULATION: selected},
    )

    table = document.tables[0]
    assert len(table.rows) == 4
    assert [row.cells[1].text for row in table.rows[1:]] == [
        "CaseT / 1",
        "CaseArea / 1",
        "CaseDuration / 1",
    ]
    assert [row.cells[0].text for row in table.rows[1:]] == [
        "Highest sustained Vₜ",
        "Cumulative stress",
        "Longest duration",
    ]


def test_sustained_sdpf_build_reuses_loaded_run_and_persists_json(tmp_path, monkeypatch) -> None:
    import numpy as np
    from openpyxl import Workbook

    from results_analysis_app.models import ScopeEntry
    from results_analysis_app import analysis_engine, sustained_sdpf, voltage_envelope
    from results_analysis_app.voltage_envelope import build_voltage_envelopes

    class DummyExcel:
        def __enter__(self):
            return object()

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(voltage_envelope, "excel_app", DummyExcel)
    monkeypatch.setattr(voltage_envelope, "autofit_workbook", lambda *_args: None)

    project = tmp_path / "Project"
    case_folder = project / "Case_folder"
    case_folder.mkdir(parents=True)
    inf_path = case_folder / "C1_r00001.inf"
    inf_path.write_text(
        "\n".join(
            [
                'PGB(1) Output Desc="MM_LGp_a" Group="MM_66_A" Units="kV"',
                'PGB(2) Output Desc="MM_LGp_b" Group="MM_66_A" Units="kV"',
                'PGB(3) Output Desc="MM_LGp_c" Group="MM_66_A" Units="kV"',
                'PGB(4) Output Desc="MM_LLp_a" Group="MM_66_A" Units="kV"',
                'PGB(5) Output Desc="MM_LLp_b" Group="MM_66_A" Units="kV"',
                'PGB(6) Output Desc="MM_LLp_c" Group="MM_66_A" Units="kV"',
            ]
        ),
        encoding="utf-8",
    )
    time = np.arange(0.0, 0.2, 0.0002)
    amplitude = np.where((time >= 0.03) & (time < 0.1), 2.0, 0.4)
    lines = ["header"]
    for index, time_s in enumerate(time):
        values = [time_s]
        values.extend(float(amplitude[index] * np.sin(2 * np.pi * 50 * time_s + phase)) for phase in (0.0, 2.1, 4.2))
        values.extend(float(amplitude[index] * np.sin(2 * np.pi * 50 * time_s + phase)) for phase in (0.0, 2.1, 4.2))
        lines.append(" ".join(f"{value:.12g}" for value in values))
    inf_path.with_name("C1_r00001_01.out").write_text("\n".join(lines), encoding="utf-8")

    workbook = Workbook()
    input_data = workbook.active
    input_data.title = "Input_Data"
    input_data.append(["Final duration", 0.2])
    mm_blocks = workbook.create_sheet("MM_blocks")
    mm_blocks.append(["Group", "Un", "Um", "SDPF_LG", "SDPF_LL", "SIWL_LG", "SIWL_LL"])
    mm_blocks.append(["MM_66_A", 66.0, 72.5, 1.0, 1.0, 1.0, 1.0])
    workbook.save(project / "Input_Data_PSCAD_test.xlsx")
    workbook.close()

    build_voltage_envelopes(
        project,
        [ScopeEntry.full()],
        ["66"],
        log=None,
        envelope_workers=1,
        build_charts=False,
        envelope_fallback_frequency=60.0,
        sustained_sdpf_settings={"enabled": True, "duration_ms": 30.0},
    )

    result_path = project / "Voltage_envelope" / "Full" / "Sustained_SDpf.json"
    assert result_path.is_file()
    payload = result_path.read_text(encoding="utf-8")
    assert '"66"' in payload
    settings = sustained_sdpf.SustainedSDPFSettings(enabled=True, duration_ms=30.0)
    cache_status = sustained_sdpf.validate_result_cache(
        sustained_sdpf.load_results(project, "Full"),
        project,
        settings,
    )
    assert cache_status.valid
    batch_outputs = analysis_engine.create_plot_batches(
        project,
        [ScopeEntry.full()],
        ["66"],
        [],
        sustained_sdpf_settings=settings.to_mapping(),
        excel_waveform_exports_enabled=False,
    )
    assert project / "Plots" / "Plot_batch" / "batch_paste_Full_Sustained_SDPF.xlsx" in batch_outputs
    summary_path = project / "Voltage_envelope" / "Full" / "Sustained_SDpf_summary.xlsx"
    assert summary_path.is_file()
    from openpyxl import load_workbook

    workbook = load_workbook(summary_path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Ranked cases", "Representative selections"]
    ranked_rows = list(workbook["Ranked cases"].values)
    context = {row[0]: row[1] for row in ranked_rows[1:6] if row and row[0]}
    assert context["Frequency (Hz)"] == 60.0
    assert context["Derived complete-cycle coverage"] == 2
    workbook.close()
